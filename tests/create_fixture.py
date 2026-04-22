"""Run once to create tests/fixtures/sample_template.xlsx."""
import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles import Color


def create_fixture():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # A1:C2 — merged region, bold + colored background
    ws.merge_cells("A1:C2")
    cell = ws["A1"]
    cell.value = "Merged Header"
    cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFFFF")
    cell.fill = PatternFill(patternType="solid", fgColor=Color(rgb="FF003366"))
    cell.alignment = Alignment(horizontal="center", vertical="center")

    # A3 — plain string
    ws["A3"].value = "Label"
    ws["A3"].font = Font(name="Calibri", size=11)

    # B3 — number with format
    ws["B3"].value = 12345.678
    ws["B3"].number_format = "0.00"
    ws["B3"].alignment = Alignment(horizontal="right")

    # C3 — formula
    ws["C3"].value = "=B3*2"

    # A4 — date
    ws["A4"].value = datetime.datetime(2024, 6, 15, 9, 30)
    ws["A4"].number_format = "DD/MM/YYYY HH:MM"

    # B4 — bold + italic + colored font
    ws["B4"].value = "Styled Text"
    ws["B4"].font = Font(bold=True, italic=True, color="FFCC0000", size=12)

    # C4 — solid fill (yellow)
    ws["C4"].value = "Yellow Cell"
    ws["C4"].fill = PatternFill(patternType="solid", fgColor=Color(rgb="FFFFFF00"))

    # A5 — all four borders
    ws["A5"].value = "Bordered"
    ws["A5"].border = Border(
        top=Side(border_style="thick", color=Color(rgb="FF000000")),
        bottom=Side(border_style="thin", color=Color(rgb="FF000000")),
        left=Side(border_style="medium", color=Color(rgb="FF0000FF")),
        right=Side(border_style="dashed", color=Color(rgb="FFFF0000")),
    )

    # B5 — wrapped text + center alignment
    ws["B5"].value = "This is a long text that should be wrapped inside the cell"
    ws["B5"].alignment = Alignment(wrap_text=True, horizontal="center", vertical="top")

    # D1:E3 — another merged region
    ws.merge_cells("D1:E3")
    ws["D1"].value = "Second Merge"
    ws["D1"].fill = PatternFill(patternType="solid", fgColor=Color(rgb="FFCCFFCC"))
    ws["D1"].alignment = Alignment(horizontal="center", vertical="center")

    # Custom column widths
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["B"].width = 18.0
    ws.column_dimensions["C"].width = 15.0

    # Custom row heights
    ws.row_dimensions[1].height = 30.0
    ws.row_dimensions[2].height = 30.0
    ws.row_dimensions[5].height = 50.0

    out = Path(__file__).parent / "fixtures" / "sample_template.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"Fixture written to {out}")


if __name__ == "__main__":
    create_fixture()
