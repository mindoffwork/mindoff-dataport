from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

from mindoff_data_export import build_template_with_data

# §1 Constants & Exceptions

polars = pytest.importorskip("polars", reason="polars not installed")

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _cell(coord: str, value):
    return {
        "coordinate": coord,
        "value": value,
        "cell_type": "string",
        "number_format": None,
        "font": {
            "name": "Calibri",
            "size": 11.0,
            "bold": False,
            "italic": False,
            "underline": None,
            "color": None,
        },
        "fill": {"bg_color": None},
        "alignment": {"horizontal": None, "vertical": None, "wrap_text": False},
        "borders": {
            "top": {"style": None, "color": None},
            "bottom": {"style": None, "color": None},
            "left": {"style": None, "color": None},
            "right": {"style": None, "color": None},
        },
        "merged": False,
        "merge_anchor": None,
    }


def _schema(
    cells: dict[str, dict], *, dims: str = "A1:C3", merges: list[str] | None = None
):
    return {
        "sheets": [
            {
                "name": "Sheet1",
                "dimensions": dims,
                "merged_regions": merges or [],
                "column_widths": {},
                "row_heights": {},
                "cells": cells,
            }
        ]
    }


def _temp_output_path(managed_tmp_dir: Path, name: str) -> Path:
    return managed_tmp_dir / name


# §4 Public Functions


def test_streaming_split_outputs_and_names(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1, 2, 3, 4, 5]})
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    paths = build_template_with_data(
        schema,
        {"Sheet1": {"rows": df}},
        str(out),
        export_mode="streaming",
        streaming_chunk_rows=2,
        max_rows_per_workbook=3,
    )

    assert isinstance(paths, list)
    assert len(paths) == 1
    zip_path = Path(paths[0])
    assert zip_path.name == "filled.zip"
    assert zip_path.exists()
    assert not (managed_tmp_dir / "filled.part001.xlsx").exists()
    assert not (managed_tmp_dir / "filled.part002.xlsx").exists()

    with ZipFile(zip_path) as zip_file:
        assert zip_file.namelist() == ["filled.part001.xlsx", "filled.part002.xlsx"]
        zip_file.extractall(path=managed_tmp_dir)

    part1 = managed_tmp_dir / "filled.part001.xlsx"
    part2 = managed_tmp_dir / "filled.part002.xlsx"
    assert part1.exists()
    assert part2.exists()

    wb1 = openpyxl.load_workbook(part1, data_only=True)
    wb2 = openpyxl.load_workbook(part2, data_only=True)
    ws1 = wb1["Sheet1"]
    ws2 = wb2["Sheet1"]
    assert [ws1["A1"].value, ws1["A2"].value, ws1["A3"].value] == [1, 2, 3]
    assert [ws2["A1"].value, ws2["A2"].value] == [4, 5]
    wb1.close()
    wb2.close()


def test_streaming_single_part_does_not_zip(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1, 2]})
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    paths = build_template_with_data(
        schema,
        {"Sheet1": {"rows": df}},
        str(out),
        export_mode="streaming",
        streaming_chunk_rows=2,
        max_rows_per_workbook=10,
    )

    assert len(paths) == 1
    assert Path(paths[0]).name == "filled.part001.xlsx"
    assert Path(paths[0]).exists()
    assert not (managed_tmp_dir / "filled.zip").exists()


def test_streaming_rejects_hug_mode(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    with pytest.raises(ValueError, match="does not support 'hug'"):
        build_template_with_data(
            schema,
            {"Sheet1": {"rows": df}},
            str(out),
            export_mode="streaming",
            column_width_mode="hug",
        )


def test_streaming_rejects_merged_regions(managed_tmp_dir: Path):
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}")},
        dims="A1:A1",
        merges=["A1:B1"],
    )
    df = polars.DataFrame({"A": [1]})
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    with pytest.raises(ValueError, match="does not support merged cells"):
        build_template_with_data(
            schema, {"Sheet1": {"rows": df}}, str(out), export_mode="streaming"
        )


def test_streaming_anchor_style_is_cloned(managed_tmp_dir: Path):
    anchor = _cell("A1", "{{rows:dataframe-content}}")
    anchor["fill"] = {"bg_color": "FFFF0000"}
    anchor["number_format"] = "0.00"
    schema = _schema({"A1": anchor}, dims="A1:A1")
    df = polars.DataFrame({"A": [1.5]})
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    paths = build_template_with_data(
        schema, {"Sheet1": {"rows": df}}, str(out), export_mode="streaming"
    )
    wb = openpyxl.load_workbook(paths[0])
    cell = wb["Sheet1"]["A1"]
    assert cell.fill.fgColor.type == "rgb"
    assert cell.fill.fgColor.rgb == "FFFF0000"
    assert cell.number_format == "0.00"
    wb.close()


def test_streaming_lazyframe_uses_collect_batches(monkeypatch, managed_tmp_dir: Path):
    lf = polars.DataFrame({"A": [1, 2]}).lazy()
    lazy_cls = type(lf)
    called = {"collect_batches": 0}
    original_batches = lazy_cls.collect_batches
    original_collect = lazy_cls.collect

    def _collect_batches(self, **kwargs):
        called["collect_batches"] += 1
        return original_batches(self, **kwargs)

    def _collect_fail(self, *args, **kwargs):
        raise AssertionError("collect() should not be called in streaming path")

    monkeypatch.setattr(lazy_cls, "collect_batches", _collect_batches)
    monkeypatch.setattr(lazy_cls, "collect", _collect_fail)

    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")
    paths = build_template_with_data(
        schema,
        {"Sheet1": {"rows": lf}},
        str(out),
        export_mode="streaming",
        streaming_chunk_rows=1,
        max_rows_per_workbook=10,
    )

    assert called["collect_batches"] >= 1
    wb = openpyxl.load_workbook(paths[0], data_only=True)
    assert [wb["Sheet1"]["A1"].value, wb["Sheet1"]["A2"].value] == [1, 2]
    wb.close()


def test_streaming_expands_dynamic_sheet_names_in_order(managed_tmp_dir: Path):
    schema = {
        "sheets": [
            {
                "name": "{{sheet_1}}",
                "dimensions": "A1:A1",
                "merged_regions": [],
                "column_widths": {},
                "row_heights": {},
                "cells": {"A1": _cell("A1", "{{name:string}}")},
            }
        ]
    }
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    paths = build_template_with_data(
        schema,
        {
            "sheet_1": {
                "Sheet Name 1": {"name": "Alpha"},
                "Sheet Name 2": {"name": "Beta"},
            }
        },
        str(out),
        export_mode="streaming",
        max_rows_per_workbook=10,
    )

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    assert wb.sheetnames == ["Sheet Name 1", "Sheet Name 2"]
    assert wb["Sheet Name 1"]["A1"].value == "Alpha"
    assert wb["Sheet Name 2"]["A1"].value == "Beta"
    wb.close()


def test_streaming_bundle_with_parquet_builds_manifest_and_foldered_zip(
    managed_tmp_dir: Path,
):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1, 2, 3, 4, 5]})
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    paths = build_template_with_data(
        schema,
        {"Sheet1": {"rows": df}},
        str(out),
        export_mode="streaming",
        streaming_chunk_rows=2,
        max_rows_per_workbook=3,
        streaming_bundle_with_parquet=True,
    )

    assert len(paths) == 1
    zip_path = Path(paths[0])
    assert zip_path.name == "filled.zip"
    assert zip_path.exists()

    with ZipFile(zip_path) as zip_file:
        names = sorted(zip_file.namelist())
        assert "manifest.json" in names
        assert "report/filled.part001.xlsx" in names
        assert "report/filled.part002.xlsx" in names
        parquet_names = [name for name in names if name.startswith("data/")]
        assert len(parquet_names) == 1

        manifest = json.loads(zip_file.read("manifest.json").decode("utf-8"))
        assert manifest["export_mode"] == "streaming"
        assert manifest["streaming_bundle_with_parquet"] is True
        assert manifest["report_files"] == ["filled.part001.xlsx", "filled.part002.xlsx"]
        assert len(manifest["data_files"]) == 1
        assert manifest["data_files"][0]["rows"] == 5
        assert manifest["data_files"][0]["columns"] == ["A"]
        assert manifest["data_files"][0]["path"] == parquet_names[0]

        zip_file.extractall(path=managed_tmp_dir)

    wb1 = openpyxl.load_workbook(managed_tmp_dir / "report" / "filled.part001.xlsx")
    wb2 = openpyxl.load_workbook(managed_tmp_dir / "report" / "filled.part002.xlsx")
    assert [wb1["Sheet1"]["A1"].value, wb1["Sheet1"]["A2"].value, wb1["Sheet1"]["A3"].value] == [1, 2, 3]
    assert [wb2["Sheet1"]["A1"].value, wb2["Sheet1"]["A2"].value] == [4, 5]
    wb1.close()
    wb2.close()

    parquet_df = polars.read_parquet(managed_tmp_dir / manifest["data_files"][0]["path"])
    assert parquet_df["A"].to_list() == [1, 2, 3, 4, 5]


def test_streaming_bundle_with_parquet_rejects_fidelity_mode(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")

    with pytest.raises(
        ValueError,
        match="streaming_bundle_with_parquet is only supported with export_mode='streaming'",
    ):
        build_template_with_data(
            schema,
            {"Sheet1": {"name": "Alice"}},
            str(out),
            export_mode="fidelity",
            streaming_bundle_with_parquet=True,
        )


def test_streaming_bundle_with_parquet_rejects_non_polars_sources(managed_tmp_dir: Path):
    pandas = pytest.importorskip("pandas", reason="pandas not installed")
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    out = _temp_output_path(managed_tmp_dir, "filled.xlsx")
    df = pandas.DataFrame({"A": [1, 2]})

    with pytest.raises(
        ValueError,
        match="supports only polars DataFrame/LazyFrame sources",
    ):
        build_template_with_data(
            schema,
            {"Sheet1": {"rows": df}},
            str(out),
            export_mode="streaming",
            streaming_bundle_with_parquet=True,
        )


# §5 Entrypoints
