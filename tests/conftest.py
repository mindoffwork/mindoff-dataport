import os
import sys
import tempfile
from pathlib import Path

import pytest

# §1 Types

# §2 Constants

# Ensure src is on path when running without editable install.
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))
PROJECT_ROOT = Path(__file__).parent.parent
FIXTURE_PATH = Path(__file__).parent / "fixtures" / "sample_template.xlsx"

# §3 Private Helpers

_ORIGINAL_OS_MKDIR = os.mkdir


def _patch_windows_mkdir_mode() -> None:
    """Work around Python 3.13 Windows ACL behavior for mode 0o700 directories."""
    if os.name != "nt" or getattr(os.mkdir, "_mindoff_mode_patch", False):
        return

    def _safe_mkdir(path, mode=0o777, *, dir_fd=None):
        safe_mode = 0o777 if mode == 0o700 else mode
        if dir_fd is None:
            return _ORIGINAL_OS_MKDIR(path, safe_mode)
        return _ORIGINAL_OS_MKDIR(path, safe_mode, dir_fd=dir_fd)

    _safe_mkdir._mindoff_mode_patch = True
    os.mkdir = _safe_mkdir


def _ensure_writable_temp_root() -> Path:
    """Prefer OS temp to avoid workspace file-change loops during test discovery."""
    system_root = Path(tempfile.gettempdir()) / "mindoff_data_export_pytest"
    try:
        system_root.mkdir(parents=True, exist_ok=True)
        temp_root = system_root
    except OSError:
        fallback_root = PROJECT_ROOT / ".tmp"
        fallback_root.mkdir(parents=True, exist_ok=True)
        temp_root = fallback_root

    # Keep Python and pytest temp discovery aligned.
    os.environ["TMP"] = str(temp_root)
    os.environ["TEMP"] = str(temp_root)
    os.environ["TMPDIR"] = str(temp_root)
    tempfile.tempdir = str(temp_root)
    return temp_root


def _configure_stable_basetemp(config) -> None:
    if config.option.basetemp:
        return
    temp_root = _ensure_writable_temp_root() / "pytest-basetemp"
    temp_root.mkdir(parents=True, exist_ok=True)
    config.option.basetemp = str(temp_root)


def _create_fixture() -> None:
    import datetime

    import openpyxl
    from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.merge_cells("A1:C2")
    cell = ws["A1"]
    cell.value = "Merged Header"
    cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFFFF")
    cell.fill = PatternFill(patternType="solid", fgColor=Color(rgb="FF003366"))
    cell.alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"].value = "Label"
    ws["A3"].font = Font(name="Calibri", size=11)

    ws["B3"].value = 12345.678
    ws["B3"].number_format = "0.00"
    ws["B3"].alignment = Alignment(horizontal="right")

    ws["C3"].value = "=B3*2"

    ws["A4"].value = datetime.datetime(2024, 6, 15, 9, 30)
    ws["A4"].number_format = "DD/MM/YYYY HH:MM"

    ws["B4"].value = "Styled Text"
    ws["B4"].font = Font(bold=True, italic=True, color="FFCC0000", size=12)

    ws["C4"].value = "Yellow Cell"
    ws["C4"].fill = PatternFill(patternType="solid", fgColor=Color(rgb="FFFFFF00"))

    ws["A5"].value = "Bordered"
    ws["A5"].border = Border(
        top=Side(border_style="thick", color=Color(rgb="FF000000")),
        bottom=Side(border_style="thin", color=Color(rgb="FF000000")),
        left=Side(border_style="medium", color=Color(rgb="FF0000FF")),
        right=Side(border_style="dashed", color=Color(rgb="FFFF0000")),
    )

    ws["B5"].value = "This is a long text that should be wrapped inside the cell"
    ws["B5"].alignment = Alignment(wrap_text=True, horizontal="center", vertical="top")

    ws.merge_cells("D1:E3")
    ws["D1"].value = "Second Merge"
    ws["D1"].fill = PatternFill(patternType="solid", fgColor=Color(rgb="FFCCFFCC"))
    ws["D1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["B"].width = 18.0
    ws.column_dimensions["C"].width = 15.0
    ws.row_dimensions[1].height = 30.0
    ws.row_dimensions[2].height = 30.0
    ws.row_dimensions[5].height = 50.0

    FIXTURE_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(FIXTURE_PATH))


def pytest_configure(config) -> None:
    _patch_windows_mkdir_mode()
    _configure_stable_basetemp(config)
    if not FIXTURE_PATH.exists():
        _create_fixture()


# §4 Public API


@pytest.fixture(scope="session")
def fixture_path() -> str:
    return str(FIXTURE_PATH)


@pytest.fixture(scope="session")
def workbook_schema(fixture_path):
    from mindoff_data_export import extract_template

    return extract_template(fixture_path)


@pytest.fixture
def managed_tmp_dir(tmp_path: Path) -> Path:
    """Provide a pytest-managed per-test temporary directory."""
    return tmp_path
