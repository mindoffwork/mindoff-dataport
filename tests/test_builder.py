import openpyxl

# §1 Types

# §2 Constants

# §3 Private Helpers

# §4 Public API


def test_build_creates_file(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    assert out.exists()


def test_built_file_is_valid_xlsx(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    wb = openpyxl.load_workbook(str(out))
    assert len(wb.worksheets) == len(workbook_schema["sheets"])
    wb.close()


def test_merged_cells_reconstructed(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    wb = openpyxl.load_workbook(str(out))
    ws = wb.worksheets[0]
    merged = [str(r) for r in ws.merged_cells.ranges]
    assert "A1:C2" in merged
    assert "D1:E3" in merged
    wb.close()


def test_column_width_reconstructed(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    wb = openpyxl.load_workbook(str(out))
    ws = wb.worksheets[0]
    assert abs(ws.column_dimensions["A"].width - 20.0) < 0.5
    wb.close()


def test_row_height_reconstructed(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    wb = openpyxl.load_workbook(str(out))
    ws = wb.worksheets[0]
    assert abs(ws.row_dimensions[1].height - 30.0) < 0.5
    wb.close()


def test_cell_value_reconstructed(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    wb = openpyxl.load_workbook(str(out), data_only=False)
    ws = wb.worksheets[0]
    assert ws["A3"].value == "Label"
    wb.close()


def test_anchor_cell_value_reconstructed(workbook_schema, managed_tmp_dir):
    from mindoff_data_export.builder import build_template

    out = managed_tmp_dir / "out.xlsx"
    build_template(workbook_schema, str(out))
    wb = openpyxl.load_workbook(str(out))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "Merged Header"
    wb.close()
