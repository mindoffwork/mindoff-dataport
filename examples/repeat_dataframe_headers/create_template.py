"""Generate template.xlsx + data.parquet for repeat_dataframe_headers example."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import polars as pl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

HERE = Path(__file__).resolve().parent
TEMPLATE_XLSX = HERE / "template.xlsx"
DATA_PARQUET = HERE / "data.parquet"


def _solid(rgb: str) -> PatternFill:
    return PatternFill(patternType="solid", fgColor=Color(rgb=rgb))


def _border() -> Border:
    side = Side(border_style="thin", color=Color(rgb="FF3A3A3A"))
    return Border(top=side, bottom=side, left=side, right=side)


def _build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Combined Header Demo"

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    ws["A1"].value = "Non-repeat dataframe sheet"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFFFF")
    ws["A1"].fill = _solid("FF1F4E78")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A1"].border = _border()
    ws.merge_cells("A1:C1")

    ws["A2"].value = "{{summary_headers:dataframe-header}}"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    ws["A2"].fill = _solid("FF2F75B5")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].border = _border()

    ws["A3"].value = "{{summary_rows:dataframe-content}}"
    ws["A3"].font = Font(name="Calibri", size=11)
    ws["A3"].fill = _solid("FFEAF2FB")
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A3"].border = _border()

    ws["A4"].value = "Footer below streamed rows (with vertical shift)"
    ws["A4"].font = Font(name="Calibri", size=10, italic=True, color="FF333333")
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_breaks.append(Break(id=3))

    repeat_ws = wb.create_sheet("Repeat Header Demo")
    repeat_ws.column_dimensions["A"].width = 20
    repeat_ws.column_dimensions["B"].width = 16
    repeat_ws.column_dimensions["C"].width = 14

    repeat_ws["A1"].value = "{{reports:repeat-start}}"
    repeat_ws["A2"].value = "Customer: {{customer_name:string}}"
    repeat_ws["A2"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFFFF")
    repeat_ws["A2"].fill = _solid("FF385723")
    repeat_ws["A2"].border = _border()
    repeat_ws.merge_cells("A2:C2")

    repeat_ws["A3"].value = "{{line_items:dataframe-header}}"
    repeat_ws["A3"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    repeat_ws["A3"].fill = _solid("FF548235")
    repeat_ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    repeat_ws["A3"].border = _border()

    repeat_ws["A4"].value = "{{line_items:dataframe-content}}"
    repeat_ws["A4"].font = Font(name="Calibri", size=11)
    repeat_ws["A4"].fill = _solid("FFE2F0D9")
    repeat_ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
    repeat_ws["A4"].border = _border()

    repeat_ws["A5"].value = "----"
    repeat_ws["A5"].font = Font(name="Calibri", size=10, color="FF666666")
    repeat_ws["A6"].value = "{{reports:repeat-end}}"
    repeat_ws.row_breaks.append(Break(id=4))
    return wb


def _build_data() -> pl.DataFrame:
    records: list[dict[str, object]] = []
    customers = [
        ("Acme", "A", 100),
        ("Globex", "G", 200),
        ("Initech", "I", 300),
        ("Umbrella", "U", 400),
        ("Wayne", "W", 500),
        ("Stark", "S", 600),
    ]
    for customer_name, prefix, base in customers:
        for idx in range(1, 81):
            records.append(
                {
                    "customer": customer_name,
                    "sku": f"{prefix}-{base + idx}",
                    "item": f"{customer_name} Item {idx}",
                    "qty": (idx % 11) + 1,
                }
            )
    return pl.DataFrame(records)


def main() -> None:
    wb = _build_workbook()
    wb.save(TEMPLATE_XLSX)
    print(f"Saved template: {TEMPLATE_XLSX}")

    df = _build_data()
    df.write_parquet(DATA_PARQUET)
    print(f"Saved data:     {DATA_PARQUET}")


if __name__ == "__main__":
    main()
