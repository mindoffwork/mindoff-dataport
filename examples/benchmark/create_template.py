"""Generates benchmark_template.xlsx for the benchmark example.

Run once before run.py:
    python examples/benchmark/create_template.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
TEMPLATE_XLSX = HERE / "benchmark_template.xlsx"

# §1. Color palette (ARGB)
_NAVY  = "FF1E3A5F"
_BLUE  = "FF2563EB"
_WHITE = "FFFFFFFF"
_GRAY  = "FFCCCCCC"


# §2. Style helpers


def _solid(rgb: str) -> PatternFill:
    return PatternFill(patternType="solid", fgColor=Color(rgb=rgb))


def _side(style: str, rgb: str = "FF000000") -> Side:
    return Side(border_style=style, color=Color(rgb=rgb))


def _hdr_cell(cell, value: str) -> None:
    cell.value = value
    cell.font = Font(name="Calibri", size=10, bold=True, color=_WHITE)
    cell.fill = _solid(_NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        top=_side("thin", _WHITE), bottom=_side("thin", _WHITE),
        left=_side("thin", _WHITE), right=_side("thin", _WHITE),
    )


def _content_cell(cell, value: str) -> None:
    cell.value = value
    cell.font = Font(name="Calibri", size=10, color=_NAVY)
    cell.fill = _solid(_WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        top=_side("thin", _GRAY), bottom=_side("thin", _GRAY),
        left=_side("thin", _GRAY), right=_side("thin", _GRAY),
    )


# §3. Workbook builder


def _build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Benchmark"

    # Column widths A–E
    for col, w in zip("ABCDE", [12.0, 28.0, 18.0, 14.0, 16.0]):
        ws.column_dimensions[col].width = w

    # ── Row 1: Title (merged, white bold on navy) ────────────────────────
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = "{{report_title:string}}"
    c.font = Font(name="Calibri", size=14, bold=True, color=_WHITE)
    c.fill = _solid(_NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36.0

    # ── Row 2: Subtitle left (date) + right (row count) ─────────────────
    ws.merge_cells("A2:C2")
    c = ws["A2"]
    c.value = "Generated: {{generated_on:date}}"
    c.font = Font(name="Calibri", size=10, color=_WHITE)
    c.fill = _solid(_BLUE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("D2:E2")
    c = ws["D2"]
    c.value = "Rows: {{row_count:number}}"
    c.font = Font(name="Calibri", size=10, color=_WHITE)
    c.fill = _solid(_BLUE)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 22.0

    # ── Row 3: Spacer ────────────────────────────────────────────────────
    ws.row_dimensions[3].height = 8.0

    # ── Row 4: Dataframe header anchor ──────────────────────────────────
    _hdr_cell(ws["A4"], "{{bench_data:dataframe-header}}")
    ws.row_dimensions[4].height = 22.0

    # ── Row 5: Dataframe content anchor ─────────────────────────────────
    _content_cell(ws["A5"], "{{bench_data:dataframe-content}}")
    ws.row_dimensions[5].height = 18.0

    return wb


# §4. Entrypoint


def main() -> None:
    HERE.mkdir(parents=True, exist_ok=True)
    wb = _build_workbook()
    wb.save(str(TEMPLATE_XLSX))
    print(f"Saved: {TEMPLATE_XLSX}")


if __name__ == "__main__":
    main()
