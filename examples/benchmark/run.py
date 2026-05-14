"""Public-facing benchmark: runtime compile+export only.

This benchmark intentionally excludes template extraction/schema creation from
timed measurements. The template is extracted once, then each benchmark run
measures compile + export only.

Produces:
    output/results.csv
    output/files/
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import statistics
import shutil
import sys
import threading
import tracemalloc
from dataclasses import dataclass
from datetime import date, timedelta
from functools import partial
from pathlib import Path
from time import perf_counter, perf_counter_ns
from typing import Any

import polars as pl


try:
    from openpyxl import Workbook as _OXLWorkbook
    from openpyxl.styles import Alignment as _Align
    from openpyxl.styles import Border as _Border
    from openpyxl.styles import Color as _Color
    from openpyxl.styles import Font as _Font
    from openpyxl.styles import PatternFill as _Fill
    from openpyxl.styles import Side as _Side

    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

try:
    import xlsxwriter as _xlsxwriter

    _HAS_XLSXWRITER = True
except ImportError:
    _HAS_XLSXWRITER = False

try:
    from reportlab.lib import colors as _rl_colors
    from reportlab.lib.pagesizes import A4 as _RL_A4
    from reportlab.platypus import SimpleDocTemplate as _RLDoc
    from reportlab.platypus import Table as _RLTable
    from reportlab.platypus import TableStyle as _RLStyle

    _HAS_REPORTLAB = True
except ImportError:
    _HAS_REPORTLAB = False

try:
    import psutil as _psutil

    _PROC = _psutil.Process()
    _HAS_PSUTIL = True
except ImportError:
    _HAS_PSUTIL = False

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as _plt
    import matplotlib.ticker as _ticker

    _HAS_MATPLOTLIB = True
except ImportError:
    _HAS_MATPLOTLIB = False

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
import mindoff_dataport as mo_dataport

# Thread-local store: mindoff bench functions write compile_s here so _timed_runs can read it.
_BENCH_CTX = threading.local()


# Â§1. Constants & Configuration
HERE = Path(__file__).resolve().parent
TEMPLATE_PATH = HERE / "benchmark_template.xlsx"
OUTPUT_DIR = HERE / "output"
ARTIFACT_DIR = OUTPUT_DIR / "files"

# Streaming modes are tested to 500K; fidelity/direct modes cap at 100K (in-memory, not their use-case).
XLSX_STREAMING_ROW_COUNTS = [1_000, 10_000, 100_000, 500_000]
XLSX_FIDELITY_ROW_COUNTS = [1_000, 10_000, 100_000]
PDF_ROW_COUNTS = [1_000, 5_000, 10_000]

RUN_TIMEOUT_S = 300
RUNS_PER_POINT = 5

_COLS = ["id", "name", "category", "value", "date"]

_HUMAN_ROW_LABELS: dict[int, str] = {
    1_000: "1K",
    5_000: "5K",
    10_000: "10K",
    100_000: "100K",
    500_000: "500K",
    1_000_000: "1M",
}

_BENCH_TITLE = "Benchmark Report"

# Method name constants — used in registries, chart series, and comparison CSV.
_M_MO_XLSX_FIDELITY = "mindoff fidelity (XLSX)"
_M_MO_XLSX_STREAM_OXL = "mindoff streaming-openpyxl (XLSX)"
_M_MO_XLSX_STREAM_XLW = "mindoff streaming-xlsxwriter (XLSX)"
_M_OXL_DIRECT = "openpyxl - manual styling"
_M_XLW_DIRECT = "xlsxwriter - manual styling"
_M_MO_PDF_FIDELITY = "mindoff fidelity (PDF)"
_M_MO_PDF_STREAM = "mindoff streaming (PDF)"
_M_RL_DIRECT = "reportlab - manual styling"

# Runtime-configurable benchmark parameters (set by CLI in main()).
ACTIVE_XLSX_STREAMING_ROW_COUNTS = XLSX_STREAMING_ROW_COUNTS
ACTIVE_XLSX_FIDELITY_ROW_COUNTS = XLSX_FIDELITY_ROW_COUNTS
ACTIVE_PDF_ROW_COUNTS = PDF_ROW_COUNTS
ACTIVE_RUN_TIMEOUT_S = RUN_TIMEOUT_S
ACTIVE_RUNS_PER_POINT = RUNS_PER_POINT


# Â§2. Data Classes
@dataclass
class BenchResult:
    method: str
    fmt: str
    rows: int
    elapsed_s: float
    peak_mb: float
    file_mb: float
    status: str = "completed"
    runs: int = ACTIVE_RUNS_PER_POINT
    compile_s: float = 0.0
    elapsed_stdev_s: float = float("nan")
    peak_stdev_mb: float = float("nan")
    output_file: str = ""


# Â§3. Private Helper Functions
def _is_valid_number(value: float) -> bool:
    return math.isfinite(value) and not math.isnan(value)


def _format_number(value: float, places: int = 1) -> str:
    if not _is_valid_number(value):
        return "-"
    return f"{value:,.{places}f}"


def _format_rows(n: int) -> str:
    return _HUMAN_ROW_LABELS.get(n, f"{n:,}")


def _make_parquet(n_rows: int) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / f"bench_data_{n_rows}.parquet"
    if path.exists():
        return path
    cats = ["Alpha", "Beta", "Gamma", "Delta"]
    base = date(2024, 1, 1)
    df = pl.DataFrame(
        {
            "id": pl.Series(list(range(1, n_rows + 1)), dtype=pl.Int32),
            "name": [f"Record {i:07d}" for i in range(1, n_rows + 1)],
            "category": [cats[i % 4] for i in range(n_rows)],
            "value": [round(100.0 + (i * 0.37 % 900), 2) for i in range(n_rows)],
            "date": [
                (base + timedelta(days=i % 365)).isoformat() for i in range(n_rows)
            ],
        }
    )
    df.write_parquet(str(path))
    print(f"  [data] generated {n_rows:,} rows -> {path.name}")
    return path


def _measure(fn) -> tuple[float, float]:
    if _HAS_PSUTIL:
        peak_rss = [_PROC.memory_info().rss]
        stop_flag = threading.Event()

        def _poll() -> None:
            while not stop_flag.is_set():
                try:
                    peak_rss[0] = max(peak_rss[0], _PROC.memory_info().rss)
                except Exception:
                    pass
                stop_flag.wait(0.05)

        baseline_rss = _PROC.memory_info().rss
        poller = threading.Thread(target=_poll, daemon=True)
        poller.start()
        t0 = perf_counter()
        fn()
        elapsed = perf_counter() - t0
        stop_flag.set()
        poller.join(timeout=1)
        peak_mb = (peak_rss[0] - baseline_rss) / 1024 / 1024
        if peak_mb < 0:
            peak_mb = peak_rss[0] / 1024 / 1024
        return elapsed, peak_mb

    if tracemalloc.is_tracing():
        tracemalloc.stop()
    tracemalloc.start()
    t0 = perf_counter()
    fn()
    elapsed = perf_counter() - t0
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    return elapsed, peak / 1024 / 1024


def _export_file_mb(tmp_dir: Path, ext: str) -> float:
    return sum(f.stat().st_size for f in tmp_dir.glob(f"*.{ext}")) / 1024 / 1024


def _safe_artifact_stem(method: str, rows: int) -> str:
    safe_method = re.sub(r"[^A-Za-z0-9]+", "_", method).strip("_").lower()
    return f"{rows:09d}_{safe_method or 'method'}"


def _artifact_path(method: str, fmt: str, rows: int, ext: str) -> Path:
    return ARTIFACT_DIR / fmt / f"{_safe_artifact_stem(method, rows)}.{ext}"


def _copy_export_outputs(tmp_dir: Path, ext: str, keep_output_path: Path) -> list[str]:
    exported = sorted(tmp_dir.glob(f"*.{ext}"))
    if not exported:
        return []
    keep_output_path.parent.mkdir(parents=True, exist_ok=True)
    if len(exported) == 1:
        shutil.copy2(exported[0], keep_output_path)
        return [str(keep_output_path)]

    saved: list[str] = []
    for index, source in enumerate(exported, start=1):
        target = keep_output_path.with_name(
            f"{keep_output_path.stem}.part{index:03d}{keep_output_path.suffix}"
        )
        shutil.copy2(source, target)
        saved.append(str(target))
    return saved


def _run_once(fn, ext: str, keep_output_path: Path | None = None) -> dict[str, object]:
    """Execute fn once in a timeout-guarded thread; return state dict."""
    state: dict[str, object] = {}

    def _target() -> None:
        try:
            scratch_root = OUTPUT_DIR / "_scratch"
            scratch_root.mkdir(parents=True, exist_ok=True)
            tmp_dir = scratch_root / f"run_{perf_counter_ns()}"
            tmp_dir.mkdir(parents=True, exist_ok=True)
            out = tmp_dir / f"output.{ext}"
            _BENCH_CTX.compile_s = 0.0
            elapsed, peak_mb = _measure(partial(fn, out))
            file_mb = _export_file_mb(tmp_dir, ext)
            compile_s = getattr(_BENCH_CTX, "compile_s", 0.0)
            saved_outputs = (
                _copy_export_outputs(tmp_dir, ext, keep_output_path)
                if keep_output_path is not None
                else []
            )
            state.update(
                elapsed=elapsed,
                peak_mb=peak_mb,
                file_mb=file_mb,
                compile_s=compile_s,
                saved_outputs=saved_outputs,
            )
            shutil.rmtree(tmp_dir, ignore_errors=True)
        except Exception as exc:
            state["error"] = str(exc)

    thread = threading.Thread(target=_target, daemon=True)
    thread.start()
    thread.join(timeout=ACTIVE_RUN_TIMEOUT_S)
    if thread.is_alive():
        state["timeout"] = True
    return state


def _timed_runs(
    fn, ext: str, keep_output_path: Path | None = None
) -> tuple[float, float, float, float, float, float, str, str]:
    elapsed_vals: list[float] = []
    peak_vals: list[float] = []
    file_vals: list[float] = []
    compile_vals: list[float] = []

    # Discard one warm-up pass to avoid cold Python import / OS file-cache effects.
    if ACTIVE_RUNS_PER_POINT >= 3:
        _run_once(fn, ext)

    for run_idx in range(ACTIVE_RUNS_PER_POINT):
        state = _run_once(
            fn,
            ext,
            keep_output_path
            if run_idx == ACTIVE_RUNS_PER_POINT - 1
            else None,
        )
        if state.get("timeout"):
            return (
                float("nan"),
                float("nan"),
                float("nan"),
                float("nan"),
                0.0,
                0.0,
                "timeout",
                f"timed out after {ACTIVE_RUN_TIMEOUT_S}s",
            )
        if "error" in state:
            return (
                float("nan"),
                float("nan"),
                float("nan"),
                float("nan"),
                0.0,
                0.0,
                "failed",
                str(state["error"]),
            )
        elapsed_vals.append(float(state["elapsed"]))
        peak_vals.append(float(state["peak_mb"]))
        file_vals.append(float(state["file_mb"]))
        compile_vals.append(float(state.get("compile_s", 0.0)))

    stdev_e = statistics.stdev(elapsed_vals) if len(elapsed_vals) >= 2 else float("nan")
    stdev_m = statistics.stdev(peak_vals) if len(peak_vals) >= 2 else float("nan")
    return (
        statistics.median(elapsed_vals),
        stdev_e,
        statistics.median(peak_vals),
        stdev_m,
        statistics.median(file_vals),
        statistics.median(compile_vals),
        "completed",
        "",
    )


# Â§4. Benchmark Targets


def _scalar_value(key: str, expected_type: str, n_rows: int) -> Any:
    lowered = key.lower()
    if expected_type == "string":
        if "title" in lowered:
            return _BENCH_TITLE
        return f"{key.replace('_', ' ').title()} sample"
    if expected_type == "number":
        if "count" in lowered or "rows" in lowered:
            return n_rows
        return max(1, n_rows // 10)
    if expected_type == "int":
        if "count" in lowered or "rows" in lowered:
            return n_rows
        return max(1, n_rows // 10)
    if expected_type == "float":
        return round(max(1, n_rows) * 1.25, 2)
    if expected_type == "boolean":
        return True
    if expected_type == "date":
        return date.today()
    raise ValueError(f"Unsupported placeholder type: {expected_type}")


def _dataframe_value(key: str, df_path: Path, n_rows: int) -> pl.LazyFrame:
    base = pl.scan_parquet(str(df_path))
    lowered = key.lower()

    if "header" in lowered:
        return base.select(
            pl.lit(f"{key} field").alias("Field"),
            pl.lit(f"{key} label").alias("Label"),
            pl.int_range(1, pl.len() + 1).alias("Order"),
        )

    if "summary" in lowered or "total" in lowered:
        return base.group_by("category").agg(
            pl.len().alias("row_total"),
            pl.col("value").sum().round(2).alias("value_total"),
        ).sort("category")

    if "detail" in lowered or "line" in lowered or "item" in lowered:
        return base.select(
            pl.col("id").alias(f"{key}_id"),
            pl.col("name").alias(f"{key}_name"),
            pl.col("value").alias(f"{key}_value"),
            pl.col("date").alias(f"{key}_date"),
        )

    return base.select(
        pl.col("id").alias(f"{key}_id"),
        pl.concat_str(
            [pl.lit(f"{key.upper()} "), pl.col("name")],
            separator="",
        ).alias(f"{key}_name"),
        pl.col("category").alias(f"{key}_category"),
        (pl.col("value") + pl.lit(len(key))).round(2).alias(f"{key}_value"),
        pl.col("date").alias(f"{key}_date"),
    )


def _record_payload(
    contract: dict[str, Any],
    df_path: Path,
    n_rows: int,
    *,
    repeat_index: int = 0,
) -> dict[str, Any]:
    payload: dict[str, Any] = {}
    for key, expected in contract.items():
        if isinstance(expected, dict):
            payload[key] = _record_payload(
                expected,
                df_path,
                n_rows,
                repeat_index=repeat_index,
            )
            continue
        if isinstance(expected, list):
            item_contract = expected[0] if expected else {}
            payload[key] = [
                _record_payload(
                    item_contract,
                    df_path,
                    min(n_rows, 10) + repeat_index + item_idx,
                    repeat_index=item_idx,
                )
                for item_idx in range(2)
            ]
            continue
        if expected == "string":
            payload[key] = (
                f"{key.replace('_', ' ').title()} sample {repeat_index + 1}"
                if repeat_index
                else _scalar_value(key, expected, n_rows)
            )
            continue
        if expected in {"number", "int"} and repeat_index:
            payload[key] = repeat_index + 1
            continue
        if expected in {"dataframe", "dataframe-header", "dataframe-content"}:
            dataframe_key = f"{key}_{repeat_index + 1}" if repeat_index else key
            payload[key] = _dataframe_value(dataframe_key, df_path, n_rows)
            continue
        payload[key] = _scalar_value(key, expected, n_rows)
    return payload


def _payload(schema, df_path: Path, n_rows: int) -> dict[str, dict[str, object]]:
    contract = mo_dataport.inputs(schema)
    payload: dict[str, dict[str, object]] = {}
    for sheet_name, sheet_contract in contract.items():
        if isinstance(sheet_contract, dict) and "*" in sheet_contract:
            payload[sheet_name] = {
                "Preview A": _record_payload(sheet_contract["*"], df_path, n_rows),
                "Preview B": _record_payload(sheet_contract["*"], df_path, n_rows // 2 or 1, repeat_index=1),
            }
            continue
        payload[sheet_name] = _record_payload(sheet_contract, df_path, n_rows)
    return payload


def _bench_mindoff_xlsx_fidelity(schema, df_path: Path, n_rows: int, out: Path) -> None:
    bundle_path = out.parent / "bundle"
    t0 = perf_counter()
    bundle = mo_dataport.compile(
        schema, _payload(schema, df_path, n_rows), bundle_path=str(bundle_path)
    )
    _BENCH_CTX.compile_s = perf_counter() - t0
    mo_dataport.export(
        bundle, str(out), export_mode="fidelity", auto_delete_bundle=False
    )


def _bench_mindoff_xlsx_streaming_openpyxl(
    schema, df_path: Path, n_rows: int, out: Path
) -> None:
    bundle_path = out.parent / "bundle"
    t0 = perf_counter()
    bundle = mo_dataport.compile(
        schema, _payload(schema, df_path, n_rows), bundle_path=str(bundle_path)
    )
    _BENCH_CTX.compile_s = perf_counter() - t0
    mo_dataport.export(
        bundle,
        str(out),
        export_mode="streaming",
        streaming_engine="openpyxl",
        streaming_chunk_rows=10_000,
        auto_delete_bundle=False,
    )


def _bench_mindoff_xlsx_streaming_xlsxwriter(
    schema, df_path: Path, n_rows: int, out: Path
) -> None:
    bundle_path = out.parent / "bundle"
    t0 = perf_counter()
    bundle = mo_dataport.compile(
        schema, _payload(schema, df_path, n_rows), bundle_path=str(bundle_path)
    )
    _BENCH_CTX.compile_s = perf_counter() - t0
    mo_dataport.export(
        bundle,
        str(out),
        export_mode="streaming",
        streaming_engine="xlsxwriter",
        streaming_chunk_rows=10_000,
        auto_delete_bundle=False,
    )


def _bench_openpyxl_direct(df_path: Path, n_rows: int, out: Path) -> None:
    df = pl.read_parquet(str(df_path))
    wb = _OXLWorkbook()
    ws = wb.active
    ws.title = "Benchmark"
    for col, w in zip("ABCDE", [12, 28, 18, 14, 16]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = _BENCH_TITLE
    c.font = _Font(name="Calibri", size=14, bold=True, color="FFFFFFFF")
    c.fill = _Fill(patternType="solid", fgColor=_Color(rgb="FF1E3A5F"))
    c.alignment = _Align(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A2:C2")
    ws["A2"].value = f"Generated: {date.today()}"
    ws["A2"].font = _Font(name="Calibri", size=10, color="FFFFFFFF")
    ws["A2"].fill = _Fill(patternType="solid", fgColor=_Color(rgb="FF2563EB"))
    ws["A2"].alignment = _Align(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("D2:E2")
    ws["D2"].value = f"Rows: {n_rows}"
    ws["D2"].font = _Font(name="Calibri", size=10, color="FFFFFFFF")
    ws["D2"].fill = _Fill(patternType="solid", fgColor=_Color(rgb="FF2563EB"))
    ws["D2"].alignment = _Align(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 8
    hdr_font = _Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
    hdr_fill = _Fill(patternType="solid", fgColor=_Color(rgb="FF1E3A5F"))
    hdr_align = _Align(horizontal="center", vertical="center")
    hdr_border = _Border(
        left=_Side(border_style="thin", color="FFFFFFFF"),
        right=_Side(border_style="thin", color="FFFFFFFF"),
        top=_Side(border_style="thin", color="FFFFFFFF"),
        bottom=_Side(border_style="thin", color="FFFFFFFF"),
    )
    for col_idx, hdr in enumerate(_COLS, 1):
        cell = ws.cell(row=4, column=col_idx, value=hdr)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = hdr_border
    ws.row_dimensions[4].height = 22
    dat_font = _Font(name="Calibri", size=10, color="FF1E3A5F")
    dat_border = _Border(
        left=_Side(border_style="thin", color="FFCCCCCC"),
        right=_Side(border_style="thin", color="FFCCCCCC"),
        top=_Side(border_style="thin", color="FFCCCCCC"),
        bottom=_Side(border_style="thin", color="FFCCCCCC"),
    )
    dat_align = _Align(horizontal="center", vertical="center")
    for row_idx, row in enumerate(df.iter_rows(), 5):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = dat_font
            cell.border = dat_border
            cell.alignment = dat_align
    wb.save(str(out))


def _bench_xlsxwriter_direct(df_path: Path, n_rows: int, out: Path) -> None:
    df = pl.read_parquet(str(df_path))
    with _xlsxwriter.Workbook(str(out)) as wb:
        ws = wb.add_worksheet("Benchmark")
        title_fmt = wb.add_format(
            {
                "bold": True,
                "font_color": "#FFFFFF",
                "bg_color": "#1E3A5F",
                "font_size": 14,
                "align": "center",
                "valign": "vcenter",
            }
        )
        sub_l_fmt = wb.add_format(
            {
                "font_color": "#FFFFFF",
                "bg_color": "#2563EB",
                "font_size": 10,
                "align": "left",
                "valign": "vcenter",
                "indent": 1,
            }
        )
        sub_r_fmt = wb.add_format(
            {
                "font_color": "#FFFFFF",
                "bg_color": "#2563EB",
                "font_size": 10,
                "align": "right",
                "valign": "vcenter",
            }
        )
        hdr_fmt = wb.add_format(
            {
                "bold": True,
                "font_color": "#FFFFFF",
                "bg_color": "#1E3A5F",
                "font_size": 10,
                "align": "center",
                "valign": "vcenter",
                "border": 1,
                "border_color": "#FFFFFF",
            }
        )
        dat_fmt = wb.add_format(
            {
                "font_color": "#1E3A5F",
                "font_size": 10,
                "align": "center",
                "valign": "vcenter",
                "border": 1,
                "border_color": "#CCCCCC",
            }
        )
        ws.set_column("A:A", 12)
        ws.set_column("B:B", 28)
        ws.set_column("C:C", 18)
        ws.set_column("D:D", 14)
        ws.set_column("E:E", 16)
        ws.merge_range("A1:E1", _BENCH_TITLE, title_fmt)
        ws.set_row(0, 36)
        ws.merge_range("A2:C2", f"Generated: {date.today()}", sub_l_fmt)
        ws.merge_range("D2:E2", f"Rows: {n_rows}", sub_r_fmt)
        ws.set_row(1, 22)
        ws.set_row(2, 8)
        for col_idx, hdr in enumerate(_COLS):
            ws.write(3, col_idx, hdr, hdr_fmt)
        ws.set_row(3, 22)
        for row_idx, row in enumerate(df.iter_rows(), 4):
            for col_idx, val in enumerate(row):
                ws.write(row_idx, col_idx, val, dat_fmt)


def _bench_mindoff_pdf_fidelity(schema, df_path: Path, n_rows: int, out: Path) -> None:
    bundle_path = out.parent / "bundle"
    t0 = perf_counter()
    bundle = mo_dataport.compile(
        schema, _payload(schema, df_path, n_rows), bundle_path=str(bundle_path)
    )
    _BENCH_CTX.compile_s = perf_counter() - t0
    mo_dataport.export(
        bundle,
        str(out),
        format="pdf",
        auto_delete_bundle=False,
    )


def _bench_mindoff_pdf_streaming(schema, df_path: Path, n_rows: int, out: Path) -> None:
    bundle_path = out.parent / "bundle"
    t0 = perf_counter()
    bundle = mo_dataport.compile(
        schema, _payload(schema, df_path, n_rows), bundle_path=str(bundle_path)
    )
    _BENCH_CTX.compile_s = perf_counter() - t0
    mo_dataport.export(
        bundle,
        str(out),
        format="pdf",
        streaming_chunk_rows=200,
        auto_delete_bundle=False,
    )


def _bench_reportlab_direct(df_path: Path, _n_rows: int, out: Path) -> None:
    df = pl.read_parquet(str(df_path))
    raw_col_widths = [w * 7.0 for w in [12.0, 28.0, 18.0, 14.0, 16.0]]
    available_width = _RL_A4[0] - 72
    scale = min(available_width / sum(raw_col_widths), 1.0)
    col_widths = [w * scale for w in raw_col_widths]
    row_heights = [36.0, 22.0, 8.0, 22.0] + [18.0] * len(df)
    data = [
        [_BENCH_TITLE, "", "", "", ""],
        [f"Generated: {date.today()}", "", "", f"Rows: {_n_rows}", ""],
        ["", "", "", "", ""],
        _COLS,
        *[list(row) for row in df.iter_rows()],
    ]
    table = _RLTable(data, colWidths=col_widths, rowHeights=row_heights, repeatRows=0)
    table.setStyle(
        _RLStyle(
            [
                ("SPAN", (0, 0), (-1, 0)),
                ("SPAN", (0, 1), (2, 1)),
                ("SPAN", (3, 1), (4, 1)),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BACKGROUND", (0, 0), (-1, 0), _rl_colors.HexColor("#1E3A5F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), _rl_colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 14),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("BACKGROUND", (0, 1), (-1, 1), _rl_colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 1), (-1, 1), _rl_colors.white),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, 1), 10),
                ("ALIGN", (0, 1), (2, 1), "LEFT"),
                ("ALIGN", (3, 1), (4, 1), "RIGHT"),
                ("LEFTPADDING", (0, 1), (2, 1), 10),
                ("BACKGROUND", (0, 3), (-1, 3), _rl_colors.HexColor("#1E3A5F")),
                ("TEXTCOLOR", (0, 3), (-1, 3), _rl_colors.white),
                ("FONTNAME", (0, 3), (-1, 3), "Helvetica-Bold"),
                ("FONTSIZE", (0, 3), (-1, 3), 10),
                ("GRID", (0, 3), (-1, 3), 0.5, _rl_colors.white),
                ("FONTNAME", (0, 4), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 4), (-1, -1), 10),
                ("TEXTCOLOR", (0, 4), (-1, -1), _rl_colors.HexColor("#1E3A5F")),
                ("GRID", (0, 4), (-1, -1), 0.5, _rl_colors.HexColor("#CCCCCC")),
            ]
        )
    )
    _RLDoc(
        str(out),
        pagesize=_RL_A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    ).build([table])


# §5. Benchmark Orchestration
# Registry tuples: (name, fn, available, row_counts)
# Streaming modes run to 500K; fidelity/direct modes cap at 100K (in-memory, not their use-case).
def _xlsx_registry(schema) -> list[tuple[str, object, bool, list[int]]]:
    return [
        (
            _M_MO_XLSX_FIDELITY,
            partial(_bench_mindoff_xlsx_fidelity, schema),
            True,
            ACTIVE_XLSX_FIDELITY_ROW_COUNTS,
        ),
        (
            _M_MO_XLSX_STREAM_OXL,
            partial(_bench_mindoff_xlsx_streaming_openpyxl, schema),
            True,
            ACTIVE_XLSX_STREAMING_ROW_COUNTS,
        ),
        (
            _M_MO_XLSX_STREAM_XLW,
            partial(_bench_mindoff_xlsx_streaming_xlsxwriter, schema),
            _HAS_XLSXWRITER,
            ACTIVE_XLSX_STREAMING_ROW_COUNTS,
        ),
        (_M_OXL_DIRECT, _bench_openpyxl_direct, _HAS_OPENPYXL, ACTIVE_XLSX_FIDELITY_ROW_COUNTS),
        (_M_XLW_DIRECT, _bench_xlsxwriter_direct, _HAS_XLSXWRITER, ACTIVE_XLSX_FIDELITY_ROW_COUNTS),
    ]


def _pdf_registry(schema) -> list[tuple[str, object, bool, list[int]]]:
    return [
        (_M_MO_PDF_FIDELITY, partial(_bench_mindoff_pdf_fidelity, schema), True, ACTIVE_PDF_ROW_COUNTS),
        (_M_MO_PDF_STREAM, partial(_bench_mindoff_pdf_streaming, schema), True, ACTIVE_PDF_ROW_COUNTS),
        (_M_RL_DIRECT, _bench_reportlab_direct, _HAS_REPORTLAB, ACTIVE_PDF_ROW_COUNTS),
    ]


def run_benchmarks(schema) -> list[BenchResult]:
    results: list[BenchResult] = []

    def _run_one(name: str, fn, fmt: str, ext: str, n: int, df_path: Path) -> None:
        artifact_path = _artifact_path(name, fmt, n, ext)
        elapsed_s, stdev_e, peak_mb, stdev_m, file_mb, compile_s, status, detail = (
            _timed_runs(partial(fn, df_path, n), ext, keep_output_path=artifact_path)
        )
        results.append(
            BenchResult(
                name,
                fmt,
                n,
                elapsed_s,
                peak_mb,
                file_mb,
                status=status,
                runs=ACTIVE_RUNS_PER_POINT,
                compile_s=compile_s,
                elapsed_stdev_s=stdev_e,
                peak_stdev_mb=stdev_m,
                output_file=str(artifact_path.relative_to(OUTPUT_DIR))
                if status == "completed"
                else "",
            )
        )
        if status == "completed":
            print(
                f"  {name:<42} {n:>10,} rows  median {elapsed_s:7.2f}s  "
                f"{peak_mb:7.1f} MB  saved {artifact_path.relative_to(OUTPUT_DIR)}"
            )
        else:
            print(f"  [{status.upper()}] {name} @ {n:,}: {detail}")

    print("\n=== XLSX runtime benchmarks (compile + export) ===")
    for name, fn, available, row_counts in _xlsx_registry(schema):
        if not available:
            print(f"  [SKIP] {name}")
            results.append(
                BenchResult(
                    name,
                    "xlsx",
                    0,
                    float("nan"),
                    float("nan"),
                    0.0,
                    status="skipped",
                    runs=0,
                )
            )
            continue
        for n in row_counts:
            _run_one(name, fn, "xlsx", "xlsx", n, _make_parquet(n))

    print("\n=== PDF runtime benchmarks (compile + export) ===")
    for name, fn, available, row_counts in _pdf_registry(schema):
        if not available:
            print(f"  [SKIP] {name}")
            results.append(
                BenchResult(
                    name,
                    "pdf",
                    0,
                    float("nan"),
                    float("nan"),
                    0.0,
                    status="skipped",
                    runs=0,
                )
            )
            continue
        for n in row_counts:
            _run_one(name, fn, "pdf", "pdf", n, _make_parquet(n))

    return results


# Â§6. Outputs
def save_csv(results: list[BenchResult]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTPUT_DIR / "results.csv"
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "method",
                "format",
                "rows",
                "elapsed_s",
                "elapsed_stdev_s",
                "peak_mb",
                "peak_stdev_mb",
                "compile_s",
                "file_mb",
                "output_file",
                "rows_per_s",
                "peak_mb_per_10k_rows",
                "file_mb_per_10k_rows",
                "status",
                "runs",
            ]
        )
        for r in results:
            rows_per_s = (
                r.rows / r.elapsed_s
                if r.rows > 0 and r.status == "completed" and r.elapsed_s > 0
                else float("nan")
            )
            peak_mb_per_10k = (
                r.peak_mb / (r.rows / 10_000)
                if r.rows > 0 and r.status == "completed"
                else float("nan")
            )
            file_mb_per_10k = (
                r.file_mb / (r.rows / 10_000)
                if r.rows > 0 and r.status == "completed"
                else float("nan")
            )
            w.writerow(
                [
                    r.method,
                    r.fmt,
                    r.rows,
                    r.elapsed_s,
                    r.elapsed_stdev_s,
                    r.peak_mb,
                    r.peak_stdev_mb,
                    r.compile_s,
                    r.file_mb,
                    r.output_file,
                    rows_per_s,
                    peak_mb_per_10k,
                    file_mb_per_10k,
                    r.status,
                    r.runs,
                ]
            )
    return out


def _mindoff_category(
    mindoff_value: float, baseline_value: float, acceptable_ratio: float = 1.15
) -> str:
    if not (_is_valid_number(mindoff_value) and _is_valid_number(baseline_value)):
        return "n/a"
    if mindoff_value <= baseline_value:
        return "winner"
    # "acceptable" means mindoff is close enough to baseline when not winning.
    if mindoff_value <= baseline_value * acceptable_ratio:
        return "acceptable"
    return "loser"


def _format_metric_pair(
    category: str, mindoff_value: float, baseline_value: float, unit: str
) -> str:
    if not (_is_valid_number(mindoff_value) and _is_valid_number(baseline_value)):
        return f"{category} (-/-)"
    return f"{category} ({mindoff_value:.3f}{unit}/{baseline_value:.3f}{unit})"


def save_comparison_csv(results: list[BenchResult]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTPUT_DIR / "comparison.csv"
    completed = [
        r for r in results if r.status == "completed" and r.rows > 0 and r.file_mb >= 0.0
    ]
    by_key = {(r.method, r.fmt, r.rows): r for r in completed}

    comparisons = [
        (
            "xlsx-fidelity-vs-openpyxl",
            (_M_MO_XLSX_FIDELITY, "xlsx", "mindoff"),
            (_M_OXL_DIRECT, "xlsx", "baseline"),
        ),
        (
            "xlsx-streaming-openpyxl-vs-openpyxl",
            (_M_MO_XLSX_STREAM_OXL, "xlsx", "mindoff"),
            (_M_OXL_DIRECT, "xlsx", "baseline"),
        ),
        (
            "xlsx-streaming-xlsxwriter-vs-xlsxwriter",
            (_M_MO_XLSX_STREAM_XLW, "xlsx", "mindoff"),
            (_M_XLW_DIRECT, "xlsx", "baseline"),
        ),
        (
            "pdf-fidelity-vs-reportlab",
            (_M_MO_PDF_FIDELITY, "pdf", "mindoff"),
            (_M_RL_DIRECT, "pdf", "baseline"),
        ),
        (
            "pdf-streaming-vs-reportlab",
            (_M_MO_PDF_STREAM, "pdf", "mindoff"),
            (_M_RL_DIRECT, "pdf", "baseline"),
        ),
    ]

    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "scenario",
                "format",
                "rows",
                "mindoff_method",
                "baseline_method",
                "speed_category",
                "memory_category",
                "file_size_category",
                "overall_category",
            ]
        )
        for scenario, left_meta, right_meta in comparisons:
            left_method, fmt, _ = left_meta
            right_method, _, _ = right_meta
            rows_in_fmt = sorted({r.rows for r in completed if r.fmt == fmt})
            for rows in rows_in_fmt:
                left = by_key.get((left_method, fmt, rows))
                right = by_key.get((right_method, fmt, rows))
                if left is None or right is None:
                    continue
                speed_category = _mindoff_category(
                    left.elapsed_s, right.elapsed_s, acceptable_ratio=1.25
                )
                memory_category = _mindoff_category(left.peak_mb, right.peak_mb)
                file_category = _mindoff_category(left.file_mb, right.file_mb)
                speed_display = _format_metric_pair(
                    speed_category, left.elapsed_s, right.elapsed_s, "s"
                )
                memory_display = _format_metric_pair(
                    memory_category, left.peak_mb, right.peak_mb, "MB"
                )
                file_display = _format_metric_pair(
                    file_category, left.file_mb, right.file_mb, "MB"
                )
                categories = (speed_category, memory_category, file_category)
                if "n/a" in categories:
                    overall = "n/a"
                elif "loser" in categories:
                    if "winner" in categories or "acceptable" in categories:
                        overall = "acceptable"
                    else:
                        overall = "loser"
                elif "acceptable" in categories:
                    overall = "acceptable"
                else:
                    overall = "winner"
                loser_count = sum(
                    category == "loser"
                    for category in categories
                )
                if loser_count >= 2:
                    overall = "loser"
                w.writerow(
                    [
                        scenario,
                        fmt,
                        _format_rows(rows),
                        left.method,
                        right.method,
                        speed_display,
                        memory_display,
                        file_display,
                        overall,
                    ]
                )
    return out


# §7. Charts
# Time panel: all Mindoff modes — demonstrates linear O(n) scaling across modes.
# Memory panel: streaming vs. raw-loop baselines only — fidelity excluded because it is an
# in-memory mode intended for smaller outputs and comparing it on memory defeats the story.
_CHART_SERIES = {
    "xlsx": {
        "time": [
            (_M_MO_XLSX_STREAM_OXL, "#2563EB", "solid",  "Mindoff streaming · openpyxl"),
            (_M_MO_XLSX_STREAM_XLW, "#7C3AED", "solid",  "Mindoff streaming · xlsxwriter"),
            (_M_MO_XLSX_FIDELITY,   "#0EA5E9", "dashed", "Mindoff fidelity (openpyxl engine)"),
        ],
        # memory: list of (group_label, [(method, color, legend_label), ...])
        # Each group is rendered as adjacent bars with a wider gap separating groups.
        "memory": [
            ("openpyxl engine", [
                (_M_MO_XLSX_STREAM_OXL, "#2563EB", "Mindoff streaming · openpyxl"),
                (_M_OXL_DIRECT,         "#DC2626", "openpyxl (raw loop)"),
            ]),
            ("xlsxwriter engine", [
                (_M_MO_XLSX_STREAM_XLW, "#7C3AED", "Mindoff streaming · xlsxwriter"),
                (_M_XLW_DIRECT,         "#EA580C", "xlsxwriter (raw loop)"),
            ]),
        ],
    },
    "pdf": {
        "time": [
            (_M_MO_PDF_STREAM,   "#2563EB", "solid",  "Mindoff streaming"),
            (_M_MO_PDF_FIDELITY, "#0EA5E9", "dashed", "Mindoff fidelity"),
        ],
        "memory": [
            ("", [
                (_M_MO_PDF_STREAM, "#2563EB", "Mindoff streaming"),
                (_M_RL_DIRECT,     "#DC2626", "ReportLab (raw loop)"),
            ]),
        ],
    },
}

def save_charts(results: list[BenchResult]) -> list[Path]:
    if not _HAS_MATPLOTLIB:
        print("\n[charts] matplotlib not installed -- skipping chart generation")
        return []

    charts_dir = OUTPUT_DIR / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    docs_dir = HERE.parents[1] / "docs" / "benchmark"
    docs_dir.mkdir(parents=True, exist_ok=True)

    completed = {
        (r.method, r.fmt, r.rows): r
        for r in results
        if r.status == "completed" and r.rows > 0
    }

    _plt.rcParams.update({
        "font.family": "sans-serif",
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.edgecolor": "#CBD5E1",
        "axes.linewidth": 0.8,
        "xtick.color": "#475569",
        "ytick.color": "#475569",
        "xtick.labelsize": 9,
        "ytick.labelsize": 9,
        "grid.color": "#E2E8F0",
        "grid.linewidth": 0.7,
        "legend.framealpha": 0.92,
        "legend.edgecolor": "#E2E8F0",
        "legend.fontsize": 8.5,
    })

    def _smart_y_fmt(v, _):
        if v >= 1000:
            return f"{v / 1000:,.1f}K"
        if v >= 10:
            return f"{v:,.0f}"
        if v >= 1:
            return f"{v:.1f}"
        return f"{v:.2f}"

    saved: list[Path] = []

    for fmt, panels in _CHART_SERIES.items():
        row_counts = sorted({r.rows for r in results if r.fmt == fmt and r.rows > 0})
        if not row_counts:
            continue

        fig, (ax_time, ax_mem) = _plt.subplots(1, 2, figsize=(13, 6.5))
        fig.patch.set_facecolor("#FFFFFF")
        fig.suptitle(
            f"mindoff-dataport  ·  {fmt.upper()} Export Performance",
            fontsize=13, fontweight="bold", color="#1E3A5F", y=0.98,
        )

        x_labels = [_format_rows(n) for n in row_counts]

        # Time panel: line chart — shows O(n) scaling trend across dataset sizes.
        ax_time.set_facecolor("#FAFAFA")
        ax_time.set_title(
            "Export Time  ·  lower is better ↓",
            fontsize=10.5, color="#1E3A5F", pad=10, fontweight="bold",
        )
        ax_time.set_xlabel("Dataset size (rows)", fontsize=9.5, color="#475569", labelpad=6)
        ax_time.set_ylabel("Wall-clock time (s)", fontsize=9.5, color="#475569", labelpad=6)
        ax_time.set_xscale("log")
        ax_time.set_xticks(row_counts)
        ax_time.set_xticklabels(x_labels)
        ax_time.xaxis.set_minor_formatter(_ticker.NullFormatter())
        ax_time.yaxis.set_major_formatter(_ticker.FuncFormatter(_smart_y_fmt))
        ax_time.grid(axis="y", alpha=0.7)

        time_plotted = False
        _time_max = 0.0
        for method, color, linestyle, label in panels["time"]:
            xs, ys, errs = [], [], []
            for n in row_counts:
                r = completed.get((method, fmt, n))
                if r is None:
                    continue
                xs.append(n)
                ys.append(r.elapsed_s)
                errs.append(r.elapsed_stdev_s if _is_valid_number(r.elapsed_stdev_s) else 0)
            if not xs:
                continue
            time_plotted = True
            _time_max = max(_time_max, max(ys))
            is_mindoff = method.startswith("mindoff")
            lw = 2.5 if is_mindoff else 1.8
            zorder = 3 if is_mindoff else 2
            marker = "o" if linestyle == "solid" else "s"
            ms = 6 if is_mindoff else 5
            ax_time.plot(
                xs, ys,
                color=color, linestyle=linestyle, linewidth=lw,
                marker=marker, markersize=ms, label=label, zorder=zorder,
                solid_capstyle="round",
            )
            lower = [max(0.0, y - e) for y, e in zip(ys, errs)]
            upper = [y + e for y, e in zip(ys, errs)]
            ax_time.fill_between(xs, lower, upper, color=color, alpha=0.11, zorder=zorder - 1)

        if time_plotted:
            ax_time.set_ylim(bottom=0, top=_time_max * 1.3)
            ax_time.legend(loc="upper left")

        # Memory panel: grouped bar chart — pairs of bars per engine, with a gap between groups.
        # panels["memory"] is list of (group_label, [(method, color, legend_label), ...]).
        mem_groups = panels["memory"]
        x_pos = list(range(len(row_counts)))
        n_groups = len(mem_groups)
        group_sizes = [len(series) for _, series in mem_groups]
        n_bars_total = sum(group_sizes)

        # Geometry: each bar is bar_w wide; groups are separated by gap_w.
        bar_w = 0.70 / max(n_bars_total + (n_groups - 1) * 0.5, 1)
        gap_w = bar_w * 1.8

        # Pre-compute per-bar center offsets from each x_pos integer.
        total_span = n_bars_total * bar_w + max(n_groups - 1, 0) * gap_w
        bar_offsets: list[float] = []
        cursor = -total_span / 2.0 + bar_w / 2.0
        for g_idx, (_, series_list) in enumerate(mem_groups):
            for _ in series_list:
                bar_offsets.append(cursor)
                cursor += bar_w
            cursor += gap_w

        # Group center offsets (used for engine label annotations).
        group_centers: list[float] = []
        idx = 0
        for _, series_list in mem_groups:
            gc = (bar_offsets[idx] + bar_offsets[idx + len(series_list) - 1]) / 2.0
            group_centers.append(gc)
            idx += len(series_list)

        ax_mem.set_facecolor("#FAFAFA")
        ax_mem.set_title(
            "Peak Memory Usage  ·  lower is better ↓",
            fontsize=10.5, color="#1E3A5F", pad=10, fontweight="bold",
        )
        ax_mem.set_xlabel("Dataset size (rows)", fontsize=9.5, color="#475569", labelpad=6)
        ax_mem.set_ylabel("Peak memory usage (MB)", fontsize=9.5, color="#475569", labelpad=6)
        ax_mem.set_xticks(x_pos)
        ax_mem.set_xticklabels(x_labels)
        ax_mem.yaxis.set_major_formatter(_ticker.FuncFormatter(_smart_y_fmt))
        ax_mem.grid(axis="y", alpha=0.7)

        mem_plotted = False
        _mem_max = 0.0
        bar_idx = 0
        for g_idx, (group_name, series_list) in enumerate(mem_groups):
            for method, color, label in series_list:
                offset = bar_offsets[bar_idx]
                bar_idx += 1
                xs_bar, ys_bar = [], []
                for j, n in enumerate(row_counts):
                    r = completed.get((method, fmt, n))
                    if r is not None and _is_valid_number(r.peak_mb):
                        xs_bar.append(x_pos[j] + offset)
                        ys_bar.append(r.peak_mb)
                if not xs_bar:
                    continue
                mem_plotted = True
                _mem_max = max(_mem_max, max(ys_bar))
                is_mindoff = method.startswith("mindoff")
                ax_mem.bar(
                    xs_bar, ys_bar, bar_w * 0.9,
                    color=color, label=label,
                    alpha=0.9 if is_mindoff else 0.7,
                    zorder=2,
                )
                for xb, yb in zip(xs_bar, ys_bar):
                    ax_mem.text(
                        xb, yb, f"({_smart_y_fmt(yb, None)})",
                        ha="center", va="bottom", fontsize=6.5,
                        color="#334155", zorder=3,
                    )

        # Engine group labels below each cluster of bars.
        for g_idx, (group_name, _) in enumerate(mem_groups):
            if group_name:
                for xp in x_pos:
                    ax_mem.annotate(
                        group_name,
                        xy=(xp + group_centers[g_idx], 0),
                        xytext=(0, -22),
                        textcoords="offset points",
                        ha="center", va="top",
                        fontsize=7, color="#64748B",
                        annotation_clip=False,
                    )

        if mem_plotted:
            ax_mem.set_ylim(bottom=0, top=_mem_max * 1.35)
            ax_mem.legend(loc="upper left", fontsize=8)

        if not (time_plotted or mem_plotted):
            _plt.close(fig)
            continue

        runs_note = (
            f"Median of {ACTIVE_RUNS_PER_POINT} runs · ±1σ confidence bands on time panel · "
            "compile + export only (extraction excluded) · "
            "python examples/benchmark/run.py to reproduce"
        )
        fig.text(0.5, 0.01, runs_note, ha="center", fontsize=7.5, color="#94A3B8",
                 fontstyle="italic")
        fig.tight_layout(rect=[0, 0.06, 1, 0.93], pad=1.5, w_pad=3.0)

        out_path = charts_dir / f"benchmark_{fmt}.png"
        fig.savefig(str(out_path), dpi=130, bbox_inches="tight", pad_inches=0.3, facecolor="#FFFFFF")
        shutil.copy2(str(out_path), str(docs_dir / out_path.name))
        _plt.close(fig)
        saved.append(out_path)
        print(f"  [chart] {out_path.relative_to(OUTPUT_DIR.parent)}  ->  docs/benchmark/{out_path.name}")

    return saved


# §8. Entrypoint
def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run public benchmark.")
    profile = parser.add_mutually_exclusive_group()
    profile.add_argument(
        "--quick",
        action="store_true",
        help="Fast dev profile (small rows, single run, shorter timeout).",
    )
    profile.add_argument(
        "--full",
        action="store_true",
        help="Publish profile (default).",
    )
    parser.add_argument(
        "--runs", type=int, default=None, help="Override runs per point."
    )
    parser.add_argument(
        "--timeout", type=int, default=None, help="Override timeout seconds per run."
    )
    return parser.parse_args()


def _apply_cli_profile(args: argparse.Namespace) -> None:
    global ACTIVE_XLSX_STREAMING_ROW_COUNTS
    global ACTIVE_XLSX_FIDELITY_ROW_COUNTS
    global ACTIVE_PDF_ROW_COUNTS
    global ACTIVE_RUNS_PER_POINT
    global ACTIVE_RUN_TIMEOUT_S

    # Default publish profile.
    ACTIVE_XLSX_STREAMING_ROW_COUNTS = XLSX_STREAMING_ROW_COUNTS
    ACTIVE_XLSX_FIDELITY_ROW_COUNTS = XLSX_FIDELITY_ROW_COUNTS
    ACTIVE_PDF_ROW_COUNTS = PDF_ROW_COUNTS
    ACTIVE_RUNS_PER_POINT = RUNS_PER_POINT
    ACTIVE_RUN_TIMEOUT_S = RUN_TIMEOUT_S

    if args.quick:
        ACTIVE_XLSX_STREAMING_ROW_COUNTS = [1_000, 10_000]
        ACTIVE_XLSX_FIDELITY_ROW_COUNTS = [1_000, 10_000]
        ACTIVE_PDF_ROW_COUNTS = [1_000, 10_000]
        ACTIVE_RUNS_PER_POINT = 1
        ACTIVE_RUN_TIMEOUT_S = 120

    if args.runs is not None:
        ACTIVE_RUNS_PER_POINT = max(1, args.runs)
    if args.timeout is not None:
        ACTIVE_RUN_TIMEOUT_S = max(30, args.timeout)


def main() -> None:
    args = _parse_args()
    _apply_cli_profile(args)

    if not TEMPLATE_PATH.exists():
        sys.exit(
            f"Template not found: {TEMPLATE_PATH}\nRun first: python examples/benchmark/create_template.py"
        )

    print(f"Template: {TEMPLATE_PATH}")
    print(f"Output:   {OUTPUT_DIR}")
    print(
        "Profile:  "
        f"xlsx_streaming={[_format_rows(n) for n in ACTIVE_XLSX_STREAMING_ROW_COUNTS]}, "
        f"xlsx_fidelity={[_format_rows(n) for n in ACTIVE_XLSX_FIDELITY_ROW_COUNTS]}, "
        f"pdf_rows={[_format_rows(n) for n in ACTIVE_PDF_ROW_COUNTS]}, "
        f"runs={ACTIVE_RUNS_PER_POINT}, "
        f"warmup={'yes' if ACTIVE_RUNS_PER_POINT >= 3 else 'no'}, "
        f"timeout={ACTIVE_RUN_TIMEOUT_S}s"
    )
    print("Preparing schema once (excluded from timed benchmark runs)...")
    schema = mo_dataport.extract(str(TEMPLATE_PATH))
    if ARTIFACT_DIR.exists():
        shutil.rmtree(ARTIFACT_DIR)

    results = run_benchmarks(schema)
    csv_path = save_csv(results)
    comparison_path = save_comparison_csv(results)
    chart_paths = save_charts(results)

    print(f"\nResults CSV:    {csv_path}")
    print(f"Comparison CSV: {comparison_path}")
    if chart_paths:
        print(f"Charts:         {', '.join(str(p) for p in chart_paths)}")


if __name__ == "__main__":
    main()
