"""Generates template.xlsx and data.parquet for the style_showcase example.

Run once before xlsx.py / pdf.py:
    python examples/style_showcase/create_template.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import polars as pl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
TEMPLATE_XLSX = HERE / "template.xlsx"
DATA_PARQUET = HERE / "data.parquet"

# §1. Color palette (ARGB)
_NAVY = "FF1F3864"
_BLUE_MID = "FF2E75B6"
_BLUE_LIGHT = "FFDCE6F1"
_WHITE = "FFFFFFFF"
_GRAY_LIGHT = "FFF5F5F5"
_TEAL_DARK = "FF00695C"
_TEAL_LIGHT = "FFE0F2F1"
_GREEN_DARK = "FF375623"
_ORANGE_DARK = "FFC55A11"
_RED = "FFCC0000"
_CREAM = "FFFFF0E0"


# §2. Style helpers


def _solid(rgb: str) -> PatternFill:
    return PatternFill(patternType="solid", fgColor=Color(rgb=rgb))


def _side(style: str, rgb: str = "FF000000") -> Side:
    return Side(border_style=style, color=Color(rgb=rgb))


def _section_header_style(cell, text: str) -> None:
    """Bold navy label on light-blue tint with left accent border."""
    cell.value = text
    cell.font = Font(name="Calibri", size=13, bold=True, color=_NAVY)
    cell.fill = _solid(_BLUE_LIGHT)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cell.border = Border(
        left=_side("thick", _NAVY),
        bottom=_side("thin", _BLUE_MID),
    )


def _kpi_label(cell, text: str, dark: str) -> None:
    cell.value = text
    cell.font = Font(name="Calibri", size=10, bold=True, color=_WHITE)
    cell.fill = _solid(dark)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        top=_side("medium", dark),
        left=_side("thin", _WHITE),
        right=_side("thin", _WHITE),
    )


def _kpi_value(cell, placeholder: str, dark: str) -> None:
    cell.value = placeholder
    cell.font = Font(name="Calibri", size=20, bold=True, color=_WHITE)
    cell.fill = _solid(dark)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        bottom=_side("medium", dark),
        left=_side("thin", _WHITE),
        right=_side("thin", _WHITE),
    )


def _df_header_cell(cell, placeholder: str) -> None:
    cell.value = placeholder
    cell.font = Font(name="Calibri", size=11, bold=True, color=_WHITE)
    cell.fill = _solid(_NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        top=_side("thin", _WHITE),
        bottom=_side("medium", _WHITE),
        left=_side("thin", _WHITE),
        right=_side("thin", _WHITE),
    )


def _df_content_cell(cell, placeholder: str) -> None:
    cell.value = placeholder
    cell.font = Font(name="Calibri", size=11, color=_NAVY)
    cell.fill = _solid(_WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(
        top=_side("thin", "FFBFBFBF"),
        bottom=_side("thin", "FFBFBFBF"),
        left=_side("thin", "FFBFBFBF"),
        right=_side("thin", "FFBFBFBF"),
    )


def _strip(ws, row: int, rgb: str, pattern: str = "solid") -> None:
    """Fill a full-width strip row."""
    for col in "ABCDEF":
        ws[f"{col}{row}"].fill = PatternFill(
            patternType=pattern,
            fgColor=Color(rgb=rgb),
            bgColor=Color(rgb=_WHITE),
        )


# §3. Workbook builder


def _build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1 Review"

    # Column widths (A–F, 6 columns)
    for col, w in zip("ABCDEF", [22.0, 13.0, 13.0, 14.0, 13.0, 14.0]):
        ws.column_dimensions[col].width = w

    # ── Row 1: Report title (merged, large white bold on navy) ──────────────
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "{{report_title:string}}"
    c.font = Font(name="Calibri", size=22, bold=True, color=_WHITE)
    c.fill = _solid(_NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = Border(bottom=_side("medium", _BLUE_MID))
    ws.row_dimensions[1].height = 44.0

    # ── Row 2: Department (left) and generated date (right) ────────────────
    ws.merge_cells("A2:C2")
    c = ws["A2"]
    c.value = "{{department:string}}"
    c.font = Font(name="Calibri", size=11, italic=True, color=_WHITE)
    c.fill = _solid(_BLUE_MID)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("D2:F2")
    c = ws["D2"]
    c.value = "Generated: {{generated_on:date}}"
    c.font = Font(name="Calibri", size=11, italic=True, color=_WHITE)
    c.fill = _solid(_BLUE_MID)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 22.0

    # ── Row 3: Pattern-fill separator strip ────────────────────────────────
    _strip(ws, 3, _BLUE_MID, pattern="gray125")
    ws.row_dimensions[3].height = 6.0

    # ── Row 4: KPI section label ───────────────────────────────────────────
    ws.merge_cells("A4:F4")
    _section_header_style(ws["A4"], "KEY METRICS")
    ws.row_dimensions[4].height = 28.0

    # ── Rows 5–6: KPI boxes (3 × 2-column blocks) ─────────────────────────
    kpis = [
        ("A5:B5", "A6:B6", "Revenue",    "{{kpi_revenue:string}}", _GREEN_DARK),
        ("C5:D5", "C6:D6", "YoY Growth", "{{kpi_growth:string}}",  _BLUE_MID),
        ("E5:F5", "E6:F6", "Net Margin", "{{kpi_margin:string}}",  _ORANGE_DARK),
    ]
    for lrng, vrng, label, ph, dark in kpis:
        ws.merge_cells(lrng)
        _kpi_label(ws[lrng.split(":")[0]], label, dark)
        ws.merge_cells(vrng)
        _kpi_value(ws[vrng.split(":")[0]], ph, dark)

    ws.row_dimensions[5].height = 20.0
    ws.row_dimensions[6].height = 40.0

    # ── Row 7: Spacer strip ────────────────────────────────────────────────
    _strip(ws, 7, _BLUE_LIGHT)
    ws.row_dimensions[7].height = 8.0

    # ── Row 8: Style reference section label ──────────────────────────────
    ws.merge_cells("A8:F8")
    _section_header_style(ws["A8"], "STYLE REFERENCE")
    ws.row_dimensions[8].height = 28.0

    # ── Row 9: Typography demos ────────────────────────────────────────────
    # A9 — Strikethrough (strike=True)
    ws["A9"].value = "Strikethrough"
    ws["A9"].font = Font(name="Calibri", size=11, strike=True, color=_RED)
    ws["A9"].fill = _solid(_GRAY_LIGHT)
    ws["A9"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A9"].border = Border(
        bottom=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # B9 — Subscript (vert_align="subscript")
    ws["B9"].value = "H₂O"
    ws["B9"].font = Font(name="Calibri", size=13, vertAlign="subscript", color=_NAVY)
    ws["B9"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B9"].border = Border(
        bottom=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # C9 — Superscript (vert_align="superscript")
    ws["C9"].value = "E=mc²"
    ws["C9"].font = Font(name="Calibri", size=13, vertAlign="superscript", color=_NAVY)
    ws["C9"].alignment = Alignment(horizontal="center", vertical="center")
    ws["C9"].border = Border(
        bottom=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # D9 — Underline (underline="single")
    ws["D9"].value = "Underlined"
    ws["D9"].font = Font(name="Calibri", size=11, underline="single", color=_NAVY)
    ws["D9"].alignment = Alignment(horizontal="center", vertical="center")
    ws["D9"].border = Border(
        bottom=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # E9 — Bold + italic
    ws["E9"].value = "Bold & Italic"
    ws["E9"].font = Font(name="Calibri", size=11, bold=True, italic=True, color=_NAVY)
    ws["E9"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E9"].border = Border(
        bottom=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # F9 — Indented text (indent=3)
    ws["F9"].value = "→ Indented"
    ws["F9"].font = Font(name="Calibri", size=11, color=_TEAL_DARK)
    ws["F9"].fill = _solid(_TEAL_LIGHT)
    ws["F9"].alignment = Alignment(horizontal="left", vertical="center", indent=3)
    ws["F9"].border = Border(bottom=_side("thin", _BLUE_LIGHT))

    ws.row_dimensions[9].height = 24.0

    # ── Row 10: Layout and border demos ───────────────────────────────────
    # A10 — Wrap text (wrap_text=True, vertical=top)
    ws["A10"].value = "Wrap: Long text that wraps across multiple lines within the cell"
    ws["A10"].font = Font(name="Calibri", size=9, color=_NAVY)
    ws["A10"].fill = _solid(_GRAY_LIGHT)
    ws["A10"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws["A10"].border = Border(
        top=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # B10 — Shrink to fit (shrink_to_fit=True)
    ws["B10"].value = "Shrink-to-fit text"
    ws["B10"].font = Font(name="Calibri", size=11, color=_NAVY)
    ws["B10"].alignment = Alignment(
        horizontal="center", vertical="center", shrinkToFit=True
    )
    ws["B10"].border = Border(
        top=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # C10 — Text rotation (text_rotation=45)
    ws["C10"].value = "Rotated 45°"
    ws["C10"].font = Font(name="Calibri", size=10, color=_NAVY)
    ws["C10"].alignment = Alignment(
        horizontal="center", vertical="bottom", textRotation=45
    )
    ws["C10"].border = Border(
        top=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # D10 — RTL reading order (reading_order=2)
    ws["D10"].value = "← RTL Text"
    ws["D10"].font = Font(name="Calibri", size=11, color=_NAVY)
    ws["D10"].fill = _solid(_CREAM)
    ws["D10"].alignment = Alignment(
        horizontal="right", vertical="center", readingOrder=2
    )
    ws["D10"].border = Border(
        top=_side("thin", _BLUE_LIGHT),
        right=_side("thin", _BLUE_LIGHT),
    )

    # E10 — Diagonal borders (diagonal_up=True, diagonal_down=True)
    ws["E10"].value = "Diagonal"
    ws["E10"].font = Font(name="Calibri", size=11, color=_NAVY)
    ws["E10"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E10"].border = Border(
        top=_side("thin", _NAVY),
        right=_side("thin", _NAVY),
        bottom=_side("thin", _NAVY),
        left=_side("thin", _NAVY),
        diagonal=_side("medium", _NAVY),
        diagonalUp=True,
        diagonalDown=True,
    )

    # F10 — Pattern fill with logical start/end borders
    ws["F10"].value = "Start / End"
    ws["F10"].font = Font(name="Calibri", size=11, bold=True, color=_NAVY)
    ws["F10"].fill = PatternFill(
        patternType="gray125",
        fgColor=Color(rgb=_NAVY),
        bgColor=Color(rgb=_BLUE_LIGHT),
    )
    ws["F10"].alignment = Alignment(horizontal="center", vertical="center")
    ws["F10"].border = Border(
        start=_side("thick", _GREEN_DARK),
        end=_side("thick", _RED),
        top=_side("thin", _NAVY),
        bottom=_side("thin", _NAVY),
    )

    ws.row_dimensions[10].height = 38.0

    # ── Row 11: Spacer strip ───────────────────────────────────────────────
    _strip(ws, 11, _BLUE_LIGHT)
    ws.row_dimensions[11].height = 8.0

    # ── Row 12: Data section label ─────────────────────────────────────────
    ws.merge_cells("A12:F12")
    _section_header_style(ws["A12"], "QUARTERLY SALES DATA")
    ws.row_dimensions[12].height = 28.0

    # ── Row 13: Dataframe header anchor ───────────────────────────────────
    _df_header_cell(ws["A13"], "{{sales:dataframe-header}}")
    ws.row_dimensions[13].height = 22.0

    # ── Row 14: Dataframe content anchor ──────────────────────────────────
    _df_content_cell(ws["A14"], "{{sales:dataframe-content}}")
    ws.row_dimensions[14].height = 18.0

    return wb


# §4. Data builder


def _build_data() -> pl.DataFrame:
    return pl.DataFrame(
        {
            "Product": [
                "Omega Suite",
                "Nexus Hub",
                "Apex Cloud",
                "Vertex AI",
                "Prism Tools",
                "Cascade DB",
                "Echo Sync",
            ],
            "Region": ["EMEA", "APAC", "Americas", "EMEA", "APAC", "Americas", "EMEA"],
            "Units": [1240, 890, 1650, 720, 480, 930, 610],
            "Revenue": [620000, 445000, 825000, 360000, 240000, 465000, 305000],
            "Growth": [0.183, 0.122, 0.215, 0.089, -0.031, 0.147, 0.094],
            "Margin": [0.314, 0.298, 0.342, 0.271, 0.255, 0.318, 0.287],
        }
    )


# §5. Entrypoint


def main() -> None:
    HERE.mkdir(parents=True, exist_ok=True)

    wb = _build_workbook()
    wb.save(str(TEMPLATE_XLSX))
    print(f"Saved template: {TEMPLATE_XLSX}")

    df = _build_data()
    df.write_parquet(str(DATA_PARQUET))
    print(f"Saved data:     {DATA_PARQUET}")


if __name__ == "__main__":
    main()
