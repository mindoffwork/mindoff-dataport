from __future__ import annotations

import hashlib
import html
import re
from pathlib import Path
from typing import Any

from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, legal, letter, landscape, portrait
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from .bundle import ReportBundle
from .renderer import _parse_dims
from .schema import CellSchema, SheetSchema
from .xlsx_renderer import (
    _content_cells,
    _data_source_map,
    _expanded_sheet_from_manifest,
    _header_cells,
    _sheet_with_overrides,
)

__all__ = ["export_report_bundle"]

# Â§1 Constants & Exceptions

_PAGE_SIZES = {
    "A4": A4,
    "LETTER": letter,
    "LEGAL": legal,
}

_BORDER_WIDTHS = {
    "hair": 0.25,
    "thin": 0.5,
    "medium": 1.0,
    "thick": 1.5,
    "dashed": 0.75,
    "dotted": 0.5,
    "double": 1.25,
}

_DEFAULT_COLUMN_WIDTH = 15.0
_DEFAULT_ROW_HEIGHT = 15.0
_EXCEL_WIDTH_TO_POINTS = 7.0
_MIN_COLUMN_WIDTH = 24.0
_MAX_COLUMN_WIDTH = 180.0

# Â§2 Classes and Sub Classes


class _FontResolver:
    def __init__(self, fonts: dict[str, Any] | None = None) -> None:
        self._families = _register_fonts(fonts or {})

    def font_name(self, font: dict[str, Any]) -> str:
        family = font.get("name")
        custom = self._families.get(str(family)) if family else None
        if custom is not None:
            return _variant_font_name(
                custom,
                bold=bool(font.get("bold")),
                italic=bool(font.get("italic")),
            )
        return _builtin_font_name(font)


# Â§3 Private Helper Functions


def _page_size(name: str, orientation: str) -> tuple[float, float]:
    size = _PAGE_SIZES.get(name.upper())
    if size is None:
        expected = ", ".join(sorted(_PAGE_SIZES))
        raise ValueError(
            f"Unsupported PDF page_size '{name}'. Expected one of: {expected}."
        )
    if orientation == "portrait":
        return portrait(size)
    if orientation == "landscape":
        return landscape(size)
    raise ValueError(
        f"Unsupported PDF orientation '{orientation}'. Expected 'portrait' or 'landscape'."
    )


def _hex_color(value: str | None) -> colors.Color | None:
    if not value:
        return None
    raw = value[-6:]
    if len(raw) != 6:
        return None
    try:
        return colors.HexColor(f"#{raw}")
    except ValueError:
        return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _builtin_font_name(font: dict[str, Any]) -> str:
    bold = bool(font.get("bold"))
    italic = bool(font.get("italic"))
    if bold and italic:
        return "Helvetica-BoldOblique"
    if bold:
        return "Helvetica-Bold"
    if italic:
        return "Helvetica-Oblique"
    return "Helvetica"


def _register_fonts(fonts: dict[str, Any]) -> dict[str, dict[str, str]]:
    families: dict[str, dict[str, str]] = {}
    for family, config in fonts.items():
        variants = _font_variants(config)
        registered: dict[str, str] = {}
        for variant, path in variants.items():
            if path is None:
                continue
            font_path = Path(path)
            if not font_path.exists():
                raise ValueError(
                    f"PDF font '{family}' variant '{variant}' file does not exist: {font_path}"
                )
            font_name = _registered_font_name(family, variant, font_path)
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
            registered[variant] = font_name

        if "regular" not in registered:
            raise ValueError(f"PDF font '{family}' requires a regular font file")
        families[str(family)] = registered
        _register_font_family(str(family), registered)
    return families


def _font_variants(config: Any) -> dict[str, str | None]:
    if isinstance(config, (str, Path)):
        return {
            "regular": str(config),
            "bold": None,
            "italic": None,
            "bold_italic": None,
        }
    if not isinstance(config, dict):
        raise TypeError(
            f"PDF font config must be a path or dict of variants, got {type(config).__name__}"
        )
    return {
        "regular": config.get("regular"),
        "bold": config.get("bold"),
        "italic": config.get("italic"),
        "bold_italic": config.get("bold_italic") or config.get("boldItalic"),
    }


def _registered_font_name(family: str, variant: str, path: Path) -> str:
    safe_family = re.sub(r"[^A-Za-z0-9]+", "-", family).strip("-") or "Font"
    safe_variant = variant.replace("_", "-")
    digest = hashlib.sha1(str(path.resolve()).encode("utf-8")).hexdigest()[:12]
    return f"Mindoff-{safe_family}-{safe_variant}-{digest}"


def _register_font_family(family: str, registered: dict[str, str]) -> None:
    regular = registered["regular"]
    pdfmetrics.registerFontFamily(
        family,
        normal=regular,
        bold=registered.get("bold", regular),
        italic=registered.get("italic", regular),
        boldItalic=registered.get("bold_italic", registered.get("bold", regular)),
    )


def _variant_font_name(
    registered: dict[str, str], *, bold: bool, italic: bool
) -> str:
    if bold and italic:
        return (
            registered.get("bold_italic")
            or registered.get("bold")
            or registered["regular"]
        )
    if bold:
        return registered.get("bold") or registered["regular"]
    if italic:
        return registered.get("italic") or registered["regular"]
    return registered["regular"]


def _alignment(value: str | None) -> str:
    if value in {"center", "centerContinuous"}:
        return "CENTER"
    if value == "right":
        return "RIGHT"
    return "LEFT"


def _paragraph_alignment(value: str | None) -> int:
    if value in {"center", "centerContinuous"}:
        return TA_CENTER
    if value == "right":
        return TA_RIGHT
    return TA_LEFT


def _vertical_alignment(value: str | None) -> str:
    if value == "center":
        return "MIDDLE"
    if value == "bottom":
        return "BOTTOM"
    return "TOP"


def _paragraph(cell: CellSchema, font_resolver: _FontResolver) -> Paragraph | str:
    text = html.escape(_cell_text(cell.get("value")), quote=False)
    if not text:
        return ""

    font = cell["font"]
    if font.get("underline"):
        text = f"<u>{text}</u>"

    font_size = float(font.get("size") or 11.0)
    color = _hex_color(font.get("color"))
    style = ParagraphStyle(
        name=f"cell-{cell['coordinate']}",
        fontName=font_resolver.font_name(font),
        fontSize=font_size,
        leading=max(font_size * 1.2, font_size + 2),
        textColor=color or colors.black,
        alignment=_paragraph_alignment(cell["alignment"].get("horizontal")),
    )
    return Paragraph(text.replace("\n", "<br/>"), style)


def _coord_indexes(coord: str) -> tuple[int, int]:
    col_letter, row_idx = coordinate_from_string(coord)
    return row_idx, column_index_from_string(col_letter)


def _is_non_anchor_merged_cell(cell: CellSchema) -> bool:
    return bool(cell.get("merged")) and cell.get("merge_anchor") not in (
        None,
        cell["coordinate"],
    )


def _expanded_cells(
    bundle: ReportBundle,
    sheet: SheetSchema,
    *,
    streaming_chunk_rows: int,
) -> tuple[SheetSchema, dict[tuple[int, int], CellSchema]]:
    source_map = _data_source_map(bundle)
    sheet = _expanded_sheet_from_manifest(sheet, source_map)
    cells: dict[tuple[int, int], CellSchema] = {}

    for cell in sheet["cells"].values():
        if _is_non_anchor_merged_cell(cell):
            continue
        cells[_coord_indexes(cell["coordinate"])] = cell

    for anchor in sheet.get("dataframe_anchors", []):
        anchor_cells = (
            _header_cells(anchor)
            if anchor["placeholder_type"] == "dataframe-headers"
            else _content_cells(
                bundle,
                source_map,
                anchor,
                batch_size=streaming_chunk_rows,
            )
        )
        for coord, cell in anchor_cells:
            cells[_coord_indexes(coord)] = cell

    return sheet, cells


def _hug_column_width(
    col_idx: int,
    min_row: int,
    max_row: int,
    cells: dict[tuple[int, int], CellSchema],
) -> float:
    max_chars = 1
    for row_idx in range(min_row, max_row + 1):
        cell = cells.get((row_idx, col_idx))
        if cell is None:
            continue
        text = _cell_text(cell.get("value"))
        longest_line = max((len(part) for part in text.splitlines()), default=0)
        font_size = float(cell["font"].get("size") or 11.0)
        factor = 1.15 if cell["font"].get("bold") else 1.0
        max_chars = max(max_chars, int(longest_line * factor * font_size / 11.0))
    return min(max(max_chars * 6.0 + 10.0, _MIN_COLUMN_WIDTH), _MAX_COLUMN_WIDTH)


def _column_widths(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    available_width: float,
) -> list[float]:
    mode = sheet.get("column_width_mode", "fixed")
    widths: list[float] = []
    if mode == "even":
        width = float(sheet.get("default_column_width", _DEFAULT_COLUMN_WIDTH))
        widths = [width * _EXCEL_WIDTH_TO_POINTS] * (max_col - min_col + 1)
    elif mode == "hug":
        widths = [
            _hug_column_width(col_idx, min_row, max_row, cells)
            for col_idx in range(min_col, max_col + 1)
        ]
    else:
        for col_idx in range(min_col, max_col + 1):
            col_letter = _col_letter(col_idx)
            width = sheet["column_widths"].get(col_letter, _DEFAULT_COLUMN_WIDTH)
            widths.append(float(width or _DEFAULT_COLUMN_WIDTH) * _EXCEL_WIDTH_TO_POINTS)

    total = sum(widths)
    if total > available_width and total > 0:
        scale = available_width / total
        widths = [max(width * scale, _MIN_COLUMN_WIDTH) for width in widths]
    return widths


def _row_heights(sheet: SheetSchema, min_row: int, max_row: int) -> list[float | None]:
    mode = sheet.get("row_height_mode", "fixed")
    if mode == "even":
        height = float(sheet.get("default_row_height", _DEFAULT_ROW_HEIGHT))
        return [height] * (max_row - min_row + 1)
    if mode == "hug":
        return [None] * (max_row - min_row + 1)
    return [
        float(sheet["row_heights"].get(str(row_idx)) or _DEFAULT_ROW_HEIGHT)
        for row_idx in range(min_row, max_row + 1)
    ]


def _col_letter(col_idx: int) -> str:
    result = ""
    while col_idx:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _table_data(
    cells: dict[tuple[int, int], CellSchema],
    font_resolver: _FontResolver,
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
) -> list[list[Any]]:
    data: list[list[Any]] = []
    for row_idx in range(min_row, max_row + 1):
        row: list[Any] = []
        for col_idx in range(min_col, max_col + 1):
            cell = cells.get((row_idx, col_idx))
            row.append(_paragraph(cell, font_resolver) if cell is not None else "")
        data.append(row)
    return data


def _style_cell(
    commands: list[tuple[Any, ...]],
    cell: CellSchema,
    row_idx: int,
    col_idx: int,
    min_row: int,
    min_col: int,
    font_resolver: _FontResolver,
) -> None:
    point = (col_idx - min_col, row_idx - min_row)
    font = cell["font"]
    fill = cell["fill"]
    alignment = cell["alignment"]

    commands.append(("FONTNAME", point, point, font_resolver.font_name(font)))
    commands.append(("FONTSIZE", point, point, float(font.get("size") or 11.0)))
    font_color = _hex_color(font.get("color"))
    if font_color is not None:
        commands.append(("TEXTCOLOR", point, point, font_color))
    fill_color = _hex_color(fill.get("bg_color"))
    if fill_color is not None:
        commands.append(("BACKGROUND", point, point, fill_color))
    commands.append(("ALIGN", point, point, _alignment(alignment.get("horizontal"))))
    commands.append(
        ("VALIGN", point, point, _vertical_alignment(alignment.get("vertical")))
    )
    commands.append(("LEFTPADDING", point, point, 4))
    commands.append(("RIGHTPADDING", point, point, 4))
    commands.append(("TOPPADDING", point, point, 3))
    commands.append(("BOTTOMPADDING", point, point, 3))

    _border(commands, point, "LINEABOVE", cell["borders"]["top"])
    _border(commands, point, "LINEBELOW", cell["borders"]["bottom"])
    _border(commands, point, "LINEBEFORE", cell["borders"]["left"])
    _border(commands, point, "LINEAFTER", cell["borders"]["right"])


def _border(
    commands: list[tuple[Any, ...]],
    point: tuple[int, int],
    command: str,
    side: dict[str, Any],
) -> None:
    style = side.get("style")
    if not style:
        return
    width = _BORDER_WIDTHS.get(style, 0.5)
    color = _hex_color(side.get("color")) or colors.black
    commands.append((command, point, point, width, color))


def _border_range(
    commands: list[tuple[Any, ...]],
    start: tuple[int, int],
    end: tuple[int, int],
    command: str,
    side: dict[str, Any],
) -> None:
    style = side.get("style")
    if not style:
        return
    width = _BORDER_WIDTHS.get(style, 0.5)
    color = _hex_color(side.get("color")) or colors.black
    commands.append((command, start, end, width, color))


def _span_commands(
    sheet: SheetSchema,
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
) -> list[tuple[Any, ...]]:
    commands: list[tuple[Any, ...]] = []
    for raw_region in sheet.get("merged_regions", []):
        region = CellRange(raw_region)
        if (
            region.min_col < min_col
            or region.min_row < min_row
            or region.max_col > max_col
            or region.max_row > max_row
        ):
            continue
        start = (region.min_col - min_col, region.min_row - min_row)
        end = (region.max_col - min_col, region.max_row - min_row)
        if start != end:
            commands.append(("SPAN", start, end))
    return commands


def _merged_border_commands(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
) -> list[tuple[Any, ...]]:
    commands: list[tuple[Any, ...]] = []
    for raw_region in sheet.get("merged_regions", []):
        region = CellRange(raw_region)
        if (
            region.min_col < min_col
            or region.min_row < min_row
            or region.max_col > max_col
            or region.max_row > max_row
        ):
            continue
        anchor = cells.get((region.min_row, region.min_col))
        if anchor is None:
            continue
        start = (region.min_col - min_col, region.min_row - min_row)
        end = (region.max_col - min_col, region.max_row - min_row)
        borders = anchor["borders"]
        _border_range(commands, start, end, "LINEABOVE", borders["top"])
        _border_range(commands, start, end, "LINEBELOW", borders["bottom"])
        _border_range(commands, start, end, "LINEBEFORE", borders["left"])
        _border_range(commands, start, end, "LINEAFTER", borders["right"])
    return commands


def _table_style(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    font_resolver: _FontResolver,
) -> TableStyle:
    commands: list[tuple[Any, ...]] = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    commands.extend(_span_commands(sheet, min_col, min_row, max_col, max_row))
    for (row_idx, col_idx), cell in cells.items():
        if min_row <= row_idx <= max_row and min_col <= col_idx <= max_col:
            _style_cell(
                commands,
                cell,
                row_idx,
                col_idx,
                min_row,
                min_col,
                font_resolver,
            )
    commands.extend(
        _merged_border_commands(sheet, cells, min_col, min_row, max_col, max_row)
    )
    return TableStyle(commands)


def _sheet_table(
    bundle: ReportBundle,
    raw_sheet: dict[str, Any],
    *,
    available_width: float,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    font_resolver: _FontResolver,
) -> Table:
    sheet = _sheet_with_overrides(
        raw_sheet,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
    )
    sheet, cells = _expanded_cells(
        bundle,
        sheet,
        streaming_chunk_rows=streaming_chunk_rows,
    )
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    data = _table_data(cells, font_resolver, min_col, min_row, max_col, max_row)
    table = Table(
        data,
        colWidths=_column_widths(
            sheet,
            cells,
            min_col,
            min_row,
            max_col,
            max_row,
            available_width,
        ),
        rowHeights=_row_heights(sheet, min_row, max_row),
        repeatRows=0,
        splitByRow=1,
    )
    table.setStyle(
        _table_style(sheet, cells, min_col, min_row, max_col, max_row, font_resolver)
    )
    return table


def _ensure_output_parent(output_path: str) -> None:
    parent = Path(output_path).parent
    if parent and not parent.exists():
        parent.mkdir(parents=True, exist_ok=True)


# Â§4 Public Functions


def export_report_bundle(
    bundle: ReportBundle,
    output_path: str,
    *,
    column_width_mode: str | None = None,
    row_height_mode: str | None = None,
    default_column_width: float | None = None,
    default_row_height: float | None = None,
    page_size: str = "A4",
    orientation: str = "portrait",
    margin: float = 36,
    streaming_chunk_rows: int = 50_000,
    fonts: dict[str, Any] | None = None,
) -> None:
    """Render a ReportBundle to a styled PDF."""
    if streaming_chunk_rows <= 0:
        raise ValueError(
            f"streaming_chunk_rows must be greater than 0, got {streaming_chunk_rows}"
        )
    if margin < 0:
        raise ValueError(f"margin must be non-negative, got {margin}")

    pagesize = _page_size(page_size, orientation)
    font_resolver = _FontResolver(fonts)
    _ensure_output_parent(output_path)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )
    available_width = pagesize[0] - (margin * 2)
    story: list[Any] = []
    for index, raw_sheet in enumerate(bundle.report["sheets"]):
        if index:
            story.append(PageBreak())
        story.append(
            _sheet_table(
                bundle,
                raw_sheet,
                available_width=available_width,
                column_width_mode=column_width_mode,
                row_height_mode=row_height_mode,
                default_column_width=default_column_width,
                default_row_height=default_row_height,
                streaming_chunk_rows=streaming_chunk_rows,
                font_resolver=font_resolver,
            )
        )
    doc.build(story)


# Â§5 Entrypoints
