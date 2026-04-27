from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import reportlab

from mindoff_dataport import mode, parquet_source
from mindoff_dataport.bundle import load_report_bundle
from mindoff_dataport.pdf_renderer import _FontResolver, _table_style

# §1 Constants & Exceptions

polars = pytest.importorskip("polars", reason="polars not installed")

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _cell(coord: str, value):
    return {
        "coordinate": coord,
        "value": value,
        "cell_type": "string",
        "number_format": None,
        "font": {
            "name": "Calibri",
            "size": 11.0,
            "bold": False,
            "italic": False,
            "underline": None,
            "color": None,
        },
        "fill": {"bg_color": None},
        "alignment": {"horizontal": None, "vertical": None, "wrap_text": False},
        "borders": {
            "top": {"style": None, "color": None},
            "bottom": {"style": None, "color": None},
            "left": {"style": None, "color": None},
            "right": {"style": None, "color": None},
        },
        "merged": False,
        "merge_anchor": None,
    }


def _schema(cells: dict[str, dict], *, dims: str = "A1:C3") -> dict:
    return {
        "sheets": [
            {
                "name": "Sheet1",
                "dimensions": dims,
                "merged_regions": [],
                "column_widths": {},
                "row_heights": {},
                "cells": cells,
            }
        ]
    }


def _schema_with_options(cells: dict[str, dict], *, dims: str = "A1:C3", **options):
    schema = _schema(cells, dims=dims)
    schema["sheets"][0].update(options)
    return schema


def _has_line_command(commands: list[tuple], command: str, width: float) -> bool:
    return any(
        item[:4] == (command, (0, 0), (1, 1), width)
        for item in commands
    )


def _vera_font_path() -> str:
    return str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")


# §4 Public Functions


def test_compile_creates_valid_report_bundle_directory(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{name:string}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:B2",
    )
    df = polars.DataFrame({"A": [1, 2], "B": [3, 4]})
    bundle_path = managed_tmp_dir / "report_bundle"

    bundle = mode.compile(
        schema,
        {"Sheet1": {"name": "Alice", "rows": df}},
        bundle_path=str(bundle_path),
    )

    assert bundle_path.exists()
    assert bundle_path.is_dir()
    assert (bundle_path / "manifest.json").exists()
    assert (bundle_path / "report.json").exists()
    assert bundle.report["sheets"][0]["cells"]["A1"]["value"] == "Alice"
    assert "A2" not in bundle.report["sheets"][0]["cells"]
    source_path = bundle.report["sheets"][0]["dataframe_anchors"][0]["source"]
    assert source_path.startswith("data/")
    assert (bundle_path / source_path).exists()

    manifest = json.loads((bundle_path / "manifest.json").read_text(encoding="utf-8"))
    report = json.loads((bundle_path / "report.json").read_text(encoding="utf-8"))

    assert manifest["bundle_format"] == "directory"
    assert manifest["output_capabilities"] == {
        "xlsx": True,
        "pdf": True,
        "image": False,
    }
    assert report["sheets"][0]["dataframe_anchors"][0]["columns"] == ["A", "B"]


def test_compile_without_bundle_path_creates_temp_directory():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})

    bundle = mode.compile(schema, {"Sheet1": {"rows": df}})

    path = Path(bundle.path)
    assert path.exists()
    assert path.is_dir()
    assert (path / "manifest.json").exists()
    paths = [source["path"] for source in bundle.manifest["dataframe_sources"]]
    assert paths == ["data/Sheet1__rows__dataframe-content.parquet"]
    assert (path / paths[0]).exists()


def test_compile_accepts_parquet_source_without_expanding_rows(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    pq.write_table(pa.table({"A": [1, 2], "B": [3, 4]}), source_path)
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "bundle"

    bundle = mode.compile(
        schema,
        {"Sheet1": {"rows": parquet_source(str(source_path), columns=["B", "A"])}},
        bundle_path=str(bundle_path),
    )

    source = bundle.manifest["dataframe_sources"][0]
    assert source["columns"] == ["B", "A"]
    assert source["rows"] == 2
    assert source["format"] == "parquet"
    assert (bundle_path / source["path"]).exists()


def test_xlsx_export_uses_bundle_data_without_expanding_report(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{name:string}}"),
            "A2": _cell("A2", "{{headers:dataframe-headers}}"),
            "A3": _cell("A3", "{{rows:dataframe-content}}"),
        },
        dims="A1:B3",
    )
    df = polars.DataFrame({"A": [1, 2], "B": [3, 4]})
    bundle = mode.compile(
        schema,
        {"Sheet1": {"name": "Alice", "headers": df, "rows": df}},
    )

    assert "A3" not in bundle.report["sheets"][0]["cells"]
    out = managed_tmp_dir / "out.xlsx"
    paths = mode.export(bundle, str(out), export_mode="streaming")

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert ws["A1"].value == "Alice"
    assert [ws["A2"].value, ws["B2"].value] == ["A", "B"]
    assert [ws["A3"].value, ws["B3"].value, ws["A4"].value, ws["B4"].value] == [
        1,
        3,
        2,
        4,
    ]
    wb.close()


def test_export_accepts_bundle_path(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "report_bundle"
    mode.compile(schema, {"Sheet1": {"name": "Alice"}}, bundle_path=str(bundle_path))
    out = managed_tmp_dir / "out.xlsx"

    mode.export(str(bundle_path), str(out))

    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb["Sheet1"]["A1"].value == "Alice"
    wb.close()


def test_xlsx_export_preserves_gridline_visibility(managed_tmp_dir: Path):
    schema = _schema_with_options(
        {"A1": _cell("A1", "No gridlines")},
        dims="A1:A1",
        show_gridlines=False,
    )
    bundle = mode.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "no-gridlines.xlsx"

    mode.export(bundle, str(out))

    wb = openpyxl.load_workbook(out)
    assert wb["Sheet1"].sheet_view.showGridLines is False
    wb.close()


def test_xlsx_export_preserves_merged_region_border(managed_tmp_dir: Path):
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thick", "color": "FF000000"},
        "left": {"style": "medium", "color": "FF000000"},
        "right": {"style": "dashed", "color": "FF000000"},
    }
    shadow = _cell("B2", "")
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema({"A1": title, "B2": shadow}, dims="A1:B2")
    schema["sheets"][0]["merged_regions"] = ["A1:B2"]
    bundle = mode.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "merged-border.xlsx"

    mode.export(bundle, str(out))

    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    assert ws["A1"].border.top.style == "thin"
    assert ws["A2"].border.bottom.style == "thick"
    assert ws["A1"].border.left.style == "medium"
    assert ws["B1"].border.right.style == "dashed"
    wb.close()


def test_export_pdf_creates_nonempty_pdf(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mode.compile(schema, {"Sheet1": {"name": "Alice"}})
    out = managed_tmp_dir / "out.pdf"

    mode.export(bundle, str(out), format="pdf")

    assert out.exists()
    assert out.stat().st_size > 0
    assert out.read_bytes().startswith(b"%PDF")


def test_export_image_is_reserved(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mode.compile(schema, {"Sheet1": {"name": "Alice"}})

    with pytest.raises(NotImplementedError, match="image"):
        mode.export(bundle, str(managed_tmp_dir / "out.png"), format="image")


def test_pdf_export_renders_parquet_backed_dataframe(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    pq.write_table(pa.table({"A": [1, 2], "B": [3, 4]}), source_path)
    schema = _schema(
        {
            "A1": _cell("A1", "{{headers:dataframe-headers}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:B2",
    )
    bundle = mode.compile(
        schema,
        {
            "Sheet1": {
                "headers": parquet_source(str(source_path), columns=["B", "A"]),
                "rows": parquet_source(str(source_path), columns=["B", "A"]),
            }
        },
    )
    out = managed_tmp_dir / "table.pdf"

    mode.export(bundle, str(out), format="pdf", streaming_chunk_rows=1)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_export_handles_merges_and_basic_styles(managed_tmp_dir: Path):
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["font"] = dict(title["font"])
    title["font"]["bold"] = True
    title["fill"] = {"bg_color": "FF003366"}
    title["alignment"] = dict(title["alignment"])
    title["alignment"]["horizontal"] = "center"
    shadow = _cell("B1", "")
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema({"A1": title, "B1": shadow}, dims="A1:B1")
    schema["sheets"][0]["merged_regions"] = ["A1:B1"]
    bundle = mode.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "styled.pdf"

    mode.export(bundle, str(out), format="pdf")

    assert out.exists()
    assert out.stat().st_size > 0


def test_pdf_table_style_applies_merged_region_border_to_span():
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thick", "color": "FF000000"},
        "left": {"style": "medium", "color": "FF000000"},
        "right": {"style": "dashed", "color": "FF000000"},
    }
    sheet = _schema({"A1": title}, dims="A1:B2")["sheets"][0]
    sheet["merged_regions"] = ["A1:B2"]

    commands = _table_style(
        sheet, {(1, 1): title}, 1, 1, 2, 2, _FontResolver()
    ).getCommands()

    assert _has_line_command(commands, "LINEABOVE", 0.5)
    assert _has_line_command(commands, "LINEBELOW", 1.5)
    assert _has_line_command(commands, "LINEBEFORE", 1.0)
    assert _has_line_command(commands, "LINEAFTER", 0.75)


def test_pdf_table_style_does_not_draw_default_grid_for_empty_cells():
    cell = _cell("A1", "Visible")
    sheet = _schema({"A1": cell}, dims="A1:B2")["sheets"][0]

    commands = _table_style(
        sheet, {(1, 1): cell}, 1, 1, 2, 2, _FontResolver()
    ).getCommands()

    assert not any(item[0] == "GRID" for item in commands)
    assert not any(
        item[0] in {"LINEABOVE", "LINEBELOW", "LINEBEFORE", "LINEAFTER"}
        for item in commands
    )


def test_pdf_table_style_uses_registered_custom_font():
    cell = _cell("A1", "Custom Font")
    cell["font"] = dict(cell["font"])
    cell["font"]["name"] = "Vera"
    sheet = _schema({"A1": cell}, dims="A1:A1")["sheets"][0]
    resolver = _FontResolver({"Vera": _vera_font_path()})

    commands = _table_style(sheet, {(1, 1): cell}, 1, 1, 1, 1, resolver).getCommands()

    font_commands = [item for item in commands if item[0] == "FONTNAME"]
    assert font_commands
    assert str(font_commands[0][3]).startswith("Mindoff-Vera-regular-")


def test_pdf_export_accepts_custom_fonts(managed_tmp_dir: Path):
    cell = _cell("A1", "Custom Font")
    cell["font"] = dict(cell["font"])
    cell["font"]["name"] = "Vera"
    schema = _schema({"A1": cell}, dims="A1:A1")
    bundle = mode.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "custom-font.pdf"

    mode.export(
        bundle,
        str(out),
        format="pdf",
        fonts={"Vera": {"regular": _vera_font_path()}},
    )

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_export_accepts_bundle_path(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "report_bundle"
    mode.compile(schema, {"Sheet1": {"name": "Alice"}}, bundle_path=str(bundle_path))
    out = managed_tmp_dir / "out.pdf"

    mode.export(str(bundle_path), str(out), format="pdf")

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_export_auto_delete_removes_bundle_after_success(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "bundle"
    bundle = mode.compile(
        schema,
        {"Sheet1": {"name": "Alice"}},
        bundle_path=str(bundle_path),
    )
    out = managed_tmp_dir / "out.pdf"

    mode.export(bundle, str(out), format="pdf", auto_delete_bundle=True)

    assert out.exists()
    assert not bundle_path.exists()


def test_pdf_export_rejects_streaming_export_mode(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mode.compile(schema, {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="PDF export does not support export_mode"):
        mode.export(
            bundle,
            str(managed_tmp_dir / "out.pdf"),
            format="pdf",
            export_mode="streaming",
        )


def test_load_rejects_malformed_bundle(managed_tmp_dir: Path):
    bad_path = managed_tmp_dir / "bad"
    bad_path.mkdir()
    (bad_path / "manifest.json").write_text("{}", encoding="utf-8")

    with pytest.raises(ValueError, match="Malformed report bundle"):
        load_report_bundle(str(bad_path))


def test_export_auto_delete_removes_bundle_after_success(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "bundle"
    bundle = mode.compile(
        schema,
        {"Sheet1": {"name": "Alice"}},
        bundle_path=str(bundle_path),
    )
    out = managed_tmp_dir / "out.xlsx"

    mode.export(bundle, str(out), auto_delete_bundle=True)

    assert out.exists()
    assert not bundle_path.exists()


def test_failed_export_preserves_bundle_directory(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})
    bundle_path = managed_tmp_dir / "bundle"
    bundle = mode.compile(
        schema,
        {"Sheet1": {"rows": df}},
        bundle_path=str(bundle_path),
    )

    with pytest.raises(ValueError, match="Fidelity XLSX export does not support"):
        mode.export(
            bundle,
            str(managed_tmp_dir / "out.xlsx"),
            auto_delete_bundle=True,
        )

    assert bundle_path.exists()


# §5 Entrypoints
