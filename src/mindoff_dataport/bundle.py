from __future__ import annotations

import datetime
import json
import shutil
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)

from .template_contract import (
    PLACEHOLDER_RE,
    _DATAFRAME_TYPES,
    _SCALAR_TYPES,
    _infer_cell_type,
    _parse_dims,
    _repeat_sections,
    _resolve_sheet_payloads,
    _substitute_scalars,
    _to_headers,
    get_template_inputs,
)
from .page_breaks import resolve_compiled_sheet_page_breaks
from .repeat import RepeatRecords, is_repeat_records
from .schema import CellSchema, SheetSchema, WorkbookSchema

__all__ = [
    "ReportBundle",
    "RepeatRecords",
    "compile_report_bundle",
    "load_report_bundle",
]

# §1. Constants & Exceptions

BUNDLE_VERSION = "1.1"
_ALIGNMENTS = frozenset({"left", "center", "right"})
_DATAFRAME_SHIFT_MODES = frozenset({"both", "horizontal", "vertical", "none"})

# §2. Classes and Sub Classes


@dataclass(frozen=True)
class ReportBundle:
    """Canonical intermediate report artifact, backed by a directory."""

    manifest: dict[str, Any]
    report: dict[str, Any]
    path: str

    def write(self, bundle_path: str) -> None:
        _copy_bundle_dir(self, bundle_path)

    @classmethod
    def load(cls, bundle_path: str) -> "ReportBundle":
        return load_report_bundle(bundle_path)


# §3. Private Helper Functions


def _copy_bundle_dir(bundle: ReportBundle, bundle_path: str) -> None:
    source = Path(bundle.path).resolve()
    target = Path(bundle_path).resolve()
    if source == target:
        _write_bundle_metadata(bundle)
        return
    if target.exists() and not target.is_dir():
        raise ValueError(f"Report bundle path must be a directory: {bundle_path}")
    shutil.copytree(source, target, dirs_exist_ok=True)


def _write_bundle_metadata(bundle: ReportBundle) -> None:
    path = Path(bundle.path)
    path.mkdir(parents=True, exist_ok=True)
    (path / "manifest.json").write_text(
        json.dumps(bundle.manifest, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    (path / "report.json").write_text(
        json.dumps(bundle.report, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def _prepare_bundle_dir(bundle_path: str | None) -> Path:
    if bundle_path is None:
        path = Path(tempfile.mkdtemp(prefix="mindoff_report_bundle_"))
    else:
        path = Path(bundle_path)
    if path.exists() and not path.is_dir():
        raise ValueError(f"Report bundle path must be a directory: {path}")
    path.mkdir(parents=True, exist_ok=True)
    (path / "data").mkdir(exist_ok=True)
    return path


def _parquet_metadata(path: Path) -> tuple[list[str], int]:
    parquet_file = pq.ParquetFile(path)
    columns = [str(name) for name in parquet_file.schema_arrow.names]
    return columns, int(parquet_file.metadata.num_rows)


def _write_source_file(
    *,
    value: Any,
    source_id: str,
    bundle_dir: Path,
    source_cache: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    cache_key = id(value)
    cached = source_cache.get(cache_key)
    if cached is not None:
        return cached

    record = _write_dataframe_source(
        value=value,
        source_id=source_id,
        bundle_dir=bundle_dir,
    )
    source_cache[cache_key] = record
    return record


def _write_dataframe_source(
    *, value: Any, source_id: str, bundle_dir: Path
) -> dict[str, Any]:
    rel_path = f"data/{source_id}.parquet"
    output_path = bundle_dir / rel_path
    module = getattr(type(value), "__module__", "") or ""
    qualname = type(value).__qualname__

    if "polars" in module and qualname == "LazyFrame":
        value.sink_parquet(output_path)
    elif "polars" in module and qualname == "DataFrame":
        value.write_parquet(output_path)
    else:
        raise TypeError(
            f"Expected a polars DataFrame or LazyFrame, got {type(value).__name__}"
        )

    columns, row_count = _parquet_metadata(output_path)
    record = {
        "id": source_id,
        "path": rel_path,
        "format": "parquet",
        "columns": columns,
        "rows": row_count,
        "file_backed": True,
    }
    return record


def _write_repeat_record_source(
    *, value: Any, source_id: str, bundle_dir: Path
) -> dict[str, Any]:
    rel_path = f"data/{source_id}.parquet"
    output_path = bundle_dir / rel_path
    module = getattr(type(value), "__module__", "") or ""
    qualname = type(value).__qualname__

    if "polars" in module and qualname == "LazyFrame":
        value.sink_parquet(output_path)
    elif "polars" in module and qualname == "DataFrame":
        value.write_parquet(output_path)
    elif isinstance(value, list):
        try:
            import polars as pl
        except ImportError as exc:  # pragma: no cover - exercised only without optional dep
            raise TypeError("repeat_records list input requires polars to be installed") from exc
        pl.DataFrame(value).write_parquet(output_path)
    elif isinstance(value, Iterable) and not isinstance(value, (str, bytes, dict)):
        _write_iterable_record_source(value, output_path)
    else:
        raise TypeError(
            f"Expected repeat_records records as a polars DataFrame/LazyFrame or list of dicts, got {type(value).__name__}"
        )

    columns, row_count = _parquet_metadata(output_path)
    return {
        "id": source_id,
        "path": rel_path,
        "format": "parquet",
        "columns": columns,
        "rows": row_count,
        "file_backed": True,
        "repeat_records": True,
    }


def _write_iterable_record_source(value: Iterable[Any], output_path: Path) -> None:
    writer: pq.ParquetWriter | None = None
    batch: list[dict[str, Any]] = []
    try:
        for item in value:
            if not isinstance(item, dict):
                raise TypeError(
                    f"repeat_records iterable items must be dicts, got {type(item).__name__}"
                )
            batch.append(item)
            if len(batch) >= 10_000:
                writer = _write_record_batch(batch, output_path, writer)
                batch = []
        if batch:
            writer = _write_record_batch(batch, output_path, writer)
        if writer is None:
            pq.write_table(pa.Table.from_pylist([]), output_path)
    finally:
        if writer is not None:
            writer.close()


def _write_record_batch(
    batch: list[dict[str, Any]],
    output_path: Path,
    writer: pq.ParquetWriter | None,
) -> pq.ParquetWriter:
    table = pa.Table.from_pylist(batch)
    if writer is None:
        writer = pq.ParquetWriter(output_path, table.schema)
    writer.write_table(table)
    return writer


def _schema_value(value: Any) -> Any:
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return value


def _unique_id(used: set[str], raw: str) -> str:
    safe = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in raw)
    safe = safe.strip("._-") or "source"
    candidate = safe
    idx = 2
    while candidate in used:
        candidate = f"{safe}_{idx}"
        idx += 1
    used.add(candidate)
    return candidate


def _coord_indexes(coord: str) -> tuple[int, int]:
    col_letter, row_idx = coordinate_from_string(coord)
    return row_idx, column_index_from_string(col_letter)


def _compile_sheet(
    *,
    sheet: SheetSchema,
    output_name: str,
    sheet_data: dict[str, Any],
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    dataframe_options: dict[str, Any],
    dataframe_shift: str,
) -> dict[str, Any]:
    repeats = _repeat_sections(sheet)
    if repeats:
        return _compile_repeat_sheet(
            sheet=sheet,
            repeats=repeats,
            output_name=output_name,
            sheet_data=sheet_data,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
            dataframe_options=dataframe_options,
            dataframe_shift=dataframe_shift,
        )

    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    static_cells: dict[str, CellSchema] = {}
    anchors: list[dict[str, Any]] = []

    for coord, cell in sheet["cells"].items():
        value = cell.get("value")
        if not isinstance(value, str):
            static_cells[coord] = cell
            continue

        full_match = PLACEHOLDER_RE.fullmatch(value.strip())
        if (
            full_match
            and full_match.group(2) in _DATAFRAME_TYPES
            and full_match.group(1) in sheet_data
        ):
            key = full_match.group(1)
            anchor_type = full_match.group(2)
            columns = _to_headers(sheet_data[key])
            column_layouts = _column_layouts(
                columns, dataframe_options.get(output_name, {}).get(key)
            )
            col_letter, row_idx = coordinate_from_string(coord)
            col_idx = column_index_from_string(col_letter)
            max_col = max(max_col, col_idx + max(_occupied_width(column_layouts) - 1, 0))

            if anchor_type == "dataframe-header":
                anchor_id = _unique_id(used_ids, f"{output_name}__{key}__{anchor_type}")
                anchors.append(
                    _dataframe_anchor(
                        anchor_id=anchor_id,
                        key=key,
                        placeholder_type=anchor_type,
                        coord=coord,
                        row_idx=row_idx,
                        col_idx=col_idx,
                        columns=columns,
                        column_layouts=column_layouts,
                        source_record=None,
                        cell=cell,
                    )
                )
                continue

            source_id_type = (
                "dataframe-content" if anchor_type == "dataframe" else anchor_type
            )
            anchor_id = _unique_id(used_ids, f"{output_name}__{key}__{source_id_type}")
            source_record = _write_source_file(
                value=sheet_data[key],
                source_id=anchor_id,
                bundle_dir=bundle_dir,
                source_cache=source_cache,
            )
            if source_record not in data_sources:
                data_sources.append(source_record)
            if anchor_type == "dataframe":
                header_id = _unique_id(
                    used_ids, f"{output_name}__{key}__dataframe-header"
                )
                anchors.append(
                    _dataframe_anchor(
                        anchor_id=header_id,
                        key=key,
                        placeholder_type="dataframe-header",
                        coord=coord,
                        row_idx=row_idx,
                        col_idx=col_idx,
                        columns=columns,
                        column_layouts=column_layouts,
                        source_record=None,
                        cell=cell,
                    )
                )
                row_idx += 1
            anchors.append(
                _dataframe_anchor(
                    anchor_id=anchor_id,
                    key=key,
                    placeholder_type="dataframe-content",
                    coord=coord,
                    row_idx=row_idx,
                    col_idx=col_idx,
                    columns=columns,
                    column_layouts=column_layouts,
                    source_record=source_record,
                    cell=cell,
                )
            )
            continue

        new_value = _substitute_scalars(value, sheet_data)
        new_cell: dict[str, Any] = dict(cell)
        new_cell["value"] = _schema_value(new_value)
        if full_match and full_match.group(2) in _SCALAR_TYPES:
            new_cell["cell_type"] = _infer_cell_type(new_value)
        static_cells[coord] = new_cell  # type: ignore[assignment]

    dimensions = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    result = dict(sheet)
    result["name"] = output_name
    result["dimensions"] = dimensions
    result["cells"] = static_cells
    result["dataframe_anchors"] = anchors
    _shift_template_content_around_dataframes(result, dataframe_shift=dataframe_shift)
    _validate_template_merges_do_not_overlap_dataframes(result)
    resolve_compiled_sheet_page_breaks(result)
    return result


def _compile_repeat_sheet(
    *,
    sheet: SheetSchema,
    repeats: list[dict[str, Any]],
    output_name: str,
    sheet_data: dict[str, Any],
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    dataframe_options: dict[str, Any],
    dataframe_shift: str,
) -> dict[str, Any]:
    for repeat in repeats:
        _validate_repeat_layout(sheet, repeat)
    min_col, min_row, max_col, _ = _parse_dims(sheet["dimensions"])
    static_cells: dict[str, CellSchema] = {}
    repeat_sections: list[dict[str, Any]] = []
    repeat_ranges = [(repeat["start_row"], repeat["end_row"]) for repeat in repeats]

    for coord, cell in sheet["cells"].items():
        row_idx, _ = _coord_indexes(coord)
        if _row_in_ranges(row_idx, repeat_ranges):
            continue
        if _is_non_anchor_merged_cell(cell):
            continue
        compiled = _compile_cell(
            cell=cell,
            payload=sheet_data,
            output_name=output_name,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
            sheet_dataframe_options=dataframe_options.get(output_name, {}),
        )
        static_cells.update(compiled["cells"])

    for repeat in repeats:
        section, section_max_col = _compile_repeat_section(
            sheet=sheet,
            repeat=repeat,
            output_name=output_name,
            sheet_data=sheet_data,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
            dataframe_options=dataframe_options,
            dataframe_shift=dataframe_shift,
        )
        repeat_sections.append(section)
        max_col = max(max_col, section_max_col)

    result = dict(sheet)
    result["name"] = output_name
    result["dimensions"] = (
        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{_last_static_row(sheet, repeats)}"
    )
    result["merged_regions"] = _static_merged_regions(sheet, repeats)
    result["cells"] = static_cells
    result["dataframe_anchors"] = []
    result["repeat_sections"] = repeat_sections
    resolve_compiled_sheet_page_breaks(result)
    return result


def _compile_repeat_section(
    *,
    sheet: SheetSchema,
    repeat: dict[str, Any],
    output_name: str,
    sheet_data: dict[str, Any],
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    dataframe_options: dict[str, Any],
    dataframe_shift: str,
) -> tuple[dict[str, Any], int]:
    repeat_key = repeat["key"]
    records = sheet_data[repeat_key]
    if is_repeat_records(records):
        return _compile_source_repeat_section(
            sheet=sheet,
            repeat=repeat,
            output_name=output_name,
            repeat_records=records,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
            dataframe_options=dataframe_options,
            dataframe_shift=dataframe_shift,
        )
    block_start = repeat["start_row"] + 1
    block_end = repeat["end_row"] - 1
    repeat_merges = _repeat_merged_regions(sheet, repeat, block_start, block_end)
    compiled_records: list[dict[str, Any]] = []
    block_height = max(block_end - block_start + 1, 0)
    max_col = _parse_dims(sheet["dimensions"])[2]
    for index, record in enumerate(records):
        record_cells: list[dict[str, Any]] = []
        record_anchors: list[dict[str, Any]] = []
        for coord, cell in sheet["cells"].items():
            row_idx, col_idx = _coord_indexes(coord)
            if row_idx < block_start or row_idx > block_end:
                continue
            if _is_non_anchor_merged_cell(cell):
                continue
            compiled = _compile_cell(
                cell=cell,
                payload=record,
                output_name=f"{output_name}__{repeat_key}_{index + 1}",
                bundle_dir=bundle_dir,
                data_sources=data_sources,
                used_ids=used_ids,
                source_cache=source_cache,
                sheet_dataframe_options=dataframe_options.get(output_name, {}),
            )
            row_offset = row_idx - block_start
            for item in compiled["cells"].values():
                record_cells.append(
                    {
                        "row_offset": row_offset,
                        "start_col": col_idx,
                        "cell": item,
                    }
                )
            for anchor in compiled["anchors"]:
                anchor = dict(anchor)
                anchor["start_row_offset"] = anchor.pop("start_row") - block_start
                record_anchors.append(anchor)
                max_col = max(
                    max_col,
                    anchor["start_col"] + max(_occupied_width(anchor["column_layouts"]) - 1, 0),
                )
        (
            record_cells,
            record_anchors,
            record_merges,
            record_block_height,
        ) = _shift_repeat_record_content_around_dataframes(
            record_cells,
            record_anchors,
            repeat_merges,
            block_height=block_height,
            dataframe_shift=dataframe_shift,
        )
        for anchor in record_anchors:
            output_range = _repeat_dataframe_output_range(anchor)
            if output_range is None:
                continue
            max_col = max(max_col, output_range.max_col)
        compiled_records.append(
            {
                "index": index,
                "block_height": record_block_height,
                "cells": record_cells,
                "dataframe_anchors": record_anchors,
                "merged_regions": record_merges,
            }
        )
    cell_templates = _compact_repeat_record_cells(compiled_records)
    section = {
            "key": repeat_key,
            "start_row": repeat["start_row"],
            "end_row": repeat["end_row"],
            "template_start_row": block_start,
            "template_end_row": block_end,
            "block_height": block_height,
            "merged_regions": repeat_merges,
            "cell_templates": cell_templates,
            "records": compiled_records,
        }
    _validate_repeat_merges_do_not_overlap_dataframes(section, dataframe_shift)
    return section, max_col


def _compile_source_repeat_section(
    *,
    sheet: SheetSchema,
    repeat: dict[str, Any],
    output_name: str,
    repeat_records: RepeatRecords,
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    dataframe_options: dict[str, Any],
    dataframe_shift: str,
) -> tuple[dict[str, Any], int]:
    repeat_key = repeat["key"]
    block_start = repeat["start_row"] + 1
    block_end = repeat["end_row"] - 1
    repeat_merges = _repeat_merged_regions(sheet, repeat, block_start, block_end)
    block_height = max(block_end - block_start + 1, 0)
    max_col = _parse_dims(sheet["dimensions"])[2]

    record_source_id = _unique_id(used_ids, f"{output_name}__{repeat_key}__records")
    record_source = _write_repeat_record_source(
        value=repeat_records.records,
        source_id=record_source_id,
        bundle_dir=bundle_dir,
    )
    data_sources.append(record_source)
    record_columns = set(record_source["columns"])
    constants = repeat_records.constants

    record_cells: list[dict[str, Any]] = []
    record_anchors: list[dict[str, Any]] = []
    for coord, cell in sheet["cells"].items():
        row_idx, col_idx = _coord_indexes(coord)
        if row_idx < block_start or row_idx > block_end:
            continue
        if _is_non_anchor_merged_cell(cell):
            continue
        compiled = _compile_source_repeat_cell(
            cell=cell,
            constants=constants,
            record_columns=record_columns,
            output_name=output_name,
            repeat_key=repeat_key,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
            sheet_dataframe_options=dataframe_options.get(output_name, {}),
        )
        row_offset = row_idx - block_start
        for item in compiled["cells"]:
            item["row_offset"] = row_offset
            item["start_col"] = col_idx
            record_cells.append(item)
        for anchor in compiled["anchors"]:
            anchor = dict(anchor)
            anchor["start_row_offset"] = anchor.pop("start_row") - block_start
            record_anchors.append(anchor)
            max_col = max(
                max_col,
                anchor["start_col"] + max(_occupied_width(anchor["column_layouts"]) - 1, 0),
            )

    shifted_cells, record_anchors, record_merges, record_block_height = (
        _shift_repeat_record_content_around_dataframes(
            record_cells,
            record_anchors,
            repeat_merges,
            block_height=block_height,
            dataframe_shift=dataframe_shift,
        )
    )
    cell_templates = _compact_source_repeat_cells(shifted_cells)
    for anchor in record_anchors:
        output_range = _repeat_dataframe_output_range(anchor)
        if output_range is not None:
            max_col = max(max_col, output_range.max_col)
    section = {
        "key": repeat_key,
        "start_row": repeat["start_row"],
        "end_row": repeat["end_row"],
        "template_start_row": block_start,
        "template_end_row": block_end,
        "block_height": block_height,
        "record_block_height": record_block_height,
        "record_count": record_source["rows"],
        "record_source": record_source["path"],
        "record_source_format": record_source["format"],
        "record_columns": record_source["columns"],
        "constants": {
            key: _schema_value(value)
            for key, value in constants.items()
            if not _is_dataframe_like(value)
        },
        "merged_regions": repeat_merges,
        "cell_templates": cell_templates,
        "record_bindings": shifted_cells,
        "dataframe_anchors": record_anchors,
        "merged_record_regions": record_merges,
    }
    _validate_repeat_merges_do_not_overlap_dataframes(section, dataframe_shift)
    return section, max_col


def _compile_source_repeat_cell(
    *,
    cell: CellSchema,
    constants: dict[str, Any],
    record_columns: set[str],
    output_name: str,
    repeat_key: str,
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    sheet_dataframe_options: dict[str, Any],
) -> dict[str, Any]:
    value = cell.get("value")
    if not isinstance(value, str):
        return {"cells": [{"cell": cell, "value": _schema_value(value)}], "anchors": []}

    full_match = PLACEHOLDER_RE.fullmatch(value.strip())
    if full_match and full_match.group(2) in _DATAFRAME_TYPES:
        key = full_match.group(1)
        if key not in constants:
            raise KeyError(
                f"Repeat section '{repeat_key}' requires dataframe '{key}' as a repeat_records constant"
            )
        anchor_type = full_match.group(2)
        row_idx, col_idx = _coord_indexes(cell["coordinate"])
        return {
            "cells": [],
            "anchors": _compile_dataframe_anchors(
                key=key,
                anchor_type=anchor_type,
                cell=cell,
                row_idx=row_idx,
                col_idx=col_idx,
                value=constants[key],
                output_name=f"{output_name}__{repeat_key}__constant",
                bundle_dir=bundle_dir,
                data_sources=data_sources,
                used_ids=used_ids,
                source_cache=source_cache,
                sheet_dataframe_options=sheet_dataframe_options,
            ),
        }

    keys = [
        match.group(1)
        for match in PLACEHOLDER_RE.finditer(value)
        if match.group(2) in _SCALAR_TYPES
    ]
    missing = [key for key in keys if key not in record_columns and key not in constants]
    if missing:
        raise KeyError(
            f"Repeat section '{repeat_key}' requires scalar column '{missing[0]}' in repeat_records"
        )
    return {
        "cells": [
            {
                "cell": cell,
                "value_template": value,
                "scalar_keys": keys,
                "full_scalar": bool(full_match and full_match.group(2) in _SCALAR_TYPES),
            }
        ],
        "anchors": [],
    }


def _is_dataframe_like(value: Any) -> bool:
    module = getattr(type(value), "__module__", "") or ""
    qualname = type(value).__qualname__
    return "polars" in module and qualname in {"DataFrame", "LazyFrame"}


def _compact_source_repeat_cells(cells: list[dict[str, Any]]) -> list[CellSchema]:
    templates: list[CellSchema] = []
    template_ids: dict[str, int] = {}
    for item in cells:
        cell = item.pop("cell")
        template = _repeat_cell_template(cell)
        key = json.dumps(template, sort_keys=True, separators=(",", ":"), default=str)
        template_id = template_ids.get(key)
        if template_id is None:
            template_id = len(templates)
            template_ids[key] = template_id
            templates.append(template)  # type: ignore[arg-type]
        item["cell_template"] = template_id
    return templates


def _compact_repeat_record_cells(records: list[dict[str, Any]]) -> list[CellSchema]:
    templates: list[CellSchema] = []
    template_ids: dict[str, int] = {}
    for record in records:
        compact_cells: list[dict[str, Any]] = []
        for item in record["cells"]:
            cell = item["cell"]
            template = _repeat_cell_template(cell)
            key = json.dumps(template, sort_keys=True, separators=(",", ":"), default=str)
            template_id = template_ids.get(key)
            if template_id is None:
                template_id = len(templates)
                template_ids[key] = template_id
                templates.append(template)  # type: ignore[arg-type]
            compact_cells.append(
                {
                    "row_offset": item["row_offset"],
                    "start_col": item["start_col"],
                    "cell_template": template_id,
                    "value": cell.get("value"),
                    "cell_type": cell.get("cell_type"),
                }
            )
        record["cells"] = compact_cells
    return templates


def _repeat_cell_template(cell: CellSchema) -> dict[str, Any]:
    template = dict(cell)
    template["coordinate"] = "A1"
    template["value"] = None
    template["cell_type"] = "empty"
    template["merged"] = False
    template["merge_anchor"] = None
    return template


def _validate_repeat_layout(sheet: SheetSchema, repeat: dict[str, Any]) -> None:
    block_start = repeat["start_row"] + 1
    block_end = repeat["end_row"] - 1
    content_rows = _repeat_content_rows(sheet, block_start, block_end)
    for region in sheet.get("merged_regions", []):
        start, end = region.split(":")
        start_row, _ = _coord_indexes(start)
        end_row, _ = _coord_indexes(end)
        if start_row < block_start or end_row > block_end:
            continue
        if any(start_row <= row_idx <= end_row for row_idx in content_rows):
            raise ValueError(
                f"Repeat section '{repeat['key']}' does not support merged cells that overlap dataframe-content rows"
            )


def _row_in_ranges(row_idx: int, ranges: list[tuple[int, int]]) -> bool:
    return any(start <= row_idx <= end for start, end in ranges)


def _is_non_anchor_merged_cell(cell: CellSchema) -> bool:
    return bool(cell.get("merged")) and cell.get("merge_anchor") not in (
        None,
        cell["coordinate"],
    )


def _repeat_content_rows(
    sheet: SheetSchema, block_start: int, block_end: int
) -> set[int]:
    rows: set[int] = set()
    for cell in sheet["cells"].values():
        row_idx, _ = _coord_indexes(cell["coordinate"])
        if row_idx < block_start or row_idx > block_end:
            continue
        value = cell.get("value")
        if not isinstance(value, str):
            continue
        match = PLACEHOLDER_RE.fullmatch(value.strip())
        if not match:
            continue
        if match.group(2) == "dataframe-content":
            rows.add(row_idx)
        elif match.group(2) == "dataframe":
            rows.add(row_idx + 1)
    return rows


def _static_merged_regions(sheet: SheetSchema, repeats: list[dict[str, Any]]) -> list[str]:
    regions: list[str] = []
    repeat_ranges = [(repeat["start_row"], repeat["end_row"]) for repeat in repeats]
    for region in sheet.get("merged_regions", []):
        start, end = region.split(":")
        start_row, _ = _coord_indexes(start)
        end_row, _ = _coord_indexes(end)
        if not any(start_row <= repeat_end and end_row >= repeat_start for repeat_start, repeat_end in repeat_ranges):
            regions.append(region)
    return regions


def _last_static_row(sheet: SheetSchema, repeats: list[dict[str, Any]]) -> int:
    _, min_row, _, max_row = _parse_dims(sheet["dimensions"])
    repeat_ranges = [(repeat["start_row"], repeat["end_row"]) for repeat in repeats]
    last = min_row
    for cell in sheet["cells"].values():
        row_idx, _ = _coord_indexes(cell["coordinate"])
        value = cell.get("value")
        if _row_in_ranges(row_idx, repeat_ranges):
            continue
        if value not in (None, ""):
            last = max(last, row_idx)
    return max(last, max(repeat["start_row"] for repeat in repeats), min(max_row, last))


def _repeat_merged_regions(
    sheet: SheetSchema, repeat: dict[str, Any], block_start: int, block_end: int
) -> list[dict[str, Any]]:
    regions: list[dict[str, Any]] = []
    for region in sheet.get("merged_regions", []):
        start, end = region.split(":")
        start_row, start_col = _coord_indexes(start)
        end_row, end_col = _coord_indexes(end)
        if start_row < block_start or end_row > block_end:
            continue
        regions.append(
            {
                "min_row_offset": start_row - block_start,
                "max_row_offset": end_row - block_start,
                "min_col": start_col,
                "max_col": end_col,
            }
        )
    return regions


def _shift_repeat_record_content_around_dataframes(
    record_cells: list[dict[str, Any]],
    record_anchors: list[dict[str, Any]],
    merged_regions: list[dict[str, Any]],
    *,
    block_height: int,
    dataframe_shift: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], int]:
    if dataframe_shift == "none":
        return record_cells, record_anchors, merged_regions, block_height
    shifted_anchors, footprints = _shift_repeat_dataframe_anchors(
        record_anchors, dataframe_shift=dataframe_shift
    )
    if not footprints:
        return record_cells, shifted_anchors, merged_regions, block_height

    merge_shifts: dict[tuple[int, int], tuple[int, int]] = {}
    shifted_merges: list[dict[str, Any]] = []
    max_row_offset = max(block_height - 1, 0)
    for merge in merged_regions:
        merge_range = CellRange(
            min_col=merge["min_col"],
            min_row=merge["min_row_offset"] + 1,
            max_col=merge["max_col"],
            max_row=merge["max_row_offset"] + 1,
        )
        row_shift, col_shift = _shift_for_range(
            merge_range, footprints, dataframe_shift=dataframe_shift
        )
        shifted = _shift_range(merge_range, row_shift, col_shift)
        shifted_merges.append(
            {
                "min_row_offset": shifted.min_row - 1,
                "max_row_offset": shifted.max_row - 1,
                "min_col": shifted.min_col,
                "max_col": shifted.max_col,
            }
        )
        max_row_offset = max(max_row_offset, shifted.max_row - 1)
        for row_idx in range(merge_range.min_row, merge_range.max_row + 1):
            for col_idx in range(merge_range.min_col, merge_range.max_col + 1):
                merge_shifts[(row_idx - 1, col_idx)] = (row_shift, col_shift)

    shifted_cells: list[dict[str, Any]] = []
    for item in record_cells:
        row_offset = int(item["row_offset"])
        start_col = int(item["start_col"])
        row_shift, col_shift = merge_shifts.get(
            (row_offset, start_col),
            _shift_for_cell(
                row_offset + 1, start_col, footprints, dataframe_shift=dataframe_shift
            ),
        )
        shifted_cell = _shift_cell(item["cell"], row_shift, col_shift)
        shifted_item = dict(item)
        shifted_item["row_offset"] = row_offset + row_shift
        shifted_item["start_col"] = start_col + col_shift
        shifted_item["cell"] = shifted_cell
        shifted_cells.append(shifted_item)
        max_row_offset = max(max_row_offset, row_offset + row_shift)

    for anchor in shifted_anchors:
        output_range = _repeat_dataframe_output_range(anchor)
        if output_range is not None:
            max_row_offset = max(max_row_offset, output_range.max_row - 1)

    return shifted_cells, shifted_anchors, shifted_merges, max_row_offset + 1


def _shift_repeat_dataframe_anchors(
    anchors: list[dict[str, Any]], *, dataframe_shift: str
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    resolved_anchors: list[dict[str, Any]] = []
    resolved_footprints: list[dict[str, Any]] = []
    for anchor in sorted(
        anchors,
        key=lambda item: (
            int(item["start_row_offset"]),
            int(item["start_col"]),
            str(item["coordinate"]),
            str(item["key"]),
            str(item["placeholder_type"]),
        ),
    ):
        original_row_offset = int(anchor["start_row_offset"])
        row_shift, col_shift = _shift_for_cell(
            original_row_offset + 1,
            int(anchor["start_col"]),
            resolved_footprints,
            dataframe_shift=dataframe_shift,
        )
        shifted_anchor = dict(anchor)
        shifted_anchor["start_row_offset"] = original_row_offset + row_shift
        shifted_anchor["start_col"] = int(anchor["start_col"]) + col_shift
        resolved_anchors.append(shifted_anchor)
        _extend_repeat_dataframe_shift_footprints(resolved_footprints, shifted_anchor, original_row_offset)
    return resolved_anchors, resolved_footprints


def _extend_repeat_dataframe_shift_footprints(
    footprints: list[dict[str, Any]], anchor: dict[str, Any], original_row_offset: int
) -> None:
    output_range = _repeat_dataframe_output_range(anchor)
    if output_range is None:
        return
    key = (anchor["coordinate"], anchor["key"], int(anchor["start_col"]))
    # Use the original (pre-shift) row as the comparison baseline so that cells
    # and merges below this anchor in the template are correctly identified even
    # after earlier anchors have already pushed this anchor's shifted row far down.
    original_start_row = original_row_offset + 1
    row_delta = output_range.max_row - output_range.min_row  # source_rows - 1
    for item in footprints:
        if item["key"] != key:
            continue
        item["start_row"] = min(item["start_row"], original_start_row)
        item["max_row"] = max(item["max_row"], original_start_row + row_delta)
        item["max_col"] = max(item["max_col"], output_range.max_col)
        item["col_delta"] = max(
            item["col_delta"], output_range.max_col - output_range.min_col
        )
        item["row_delta"] = item["max_row"] - item["start_row"]
        return
    footprints.append(
        {
            "key": key,
            "start_row": original_start_row,
            "start_col": output_range.min_col,
            "max_row": original_start_row + row_delta,
            "max_col": output_range.max_col,
            "row_delta": row_delta,
            "col_delta": output_range.max_col - output_range.min_col,
        }
    )
def _compile_cell(
    *,
    cell: CellSchema,
    payload: dict[str, Any],
    output_name: str,
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    sheet_dataframe_options: dict[str, Any],
) -> dict[str, Any]:
    value = cell.get("value")
    if not isinstance(value, str):
        return {"cells": {cell["coordinate"]: cell}, "anchors": []}

    full_match = PLACEHOLDER_RE.fullmatch(value.strip())
    if (
        full_match
        and full_match.group(2) in _DATAFRAME_TYPES
        and full_match.group(1) in payload
    ):
        key = full_match.group(1)
        anchor_type = full_match.group(2)
        row_idx, col_idx = _coord_indexes(cell["coordinate"])
        return {
            "cells": {},
            "anchors": _compile_dataframe_anchors(
                key=key,
                anchor_type=anchor_type,
                cell=cell,
                row_idx=row_idx,
                col_idx=col_idx,
                value=payload[key],
                output_name=output_name,
                bundle_dir=bundle_dir,
                data_sources=data_sources,
                used_ids=used_ids,
                source_cache=source_cache,
                sheet_dataframe_options=sheet_dataframe_options,
            ),
        }

    new_value = _substitute_scalars(value, payload)
    new_cell: dict[str, Any] = dict(cell)
    new_cell["value"] = _schema_value(new_value)
    if full_match and full_match.group(2) in _SCALAR_TYPES:
        new_cell["cell_type"] = _infer_cell_type(new_value)
    return {"cells": {cell["coordinate"]: new_cell}, "anchors": []}


def _compile_dataframe_anchors(
    *,
    key: str,
    anchor_type: str,
    cell: CellSchema,
    row_idx: int,
    col_idx: int,
    value: Any,
    output_name: str,
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[int, dict[str, Any]],
    sheet_dataframe_options: dict[str, Any],
) -> list[dict[str, Any]]:
    columns = _to_headers(value)
    column_layouts = _column_layouts(columns, sheet_dataframe_options.get(key))
    if anchor_type == "dataframe-header":
        anchor_id = _unique_id(used_ids, f"{output_name}__{key}__{anchor_type}")
        return [
            _dataframe_anchor(
                anchor_id=anchor_id,
                key=key,
                placeholder_type=anchor_type,
                coord=cell["coordinate"],
                row_idx=row_idx,
                col_idx=col_idx,
                columns=columns,
                column_layouts=column_layouts,
                source_record=None,
                cell=cell,
            )
        ]

    source_id_type = "dataframe-content" if anchor_type == "dataframe" else anchor_type
    anchor_id = _unique_id(used_ids, f"{output_name}__{key}__{source_id_type}")
    source_record = _write_source_file(
        value=value,
        source_id=anchor_id,
        bundle_dir=bundle_dir,
        source_cache=source_cache,
    )
    if source_record not in data_sources:
        data_sources.append(source_record)

    anchors: list[dict[str, Any]] = []
    if anchor_type == "dataframe":
        header_id = _unique_id(used_ids, f"{output_name}__{key}__dataframe-header")
        anchors.append(
            _dataframe_anchor(
                anchor_id=header_id,
                key=key,
                placeholder_type="dataframe-header",
                coord=cell["coordinate"],
                row_idx=row_idx,
                col_idx=col_idx,
                columns=columns,
                column_layouts=column_layouts,
                source_record=None,
                cell=cell,
            )
        )
        row_idx += 1
    anchors.append(
        _dataframe_anchor(
            anchor_id=anchor_id,
            key=key,
            placeholder_type="dataframe-content",
            coord=cell["coordinate"],
            row_idx=row_idx,
            col_idx=col_idx,
            columns=columns,
            column_layouts=column_layouts,
            source_record=source_record,
            cell=cell,
        )
    )
    return anchors


def _dataframe_anchor(
    *,
    anchor_id: str,
    key: str,
    placeholder_type: str,
    coord: str,
    row_idx: int,
    col_idx: int,
    columns: list[str],
    column_layouts: list[dict[str, Any]],
    source_record: dict[str, Any] | None,
    cell: CellSchema,
) -> dict[str, Any]:
    return {
        "id": anchor_id,
        "key": key,
        "placeholder_type": placeholder_type,
        "coordinate": coord,
        "start_row": row_idx,
        "start_col": col_idx,
        "columns": columns,
        "column_layouts": column_layouts,
        "source": source_record["path"] if source_record else None,
        "source_format": source_record["format"] if source_record else None,
        "source_rows": source_record["rows"] if source_record else None,
        "cell": cell,
    }


def _column_layouts(
    columns: list[str], options: dict[str, Any] | None
) -> list[dict[str, Any]]:
    configured = (options or {}).get("columns", {}) if isinstance(options, dict) else {}
    layouts: list[dict[str, Any]] = []
    offset = 0
    for column in columns:
        column_options = configured.get(column, {}) if isinstance(configured, dict) else {}
        occupation = column_options.get("occupation", 1)
        if not isinstance(occupation, int) or occupation <= 0:
            raise ValueError(
                f"dataframe column '{column}' occupation must be a positive integer"
            )
        alignment = column_options.get("alignment")
        if alignment is not None and alignment not in _ALIGNMENTS:
            raise ValueError(
                f"dataframe column '{column}' alignment must be one of left, center, right"
            )
        layout: dict[str, Any] = {
            "name": column,
            "start_col_offset": offset,
            "occupation": occupation,
        }
        if alignment is not None:
            layout["alignment"] = alignment
        layouts.append(layout)
        offset += occupation
    return layouts


def _occupied_width(column_layouts: list[dict[str, Any]]) -> int:
    if not column_layouts:
        return 0
    last = column_layouts[-1]
    return int(last["start_col_offset"]) + int(last["occupation"])


def _shift_template_content_around_dataframes(
    sheet: dict[str, Any], *, dataframe_shift: str
) -> None:
    if dataframe_shift == "none":
        return
    footprints = _dataframe_shift_footprints(sheet.get("dataframe_anchors", []))
    if not footprints:
        return

    _validate_dataframe_anchors_are_not_template_merged(sheet, footprints)
    merge_shifts: dict[tuple[int, int], tuple[int, int]] = {}
    shifted_merges: list[str] = []
    for region in sheet.get("merged_regions", []):
        merge_range = CellRange(region)
        row_shift, col_shift = _shift_for_range(
            merge_range, footprints, dataframe_shift=dataframe_shift
        )
        shifted = _shift_range(merge_range, row_shift, col_shift)
        shifted_merges.append(str(shifted))
        for row_idx in range(merge_range.min_row, merge_range.max_row + 1):
            for col_idx in range(merge_range.min_col, merge_range.max_col + 1):
                merge_shifts[(row_idx, col_idx)] = (row_shift, col_shift)

    shifted_cells: dict[str, CellSchema] = {}
    for cell in sheet["cells"].values():
        row_idx, col_idx = _coord_indexes(cell["coordinate"])
        row_shift, col_shift = merge_shifts.get(
            (row_idx, col_idx),
            _shift_for_cell(
                row_idx, col_idx, footprints, dataframe_shift=dataframe_shift
            ),
        )
        new_cell = _shift_cell(cell, row_shift, col_shift)
        shifted_cells[new_cell["coordinate"]] = new_cell

    sheet["merged_regions"] = shifted_merges
    sheet["cells"] = shifted_cells
    _refresh_dimensions(sheet)


def _dataframe_shift_footprints(
    anchors: list[dict[str, Any]],
) -> list[dict[str, int]]:
    grouped: dict[tuple[str, str, int], dict[str, int]] = {}
    for anchor in anchors:
        output_range = _dataframe_output_range(anchor)
        if output_range is None:
            continue
        key = (anchor["coordinate"], anchor["key"], int(anchor["start_col"]))
        item = grouped.get(key)
        if item is None:
            grouped[key] = {
                "start_row": output_range.min_row,
                "start_col": output_range.min_col,
                "max_row": output_range.max_row,
                "max_col": output_range.max_col,
                "row_delta": 0,
                "col_delta": output_range.max_col - output_range.min_col,
            }
            continue
        item["start_row"] = min(item["start_row"], output_range.min_row)
        item["max_row"] = max(item["max_row"], output_range.max_row)
        item["max_col"] = max(item["max_col"], output_range.max_col)
        item["col_delta"] = max(
            item["col_delta"], output_range.max_col - output_range.min_col
        )

    footprints = list(grouped.values())
    for item in footprints:
        item["row_delta"] = item["max_row"] - item["start_row"]
    return footprints


def _validate_dataframe_anchors_are_not_template_merged(
    sheet: dict[str, Any], footprints: list[dict[str, int]]
) -> None:
    if not sheet.get("merged_regions"):
        return
    anchors_with_ranges = [
        (anchor, output_range)
        for anchor in sheet.get("dataframe_anchors", [])
        for output_range in [_dataframe_output_range(anchor)]
        if output_range is not None
    ]
    for region in sheet["merged_regions"]:
        merge_range = CellRange(region)
        for anchor, output_range in anchors_with_ranges:
            # Only reject merges that cover the anchor cell itself (the
            # placeholder position). Rows/columns outside the anchor cell are
            # template content that dataframe_shift will move out of the way;
            # the post-shift validator catches any remaining overlaps.
            anchor_cell = CellRange(
                min_col=output_range.min_col,
                min_row=output_range.min_row,
                max_col=output_range.min_col,
                max_row=output_range.min_row,
            )
            if _ranges_overlap(merge_range, anchor_cell):
                raise ValueError(_merged_dataframe_overlap_error(anchor, merge_range, output_range))


def _shift_for_range(
    region: CellRange, footprints: list[dict[str, int]], *, dataframe_shift: str
) -> tuple[int, int]:
    row_shift = 0
    col_shift = 0
    for footprint in footprints:
        apply_h = dataframe_shift in {"both", "horizontal"} and _range_is_right_of_footprint(
            region, footprint
        )
        apply_v = dataframe_shift in {"both", "vertical"} and _range_is_below_footprint(
            region, footprint
        )
        # Diagonal gap: a range strictly past both extents of the footprint
        # (min_col > max_col AND min_row > max_row) has no row overlap for the
        # horizontal check and no col overlap for the vertical check, so both
        # conditions above fail.  Apply whichever shifts are enabled so that
        # template content diagonally past the footprint stays in the correct
        # relative position after the dataframe expands.  This also prevents
        # overlaps caused by a straddling merge (min_col ≤ max_col but
        # max_col > footprint["max_col"]) shifting into tall merges that sit
        # entirely to the right of max_col but still below max_row.
        if region.min_col > footprint["max_col"] and region.min_row > footprint["max_row"]:
            if dataframe_shift in {"both", "horizontal"}:
                apply_h = True
            if dataframe_shift in {"both", "vertical"}:
                apply_v = True
        if apply_h:
            col_shift += footprint["col_delta"]
        if apply_v:
            row_shift += footprint["row_delta"]
    return row_shift, col_shift


def _shift_for_cell(
    row_idx: int,
    col_idx: int,
    footprints: list[dict[str, int]],
    *,
    dataframe_shift: str,
) -> tuple[int, int]:
    row_shift = 0
    col_shift = 0
    for footprint in footprints:
        apply_h = (
            dataframe_shift in {"both", "horizontal"}
            and footprint["start_row"] <= row_idx <= footprint["max_row"]
            and col_idx > footprint["start_col"]
        )
        apply_v = (
            dataframe_shift in {"both", "vertical"}
            and footprint["start_col"] <= col_idx <= footprint["max_col"]
            and row_idx > footprint["start_row"]
        )
        # Diagonal gap: same as _shift_for_range — apply enabled shifts for
        # cells strictly past both extents of the footprint.
        if col_idx > footprint["max_col"] and row_idx > footprint["max_row"]:
            if dataframe_shift in {"both", "horizontal"}:
                apply_h = True
            if dataframe_shift in {"both", "vertical"}:
                apply_v = True
        if apply_h:
            col_shift += footprint["col_delta"]
        if apply_v:
            row_shift += footprint["row_delta"]
    return row_shift, col_shift


def _range_is_right_of_footprint(
    region: CellRange, footprint: dict[str, int]
) -> bool:
    return (
        region.min_row <= footprint["max_row"]
        and region.max_row >= footprint["start_row"]
        and region.min_col > footprint["start_col"]
    )


def _range_is_below_footprint(
    region: CellRange, footprint: dict[str, int]
) -> bool:
    return (
        region.min_col <= footprint["max_col"]
        and region.max_col >= footprint["start_col"]
        and region.min_row > footprint["start_row"]
    )


def _shift_range(region: CellRange, row_shift: int, col_shift: int) -> CellRange:
    return CellRange(
        min_col=region.min_col + col_shift,
        min_row=region.min_row + row_shift,
        max_col=region.max_col + col_shift,
        max_row=region.max_row + row_shift,
    )


def _shift_cell(cell: CellSchema, row_shift: int, col_shift: int) -> CellSchema:
    if row_shift == 0 and col_shift == 0:
        return cell
    row_idx, col_idx = _coord_indexes(cell["coordinate"])
    new_cell: dict[str, Any] = dict(cell)
    new_cell["coordinate"] = f"{get_column_letter(col_idx + col_shift)}{row_idx + row_shift}"
    merge_anchor = cell.get("merge_anchor")
    if merge_anchor is not None:
        anchor_row, anchor_col = _coord_indexes(merge_anchor)
        new_cell["merge_anchor"] = (
            f"{get_column_letter(anchor_col + col_shift)}{anchor_row + row_shift}"
        )
    return new_cell  # type: ignore[return-value]


def _refresh_dimensions(sheet: dict[str, Any]) -> None:
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    for coord in sheet["cells"]:
        row_idx, col_idx = _coord_indexes(coord)
        max_col = max(max_col, col_idx)
        max_row = max(max_row, row_idx)
    for region in sheet.get("merged_regions", []):
        merge_range = CellRange(region)
        max_col = max(max_col, merge_range.max_col)
        max_row = max(max_row, merge_range.max_row)
    for anchor in sheet.get("dataframe_anchors", []):
        output_range = _dataframe_output_range(anchor)
        if output_range is None:
            continue
        max_col = max(max_col, output_range.max_col)
        max_row = max(max_row, output_range.max_row)
    sheet["dimensions"] = (
        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    )


def _validate_template_merges_do_not_overlap_dataframes(sheet: dict[str, Any]) -> None:
    if not sheet.get("merged_regions"):
        return
    anchors_with_ranges = [
        (anchor, output_range)
        for anchor in sheet.get("dataframe_anchors", [])
        for output_range in [_dataframe_output_range(anchor)]
        if output_range is not None
    ]
    if not anchors_with_ranges:
        return
    for region in sheet["merged_regions"]:
        merge_range = CellRange(region)
        for anchor, output_range in anchors_with_ranges:
            if _ranges_overlap(merge_range, output_range):
                raise ValueError(_merged_dataframe_overlap_error(anchor, merge_range, output_range))


def _validate_repeat_merges_do_not_overlap_dataframes(section: dict[str, Any], dataframe_shift: str = "both") -> None:
    if section.get("record_source"):
        record = {
            "dataframe_anchors": section.get("dataframe_anchors", []),
            "merged_regions": section.get(
                "merged_record_regions",
                section.get("merged_regions", []),
            ),
        }
        records = [record]
    else:
        records = section["records"]
    for record in records:
        merges = record.get("merged_regions", section.get("merged_regions", []))
        if not merges:
            continue
        merge_ranges = [
            CellRange(
                min_col=merge["min_col"],
                min_row=merge["min_row_offset"] + 1,
                max_col=merge["max_col"],
                max_row=merge["max_row_offset"] + 1,
            )
            for merge in merges
        ]
        for anchor in record.get("dataframe_anchors", []):
            output_range = _repeat_dataframe_output_range(anchor)
            if output_range is None:
                continue
            for merge_range in merge_ranges:
                if dataframe_shift == "none":
                    # Without dataframe shift, merges cannot overlap the output range at all.
                    if _ranges_overlap(merge_range, output_range):
                        raise ValueError(_merged_dataframe_overlap_error(anchor, merge_range, output_range))
                else:
                    # With dataframe shift, only reject merges that overlap the anchor cell itself.
                    # Merges below the anchor have been shifted by _shift_repeat_record_content_around_dataframes
                    # and should not overlap the output range if the template is valid.
                    anchor_cell = CellRange(
                        min_col=output_range.min_col,
                        min_row=output_range.min_row,
                        max_col=output_range.min_col,
                        max_row=output_range.min_row,
                    )
                    if _ranges_overlap(merge_range, anchor_cell):
                        raise ValueError(_merged_dataframe_overlap_error(anchor, merge_range, output_range))


def _merged_dataframe_overlap_error(
    anchor: dict[str, Any], merge_range: CellRange, output_range: CellRange
) -> str:
    return (
        "Template merged regions must not overlap dataframe output ranges: "
        f"dataframe '{anchor['key']}' ({anchor['placeholder_type']}) overlaps "
        f"merged region {merge_range} with output range {output_range}"
    )


def _dataframe_output_range(anchor: dict[str, Any]) -> CellRange | None:
    width = _occupied_width(anchor.get("column_layouts", []))
    if width <= 0:
        return None
    row_count = 1
    if anchor["placeholder_type"] == "dataframe-content":
        source_rows = anchor.get("source_rows", None)
        if source_rows is None and anchor.get("source"):
            source_rows = 1
        row_count = int(source_rows or 0)
    if row_count <= 0:
        return None
    return CellRange(
        min_col=anchor["start_col"],
        min_row=anchor["start_row"],
        max_col=anchor["start_col"] + width - 1,
        max_row=anchor["start_row"] + row_count - 1,
    )


def _repeat_dataframe_output_range(anchor: dict[str, Any]) -> CellRange | None:
    width = _occupied_width(anchor.get("column_layouts", []))
    if width <= 0:
        return None
    row_count = 1
    if anchor["placeholder_type"] == "dataframe-content":
        row_count = int(anchor.get("source_rows") or 0)
    if row_count <= 0:
        return None
    start_row = anchor["start_row_offset"] + 1
    return CellRange(
        min_col=anchor["start_col"],
        min_row=start_row,
        max_col=anchor["start_col"] + width - 1,
        max_row=start_row + row_count - 1,
    )


def _ranges_overlap(left: CellRange, right: CellRange) -> bool:
    return (
        left.min_col <= right.max_col
        and left.max_col >= right.min_col
        and left.min_row <= right.max_row
        and left.max_row >= right.min_row
    )


def _validate_dataframe_shift_mode(dataframe_shift: str) -> None:
    if dataframe_shift not in _DATAFRAME_SHIFT_MODES:
        raise ValueError(
            "dataframe_shift must be one of 'both', 'horizontal', 'vertical', or 'none'"
        )


# §4. Public Functions


def compile_report_bundle(
    template: WorkbookSchema,
    data: dict[str, Any],
    bundle_path: str | None = None,
    dataframe_options: dict[str, Any] | None = None,
    dataframe_shift: str = "both",
) -> ReportBundle:
    """Validate *data* against *template* and produce a directory ReportBundle."""
    _validate_dataframe_shift_mode(dataframe_shift)
    bundle_dir = _prepare_bundle_dir(bundle_path)
    data_sources: list[dict[str, Any]] = []
    used_ids: set[str] = set()
    source_cache: dict[int, dict[str, Any]] = {}

    sheets = [
        _compile_sheet(
            sheet=sheet,
            output_name=output_name,
            sheet_data=sheet_data,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
            dataframe_options=dataframe_options or {},
            dataframe_shift=dataframe_shift,
        )
        for sheet, output_name, sheet_data, _ in _resolve_sheet_payloads(template, data)
    ]

    report = {"version": BUNDLE_VERSION, "sheets": sheets, "assets": []}
    if template.get("theme_colors"):
        report["theme_colors"] = template["theme_colors"]
    manifest = {
        "version": BUNDLE_VERSION,
        "bundle_format": "directory",
        "created_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "inputs": get_template_inputs(template),
        "sheets": [
            {
                "name": sheet["name"],
                "dataframe_anchors": [
                    anchor["id"] for anchor in sheet.get("dataframe_anchors", [])
                ],
            }
            for sheet in sheets
        ],
        "placeholder_resolution": {
            "scalars": "resolved in report.json",
            "dataframes": "stored as dataframe anchors pointing to data/*.parquet",
        },
        "dataframe_sources": data_sources,
        "assets": [],
        "output_capabilities": {"xlsx": True, "pdf": True, "image": False},
    }
    if template.get("theme_colors"):
        manifest["theme_colors"] = template["theme_colors"]
    bundle = ReportBundle(manifest=manifest, report=report, path=str(bundle_dir))
    _write_bundle_metadata(bundle)
    return bundle


def load_report_bundle(bundle_path: str) -> ReportBundle:
    path = Path(bundle_path)
    if not path.exists():
        raise FileNotFoundError(f"Report bundle not found: {bundle_path}")
    if not path.is_dir():
        raise ValueError(f"Report bundle must be a directory: {bundle_path}")

    manifest_path = path / "manifest.json"
    report_path = path / "report.json"
    if not manifest_path.exists() or not report_path.exists():
        raise ValueError(
            "Malformed report bundle: expected manifest.json and report.json"
        )

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    report = json.loads(report_path.read_text(encoding="utf-8"))
    if manifest.get("version") not in {BUNDLE_VERSION, "1.0"}:
        raise ValueError(
            f"Unsupported report bundle version: {manifest.get('version')!r}"
        )
    if manifest.get("bundle_format") != "directory":
        raise ValueError(
            f"Unsupported report bundle format: {manifest.get('bundle_format')!r}"
        )
    return ReportBundle(manifest=manifest, report=report, path=str(path))


# §5. Entrypoints
