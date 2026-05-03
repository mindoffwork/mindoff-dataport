from __future__ import annotations

import hashlib
import html
import re
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, legal, letter, landscape, portrait
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from ..compile.bundle import ReportBundle
from ..compile.template_contract import _infer_cell_type, _parse_dims
from ..schema import CellSchema, SheetSchema
from .xlsx_renderer import (
    _anchor_layouts,
    _anchor_row_merges,
    _data_source_map,
    _expanded_cells,
    _expanded_sheet_from_manifest,
    _header_cells,
    _layout_alignment,
    _occupied_width,
    _repeat_row_stream,
    _repeat_record_rows,
    _source_rows,
    _static_row_merges,
    _sheet_with_overrides,
)

__all__ = ["export_report_bundle"]

# §1. Constants & Exceptions

_PAGE_SIZES = {
    "A4": A4,
    "LETTER": letter,
    "LEGAL": legal,
}

_BORDER_WIDTHS = {
    "hair": 0.25,
    "thin": 0.5,
    "medium": 1.0,
    "thick": 1.5,
    "dashed": 0.75,
    "dotted": 0.5,
    "double": 1.25,
}

_DEFAULT_COLUMN_WIDTH = 15.0
_DEFAULT_ROW_HEIGHT = 15.0

# Default Office theme base colors (openpyxl index order: lt1, dk1, lt2, dk2, accent1-6).
# Used as a fallback when the workbook theme is not available in the PDF render path.
_DEFAULT_THEME_COLORS: list[tuple[int, int, int]] = [
    (255, 255, 255),  # 0  lt1      White
    (0, 0, 0),        # 1  dk1      Black
    (238, 236, 225),  # 2  lt2      #EEECE1
    (31, 73, 125),    # 3  dk2      #1F497D
    (79, 129, 189),   # 4  accent1  #4F81BD
    (192, 80, 77),    # 5  accent2  #C0504D
    (155, 187, 89),   # 6  accent3  #9BBB59
    (128, 100, 162),  # 7  accent4  #8064A2
    (75, 172, 198),   # 8  accent5  #4BACC6
    (247, 150, 70),   # 9  accent6  #F79646
]
_EXCEL_WIDTH_TO_POINTS = 7.0
_MIN_COLUMN_WIDTH = 24.0
_MAX_COLUMN_WIDTH = 180.0

# §2. Classes and Sub Classes


class _FontResolver:
    def __init__(self, fonts: dict[str, Any] | None = None) -> None:
        self._families = _register_fonts(fonts or {})

    def font_name(self, font: dict[str, Any]) -> str:
        family = font.get("name")
        custom = self._families.get(str(family)) if family else None
        if custom is not None:
            return _variant_font_name(
                custom,
                bold=bool(font.get("bold")),
                italic=bool(font.get("italic")),
            )
        return _builtin_font_name(font)


class _LazyFlowables:
    def __init__(self, flowables: Iterable[Any]) -> None:
        self._iterator = iter(flowables)
        self._buffer: list[Any] = []
        self._done = False

    def _fill(self) -> None:
        if self._buffer or self._done:
            return
        self._append_next()

    def _fill_to(self, index: int) -> None:
        while len(self._buffer) <= index and not self._done:
            self._append_next()

    def _append_next(self) -> None:
        try:
            self._buffer.append(next(self._iterator))
        except StopIteration:
            self._done = True

    def __len__(self) -> int:
        self._fill()
        return 1 if self._buffer else 0

    def __getitem__(self, index):
        if isinstance(index, slice):
            if index.stop is not None and index.stop > 0:
                self._fill_to(index.stop - 1)
            else:
                self._fill()
            return self._buffer[index]
        self._fill_to(index)
        return self._buffer[index]

    def __delitem__(self, index) -> None:
        self._fill()
        if isinstance(index, slice):
            del self._buffer[index]
        else:
            del self._buffer[index]

    def __setitem__(self, index, value) -> None:
        self._fill()
        self._buffer[index] = value

    def insert(self, index: int, value) -> None:
        self._fill()
        self._buffer.insert(index, value)


# §3. Private Helper Functions


def _page_size(name: str, orientation: str) -> tuple[float, float]:
    size = _PAGE_SIZES.get(name.upper())
    if size is None:
        expected = ", ".join(sorted(_PAGE_SIZES))
        raise ValueError(
            f"Unsupported PDF page_size '{name}'. Expected one of: {expected}."
        )
    if orientation == "portrait":
        return portrait(size)
    if orientation == "landscape":
        return landscape(size)
    raise ValueError(
        f"Unsupported PDF orientation '{orientation}'. Expected 'portrait' or 'landscape'."
    )


def _apply_tint(rgb: tuple[int, int, int], tint: float) -> tuple[int, int, int]:
    r, g, b = rgb
    if tint > 0:
        return (int(r + (255 - r) * tint), int(g + (255 - g) * tint), int(b + (255 - b) * tint))
    if tint < 0:
        return (int(r * (1 + tint)), int(g * (1 + tint)), int(b * (1 + tint)))
    return r, g, b


def _hex_color(value: str | None) -> colors.Color | None:
    if not value:
        return None
    if value.startswith("theme:"):
        parts = value.split(":")
        if len(parts) != 3:
            return None
        try:
            idx, tint = int(parts[1]), float(parts[2])
        except ValueError:
            return None
        if idx >= len(_DEFAULT_THEME_COLORS):
            return None
        r, g, b = _apply_tint(_DEFAULT_THEME_COLORS[idx], tint)
        return colors.Color(r / 255, g / 255, b / 255)
    raw = value[-6:]
    if len(raw) != 6:
        return None
    try:
        return colors.HexColor(f"#{raw}")
    except ValueError:
        return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _builtin_font_name(font: dict[str, Any]) -> str:
    bold = bool(font.get("bold"))
    italic = bool(font.get("italic"))
    if bold and italic:
        return "Helvetica-BoldOblique"
    if bold:
        return "Helvetica-Bold"
    if italic:
        return "Helvetica-Oblique"
    return "Helvetica"


def _register_fonts(fonts: dict[str, Any]) -> dict[str, dict[str, str]]:
    families: dict[str, dict[str, str]] = {}
    for family, config in fonts.items():
        variants = _font_variants(config)
        registered: dict[str, str] = {}
        for variant, path in variants.items():
            if path is None:
                continue
            font_path = Path(path)
            if not font_path.exists():
                raise ValueError(
                    f"PDF font '{family}' variant '{variant}' file does not exist: {font_path}"
                )
            font_name = _registered_font_name(family, variant, font_path)
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
            registered[variant] = font_name

        if "regular" not in registered:
            raise ValueError(f"PDF font '{family}' requires a regular font file")
        families[str(family)] = registered
        _register_font_family(str(family), registered)
    return families


def _font_variants(config: Any) -> dict[str, str | None]:
    if isinstance(config, (str, Path)):
        return {
            "regular": str(config),
            "bold": None,
            "italic": None,
            "bold_italic": None,
        }
    if not isinstance(config, dict):
        raise TypeError(
            f"PDF font config must be a path or dict of variants, got {type(config).__name__}"
        )
    return {
        "regular": config.get("regular"),
        "bold": config.get("bold"),
        "italic": config.get("italic"),
        "bold_italic": config.get("bold_italic") or config.get("boldItalic"),
    }


def _registered_font_name(family: str, variant: str, path: Path) -> str:
    safe_family = re.sub(r"[^A-Za-z0-9]+", "-", family).strip("-") or "Font"
    safe_variant = variant.replace("_", "-")
    digest = hashlib.sha1(str(path.resolve()).encode("utf-8")).hexdigest()[:12]
    return f"Mindoff-{safe_family}-{safe_variant}-{digest}"


def _register_font_family(family: str, registered: dict[str, str]) -> None:
    regular = registered["regular"]
    pdfmetrics.registerFontFamily(
        family,
        normal=regular,
        bold=registered.get("bold", regular),
        italic=registered.get("italic", regular),
        boldItalic=registered.get("bold_italic", registered.get("bold", regular)),
    )


def _variant_font_name(
    registered: dict[str, str], *, bold: bool, italic: bool
) -> str:
    if bold and italic:
        return (
            registered.get("bold_italic")
            or registered.get("bold")
            or registered["regular"]
        )
    if bold:
        return registered.get("bold") or registered["regular"]
    if italic:
        return registered.get("italic") or registered["regular"]
    return registered["regular"]


def _alignment(value: str | None) -> str:
    if value in {"center", "centerContinuous"}:
        return "CENTER"
    if value == "right":
        return "RIGHT"
    return "LEFT"


def _paragraph_alignment(value: str | None) -> int:
    if value in {"center", "centerContinuous"}:
        return TA_CENTER
    if value == "right":
        return TA_RIGHT
    if value in {"justify", "distributed"}:
        return TA_JUSTIFY
    return TA_LEFT


def _pdf_fill_color(fill: dict) -> "colors.Color | None":
    """Return the display background color for a fill, or None for no fill."""
    pt = fill.get("pattern_type")
    # Backward compat: old schema used only bg_color for solid fills.
    legacy_bg = fill.get("bg_color") if not fill.get("fg_color") else None
    if not pt and not legacy_bg:
        return None
    if not pt or pt == "solid":
        return _hex_color(fill.get("fg_color") or legacy_bg)
    # Patterned fills: prefer bg_color unless it is white/absent, then fall back to fg_color.
    bg_c = _hex_color(fill.get("bg_color"))
    fg_c = _hex_color(fill.get("fg_color"))
    white = colors.HexColor("#FFFFFF")
    return bg_c if (bg_c is not None and bg_c != white) else (fg_c or bg_c)


def _indent_points(alignment: dict) -> float:
    """Convert indent + relative_indent to PDF padding points (1 unit ≈ 7 pt)."""
    indent = int(alignment.get("indent") or 0)
    rel = int(alignment.get("relative_indent") or 0)
    total = max(0, indent + rel)
    return total * 7.0


def _effective_left_border(borders: dict, is_rtl: bool) -> dict:
    """Resolve left edge border, merging directional start/end for RTL awareness."""
    primary = "right" if is_rtl else "left"
    secondary = "end" if is_rtl else "start"
    side = borders.get(primary, {})
    return side if side.get("style") else borders.get(secondary, {"style": None, "color": None})


def _effective_right_border(borders: dict, is_rtl: bool) -> dict:
    """Resolve right edge border, merging directional start/end for RTL awareness."""
    primary = "left" if is_rtl else "right"
    secondary = "start" if is_rtl else "end"
    side = borders.get(primary, {})
    return side if side.get("style") else borders.get(secondary, {"style": None, "color": None})


def _vertical_alignment(value: str | None) -> str:
    if value == "center":
        return "MIDDLE"
    if value == "bottom":
        return "BOTTOM"
    return "TOP"


def _paragraph(cell: CellSchema, font_resolver: _FontResolver) -> Paragraph | str:
    text = html.escape(_cell_text(cell.get("value")), quote=False)
    if not text:
        return ""

    font = cell["font"]
    text = _apply_inline_markup(text, font)

    font_size = float(font.get("size") or 11.0)
    color = _hex_color(font.get("color"))
    style = ParagraphStyle(
        name=f"cell-{cell['coordinate']}",
        fontName=font_resolver.font_name(font),
        fontSize=font_size,
        leading=max(font_size * 1.2, font_size + 2),
        textColor=color or colors.black,
        alignment=_paragraph_alignment(cell["alignment"].get("horizontal")),
    )
    wrap = cell["alignment"].get("wrap_text", False)
    text_body = text.replace("\n", "<br/>") if wrap else text.replace("\n", " ")
    return Paragraph(text_body, style)


def _apply_inline_markup(text: str, font: dict) -> str:
    """Wrap text in ReportLab XML markup tags for underline, strike, vert_align."""
    if font.get("underline"):
        text = f"<u>{text}</u>"
    if font.get("strike"):
        text = f"<strike>{text}</strike>"
    vert = font.get("vert_align")
    if vert == "superscript":
        text = f"<super>{text}</super>"
    elif vert == "subscript":
        text = f"<sub>{text}</sub>"
    return text


def _coord_indexes(coord: str) -> tuple[int, int]:
    col_letter, row_idx = coordinate_from_string(coord)
    return row_idx, column_index_from_string(col_letter)


def _is_non_anchor_merged_cell(cell: CellSchema) -> bool:
    return bool(cell.get("merged")) and cell.get("merge_anchor") not in (
        None,
        cell["coordinate"],
    )


def _hug_column_width(
    col_idx: int,
    min_row: int,
    max_row: int,
    cells: dict[tuple[int, int], CellSchema],
) -> float:
    max_chars = 1
    for row_idx in range(min_row, max_row + 1):
        cell = cells.get((row_idx, col_idx))
        if cell is None:
            continue
        text = _cell_text(cell.get("value"))
        longest_line = max((len(part) for part in text.splitlines()), default=0)
        font_size = float(cell["font"].get("size") or 11.0)
        factor = 1.15 if cell["font"].get("bold") else 1.0
        max_chars = max(max_chars, int(longest_line * factor * font_size / 11.0))
    return min(max(max_chars * 6.0 + 10.0, _MIN_COLUMN_WIDTH), _MAX_COLUMN_WIDTH)


def _column_widths(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    available_width: float,
) -> list[float]:
    mode = sheet.get("column_width_mode", "fixed")
    widths: list[float] = []
    if mode == "even":
        width = float(sheet.get("default_column_width", _DEFAULT_COLUMN_WIDTH))
        widths = [width * _EXCEL_WIDTH_TO_POINTS] * (max_col - min_col + 1)
    elif mode == "hug":
        widths = [
            _hug_column_width(col_idx, min_row, max_row, cells)
            for col_idx in range(min_col, max_col + 1)
        ]
    else:
        for col_idx in range(min_col, max_col + 1):
            col_letter = _col_letter(col_idx)
            width = sheet["column_widths"].get(col_letter, _DEFAULT_COLUMN_WIDTH)
            widths.append(float(width or _DEFAULT_COLUMN_WIDTH) * _EXCEL_WIDTH_TO_POINTS)

    total = sum(widths)
    if total > available_width and total > 0:
        scale = available_width / total
        widths = [max(width * scale, _MIN_COLUMN_WIDTH) for width in widths]
    return widths


def _row_heights(sheet: SheetSchema, min_row: int, max_row: int) -> list[float | None]:
    mode = sheet.get("row_height_mode", "fixed")
    if mode == "even":
        height = float(sheet.get("default_row_height", _DEFAULT_ROW_HEIGHT))
        return [height] * (max_row - min_row + 1)
    if mode == "hug":
        return [None] * (max_row - min_row + 1)
    return [
        float(sheet["row_heights"].get(str(row_idx)) or _DEFAULT_ROW_HEIGHT)
        for row_idx in range(min_row, max_row + 1)
    ]


def _col_letter(col_idx: int) -> str:
    result = ""
    while col_idx:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _table_data(
    cells: dict[tuple[int, int], CellSchema],
    font_resolver: _FontResolver,
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
) -> list[list[Any]]:
    data: list[list[Any]] = []
    for row_idx in range(min_row, max_row + 1):
        row: list[Any] = []
        for col_idx in range(min_col, max_col + 1):
            cell = cells.get((row_idx, col_idx))
            row.append(_paragraph(cell, font_resolver) if cell is not None else "")
        data.append(row)
    return data


def _style_cell(
    commands: list[tuple[Any, ...]],
    cell: CellSchema,
    row_idx: int,
    col_idx: int,
    min_row: int,
    min_col: int,
    font_resolver: _FontResolver,
) -> None:
    point = (col_idx - min_col, row_idx - min_row)
    font = cell["font"]
    alignment = cell["alignment"]

    commands.append(("FONTNAME", point, point, font_resolver.font_name(font)))
    commands.append(("FONTSIZE", point, point, float(font.get("size") or 11.0)))
    font_color = _hex_color(font.get("color"))
    if font_color is not None:
        commands.append(("TEXTCOLOR", point, point, font_color))
    fill_color = _pdf_fill_color(cell["fill"])
    if fill_color is not None:
        commands.append(("BACKGROUND", point, point, fill_color))
    commands.append(("ALIGN", point, point, _alignment(alignment.get("horizontal"))))
    commands.append(
        ("VALIGN", point, point, _vertical_alignment(alignment.get("vertical")))
    )
    _apply_cell_padding(commands, point, alignment)
    _apply_cell_borders(commands, point, cell["borders"], alignment)


def _apply_cell_padding(
    commands: list[tuple[Any, ...]], point: tuple[int, int], alignment: dict
) -> None:
    indent_pts = _indent_points(alignment)
    is_rtl = int(alignment.get("reading_order") or 0) == 2
    commands.append(("LEFTPADDING", point, point, max(indent_pts, 3) if not is_rtl else 3))
    commands.append(("RIGHTPADDING", point, point, max(indent_pts, 3) if is_rtl else 3))
    commands.append(("TOPPADDING", point, point, 2))
    commands.append(("BOTTOMPADDING", point, point, 2))


def _apply_cell_borders(
    commands: list[tuple[Any, ...]],
    point: tuple[int, int],
    borders: dict,
    alignment: dict,
) -> None:
    is_rtl = int(alignment.get("reading_order") or 0) == 2
    _border(commands, point, "LINEABOVE", borders["top"])
    _border(commands, point, "LINEBELOW", borders["bottom"])
    _border(commands, point, "LINEBEFORE", _effective_left_border(borders, is_rtl))
    _border(commands, point, "LINEAFTER", _effective_right_border(borders, is_rtl))


def _border(
    commands: list[tuple[Any, ...]],
    point: tuple[int, int],
    command: str,
    side: dict[str, Any],
) -> None:
    style = side.get("style")
    if not style:
        return
    width = _BORDER_WIDTHS.get(style, 0.5)
    color = _hex_color(side.get("color")) or colors.black
    commands.append((command, point, point, width, color))


def _border_range(
    commands: list[tuple[Any, ...]],
    start: tuple[int, int],
    end: tuple[int, int],
    command: str,
    side: dict[str, Any],
) -> None:
    style = side.get("style")
    if not style:
        return
    width = _BORDER_WIDTHS.get(style, 0.5)
    color = _hex_color(side.get("color")) or colors.black
    commands.append((command, start, end, width, color))


def _span_commands(
    sheet: SheetSchema,
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
) -> list[tuple[Any, ...]]:
    commands: list[tuple[Any, ...]] = []
    for raw_region in sheet.get("merged_regions", []):
        region = CellRange(raw_region)
        if (
            region.min_col < min_col
            or region.min_row < min_row
            or region.max_col > max_col
            or region.max_row > max_row
        ):
            continue
        start = (region.min_col - min_col, region.min_row - min_row)
        end = (region.max_col - min_col, region.max_row - min_row)
        if start != end:
            commands.append(("SPAN", start, end))
    return commands


def _merged_border_commands(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
) -> list[tuple[Any, ...]]:
    commands: list[tuple[Any, ...]] = []
    for raw_region in sheet.get("merged_regions", []):
        region = CellRange(raw_region)
        if (
            region.min_col < min_col
            or region.min_row < min_row
            or region.max_col > max_col
            or region.max_row > max_row
        ):
            continue
        anchor = cells.get((region.min_row, region.min_col))
        if anchor is None:
            continue
        start = (region.min_col - min_col, region.min_row - min_row)
        end = (region.max_col - min_col, region.max_row - min_row)
        borders = anchor["borders"]
        is_rtl = int(anchor["alignment"].get("reading_order") or 0) == 2
        _border_range(commands, start, end, "LINEABOVE", borders["top"])
        _border_range(commands, start, end, "LINEBELOW", borders["bottom"])
        _border_range(commands, start, end, "LINEBEFORE", _effective_left_border(borders, is_rtl))
        _border_range(commands, start, end, "LINEAFTER", _effective_right_border(borders, is_rtl))
    return commands


def _merged_style_commands(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    font_resolver: _FontResolver,
) -> list[tuple[Any, ...]]:
    commands: list[tuple[Any, ...]] = []
    for raw_region in sheet.get("merged_regions", []):
        region = CellRange(raw_region)
        if (
            region.min_col < min_col
            or region.min_row < min_row
            or region.max_col > max_col
            or region.max_row > max_row
        ):
            continue
        anchor = cells.get((region.min_row, region.min_col))
        if anchor is None:
            continue
        start = (region.min_col - min_col, region.min_row - min_row)
        end = (region.max_col - min_col, region.max_row - min_row)
        _merged_region_style(commands, anchor, start, end, font_resolver)
    return commands


def _merged_region_style(
    commands: list[tuple[Any, ...]],
    anchor: CellSchema,
    start: tuple[int, int],
    end: tuple[int, int],
    font_resolver: _FontResolver,
) -> None:
    font = anchor["font"]
    alignment = anchor["alignment"]
    commands.append(("FONTNAME", start, end, font_resolver.font_name(font)))
    commands.append(("FONTSIZE", start, end, float(font.get("size") or 11.0)))
    font_color = _hex_color(font.get("color"))
    if font_color is not None:
        commands.append(("TEXTCOLOR", start, end, font_color))
    fill_color = _pdf_fill_color(anchor["fill"])
    if fill_color is not None:
        commands.append(("BACKGROUND", start, end, fill_color))
    commands.append(("ALIGN", start, end, _alignment(alignment.get("horizontal"))))
    commands.append(("VALIGN", start, end, _vertical_alignment(alignment.get("vertical"))))
    indent_pts = _indent_points(alignment)
    is_rtl = int(alignment.get("reading_order") or 0) == 2
    commands.append(("LEFTPADDING", start, end, indent_pts if not is_rtl else 0))
    commands.append(("RIGHTPADDING", start, end, indent_pts if is_rtl else 0))
    commands.append(("TOPPADDING", start, end, 0))
    commands.append(("BOTTOMPADDING", start, end, 0))


def _table_style(
    sheet: SheetSchema,
    cells: dict[tuple[int, int], CellSchema],
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    font_resolver: _FontResolver,
) -> TableStyle:
    commands: list[tuple[Any, ...]] = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    commands.extend(_span_commands(sheet, min_col, min_row, max_col, max_row))
    for (row_idx, col_idx), cell in cells.items():
        if min_row <= row_idx <= max_row and min_col <= col_idx <= max_col:
            _style_cell(
                commands,
                cell,
                row_idx,
                col_idx,
                min_row,
                min_col,
                font_resolver,
            )
    commands.extend(
        _merged_style_commands(
            sheet, cells, min_col, min_row, max_col, max_row, font_resolver
        )
    )
    commands.extend(
        _merged_border_commands(sheet, cells, min_col, min_row, max_col, max_row)
    )
    return TableStyle(commands)


def _sheet_table(
    bundle: ReportBundle,
    raw_sheet: dict[str, Any],
    *,
    available_width: float,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    font_resolver: _FontResolver,
) -> Table:
    sheet = _sheet_with_overrides(
        raw_sheet,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
    )
    sheet, cells = _expanded_cells(
        bundle,
        sheet,
        streaming_chunk_rows=streaming_chunk_rows,
    )
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    data = _table_data(cells, font_resolver, min_col, min_row, max_col, max_row)
    table = Table(
        data,
        colWidths=_column_widths(
            sheet,
            cells,
            min_col,
            min_row,
            max_col,
            max_row,
            available_width,
        ),
        rowHeights=_row_heights(sheet, min_row, max_row),
        repeatRows=0,
        splitByRow=1,
    )
    table.setStyle(
        _table_style(sheet, cells, min_col, min_row, max_col, max_row, font_resolver)
    )
    return table


def _has_dataframe_content(sheet: SheetSchema) -> bool:
    return any(
        anchor["placeholder_type"] == "dataframe-content"
        for anchor in sheet.get("dataframe_anchors", [])
    )


def _validate_pdf_dataframe_sizing(
    sheet: SheetSchema,
    *,
    column_width_mode: str | None,
) -> None:
    if not _has_dataframe_content(sheet):
        return
    col_mode = column_width_mode or sheet.get("column_width_mode", "fixed")
    if col_mode == "hug":
        raise ValueError(
            "PDF export does not support column_width_mode='hug' for dataframe-content sheets. "
            "Use 'fixed' or 'even' — column hug requires buffering all rows before sizing."
        )


def _dataframe_pdf_rows(
    bundle: ReportBundle,
    source_map: dict[str, dict[str, Any]],
    sheet: SheetSchema,
    *,
    batch_size: int,
) -> Iterator[dict[str, Any]]:
    sheet = _expanded_sheet_from_manifest(sheet, source_map)
    _, min_row, _, max_row = _parse_dims(sheet["dimensions"])

    static_by_row: dict[int, dict[int, CellSchema]] = {}
    merge_by_row: dict[int, list[dict[str, Any]]] = {}
    for cell in sheet["cells"].values():
        if _is_non_anchor_merged_cell(cell):
            continue
        row_idx, col_idx = _coord_indexes(cell["coordinate"])
        static_by_row.setdefault(row_idx, {})[col_idx] = cell
    for row_idx in range(min_row, max_row + 1):
        merge_by_row[row_idx] = _static_row_merges(sheet, row_idx)

    content_states: list[dict[str, Any]] = []
    for anchor in sheet.get("dataframe_anchors", []):
        if anchor["placeholder_type"] == "dataframe-header":
            row_cells = static_by_row.setdefault(anchor["start_row"], {})
            for _, cell in _header_cells(anchor):
                _, col_idx = _coord_indexes(cell["coordinate"])
                row_cells[col_idx] = cell
            merge_by_row.setdefault(anchor["start_row"], []).extend(
                _anchor_row_merges(anchor, anchor["start_row"])
            )
            continue

        source = source_map[anchor["source"]]
        content_states.append(
            {
                "anchor": anchor,
                "rows": _source_rows(bundle, source, batch_size=batch_size),
                "remaining": int(source.get("rows", 0)),
            }
        )

    for row_idx in range(min_row, max_row + 1):
        row_cells = dict(static_by_row.get(row_idx, {}))
        row_merges = list(merge_by_row.get(row_idx, []))
        for state in content_states:
            anchor = state["anchor"]
            if row_idx < anchor["start_row"] or state["remaining"] <= 0:
                continue
            try:
                row_values = next(state["rows"])
            except StopIteration:
                state["remaining"] = 0
                continue
            state["remaining"] -= 1
            for layout, value in zip(_anchor_layouts(anchor), row_values):
                col_idx = anchor["start_col"] + layout["start_col_offset"]
                content_cell = dict(anchor["cell"])
                content_cell["value"] = value
                content_cell["cell_type"] = _infer_cell_type(value)
                content_cell["alignment"] = _layout_alignment(anchor["cell"], layout)
                row_cells[col_idx] = content_cell  # type: ignore[assignment]
            row_merges.extend(_anchor_row_merges(anchor, row_idx))
        yield {"cells": row_cells, "merges": row_merges, "row_idx": row_idx}


def _chunked_row_tables(
    sheet: SheetSchema,
    row_items: Iterator[dict[str, Any]],
    min_col: int,
    max_col: int,
    available_width: float,
    font_resolver: _FontResolver,
    *,
    streaming_chunk_rows: int,
    chunk_row_height: float | None = None,
) -> Iterator[Table]:
    chunk: list[dict[str, Any]] = []
    for row_item in row_items:
        if (
            chunk
            and _chunk_merges_fit(chunk)
            and len(chunk) + _row_item_merge_height(row_item) > streaming_chunk_rows
        ):
            yield _row_chunk_table(
                sheet,
                chunk,
                min_col,
                max_col,
                available_width,
                font_resolver,
                per_row_heights=_per_row_heights(sheet, chunk, chunk_row_height),
            )
            chunk = []
        chunk.append(row_item)
        if len(chunk) >= streaming_chunk_rows and _chunk_merges_fit(chunk):
            yield _row_chunk_table(
                sheet,
                chunk,
                min_col,
                max_col,
                available_width,
                font_resolver,
                per_row_heights=_per_row_heights(sheet, chunk, chunk_row_height),
            )
            chunk = []
    if chunk:
        yield _row_chunk_table(
            sheet,
            chunk,
            min_col,
            max_col,
            available_width,
            font_resolver,
            per_row_heights=_per_row_heights(sheet, chunk, chunk_row_height),
        )


def _resolved_row_page_breaks(sheet: SheetSchema) -> set[int]:
    return {
        int(break_idx)
        for break_idx in sheet.get(
            "resolved_row_page_breaks",
            sheet.get("row_page_breaks", []),
        )
    }


def _chunked_row_flowables(
    sheet: SheetSchema,
    row_items: Iterator[dict[str, Any]],
    min_col: int,
    max_col: int,
    available_width: float,
    font_resolver: _FontResolver,
    *,
    streaming_chunk_rows: int,
    page_breaks: set[int] | None = None,
    chunk_row_height: float | None = None,
) -> Iterator[Any]:
    chunk: list[dict[str, Any]] = []
    previous_row_idx: int | None = None
    manual_breaks = page_breaks or set()
    for row_item in row_items:
        row_idx = int(row_item.get("row_idx", 0) or 0)
        if chunk and previous_row_idx in manual_breaks:
            yield _row_chunk_table(
                sheet,
                chunk,
                min_col,
                max_col,
                available_width,
                font_resolver,
                per_row_heights=_per_row_heights(sheet, chunk, chunk_row_height),
            )
            yield PageBreak()
            chunk = []
        if (
            chunk
            and _chunk_merges_fit(chunk)
            and len(chunk) + _row_item_merge_height(row_item) > streaming_chunk_rows
        ):
            yield _row_chunk_table(
                sheet,
                chunk,
                min_col,
                max_col,
                available_width,
                font_resolver,
                per_row_heights=_per_row_heights(sheet, chunk, chunk_row_height),
            )
            chunk = []
        chunk.append(row_item)
        previous_row_idx = row_idx
    if chunk:
        yield _row_chunk_table(
            sheet,
            chunk,
            min_col,
            max_col,
            available_width,
            font_resolver,
            per_row_heights=_per_row_heights(sheet, chunk, chunk_row_height),
        )


def _dataframe_sheet_flowables(
    bundle: ReportBundle,
    raw_sheet: dict[str, Any],
    *,
    available_width: float,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    font_resolver: _FontResolver,
) -> Iterator[Table]:
    sheet = _sheet_with_overrides(
        raw_sheet,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
    )
    _validate_pdf_dataframe_sizing(
        sheet,
        column_width_mode=column_width_mode,
    )
    source_map = _data_source_map(bundle)
    sheet = _expanded_sheet_from_manifest(sheet, source_map)
    min_col, _, max_col, _ = _parse_dims(sheet["dimensions"])
    chunk_row_height = _content_anchor_row_height(sheet)
    yield from _chunked_row_flowables(
        sheet,
        _dataframe_pdf_rows(
            bundle,
            source_map,
            sheet,
            batch_size=streaming_chunk_rows,
        ),
        min_col,
        max_col,
        available_width,
        font_resolver,
        streaming_chunk_rows=streaming_chunk_rows,
        page_breaks=_resolved_row_page_breaks(sheet),
        chunk_row_height=chunk_row_height,
    )


def _repeat_sheet_flowables(
    bundle: ReportBundle,
    raw_sheet: dict[str, Any],
    *,
    available_width: float,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    font_resolver: _FontResolver,
) -> Iterator[Any]:
    sheet = _sheet_with_overrides(
        raw_sheet,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
    )
    source_map = _data_source_map(bundle)
    min_col, _, max_col, _ = _parse_dims(sheet["dimensions"])
    max_col = _repeat_max_col(sheet, max_col)
    yield from _chunked_row_flowables(
        sheet,
        _repeat_row_stream(
            bundle,
            source_map,
            sheet,
            streaming_chunk_rows=streaming_chunk_rows,
        ),
        min_col,
        max_col,
        available_width,
        font_resolver,
        streaming_chunk_rows=streaming_chunk_rows,
        page_breaks=_resolved_row_page_breaks(sheet),
    )


def _static_sheet_flowables(
    bundle: ReportBundle,
    raw_sheet: dict[str, Any],
    *,
    available_width: float,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    font_resolver: _FontResolver,
) -> Iterator[Any]:
    del bundle
    sheet = _sheet_with_overrides(
        raw_sheet,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
    )
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    segment_rows = max(max_row - min_row + 1, 1)
    yield from _chunked_row_flowables(
        sheet,
        _static_pdf_rows(sheet, min_row, max_row),
        min_col,
        max_col,
        available_width,
        font_resolver,
        streaming_chunk_rows=max(segment_rows, streaming_chunk_rows),
        page_breaks=_resolved_row_page_breaks(sheet),
    )


def _sheet_flowables(
    bundle: ReportBundle,
    raw_sheet: dict[str, Any],
    *,
    available_width: float,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    font_resolver: _FontResolver,
) -> Iterator[Any]:
    if raw_sheet.get("repeat_sections"):
        yield from _repeat_sheet_flowables(
            bundle,
            raw_sheet,
            available_width=available_width,
            column_width_mode=column_width_mode,
            row_height_mode=row_height_mode,
            default_column_width=default_column_width,
            default_row_height=default_row_height,
            streaming_chunk_rows=streaming_chunk_rows,
            font_resolver=font_resolver,
        )
        return
    if _has_dataframe_content(raw_sheet):
        yield from _dataframe_sheet_flowables(
            bundle,
            raw_sheet,
            available_width=available_width,
            column_width_mode=column_width_mode,
            row_height_mode=row_height_mode,
            default_column_width=default_column_width,
            default_row_height=default_row_height,
            streaming_chunk_rows=streaming_chunk_rows,
            font_resolver=font_resolver,
        )
        return
    yield from _static_sheet_flowables(
        bundle,
        raw_sheet,
        available_width=available_width,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
        streaming_chunk_rows=streaming_chunk_rows,
        font_resolver=font_resolver,
    )


def _repeat_max_col(sheet: SheetSchema, current: int) -> int:
    max_col = current
    for section in sheet.get("repeat_sections", []):
        for record in section["records"]:
            for item in record["cells"]:
                max_col = max(max_col, item["start_col"])
            for anchor in record["dataframe_anchors"]:
                max_col = max(
                    max_col,
                    anchor["start_col"] + max(_occupied_width(anchor) - 1, 0),
                )
    return max_col


def _row_item_merge_height(row_item: dict[str, Any]) -> int:
    height = 1
    for merge in row_item.get("merges", []):
        height = max(height, merge["max_row_offset"] - merge["min_row_offset"] + 1)
    return height


def _chunk_merges_fit(rows: list[dict[str, Any]]) -> bool:
    for row_offset, row_item in enumerate(rows, start=1):
        for merge in row_item.get("merges", []):
            max_row = row_offset + merge["max_row_offset"] - merge["min_row_offset"]
            if max_row > len(rows):
                return False
    return True


def _static_pdf_rows(
    sheet: SheetSchema, start_row: int, end_row: int
) -> Iterator[dict[str, Any]]:
    if end_row < start_row:
        return
    for row_idx in range(start_row, end_row + 1):
        row_cells: dict[int, CellSchema] = {}
        for cell in sheet["cells"].values():
            cell_row, col_idx = _coord_indexes(cell["coordinate"])
            if cell_row == row_idx:
                row_cells[col_idx] = cell
        yield {"cells": row_cells, "merges": _static_row_merges(sheet, row_idx), "row_idx": row_idx}


def _per_row_heights(
    sheet: SheetSchema,
    rows: list[dict[str, Any]],
    data_row_height: float | None,
) -> list[float] | None:
    """Build a per-row height list, using explicit template heights for static rows
    and data_row_height for generated content rows (no explicit schema entry)."""
    mode = sheet.get("row_height_mode", "fixed")
    if mode == "hug":
        return None
    row_heights_map = sheet.get("row_heights", {})
    default = float(sheet.get("default_row_height", _DEFAULT_ROW_HEIGHT))
    out: list[float] = []
    for item in rows:
        row_idx = item.get("row_idx")
        if row_idx is not None:
            h = row_heights_map.get(str(row_idx))
            out.append(float(h) if h else (data_row_height or default))
        else:
            out.append(data_row_height or default)
    return out


def _row_chunk_table(
    sheet: SheetSchema,
    rows: list[dict[str, Any]],
    min_col: int,
    max_col: int,
    available_width: float,
    font_resolver: _FontResolver,
    *,
    per_row_heights: list[float] | None = None,
) -> Table:
    cells: dict[tuple[int, int], CellSchema] = {}
    merged_regions: list[str] = []
    for row_offset, row_item in enumerate(rows, start=1):
        row_map = row_item["cells"]
        for col_idx, cell in row_map.items():
            adjusted = dict(cell)
            adjusted["coordinate"] = f"{_col_letter(col_idx)}{row_offset}"
            cells[(row_offset, col_idx)] = adjusted  # type: ignore[assignment]
        for merge in row_item.get("merges", []):
            max_row = row_offset + merge["max_row_offset"] - merge["min_row_offset"]
            if max_row <= len(rows):
                merged_regions.append(
                    f"{_col_letter(merge['min_col'])}{row_offset}:"
                    f"{_col_letter(merge['max_col'])}{max_row}"
                )
    data = _table_data(cells, font_resolver, min_col, 1, max_col, len(rows))
    chunk_sheet = dict(sheet)
    chunk_sheet["merged_regions"] = merged_regions
    table = Table(
        data,
        colWidths=_column_widths(
            chunk_sheet,
            cells,
            min_col,
            1,
            max_col,
            len(rows),
            available_width,
        ),
        rowHeights=per_row_heights,
        repeatRows=0,
        splitByRow=1,
    )
    table.setStyle(_table_style(chunk_sheet, cells, min_col, 1, max_col, len(rows), font_resolver))
    return table


def _content_anchor_row_height(sheet: SheetSchema) -> float | None:
    """Return the fixed row height for dataframe-content rows, or None to auto-size."""
    mode = sheet.get("row_height_mode", "fixed")
    if mode == "hug":
        return None
    anchor = next(
        (a for a in sheet.get("dataframe_anchors", []) if a["placeholder_type"] == "dataframe-content"),
        None,
    )
    if anchor is None:
        return None
    if mode == "even":
        return float(sheet.get("default_row_height", _DEFAULT_ROW_HEIGHT))
    return float(sheet["row_heights"].get(str(anchor["start_row"])) or _DEFAULT_ROW_HEIGHT)


def _ensure_output_parent(output_path: str) -> None:
    parent = Path(output_path).parent
    if parent and not parent.exists():
        parent.mkdir(parents=True, exist_ok=True)


# §4. Public Functions


def export_report_bundle(
    bundle: ReportBundle,
    output_path: str,
    *,
    column_width_mode: str | None = None,
    row_height_mode: str | None = None,
    default_column_width: float | None = None,
    default_row_height: float | None = None,
    page_size: str = "A4",
    orientation: str = "portrait",
    margin: float = 36,
    streaming_chunk_rows: int = 50_000,
    fonts: dict[str, Any] | None = None,
) -> None:
    """Render a ReportBundle to a styled PDF."""
    if streaming_chunk_rows <= 0:
        raise ValueError(
            f"streaming_chunk_rows must be greater than 0, got {streaming_chunk_rows}"
        )
    if margin < 0:
        raise ValueError(f"margin must be non-negative, got {margin}")

    pagesize = _page_size(page_size, orientation)
    font_resolver = _FontResolver(fonts)
    _ensure_output_parent(output_path)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )
    available_width = pagesize[0] - (margin * 2)
    def flowables() -> Iterator[Any]:
        for index, raw_sheet in enumerate(bundle.report["sheets"]):
            if index:
                yield PageBreak()
            yield from _sheet_flowables(
                bundle,
                raw_sheet,
                available_width=available_width,
                column_width_mode=column_width_mode,
                row_height_mode=row_height_mode,
                default_column_width=default_column_width,
                default_row_height=default_row_height,
                streaming_chunk_rows=streaming_chunk_rows,
                font_resolver=font_resolver,
            )

    doc.build(_LazyFlowables(flowables()))


# §5. Entrypoints
