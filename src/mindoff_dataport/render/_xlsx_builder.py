from __future__ import annotations

import datetime
import functools

from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)
from openpyxl.worksheet.worksheet import Worksheet

from ..schema import (
    AlignmentSchema,
    CellBorders,
    CellSchema,
    FillSchema,
    FontSchema,
    SheetSchema,
)
from ._style_conversion import argb_to_color, dict_to_border_side

# §1. Constants & Exceptions

_BORDER_SIDE_KEYS = (
    "top", "bottom", "left", "right",
    "start", "end", "horizontal", "vertical", "diagonal",
)

# §2. Classes and Sub Classes

# §3. Private Helper Functions


def _apply_dimensions(ws: Worksheet, schema: SheetSchema) -> None:
    _apply_column_dimensions(ws, schema)
    _apply_row_dimensions(ws, schema)


def _apply_column_dimensions(ws: Worksheet, schema: SheetSchema) -> None:
    col_mode = schema.get("column_width_mode", "fixed")
    if col_mode == "fixed":
        for col_letter, width in schema["column_widths"].items():
            if width is not None:
                ws.column_dimensions[col_letter].width = width
    elif col_mode == "even":
        width = schema.get("default_column_width", 15.0)
        for col_letter in _cols_in_range(schema["dimensions"]):
            ws.column_dimensions[col_letter].width = width


def _apply_row_dimensions(ws: Worksheet, schema: SheetSchema) -> None:
    row_mode = schema.get("row_height_mode", "fixed")
    if row_mode == "fixed":
        for row_str, height in schema["row_heights"].items():
            if height is not None:
                ws.row_dimensions[int(row_str)].height = height
    elif row_mode == "even":
        height = schema.get("default_row_height", 15.0)
        for row_idx in _rows_in_range(schema["dimensions"]):
            ws.row_dimensions[row_idx].height = height


def _cols_in_range(dimensions: str) -> list[str]:
    """Parse 'A1:F9' into ['A', 'B', ..., 'F']."""
    start, end = (dimensions.split(":") + [dimensions])[:2]
    min_col = column_index_from_string(coordinate_from_string(start)[0])
    max_col = column_index_from_string(coordinate_from_string(end)[0])
    return [get_column_letter(i) for i in range(min_col, max_col + 1)]


def _rows_in_range(dimensions: str) -> list[int]:
    """Parse 'A1:F9' into [1, 2, ..., 9]."""
    start, end = (dimensions.split(":") + [dimensions])[:2]
    min_row = coordinate_from_string(start)[1]
    max_row = coordinate_from_string(end)[1]
    return list(range(min_row, max_row + 1))


def _apply_hug_columns(ws: Worksheet, schema: SheetSchema) -> None:
    """Set each column width to fit its widest cell content (approximate)."""
    for col_letter in _cols_in_range(schema["dimensions"]):
        max_len = 0
        for cell in ws[col_letter]:
            if cell.value is not None:
                length = len(str(cell.value))
                if getattr(cell.font, "bold", False):
                    length = int(length * 1.2)
                max_len = max(max_len, length)
        if max_len > 0:
            ws.column_dimensions[col_letter].width = max_len + 2


def _apply_hug_rows(ws: Worksheet, schema: SheetSchema) -> None:
    """Set each row height from the largest font size in that row."""
    for row_idx in _rows_in_range(schema["dimensions"]):
        max_font_size = 11.0
        for cell in ws[row_idx]:
            size = getattr(cell.font, "size", None)
            if size:
                max_font_size = max(max_font_size, float(size))
        ws.row_dimensions[row_idx].height = max_font_size * 1.5


def _apply_cell_value(cell, schema: CellSchema) -> None:
    if schema["cell_type"] == "date" and isinstance(schema["value"], str):
        cell.value = datetime.datetime.fromisoformat(schema["value"])
    else:
        cell.value = schema["value"]


def _apply_cell_styles(cell, schema: CellSchema) -> None:
    cell.font = _build_font(schema["font"])
    cell.fill = _build_fill(schema["fill"])
    cell.alignment = _build_alignment(schema["alignment"])
    cell.border = _build_border(schema["borders"])
    if schema["number_format"]:
        cell.number_format = schema["number_format"]


def _freeze(d: dict) -> tuple:
    """Convert nested dicts to a stable hashable key for style caches."""
    return tuple(
        (k, _freeze(v) if isinstance(v, dict) else v) for k, v in sorted(d.items())
    )


@functools.lru_cache(maxsize=512)
def _cached_font(key: tuple) -> Font:
    d = dict(key)
    color = argb_to_color(d.get("color"))
    kwargs: dict = {
        "name": d.get("name"),
        "size": d.get("size"),
        "bold": d.get("bold", False),
        "italic": d.get("italic", False),
        "underline": d.get("underline"),
        "strike": d.get("strike", False),
    }
    if d.get("vert_align") is not None:
        kwargs["vertAlign"] = d["vert_align"]
    if color is not None:
        kwargs["color"] = color
    return Font(**kwargs)


@functools.lru_cache(maxsize=512)
def _cached_fill(key: tuple) -> PatternFill:
    d = dict(key)
    pt = d.get("pattern_type") or d.get("bg_color") and "solid"
    if not pt:
        return PatternFill()
    fg = argb_to_color(d.get("fg_color") or d.get("bg_color"))
    bg = argb_to_color(d.get("bg_color") if d.get("fg_color") else None)
    kwargs: dict = {"patternType": pt}
    if fg is not None:
        kwargs["fgColor"] = fg
    if bg is not None:
        kwargs["bgColor"] = bg
    return PatternFill(**kwargs)


@functools.lru_cache(maxsize=512)
def _cached_alignment(key: tuple) -> Alignment:
    d = dict(key)
    kwargs: dict = {
        "horizontal": d.get("horizontal"),
        "vertical": d.get("vertical"),
        "wrap_text": d.get("wrap_text", False),
    }
    if d.get("indent") is not None:
        kwargs["indent"] = d["indent"]
    if d.get("relative_indent") is not None:
        kwargs["relativeIndent"] = d["relative_indent"]
    if d.get("text_rotation") is not None:
        kwargs["textRotation"] = d["text_rotation"]
    if d.get("shrink_to_fit"):
        kwargs["shrinkToFit"] = d["shrink_to_fit"]
    if d.get("reading_order") is not None:
        kwargs["readingOrder"] = d["reading_order"]
    return Alignment(**kwargs)


_FROZEN_EMPTY_SIDE = (("color", None), ("style", None))


def _has_border_side(data: dict) -> bool:
    return bool(data.get("style") or data.get("color"))


@functools.lru_cache(maxsize=512)
def _cached_border(key: tuple) -> Border:
    d = dict(key)
    kwargs = {
        k: dict_to_border_side(dict(d.get(k, _FROZEN_EMPTY_SIDE)))
        for k in ("top", "bottom", "left", "right")
    }
    for key_name in ("start", "end", "horizontal", "vertical", "diagonal"):
        side = dict(d.get(key_name, _FROZEN_EMPTY_SIDE))
        if _has_border_side(side):
            kwargs[key_name] = dict_to_border_side(side)
    kwargs["diagonalUp"] = bool(d.get("diagonal_up", False))
    kwargs["diagonalDown"] = bool(d.get("diagonal_down", False))
    kwargs["outline"] = bool(d.get("outline", True))
    return Border(**kwargs)


def _build_font(schema: FontSchema) -> Font:
    return _cached_font(_freeze(schema))


def _build_fill(schema: FillSchema) -> PatternFill:
    return _cached_fill(_freeze(schema))


def _build_alignment(schema: AlignmentSchema) -> Alignment:
    return _cached_alignment(_freeze(schema))


def _build_border(schema: CellBorders) -> Border:
    return _cached_border(_freeze(schema))


# §4. Public Functions


# §5. Entrypoints
