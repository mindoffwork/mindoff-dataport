from __future__ import annotations

import datetime

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..schema import (
    AlignmentSchema,
    BorderSide,
    CellBorders,
    CellSchema,
    CellType,
    FillSchema,
    FontSchema,
    SheetSchema,
    WorkbookSchema,
)
from ._page_breaks import extract_manual_breaks
from ..render._style_conversion import border_side_to_dict, normalize_color

# §1. Constants & Exceptions

# Shared defaults used for merged-cell stubs (never mutated downstream).
_EMPTY_BORDER_SIDE: BorderSide = {"style": None, "color": None}
_EMPTY_FONT: FontSchema = {
    "name": None,
    "size": None,
    "bold": False,
    "italic": False,
    "underline": None,
    "strike": False,
    "vert_align": None,
    "color": None,
}
_EMPTY_FILL: FillSchema = {"pattern_type": None, "fg_color": None, "bg_color": None}
_EMPTY_ALIGNMENT: AlignmentSchema = {
    "horizontal": None,
    "vertical": None,
    "wrap_text": False,
    "indent": None,
    "relative_indent": None,
    "text_rotation": None,
    "shrink_to_fit": False,
    "reading_order": None,
}
_EMPTY_BORDERS: CellBorders = {
    "top": _EMPTY_BORDER_SIDE,
    "bottom": _EMPTY_BORDER_SIDE,
    "left": _EMPTY_BORDER_SIDE,
    "right": _EMPTY_BORDER_SIDE,
    "start": _EMPTY_BORDER_SIDE,
    "end": _EMPTY_BORDER_SIDE,
    "horizontal": _EMPTY_BORDER_SIDE,
    "vertical": _EMPTY_BORDER_SIDE,
    "diagonal": _EMPTY_BORDER_SIDE,
    "diagonal_up": False,
    "diagonal_down": False,
    "outline": True,
}

# §2. Classes and Sub Classes

# §3. Private Helper Functions


def _extract_sheet(ws: Worksheet) -> SheetSchema:
    merged_regions: list[str] = [str(mr) for mr in ws.merged_cells.ranges]
    merge_map = _build_merge_map(ws)

    cells: dict[str, CellSchema] = {}
    for row in ws.iter_rows():
        for cell in row:
            cells[cell.coordinate] = _extract_cell(cell, merge_map)
    _apply_merged_region_borders(ws, cells)

    column_widths = _extract_column_widths(ws)
    row_heights = {
        str(row_idx): ws.row_dimensions[row_idx].height
        for row_idx in ws.row_dimensions
        if ws.row_dimensions[row_idx].height is not None
    }

    result: SheetSchema = {
        "name": ws.title,
        "dimensions": ws.dimensions,
        "merged_regions": merged_regions,
        "column_widths": column_widths,
        "row_heights": row_heights,
        "show_gridlines": bool(ws.sheet_view.showGridLines),
        "cells": cells,
    }
    row_page_breaks = extract_manual_breaks(ws.row_breaks)
    column_page_breaks = extract_manual_breaks(ws.col_breaks)
    if row_page_breaks:
        result["row_page_breaks"] = row_page_breaks
    if column_page_breaks:
        result["column_page_breaks"] = column_page_breaks
    return result


def _build_merge_map(ws: Worksheet) -> dict[str, str]:
    """Map each coordinate inside merged regions to its anchor coordinate."""
    result: dict[str, str] = {}
    for merged_range in ws.merged_cells.ranges:
        anchor = merged_range.coord.split(":")[0]
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                result[f"{get_column_letter(col)}{row}"] = anchor
    return result


def _extract_column_widths(ws: Worksheet) -> dict[str, float]:
    widths: dict[str, float] = {}
    for col_letter, dimension in ws.column_dimensions.items():
        width = dimension.width
        if width is None:
            continue
        min_col = dimension.min or 1
        max_col = dimension.max or min_col
        for col_idx in range(min_col, max_col + 1):
            widths[get_column_letter(col_idx)] = width
    return widths


def _extract_cell(cell: Cell | MergedCell, merge_map: dict[str, str]) -> CellSchema:
    is_merged = cell.coordinate in merge_map
    anchor = merge_map.get(cell.coordinate)

    if isinstance(cell, MergedCell):
        return {
            "coordinate": cell.coordinate,
            "value": None,
            "cell_type": "empty",
            "number_format": None,
            "font": _EMPTY_FONT,
            "fill": _EMPTY_FILL,
            "alignment": _EMPTY_ALIGNMENT,
            "borders": _EMPTY_BORDERS,
            "merged": True,
            "merge_anchor": anchor,
        }

    return {
        "coordinate": cell.coordinate,
        "value": _serialize_value(cell.value),
        "cell_type": _infer_cell_type(cell),
        "number_format": (
            cell.number_format if cell.number_format not in (None, "General") else None
        ),
        "font": _extract_font(cell),
        "fill": _extract_fill(cell),
        "alignment": _extract_alignment(cell),
        "borders": _extract_borders(cell),
        "merged": is_merged,
        "merge_anchor": anchor,
    }


def _infer_cell_type(cell: Cell) -> CellType:
    if cell.value is None:
        return "empty"
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "n":
        return "number"
    if cell.data_type == "d" or isinstance(
        cell.value, (datetime.datetime, datetime.date)
    ):
        return "date"
    return "string"


def _serialize_value(value: object) -> object:
    if isinstance(value, datetime.datetime):
        return value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return value


def _extract_font(cell: Cell) -> FontSchema:
    f = cell.font
    return {
        "name": f.name,
        "size": float(f.size) if f.size is not None else None,
        "bold": bool(f.bold),
        "italic": bool(f.italic),
        "underline": f.underline,
        "strike": bool(f.strike),
        "vert_align": f.vertAlign,
        "color": normalize_color(f.color),
    }


def _extract_fill(cell: Cell) -> FillSchema:
    fill = cell.fill
    pt = getattr(fill, "patternType", None) or getattr(fill, "fill_type", None)
    if pt and pt != "none":
        return {
            "pattern_type": pt,
            "fg_color": normalize_color(fill.fgColor) if hasattr(fill, "fgColor") else None,
            "bg_color": normalize_color(fill.bgColor) if hasattr(fill, "bgColor") else None,
        }
    return {"pattern_type": None, "fg_color": None, "bg_color": None}


def _extract_alignment(cell: Cell) -> AlignmentSchema:
    a = cell.alignment
    indent = a.indent
    rel = a.relativeIndent
    rot = a.textRotation
    order = a.readingOrder
    return {
        "horizontal": a.horizontal,
        "vertical": a.vertical,
        "wrap_text": bool(a.wrap_text),
        "indent": int(indent) if indent else None,
        "relative_indent": int(rel) if rel else None,
        "text_rotation": int(rot) if rot else None,
        "shrink_to_fit": bool(a.shrink_to_fit),
        "reading_order": int(order) if order else None,
    }


def _extract_borders(cell: Cell) -> CellBorders:
    b = cell.border
    return {
        "top": border_side_to_dict(b.top),
        "bottom": border_side_to_dict(b.bottom),
        "left": border_side_to_dict(b.left),
        "right": border_side_to_dict(b.right),
        "start": border_side_to_dict(b.start),
        "end": border_side_to_dict(b.end),
        "horizontal": border_side_to_dict(b.horizontal),
        "vertical": border_side_to_dict(b.vertical),
        "diagonal": border_side_to_dict(b.diagonal),
        "diagonal_up": bool(getattr(b, "diagonalUp", False)),
        "diagonal_down": bool(getattr(b, "diagonalDown", False)),
        "outline": bool(getattr(b, "outline", True)),
    }


def _apply_merged_region_borders(
    ws: Worksheet, cells: dict[str, CellSchema]
) -> None:
    for merged_range in ws.merged_cells.ranges:
        anchor = cells.get(merged_range.coord.split(":")[0])
        if anchor is None:
            continue
        borders = dict(anchor["borders"])
        edges = {
            "top": _merged_edge_side(ws, merged_range.min_row, merged_range, "top"),
            "bottom": _merged_edge_side(
                ws, merged_range.max_row, merged_range, "bottom"
            ),
            "left": _merged_edge_side(ws, merged_range.min_col, merged_range, "left"),
            "right": _merged_edge_side(
                ws, merged_range.max_col, merged_range, "right"
            ),
        }
        for side, edge in edges.items():
            if _has_border(edge):
                borders[side] = edge
        anchor["borders"] = borders  # type: ignore[assignment]


def _merged_edge_side(
    ws: Worksheet, edge_index: int, merged_range, side: str
) -> BorderSide:
    if side in {"top", "bottom"}:
        coords = (
            (edge_index, col)
            for col in range(merged_range.min_col, merged_range.max_col + 1)
        )
    else:
        coords = (
            (row, edge_index)
            for row in range(merged_range.min_row, merged_range.max_row + 1)
        )
    for row, col in coords:
        border_side = getattr(ws.cell(row=row, column=col).border, side)
        result = border_side_to_dict(border_side)
        if _has_border(result):
            return result
    return dict(_EMPTY_BORDER_SIDE)


def _has_border(side: BorderSide) -> bool:
    return bool(side.get("style"))


# §4. Public Functions


def extract_template(path: str) -> WorkbookSchema:
    """Load .xlsx at *path* and return a WorkbookSchema dict."""
    wb = openpyxl.load_workbook(path, data_only=False)
    return {"sheets": [_extract_sheet(ws) for ws in wb.worksheets]}


# §5. Entrypoints
