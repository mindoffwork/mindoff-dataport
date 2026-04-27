from __future__ import annotations

import datetime
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterator
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)

from .builder import _build_alignment, _build_border, _build_fill, _build_font
from .renderer import (
    PLACEHOLDER_RE,
    _assert_dataframe,
    _infer_cell_type,
    _parse_dims,
    _resolve_sheet_payloads,
    _substitute_scalars,
    _to_headers,
)
from .schema import CellSchema, SheetSchema, WorkbookSchema

__all__ = ["build_template_streaming_with_data"]

# §1 Constants & Exceptions

MAX_EXCEL_ROWS = 1_048_576

# §2 Classes and Sub Classes


@dataclass(frozen=True)
class _AnchorStyle:
    font: Any
    fill: Any
    alignment: Any
    border: Any
    number_format: str | None


@dataclass
class _SheetPlan:
    sheet: SheetSchema
    static_cells: dict[tuple[int, int], CellSchema]
    anchors: list["_StreamingAnchor"]
    min_col: int
    min_row: int
    max_col: int
    max_row: int


@dataclass
class _StreamingAnchor:
    id: str
    key: str
    sheet_name: str
    start_row: int
    start_col: int
    columns: list[str]
    style: _AnchorStyle
    source: Any
    on_batch: Callable[[Any], None] | None
    current_batch: list[tuple[Any, ...]]
    current_batch_idx: int = 0
    exhausted: bool = False
    batches: Iterator[Any] | None = None
    slice_offset: int = 0
    fallback_slice: bool = False
    non_lazy_consumed: bool = False

    def next_row(self, chunk_size: int) -> tuple[Any, ...] | None:
        if self.exhausted:
            return None

        while True:
            if self.current_batch_idx < len(self.current_batch):
                row = self.current_batch[self.current_batch_idx]
                self.current_batch_idx += 1
                return row

            self.current_batch = []
            self.current_batch_idx = 0
            batch = self._next_batch(chunk_size)
            if batch is None:
                self.exhausted = True
                return None
            self.current_batch = batch

    def _next_batch(self, chunk_size: int) -> list[tuple[Any, ...]] | None:
        module = getattr(type(self.source), "__module__", "") or ""
        qualname = type(self.source).__qualname__

        if "polars" in module and qualname == "LazyFrame":
            return self._next_lazy_polars_batch(chunk_size)

        if "polars" in module:
            if self.non_lazy_consumed:
                return None
            if self.on_batch is not None:
                self.on_batch(self.source)
            rows = self.source.rows()
            self.non_lazy_consumed = True
            return [tuple(row) for row in rows] if rows else None

        if "pandas" in module and "DataFrame" in qualname:
            if self.non_lazy_consumed:
                return None
            rows = list(self.source.itertuples(index=False))
            self.non_lazy_consumed = True
            return [tuple(row) for row in rows] if rows else None

        raise TypeError(
            f"'{self.key}' expected a polars/pandas DataFrame or LazyFrame, got {type(self.source).__name__}"
        )

    def _next_lazy_polars_batch(self, chunk_size: int) -> list[tuple[Any, ...]] | None:
        if self.fallback_slice:
            sliced = self.source.slice(self.slice_offset, chunk_size).collect()
            rows = sliced.rows()
            if not rows:
                return None
            self.slice_offset += len(rows)
            if self.on_batch is not None:
                self.on_batch(sliced)
            return [tuple(row) for row in rows]

        if self.batches is None:
            try:
                self.batches = self.source.collect_batches(
                    chunk_size=chunk_size, maintain_order=True
                )
            except Exception:
                self.fallback_slice = True
                return self._next_lazy_polars_batch(chunk_size)

        try:
            batch_df = next(self.batches)
        except StopIteration:
            return None
        except Exception:
            self.fallback_slice = True
            return self._next_lazy_polars_batch(chunk_size)

        rows = batch_df.rows()
        self.slice_offset += len(rows)
        if self.on_batch is not None:
            self.on_batch(batch_df)
        return [tuple(row) for row in rows] if rows else None


@dataclass
class _ParquetSink:
    anchor_id: str
    file_path: Path
    columns: list[str]
    row_count: int = 0
    _writer: Any = None

    def write_batch(self, batch_df: Any) -> None:
        try:
            import pyarrow.parquet as pq  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ValueError(
                "streaming_bundle_with_parquet requires optional dependency 'pyarrow'."
            ) from exc

        table = batch_df.to_arrow()
        if self._writer is None:
            self._writer = pq.ParquetWriter(str(self.file_path), table.schema)
        self._writer.write_table(table)
        self.row_count += batch_df.height

    def finalize(self) -> None:
        if self._writer is not None:
            self._writer.close()
            self._writer = None
            return
        self._write_empty()

    def _write_empty(self) -> None:
        try:
            import polars as pl  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ValueError(
                "streaming_bundle_with_parquet requires optional dependency 'polars'."
            ) from exc

        schema = {col: pl.String for col in self.columns}
        pl.DataFrame(schema=schema).write_parquet(self.file_path)


# §3 Private Helper Functions


def _validate_streaming_modes(
    schema: WorkbookSchema, column_width_mode: str | None, row_height_mode: str | None
) -> None:
    for sheet in schema["sheets"]:
        col_mode = column_width_mode or sheet.get("column_width_mode", "fixed")
        row_mode = row_height_mode or sheet.get("row_height_mode", "fixed")
        if col_mode == "hug" or row_mode == "hug":
            raise ValueError(
                "Streaming mode does not support 'hug' sizing. Use 'fixed' or 'even', "
                "or switch to export_mode='fidelity'."
            )
        if sheet.get("merged_regions"):
            raise ValueError(
                "Streaming mode does not support merged cells. Use export_mode='fidelity' "
                "for templates with merges."
            )


def _prepare_streaming_plans(
    *,
    schema: WorkbookSchema,
    data: dict[str, Any],
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    on_anchor_batch: Callable[[_StreamingAnchor], Callable[[Any], None] | None] | None = None,
) -> list[_SheetPlan]:
    plans: list[_SheetPlan] = []
    used_anchor_ids: set[str] = set()
    for sheet, output_name, sheet_data, _ in _resolve_sheet_payloads(schema, data):

        merged = dict(sheet)
        merged["name"] = output_name
        if column_width_mode is not None:
            merged["column_width_mode"] = column_width_mode
        if row_height_mode is not None:
            merged["row_height_mode"] = row_height_mode
        if default_column_width is not None:
            merged["default_column_width"] = default_column_width
        if default_row_height is not None:
            merged["default_row_height"] = default_row_height

        plans.append(
            _plan_sheet(
                merged,
                sheet_data,
                used_anchor_ids=used_anchor_ids,
                on_anchor_batch=on_anchor_batch,
            )
        )
    return plans


def _plan_sheet(
    sheet: SheetSchema,
    sheet_data: dict[str, Any],
    *,
    used_anchor_ids: set[str],
    on_anchor_batch: Callable[[_StreamingAnchor], Callable[[Any], None] | None] | None,
) -> _SheetPlan:
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    static_cells: dict[tuple[int, int], CellSchema] = {}
    anchors: list[_StreamingAnchor] = []

    for coord, cell in sheet["cells"].items():
        col_letter, row_idx = coordinate_from_string(coord)
        col_idx = column_index_from_string(col_letter)
        value = cell.get("value")
        if not isinstance(value, str):
            static_cells[(row_idx, col_idx)] = cell
            continue

        full_match = PLACEHOLDER_RE.fullmatch(value.strip())
        if (
            full_match
            and full_match.group(1) in sheet_data
            and full_match.group(2) == "dataframe-content"
        ):
            key = full_match.group(1)
            source = sheet_data[key]
            _assert_dataframe(key, source)
            columns = _headers_from_source(source)
            style = _AnchorStyle(
                font=_build_font(cell["font"]),
                fill=_build_fill(cell["fill"]),
                alignment=_build_alignment(cell["alignment"]),
                border=_build_border(cell["borders"]),
                number_format=cell["number_format"],
            )
            anchor_id = _unique_anchor_id(
                used=used_anchor_ids,
                value=f"{sheet['name']}__{key}",
            )
            anchor = _StreamingAnchor(
                id=anchor_id,
                key=key,
                sheet_name=sheet["name"],
                start_row=row_idx,
                start_col=col_idx,
                columns=columns,
                style=style,
                source=source,
                on_batch=None,
                current_batch=[],
            )
            if on_anchor_batch is not None:
                anchor.on_batch = on_anchor_batch(anchor)
            anchors.append(anchor)
            max_col = max(max_col, col_idx + max(len(columns) - 1, 0))
            continue

        if (
            full_match
            and full_match.group(1) in sheet_data
            and full_match.group(2) == "dataframe-headers"
        ):
            key = full_match.group(1)
            headers = _to_headers(sheet_data[key])
            bold_font = dict(cell["font"])
            bold_font["bold"] = True
            for offset, header in enumerate(headers):
                new_col = col_idx + offset
                new_cell = dict(cell)
                new_cell["coordinate"] = f"{get_column_letter(new_col)}{row_idx}"
                new_cell["value"] = str(header)
                new_cell["cell_type"] = "string"
                new_cell["font"] = bold_font
                static_cells[(row_idx, new_col)] = new_cell  # type: ignore[assignment]
            max_col = max(max_col, col_idx + max(len(headers) - 1, 0))
            continue

        new_val = _substitute_scalars(value, sheet_data)
        new_cell = dict(cell)
        new_cell["value"] = new_val
        if full_match:
            new_cell["cell_type"] = _infer_cell_type(new_val)
        static_cells[(row_idx, col_idx)] = new_cell  # type: ignore[assignment]

    if len(anchors) > 1:
        raise ValueError(
            "Streaming mode currently supports one dataframe-content placeholder per sheet."
        )
    return _SheetPlan(
        sheet=sheet,
        static_cells=static_cells,
        anchors=anchors,
        min_col=min_col,
        min_row=min_row,
        max_col=max_col,
        max_row=max_row,
    )


def _headers_from_source(source: Any) -> list[str]:
    module = getattr(type(source), "__module__", "") or ""
    qualname = type(source).__qualname__
    if "polars" in module:
        if qualname == "LazyFrame":
            return [str(col) for col in source.collect_schema().names()]
        return [str(col) for col in source.columns]
    if "pandas" in module and "DataFrame" in qualname:
        return [str(col) for col in source.columns]
    raise TypeError(
        f"Expected a polars or pandas DataFrame/LazyFrame, got {type(source).__name__}"
    )


def _is_polars_source(source: Any) -> bool:
    module = getattr(type(source), "__module__", "") or ""
    return "polars" in module


def _unique_anchor_id(*, used: set[str], value: str) -> str:
    base = re.sub(r"[^0-9A-Za-z._-]+", "_", value.strip()).strip("._-")
    if not base:
        base = "anchor"
    candidate = base
    i = 2
    while candidate in used:
        candidate = f"{base}_{i}"
        i += 1
    used.add(candidate)
    return candidate


def _first_active_anchor(plans: list[_SheetPlan]) -> _StreamingAnchor | None:
    for plan in plans:
        for anchor in plan.anchors:
            if not anchor.exhausted:
                return anchor
    return None


def _write_sheet_chunk(
    ws,
    plan: _SheetPlan,
    active_anchor: _StreamingAnchor | None,
    max_rows_per_workbook: int,
    rows_budget: int,
    fetch_chunk_size: int,
) -> int:
    anchor = plan.anchors[0] if plan.anchors else None
    max_col = plan.max_col
    effective_max_row = min(plan.max_row, max_rows_per_workbook)
    if anchor is not None and anchor is active_anchor:
        max_col = max(max_col, anchor.start_col + max(len(anchor.columns) - 1, 0))
    effective_max_row = max(effective_max_row, plan.min_row)
    col_count = max_col - plan.min_col + 1
    written_data_rows = 0

    for row_idx in range(plan.min_row, effective_max_row + 1):
        row_cells: list[Any] = [None] * col_count
        for col_idx in range(plan.min_col, max_col + 1):
            idx = col_idx - plan.min_col
            static = plan.static_cells.get((row_idx, col_idx))
            if static is not None:
                row_cells[idx] = _write_only_cell(ws, static)

        if (
            anchor is not None
            and anchor is active_anchor
            and row_idx >= anchor.start_row
        ):
            if written_data_rows < rows_budget:
                row_values = anchor.next_row(fetch_chunk_size)
            else:
                row_values = None
            if row_values is not None:
                for i, value in enumerate(row_values):
                    col_idx = anchor.start_col + i
                    if col_idx < plan.min_col or col_idx > max_col:
                        continue
                    idx = col_idx - plan.min_col
                    row_cells[idx] = _anchor_cell(ws, anchor.style, value)
                written_data_rows += 1

        ws.append(row_cells)

    if anchor is not None and anchor is active_anchor:
        row_idx = max(effective_max_row + 1, anchor.start_row)
        while written_data_rows < rows_budget and row_idx <= max_rows_per_workbook:
            row_values = anchor.next_row(fetch_chunk_size)
            if row_values is None:
                break
            row_cells: list[Any] = [None] * col_count
            for i, value in enumerate(row_values):
                col_idx = anchor.start_col + i
                if col_idx < plan.min_col or col_idx > max_col:
                    continue
                idx = col_idx - plan.min_col
                row_cells[idx] = _anchor_cell(ws, anchor.style, value)
            ws.append(row_cells)
            written_data_rows += 1
            row_idx += 1
    return written_data_rows


def _anchor_cell(ws, style: _AnchorStyle, value: Any) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.font = style.font
    cell.fill = style.fill
    cell.alignment = style.alignment
    cell.border = style.border
    if style.number_format:
        cell.number_format = style.number_format
    return cell


def _write_only_cell(ws, schema: CellSchema) -> WriteOnlyCell:
    value = schema["value"]
    if schema["cell_type"] == "date" and isinstance(value, str):
        value = datetime.datetime.fromisoformat(value)
    cell = WriteOnlyCell(ws, value=value)
    cell.font = _build_font(schema["font"])
    cell.fill = _build_fill(schema["fill"])
    cell.alignment = _build_alignment(schema["alignment"])
    cell.border = _build_border(schema["borders"])
    if schema["number_format"]:
        cell.number_format = schema["number_format"]
    return cell


def _apply_dimensions_write_only(ws, schema: SheetSchema) -> None:
    col_mode = schema.get("column_width_mode", "fixed")
    row_mode = schema.get("row_height_mode", "fixed")
    min_col, min_row, max_col, max_row = _parse_dims(schema["dimensions"])

    if col_mode == "fixed":
        for col_letter, width in schema["column_widths"].items():
            if width is not None:
                ws.column_dimensions[col_letter].width = width
    elif col_mode == "even":
        width = schema.get("default_column_width", 15.0)
        for col_idx in range(min_col, max_col + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    if row_mode == "fixed":
        for row_str, height in schema["row_heights"].items():
            if height is not None:
                ws.row_dimensions[int(row_str)].height = height
    elif row_mode == "even":
        height = schema.get("default_row_height", 15.0)
        for row_idx in range(min_row, max_row + 1):
            ws.row_dimensions[row_idx].height = height


def _part_path(base_output_path: str, part: int) -> str:
    p = Path(base_output_path)
    return str(p.with_name(f"{p.stem}.part{part:03d}{p.suffix}"))


def _parquet_path(base_output_path: str, anchor_id: str) -> Path:
    p = Path(base_output_path)
    return p.with_name(f"{p.stem}.{anchor_id}.parquet")


def _manifest_path(base_output_path: str) -> Path:
    p = Path(base_output_path)
    return p.with_name(f"{p.stem}.manifest.json")


def _bundle_parts_if_needed(base_output_path: str, part_paths: list[str]) -> list[str]:
    if len(part_paths) <= 1:
        return part_paths

    output = Path(base_output_path)
    zip_path = output.with_name(f"{output.stem}.zip")
    try:
        with ZipFile(zip_path, mode="w", compression=ZIP_DEFLATED) as zip_file:
            for part_path in part_paths:
                part = Path(part_path)
                zip_file.write(part, arcname=part.name)
    except Exception:
        if zip_path.exists():
            zip_path.unlink()
        raise

    for part_path in part_paths:
        Path(part_path).unlink()
    return [str(zip_path)]


def _bundle_hybrid_output(
    *,
    base_output_path: str,
    report_paths: list[str],
    parquet_sinks: list[_ParquetSink],
    manifest_path: Path,
) -> list[str]:
    output = Path(base_output_path)
    zip_path = output.with_name(f"{output.stem}.zip")
    try:
        with ZipFile(zip_path, mode="w", compression=ZIP_DEFLATED) as zip_file:
            zip_file.write(manifest_path, arcname="manifest.json")
            for report_path in report_paths:
                report = Path(report_path)
                zip_file.write(report, arcname=f"report/{report.name}")
            for sink in parquet_sinks:
                zip_file.write(sink.file_path, arcname=f"data/{sink.file_path.name}")
    except Exception:
        if zip_path.exists():
            zip_path.unlink()
        raise
    finally:
        if manifest_path.exists():
            manifest_path.unlink()
        for report_path in report_paths:
            report = Path(report_path)
            if report.exists():
                report.unlink()
        for sink in parquet_sinks:
            if sink.file_path.exists():
                sink.file_path.unlink()
    return [str(zip_path)]


# §4 Public Functions


def build_template_streaming_with_data(
    schema: WorkbookSchema,
    data: dict[str, Any],
    output_path: str,
    *,
    column_width_mode: str | None = None,
    row_height_mode: str | None = None,
    default_column_width: float | None = None,
    default_row_height: float | None = None,
    streaming_chunk_rows: int = 50_000,
    max_rows_per_workbook: int = MAX_EXCEL_ROWS,
    streaming_bundle_with_parquet: bool = False,
) -> list[str]:
    _validate_streaming_modes(schema, column_width_mode, row_height_mode)
    if max_rows_per_workbook <= 0 or max_rows_per_workbook > MAX_EXCEL_ROWS:
        raise ValueError(
            f"max_rows_per_workbook must be between 1 and {MAX_EXCEL_ROWS}, got {max_rows_per_workbook}"
        )
    if streaming_chunk_rows <= 0:
        raise ValueError(
            f"streaming_chunk_rows must be > 0, got {streaming_chunk_rows}"
        )

    parquet_sink_by_anchor_id: dict[str, _ParquetSink] = {}

    def _on_anchor_batch(anchor: _StreamingAnchor) -> Callable[[Any], None] | None:
        if not streaming_bundle_with_parquet:
            return None
        if not _is_polars_source(anchor.source):
            raise ValueError(
                "streaming_bundle_with_parquet supports only polars DataFrame/LazyFrame sources."
            )

        sink = _ParquetSink(
            anchor_id=anchor.id,
            file_path=_parquet_path(output_path, anchor.id),
            columns=[str(col) for col in anchor.columns],
        )
        parquet_sink_by_anchor_id[anchor.id] = sink
        return sink.write_batch

    plans = _prepare_streaming_plans(
        schema=schema,
        data=data,
        column_width_mode=column_width_mode,
        row_height_mode=row_height_mode,
        default_column_width=default_column_width,
        default_row_height=default_row_height,
        on_anchor_batch=_on_anchor_batch if streaming_bundle_with_parquet else None,
    )

    output_paths: list[str] = []
    part = 1
    while True:
        active_anchor = _first_active_anchor(plans)
        if active_anchor is None and part > 1:
            break

        rows_budget = (
            max_rows_per_workbook - active_anchor.start_row + 1 if active_anchor else 0
        )
        if rows_budget <= 0 and active_anchor is not None:
            raise ValueError(
                f"Anchor '{active_anchor.key}' starts at row {active_anchor.start_row}, "
                f"which exceeds max_rows_per_workbook={max_rows_per_workbook}."
            )
        wb = openpyxl.Workbook(write_only=True)
        active_written_rows = 0
        for plan in plans:
            ws = wb.create_sheet(title=plan.sheet["name"])
            _apply_dimensions_write_only(ws, plan.sheet)
            written = _write_sheet_chunk(
                ws=ws,
                plan=plan,
                active_anchor=active_anchor,
                max_rows_per_workbook=max_rows_per_workbook,
                rows_budget=rows_budget,
                fetch_chunk_size=streaming_chunk_rows,
            )
            if written > 0:
                active_written_rows = written

        final_path = _part_path(output_path, part)
        wb.save(final_path)
        output_paths.append(final_path)
        part += 1

        if active_anchor is None:
            break
        if active_written_rows == 0 and active_anchor.exhausted:
            continue

    if not output_paths:
        wb = openpyxl.Workbook(write_only=True)
        final_path = _part_path(output_path, 1)
        wb.save(final_path)
        output_paths.append(final_path)

    if not streaming_bundle_with_parquet:
        return _bundle_parts_if_needed(output_path, output_paths)

    parquet_sinks = list(parquet_sink_by_anchor_id.values())
    for sink in parquet_sinks:
        sink.finalize()

    manifest = {
        "timestamp_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "export_mode": "streaming",
        "streaming_bundle_with_parquet": True,
        "report_files": [Path(path).name for path in output_paths],
        "data_files": [
            {
                "anchor_id": sink.anchor_id,
                "path": f"data/{sink.file_path.name}",
                "rows": sink.row_count,
                "columns": sink.columns,
            }
            for sink in parquet_sinks
        ],
    }
    manifest_path = _manifest_path(output_path)
    manifest_path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return _bundle_hybrid_output(
        base_output_path=output_path,
        report_paths=output_paths,
        parquet_sinks=parquet_sinks,
        manifest_path=manifest_path,
    )


# §5 Entrypoints
