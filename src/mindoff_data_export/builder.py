from __future__ import annotations

import datetime
import functools

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)
from openpyxl.worksheet.worksheet import Worksheet

from .schema import (
    AlignmentSchema,
    CellBorders,
    CellSchema,
    FillSchema,
    FontSchema,
    SheetSchema,
    WorkbookSchema,
)
from .utils import argb_to_color, dict_to_border_side

# §1 Constants & Exceptions

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _build_sheet(ws: Worksheet, schema: SheetSchema) -> None:
    col_mode = schema.get("column_width_mode", "fixed")
    row_mode = schema.get("row_height_mode", "fixed")

    # fixed/even dimensions are applied before cells; hug runs after cell values exist.
    _apply_dimensions(ws, schema)
    _apply_cells(ws, schema)
    _apply_merges(ws, schema)

    if col_mode == "hug":
        _apply_hug_columns(ws, schema)
    if row_mode == "hug":
        _apply_hug_rows(ws, schema)


def _apply_dimensions(ws: Worksheet, schema: SheetSchema) -> None:
    col_mode = schema.get("column_width_mode", "fixed")
    row_mode = schema.get("row_height_mode", "fixed")

    if col_mode == "fixed":
        for col_letter, width in schema["column_widths"].items():
            if width is not None:
                ws.column_dimensions[col_letter].width = width
    elif col_mode == "even":
        width = schema.get("default_column_width", 15.0)
        for col_letter in _cols_in_range(schema["dimensions"]):
            ws.column_dimensions[col_letter].width = width

    if row_mode == "fixed":
        for row_str, height in schema["row_heights"].items():
            if height is not None:
                ws.row_dimensions[int(row_str)].height = height
    elif row_mode == "even":
        height = schema.get("default_row_height", 15.0)
        for row_idx in _rows_in_range(schema["dimensions"]):
            ws.row_dimensions[row_idx].height = height


def _cols_in_range(dimensions: str) -> list[str]:
    """Parse 'A1:F9' into ['A', 'B', ..., 'F']."""
    start, end = (dimensions.split(":") + [dimensions])[:2]
    min_col = column_index_from_string(coordinate_from_string(start)[0])
    max_col = column_index_from_string(coordinate_from_string(end)[0])
    return [get_column_letter(i) for i in range(min_col, max_col + 1)]


def _rows_in_range(dimensions: str) -> list[int]:
    """Parse 'A1:F9' into [1, 2, ..., 9]."""
    start, end = (dimensions.split(":") + [dimensions])[:2]
    min_row = coordinate_from_string(start)[1]
    max_row = coordinate_from_string(end)[1]
    return list(range(min_row, max_row + 1))


def _apply_hug_columns(ws: Worksheet, schema: SheetSchema) -> None:
    """Set each column width to fit its widest cell content (approximate)."""
    for col_letter in _cols_in_range(schema["dimensions"]):
        max_len = 0
        for cell in ws[col_letter]:
            if cell.value is not None:
                length = len(str(cell.value))
                if getattr(cell.font, "bold", False):
                    length = int(length * 1.2)
                max_len = max(max_len, length)
        if max_len > 0:
            ws.column_dimensions[col_letter].width = max_len + 2


def _apply_hug_rows(ws: Worksheet, schema: SheetSchema) -> None:
    """Set each row height from the largest font size in that row."""
    for row_idx in _rows_in_range(schema["dimensions"]):
        max_font_size = 11.0
        for cell in ws[row_idx]:
            size = getattr(cell.font, "size", None)
            if size:
                max_font_size = max(max_font_size, float(size))
        ws.row_dimensions[row_idx].height = max_font_size * 1.5


def _apply_cells(ws: Worksheet, schema: SheetSchema) -> None:
    for coord, cell_schema in schema["cells"].items():
        cell = ws[coord]
        _apply_cell_value(cell, cell_schema)
        _apply_cell_styles(cell, cell_schema)


def _apply_cell_value(cell, schema: CellSchema) -> None:
    if schema["cell_type"] == "date" and isinstance(schema["value"], str):
        cell.value = datetime.datetime.fromisoformat(schema["value"])
    else:
        cell.value = schema["value"]


def _apply_cell_styles(cell, schema: CellSchema) -> None:
    cell.font = _build_font(schema["font"])
    cell.fill = _build_fill(schema["fill"])
    cell.alignment = _build_alignment(schema["alignment"])
    cell.border = _build_border(schema["borders"])
    if schema["number_format"]:
        cell.number_format = schema["number_format"]


def _apply_merges(ws: Worksheet, schema: SheetSchema) -> None:
    for region_str in schema["merged_regions"]:
        ws.merge_cells(region_str)


def _freeze(d: dict) -> tuple:
    """Convert nested dicts to a stable hashable key for style caches."""
    return tuple(
        (k, _freeze(v) if isinstance(v, dict) else v) for k, v in sorted(d.items())
    )


@functools.lru_cache(maxsize=512)
def _cached_font(key: tuple) -> Font:
    d = dict(key)
    color = argb_to_color(d["color"])
    kwargs = {k: d[k] for k in ("name", "size", "bold", "italic", "underline")}
    if color is not None:
        kwargs["color"] = color
    return Font(**kwargs)


@functools.lru_cache(maxsize=512)
def _cached_fill(key: tuple) -> PatternFill:
    d = dict(key)
    if d["bg_color"] is None:
        return PatternFill()
    color = argb_to_color(d["bg_color"])
    return PatternFill(patternType="solid", fgColor=color)


@functools.lru_cache(maxsize=512)
def _cached_alignment(key: tuple) -> Alignment:
    d = dict(key)
    return Alignment(
        horizontal=d["horizontal"], vertical=d["vertical"], wrap_text=d["wrap_text"]
    )


@functools.lru_cache(maxsize=512)
def _cached_border(key: tuple) -> Border:
    sides = {k: dict(v) for k, v in key}
    return Border(
        top=dict_to_border_side(sides["top"]),
        bottom=dict_to_border_side(sides["bottom"]),
        left=dict_to_border_side(sides["left"]),
        right=dict_to_border_side(sides["right"]),
    )


def _build_font(schema: FontSchema) -> Font:
    return _cached_font(_freeze(schema))


def _build_fill(schema: FillSchema) -> PatternFill:
    return _cached_fill(_freeze(schema))


def _build_alignment(schema: AlignmentSchema) -> Alignment:
    return _cached_alignment(_freeze(schema))


def _build_border(schema: CellBorders) -> Border:
    return _cached_border(_freeze(schema))


# §4 Public Functions


def build_template(
    schema: WorkbookSchema,
    output_path: str,
    *,
    column_width_mode: str | None = None,
    row_height_mode: str | None = None,
    default_column_width: float | None = None,
    default_row_height: float | None = None,
) -> None:
    """Reconstruct a workbook from schema and write it to *output_path*."""
    overrides = {
        key: value
        for key, value in {
            "column_width_mode": column_width_mode,
            "row_height_mode": row_height_mode,
            "default_column_width": default_column_width,
            "default_row_height": default_row_height,
        }.items()
        if value is not None
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_schema in schema["sheets"]:
        if overrides:
            sheet_schema = {**sheet_schema, **overrides}
        ws = wb.create_sheet(title=sheet_schema["name"])
        _build_sheet(ws, sheet_schema)

    wb.save(output_path)


# §5 Entrypoints
