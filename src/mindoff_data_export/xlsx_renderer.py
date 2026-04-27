from __future__ import annotations

import datetime
import shutil
from pathlib import Path
from typing import Any, Iterable, Iterator
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pyarrow.parquet as pq
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Border
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)

from .builder import (
    _apply_cell_styles,
    _apply_cell_value,
    _apply_dimensions,
    _apply_hug_columns,
    _apply_hug_rows,
    _build_alignment,
    _build_border,
    _build_fill,
    _build_font,
)
from .bundle import ReportBundle, load_report_bundle
from .renderer import _infer_cell_type, _parse_dims
from .schema import CellSchema, SheetSchema

__all__ = ["export_report_bundle"]

# §1 Constants & Exceptions

MAX_EXCEL_ROWS = 1_048_576

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _coerce_bundle(bundle_or_path: ReportBundle | str) -> ReportBundle:
    if isinstance(bundle_or_path, ReportBundle):
        return bundle_or_path
    return load_report_bundle(str(bundle_or_path))


def _data_source_map(bundle: ReportBundle) -> dict[str, dict[str, Any]]:
    return {
        source["path"]: source for source in bundle.manifest.get("dataframe_sources", [])
    }


def _source_rows(
    bundle: ReportBundle, source: dict[str, Any], *, batch_size: int
) -> Iterator[tuple[Any, ...]]:
    path = source["path"]
    if source["format"] == "parquet":
        yield from _parquet_rows(_source_path(bundle, path), source["columns"], batch_size)
        return
    raise ValueError(f"Unsupported dataframe source format: {source['format']!r}")


def _source_path(bundle: ReportBundle, relative_path: str) -> Path:
    path = Path(bundle.path) / relative_path
    if not path.exists():
        raise ValueError(f"Report bundle is missing dataframe source '{relative_path}'")
    return path


def _parquet_rows(
    path: Path, columns: list[str], batch_size: int
) -> Iterator[tuple[Any, ...]]:
    parquet_file = pq.ParquetFile(path)
    for batch in parquet_file.iter_batches(batch_size=batch_size, columns=columns):
        for row in batch.to_pylist():
            yield tuple(row.get(col) for col in columns)


def _header_cells(anchor: dict[str, Any]) -> Iterable[tuple[str, CellSchema]]:
    cell = anchor["cell"]
    start_row = anchor["start_row"]
    start_col = anchor["start_col"]
    bold_font = dict(cell["font"])
    bold_font["bold"] = True
    for offset, header in enumerate(anchor["columns"]):
        coord = f"{get_column_letter(start_col + offset)}{start_row}"
        header_cell = dict(cell)
        header_cell["coordinate"] = coord
        header_cell["value"] = str(header)
        header_cell["cell_type"] = "string"
        header_cell["font"] = bold_font
        yield coord, header_cell  # type: ignore[misc]


def _content_cells(
    bundle: ReportBundle,
    source_map: dict[str, dict[str, Any]],
    anchor: dict[str, Any],
    *,
    batch_size: int,
) -> Iterator[tuple[str, CellSchema]]:
    cell = anchor["cell"]
    start_row = anchor["start_row"]
    start_col = anchor["start_col"]
    source = source_map[anchor["source"]]
    for row_offset, row in enumerate(_source_rows(bundle, source, batch_size=batch_size)):
        for col_offset, value in enumerate(row):
            coord = f"{get_column_letter(start_col + col_offset)}{start_row + row_offset}"
            content_cell = dict(cell)
            content_cell["coordinate"] = coord
            content_cell["value"] = value
            content_cell["cell_type"] = _infer_cell_type(value)
            yield coord, content_cell  # type: ignore[misc]


def _render_fidelity(
    bundle: ReportBundle,
    output_path: str,
    *,
    column_width_mode: str | None = None,
    row_height_mode: str | None = None,
    default_column_width: float | None = None,
    default_row_height: float | None = None,
) -> None:
    if bundle.manifest.get("dataframe_sources"):
        raise ValueError(
            "Fidelity XLSX export does not support file-backed dataframe sources. Use export_mode='streaming'."
        )

    source_map = _data_source_map(bundle)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for raw_sheet in bundle.report["sheets"]:
        sheet = _sheet_with_overrides(
            raw_sheet,
            column_width_mode=column_width_mode,
            row_height_mode=row_height_mode,
            default_column_width=default_column_width,
            default_row_height=default_row_height,
        )
        sheet = _expanded_sheet_from_manifest(sheet, source_map)
        ws = wb.create_sheet(title=sheet["name"])
        _apply_sheet_view(ws, sheet)
        _apply_dimensions(ws, sheet)
        for cell_schema in sheet["cells"].values():
            cell = ws[cell_schema["coordinate"]]
            _apply_cell_value(cell, cell_schema)
            _apply_cell_styles(cell, cell_schema)

        max_row = _parse_dims(sheet["dimensions"])[3]
        max_col = _parse_dims(sheet["dimensions"])[2]
        for anchor in sheet.get("dataframe_anchors", []):
            anchor_cells = (
                _header_cells(anchor)
                if anchor["placeholder_type"] == "dataframe-headers"
                else _content_cells(bundle, source_map, anchor, batch_size=50_000)
            )
            for coord, cell_schema in anchor_cells:
                cell = ws[coord]
                _apply_cell_value(cell, cell_schema)
                _apply_cell_styles(cell, cell_schema)
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)

        for region in sheet["merged_regions"]:
            ws.merge_cells(region)
        _apply_merged_region_borders(ws, sheet)

        if sheet.get("column_width_mode", "fixed") == "hug":
            _apply_hug_columns(ws, _dimensioned_sheet(sheet, max_col, max_row))
        if sheet.get("row_height_mode", "fixed") == "hug":
            _apply_hug_rows(ws, _dimensioned_sheet(sheet, max_col, max_row))

    wb.save(output_path)


def _dimensioned_sheet(
    sheet: SheetSchema, max_col: int, max_row: int
) -> SheetSchema:
    result = dict(sheet)
    min_col, min_row, _, _ = _parse_dims(sheet["dimensions"])
    result["dimensions"] = (
        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    )
    return result  # type: ignore[return-value]


def _expanded_sheet_from_manifest(
    sheet: SheetSchema, source_map: dict[str, dict[str, Any]]
) -> SheetSchema:
    _, _, max_col, max_row = _parse_dims(sheet["dimensions"])
    for anchor in sheet.get("dataframe_anchors", []):
        max_col = max(max_col, anchor["start_col"] + max(len(anchor["columns"]) - 1, 0))
        if anchor["placeholder_type"] == "dataframe-content":
            source = source_map[anchor["source"]]
            max_row = max(max_row, anchor["start_row"] + max(source.get("rows", 0) - 1, 0))
        else:
            max_row = max(max_row, anchor["start_row"])
    return _dimensioned_sheet(sheet, max_col, max_row)


def _sheet_with_overrides(
    sheet: dict[str, Any],
    *,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
) -> SheetSchema:
    result = dict(sheet)
    for key, value in {
        "column_width_mode": column_width_mode,
        "row_height_mode": row_height_mode,
        "default_column_width": default_column_width,
        "default_row_height": default_row_height,
    }.items():
        if value is not None:
            result[key] = value
    return result  # type: ignore[return-value]


def _validate_streaming(
    bundle: ReportBundle,
    source_map: dict[str, dict[str, Any]],
    column_width_mode: str | None,
    row_height_mode: str | None,
) -> None:
    for sheet in bundle.report["sheets"]:
        col_mode = column_width_mode or sheet.get("column_width_mode", "fixed")
        row_mode = row_height_mode or sheet.get("row_height_mode", "fixed")
        if col_mode == "hug" or row_mode == "hug":
            raise ValueError(
                "Streaming XLSX export does not support 'hug' sizing. Use fixed/even sizing or export_mode='fidelity'."
            )
        content_anchors = [
            anchor
            for anchor in sheet.get("dataframe_anchors", [])
            if anchor["placeholder_type"] == "dataframe-content"
        ]
        if len(content_anchors) > 1:
            raise ValueError(
                "Streaming XLSX export currently supports one dataframe-content placeholder per sheet."
            )
        _validate_streaming_merges(sheet, source_map, content_anchors)


def _validate_streaming_merges(
    sheet: SheetSchema,
    source_map: dict[str, dict[str, Any]],
    content_anchors: list[dict[str, Any]],
) -> None:
    if not sheet.get("merged_regions") or not content_anchors:
        return

    content_ranges = [
        _dataframe_content_range(anchor, source_map[anchor["source"]])
        for anchor in content_anchors
    ]
    for region in sheet["merged_regions"]:
        merge_range = CellRange(region)
        for content_range in content_ranges:
            if content_range is not None and _ranges_overlap(merge_range, content_range):
                raise ValueError(
                    "Streaming XLSX export does not support merged cells that intersect dataframe-content output. "
                    "Move the merge outside streamed rows or use export_mode='fidelity'."
                )


def _dataframe_content_range(
    anchor: dict[str, Any], source: dict[str, Any]
) -> CellRange | None:
    row_count = int(source.get("rows", 0))
    col_count = len(anchor["columns"])
    if row_count <= 0 or col_count <= 0:
        return None
    return CellRange(
        min_col=anchor["start_col"],
        min_row=anchor["start_row"],
        max_col=anchor["start_col"] + col_count - 1,
        max_row=anchor["start_row"] + row_count - 1,
    )


def _ranges_overlap(left: CellRange, right: CellRange) -> bool:
    return (
        left.min_col <= right.max_col
        and left.max_col >= right.min_col
        and left.min_row <= right.max_row
        and left.max_row >= right.min_row
    )


def _render_streaming(
    bundle: ReportBundle,
    output_path: str,
    *,
    column_width_mode: str | None,
    row_height_mode: str | None,
    default_column_width: float | None,
    default_row_height: float | None,
    streaming_chunk_rows: int,
    max_rows_per_workbook: int,
) -> list[str]:
    if streaming_chunk_rows <= 0:
        raise ValueError(
            f"streaming_chunk_rows must be greater than 0, got {streaming_chunk_rows}"
        )
    if max_rows_per_workbook <= 0 or max_rows_per_workbook > MAX_EXCEL_ROWS:
        raise ValueError(
            f"max_rows_per_workbook must be between 1 and {MAX_EXCEL_ROWS}, got {max_rows_per_workbook}"
        )
    source_map = _data_source_map(bundle)
    _validate_streaming(bundle, source_map, column_width_mode, row_height_mode)
    plans = [
        _streaming_plan(
            bundle,
            source_map,
            _sheet_with_overrides(
                sheet,
                column_width_mode=column_width_mode,
                row_height_mode=row_height_mode,
                default_column_width=default_column_width,
                default_row_height=default_row_height,
            ),
            streaming_chunk_rows=streaming_chunk_rows,
        )
        for sheet in bundle.report["sheets"]
    ]
    content_iters = [
        plan["content"]
        for plan in plans
        if plan["content"] is not None
    ]

    output_paths: list[str] = []
    part = 1
    while part == 1 or any(not item["exhausted"] for item in content_iters):
        wb = openpyxl.Workbook(write_only=True)
        wrote_any = False
        for plan in plans:
            ws = wb.create_sheet(title=plan["sheet"]["name"])
            _apply_sheet_view(ws, plan["sheet"])
            _apply_streaming_dimensions(ws, plan["sheet"])
            _apply_streaming_merges(ws, plan["sheet"])
            wrote = _write_streaming_sheet(ws, plan, max_rows_per_workbook)
            wrote_any = wrote_any or wrote
        if part > 1 and not wrote_any:
            break
        final_path = _part_path(output_path, part)
        wb.save(final_path)
        output_paths.append(final_path)
        part += 1

    return _bundle_parts_if_needed(output_path, output_paths)


def _streaming_plan(
    bundle: ReportBundle,
    source_map: dict[str, dict[str, Any]],
    sheet: SheetSchema,
    *,
    streaming_chunk_rows: int,
) -> dict[str, Any]:
    static: dict[tuple[int, int], CellSchema] = {}
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    content = None

    for cell in sheet["cells"].values():
        if _is_non_anchor_merged_cell(cell):
            continue
        col_letter, row_idx = coordinate_from_string(cell["coordinate"])
        static[(row_idx, column_index_from_string(col_letter))] = cell
    static.update(_merged_region_edge_schemas(sheet))

    for anchor in sheet.get("dataframe_anchors", []):
        if anchor["placeholder_type"] == "dataframe-headers":
            for coord, cell in _header_cells(anchor):
                row_idx = int("".join(ch for ch in coord if ch.isdigit()))
                col_idx = column_index_from_string(
                    "".join(ch for ch in coord if ch.isalpha())
                )
                static[(row_idx, col_idx)] = cell
            continue

        source = source_map[anchor["source"]]
        content = {
            "anchor": anchor,
            "rows": _source_rows(
                bundle,
                source,
                batch_size=streaming_chunk_rows,
            ),
            "exhausted": False,
        }
        max_col = max(max_col, anchor["start_col"] + max(len(anchor["columns"]) - 1, 0))

    return {
        "sheet": sheet,
        "static": static,
        "content": content,
        "min_col": min_col,
        "min_row": min_row,
        "max_col": max_col,
        "max_row": max_row,
    }


def _is_non_anchor_merged_cell(cell: CellSchema) -> bool:
    return bool(cell.get("merged")) and cell.get("merge_anchor") not in (
        None,
        cell["coordinate"],
    )


def _write_streaming_sheet(ws, plan: dict[str, Any], max_rows_per_workbook: int) -> bool:
    content = plan["content"]
    anchor = content["anchor"] if content is not None else None
    max_col = plan["max_col"]
    if anchor is not None:
        max_col = max(max_col, anchor["start_col"] + max(len(anchor["columns"]) - 1, 0))
    col_count = max_col - plan["min_col"] + 1
    wrote_content = False

    for row_idx in range(plan["min_row"], max(plan["max_row"], plan["min_row"]) + 1):
        if row_idx > max_rows_per_workbook:
            break
        row_cells = [None] * col_count
        for col_idx in range(plan["min_col"], max_col + 1):
            static = plan["static"].get((row_idx, col_idx))
            if static is not None:
                row_cells[col_idx - plan["min_col"]] = _write_only_cell(ws, static)
        if anchor is not None and row_idx >= anchor["start_row"]:
            wrote_content = _fill_streaming_row(
                ws, row_cells, plan, anchor, content
            ) or wrote_content
        ws.append(row_cells)

    if anchor is not None:
        row_idx = max(plan["max_row"] + 1, anchor["start_row"])
        while row_idx <= max_rows_per_workbook and not content["exhausted"]:
            row_cells = [None] * col_count
            wrote_content = _fill_streaming_row(
                ws, row_cells, plan, anchor, content
            ) or wrote_content
            if content["exhausted"]:
                break
            ws.append(row_cells)
            row_idx += 1
    return wrote_content


def _fill_streaming_row(
    ws, row_cells: list[Any], plan: dict[str, Any], anchor: dict[str, Any], content: dict[str, Any]
) -> bool:
    try:
        row_values = next(content["rows"])
    except StopIteration:
        content["exhausted"] = True
        return False

    style_cell = anchor["cell"]
    for offset, value in enumerate(row_values):
        col_idx = anchor["start_col"] + offset
        if col_idx < plan["min_col"] or col_idx > plan["max_col"]:
            continue
        row_cells[col_idx - plan["min_col"]] = _styled_write_only_cell(
            ws, style_cell, value
        )
    return True


def _styled_write_only_cell(ws, schema: CellSchema, value: Any) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.font = _build_font(schema["font"])
    cell.fill = _build_fill(schema["fill"])
    cell.alignment = _build_alignment(schema["alignment"])
    cell.border = _build_border(schema["borders"])
    if schema["number_format"]:
        cell.number_format = schema["number_format"]
    return cell


def _write_only_cell(ws, schema: CellSchema) -> WriteOnlyCell:
    value = schema["value"]
    if schema["cell_type"] == "date" and isinstance(value, str):
        value = datetime.datetime.fromisoformat(value)
    cell = _styled_write_only_cell(ws, schema, value)
    return cell


def _apply_merged_region_borders(ws, sheet: SheetSchema) -> None:
    for region in sheet.get("merged_regions", []):
        cell_range = CellRange(region)
        anchor = sheet["cells"].get(
            f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}"
        )
        if anchor is None:
            continue
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                border = _edge_border(anchor["borders"], cell_range, row, col)
                if border is not None:
                    ws.cell(row=row, column=col).border = border


def _merged_region_edge_schemas(
    sheet: SheetSchema,
) -> dict[tuple[int, int], CellSchema]:
    result: dict[tuple[int, int], CellSchema] = {}
    for region in sheet.get("merged_regions", []):
        cell_range = CellRange(region)
        anchor = sheet["cells"].get(
            f"{get_column_letter(cell_range.min_col)}{cell_range.min_row}"
        )
        if anchor is None:
            continue
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                border_schema = _edge_border_schema(
                    anchor["borders"], cell_range, row, col
                )
                if border_schema is None:
                    continue
                coord = f"{get_column_letter(col)}{row}"
                edge_cell = dict(anchor)
                edge_cell["coordinate"] = coord
                edge_cell["value"] = (
                    anchor["value"] if coord == anchor["coordinate"] else None
                )
                edge_cell["cell_type"] = (
                    anchor["cell_type"] if coord == anchor["coordinate"] else "empty"
                )
                edge_cell["borders"] = border_schema
                result[(row, col)] = edge_cell  # type: ignore[assignment]
    return result


def _edge_border(
    borders: dict[str, dict[str, Any]], cell_range: CellRange, row: int, col: int
) -> Border | None:
    schema = _edge_border_schema(borders, cell_range, row, col)
    return _build_border(schema) if schema is not None else None


def _edge_border_schema(
    borders: dict[str, dict[str, Any]], cell_range: CellRange, row: int, col: int
) -> dict[str, dict[str, Any]] | None:
    empty = {"style": None, "color": None}
    result = {
        "top": dict(empty),
        "bottom": dict(empty),
        "left": dict(empty),
        "right": dict(empty),
    }
    if row == cell_range.min_row:
        result["top"] = dict(borders["top"])
    if row == cell_range.max_row:
        result["bottom"] = dict(borders["bottom"])
    if col == cell_range.min_col:
        result["left"] = dict(borders["left"])
    if col == cell_range.max_col:
        result["right"] = dict(borders["right"])
    return result if any(side.get("style") for side in result.values()) else None


def _apply_streaming_merges(ws, schema: SheetSchema) -> None:
    for region in schema.get("merged_regions", []):
        ws.merged_cells.add(CellRange(region))


def _apply_sheet_view(ws, schema: SheetSchema) -> None:
    ws.sheet_view.showGridLines = schema.get("show_gridlines", True)


def _apply_streaming_dimensions(ws, schema: SheetSchema) -> None:
    col_mode = schema.get("column_width_mode", "fixed")
    row_mode = schema.get("row_height_mode", "fixed")
    min_col, min_row, max_col, max_row = _parse_dims(schema["dimensions"])

    if col_mode == "fixed":
        for col_letter, width in schema["column_widths"].items():
            if width is not None:
                ws.column_dimensions[col_letter].width = width
    elif col_mode == "even":
        width = schema.get("default_column_width", 15.0)
        for col_idx in range(min_col, max_col + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    if row_mode == "fixed":
        for row_str, height in schema["row_heights"].items():
            if height is not None:
                ws.row_dimensions[int(row_str)].height = height
    elif row_mode == "even":
        height = schema.get("default_row_height", 15.0)
        for row_idx in range(min_row, max_row + 1):
            ws.row_dimensions[row_idx].height = height


def _part_path(base_output_path: str, part: int) -> str:
    p = Path(base_output_path)
    return str(p.with_name(f"{p.stem}.part{part:03d}{p.suffix}"))


def _bundle_parts_if_needed(base_output_path: str, part_paths: list[str]) -> list[str]:
    if len(part_paths) <= 1:
        return part_paths

    output = Path(base_output_path)
    zip_path = output.with_name(f"{output.stem}.zip")
    with ZipFile(zip_path, mode="w", compression=ZIP_DEFLATED) as zip_file:
        for part_path in part_paths:
            part = Path(part_path)
            zip_file.write(part, arcname=part.name)

    for part_path in part_paths:
        Path(part_path).unlink()
    return [str(zip_path)]


def _delete_bundle_dir(bundle: ReportBundle) -> None:
    path = Path(bundle.path)
    if not path.is_dir():
        return
    if not (path / "manifest.json").exists() or not (path / "report.json").exists():
        raise ValueError(f"Refusing to delete malformed report bundle directory: {path}")
    shutil.rmtree(path)


# §4 Public Functions


def export_report_bundle(
    bundle_or_path: ReportBundle | str,
    output_path: str,
    *,
    format: str = "xlsx",
    export_mode: str = "fidelity",
    column_width_mode: str | None = None,
    row_height_mode: str | None = None,
    default_column_width: float | None = None,
    default_row_height: float | None = None,
    streaming_chunk_rows: int = 50_000,
    max_rows_per_workbook: int = MAX_EXCEL_ROWS,
    auto_delete_bundle: bool = False,
    **options: Any,
) -> None | list[str]:
    """Render a ReportBundle to the requested output format."""
    if format == "image":
        raise NotImplementedError(
            "ReportBundle export format 'image' is reserved but not implemented in v1."
        )
    if format not in {"xlsx", "pdf"}:
        raise ValueError(
            f"Unsupported report export format '{format}'. Expected 'xlsx', 'pdf', or 'image'."
        )

    bundle = _coerce_bundle(bundle_or_path)
    result: None | list[str]
    if format == "pdf":
        if export_mode != "fidelity":
            raise ValueError(
                f"PDF export does not support export_mode '{export_mode}'. PDF output paginates automatically."
            )
        from .pdf_renderer import export_report_bundle as _export_pdf_report_bundle

        _export_pdf_report_bundle(
            bundle,
            output_path,
            column_width_mode=column_width_mode,
            row_height_mode=row_height_mode,
            default_column_width=default_column_width,
            default_row_height=default_row_height,
            page_size=options.get("page_size", "A4"),
            orientation=options.get("orientation", "portrait"),
            margin=options.get("margin", 36),
            streaming_chunk_rows=streaming_chunk_rows,
            fonts=options.get("fonts"),
        )
        result = None
    elif export_mode == "streaming":
        result = _render_streaming(
            bundle,
            output_path,
            column_width_mode=column_width_mode,
            row_height_mode=row_height_mode,
            default_column_width=default_column_width,
            default_row_height=default_row_height,
            streaming_chunk_rows=streaming_chunk_rows,
            max_rows_per_workbook=max_rows_per_workbook,
        )
    elif export_mode == "fidelity":
        _render_fidelity(
            bundle,
            output_path,
            column_width_mode=column_width_mode,
            row_height_mode=row_height_mode,
            default_column_width=default_column_width,
            default_row_height=default_row_height,
        )
        result = None
    else:
        raise ValueError(
            f"Unsupported export_mode '{export_mode}'. Expected 'fidelity' or 'streaming'."
        )

    if auto_delete_bundle:
        _delete_bundle_dir(bundle)
    return result


# §5 Entrypoints
