from __future__ import annotations

import datetime

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .schema import (
    AlignmentSchema,
    BorderSide,
    CellBorders,
    CellSchema,
    CellType,
    FillSchema,
    FontSchema,
    SheetSchema,
    WorkbookSchema,
)
from .utils import border_side_to_dict, normalize_color

# §1 Types

# §2 Constants

# Shared defaults used for merged-cell stubs (never mutated downstream).
_EMPTY_BORDER_SIDE: BorderSide = {"style": None, "color": None}
_EMPTY_FONT: FontSchema = {
    "name": None,
    "size": None,
    "bold": False,
    "italic": False,
    "underline": None,
    "color": None,
}
_EMPTY_ALIGNMENT: AlignmentSchema = {"horizontal": None, "vertical": None, "wrap_text": False}
_EMPTY_BORDERS: CellBorders = {
    "top": _EMPTY_BORDER_SIDE,
    "bottom": _EMPTY_BORDER_SIDE,
    "left": _EMPTY_BORDER_SIDE,
    "right": _EMPTY_BORDER_SIDE,
}

# §3 Private Helpers


def _extract_sheet(ws: Worksheet) -> SheetSchema:
    merged_regions: list[str] = [str(mr) for mr in ws.merged_cells.ranges]
    merge_map = _build_merge_map(ws)

    cells: dict[str, CellSchema] = {}
    for row in ws.iter_rows():
        for cell in row:
            cells[cell.coordinate] = _extract_cell(cell, merge_map)

    column_widths = {
        col: ws.column_dimensions[col].width
        for col in ws.column_dimensions
        if ws.column_dimensions[col].width is not None
    }
    row_heights = {
        str(row_idx): ws.row_dimensions[row_idx].height
        for row_idx in ws.row_dimensions
        if ws.row_dimensions[row_idx].height is not None
    }

    return {
        "name": ws.title,
        "dimensions": ws.dimensions,
        "merged_regions": merged_regions,
        "column_widths": column_widths,
        "row_heights": row_heights,
        "cells": cells,
    }


def _build_merge_map(ws: Worksheet) -> dict[str, str]:
    """Map each coordinate inside merged regions to its anchor coordinate."""
    result: dict[str, str] = {}
    for merged_range in ws.merged_cells.ranges:
        anchor = merged_range.coord.split(":")[0]
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                result[f"{get_column_letter(col)}{row}"] = anchor
    return result


def _extract_cell(cell: Cell | MergedCell, merge_map: dict[str, str]) -> CellSchema:
    is_merged = cell.coordinate in merge_map
    anchor = merge_map.get(cell.coordinate)

    if isinstance(cell, MergedCell):
        return {
            "coordinate": cell.coordinate,
            "value": None,
            "cell_type": "empty",
            "number_format": None,
            "font": _EMPTY_FONT,
            "fill": {"bg_color": None},
            "alignment": _EMPTY_ALIGNMENT,
            "borders": _EMPTY_BORDERS,
            "merged": True,
            "merge_anchor": anchor,
        }

    return {
        "coordinate": cell.coordinate,
        "value": _serialize_value(cell.value),
        "cell_type": _infer_cell_type(cell),
        "number_format": cell.number_format if cell.number_format not in (None, "General") else None,
        "font": _extract_font(cell),
        "fill": _extract_fill(cell),
        "alignment": _extract_alignment(cell),
        "borders": _extract_borders(cell),
        "merged": is_merged,
        "merge_anchor": anchor,
    }


def _infer_cell_type(cell: Cell) -> CellType:
    if cell.value is None:
        return "empty"
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "n":
        return "number"
    if cell.data_type == "d" or isinstance(cell.value, (datetime.datetime, datetime.date)):
        return "date"
    return "string"


def _serialize_value(value: object) -> object:
    if isinstance(value, datetime.datetime):
        return value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return value


def _extract_font(cell: Cell) -> FontSchema:
    f = cell.font
    return {
        "name": f.name,
        "size": float(f.size) if f.size is not None else None,
        "bold": bool(f.bold),
        "italic": bool(f.italic),
        "underline": f.underline,
        "color": normalize_color(f.color),
    }


def _extract_fill(cell: Cell) -> FillSchema:
    fill = cell.fill
    fill_type = getattr(fill, "patternType", None) or getattr(fill, "fill_type", None)
    if fill_type == "solid":
        return {"bg_color": normalize_color(fill.fgColor)}
    return {"bg_color": None}


def _extract_alignment(cell: Cell) -> AlignmentSchema:
    a = cell.alignment
    return {
        "horizontal": a.horizontal,
        "vertical": a.vertical,
        "wrap_text": bool(a.wrap_text),
    }


def _extract_borders(cell: Cell) -> CellBorders:
    b = cell.border
    return {
        "top": border_side_to_dict(b.top),
        "bottom": border_side_to_dict(b.bottom),
        "left": border_side_to_dict(b.left),
        "right": border_side_to_dict(b.right),
    }


# §4 Public API


def extract_template(path: str) -> WorkbookSchema:
    """Load .xlsx at *path* and return a WorkbookSchema dict."""
    wb = openpyxl.load_workbook(path, data_only=False)
    return {"sheets": [_extract_sheet(ws) for ws in wb.worksheets]}
