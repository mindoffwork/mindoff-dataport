from __future__ import annotations

import datetime
import re
from typing import Any

from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
)

from .schema import (
    AlignmentSchema,
    CellBorders,
    FillSchema,
    FontSchema,
    SheetSchema,
    WorkbookSchema,
)

# §1 Constants & Exceptions

# Matches {{key:type}} where type may include hyphens (for example dataframe-headers).
PLACEHOLDER_RE = re.compile(r"\{\{(\w+):([\w-]+)\}\}")
SHEET_NAME_PLACEHOLDER_RE = re.compile(r"\{\{(\w+)\}\}")

_DATAFRAME_TYPES = frozenset(["dataframe-headers", "dataframe-content"])
_SCALAR_TYPES = frozenset(["string", "number", "int", "float", "date", "boolean"])
_ALL_TYPES = _SCALAR_TYPES | _DATAFRAME_TYPES

_EMPTY_BORDER_SIDE = {"style": None, "color": None}
_DEFAULT_FONT: FontSchema = {
    "name": "Calibri",
    "size": 11.0,
    "bold": False,
    "italic": False,
    "underline": None,
    "color": None,
}
_DEFAULT_FILL: FillSchema = {"bg_color": None}
_DEFAULT_ALIGNMENT: AlignmentSchema = {
    "horizontal": None,
    "vertical": None,
    "wrap_text": False,
}
_DEFAULT_BORDERS: CellBorders = {
    "top": _EMPTY_BORDER_SIDE,
    "bottom": _EMPTY_BORDER_SIDE,
    "left": _EMPTY_BORDER_SIDE,
    "right": _EMPTY_BORDER_SIDE,
}

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _validate_data(
    placeholders: dict[str, str], data: dict[str, Any], scope_label: str
) -> None:
    for key, expected_type in placeholders.items():
        if key not in data:
            raise KeyError(
                f"Sheet '{scope_label}' requires '{key}' (type: {expected_type}) "
                "but it was not provided in data"
            )
        _check_type(key, data[key], expected_type)


def _check_type(key: str, value: Any, expected: str) -> None:
    if expected == "dataframe-content":
        _assert_dataframe(key, value)
        return
    if expected == "dataframe-headers":
        _assert_headers_input(key, value)
        return

    type_map: dict[str, tuple] = {
        "string": (str,),
        "number": (int, float),
        "int": (int,),
        "float": (float,),
        "boolean": (bool,),
        "date": (str, datetime.datetime, datetime.date),
    }
    allowed = type_map.get(expected)
    if allowed and not isinstance(value, allowed):
        raise TypeError(
            f"'{key}' expected type '{expected}', got {type(value).__name__}"
        )


def _assert_dataframe(key: str, value: Any) -> None:
    module = getattr(type(value), "__module__", "") or ""
    if "polars" in module or "pandas" in module:
        return
    if _is_parquet_source(value):
        return
    raise TypeError(
        f"'{key}' expected a polars/pandas DataFrame, LazyFrame, or ParquetSource, got {type(value).__name__}"
    )


def _assert_headers_input(key: str, value: Any) -> None:
    module = getattr(type(value), "__module__", "") or ""
    if "polars" in module or "pandas" in module:
        return
    if _is_parquet_source(value):
        return
    if isinstance(value, list):
        return
    raise TypeError(
        f"'{key}' expected a polars/pandas DataFrame/LazyFrame, ParquetSource, or list of header values, got {type(value).__name__}"
    )


def _is_parquet_source(value: Any) -> bool:
    return (
        getattr(type(value), "__module__", "") == "mindoff_data_export.bundle"
        and type(value).__qualname__ == "ParquetSource"
    )


def _merge_placeholder_types(
    base: dict[str, str], incoming: dict[str, str], scope_label: str
) -> dict[str, str]:
    merged: dict[str, str] = dict(base)
    for key, type_ in incoming.items():
        existing = merged.get(key)
        if existing is not None and existing != type_:
            raise ValueError(
                f"Conflicting placeholder types for '{key}' in scope '{scope_label}': "
                f"'{existing}' vs '{type_}'"
            )
        merged[key] = type_
    return merged


def _collect_sheet_placeholders(sheet: SheetSchema) -> dict[str, str]:
    found: dict[str, str] = {}
    scope_label = sheet["name"]
    for cell_schema in sheet["cells"].values():
        value = cell_schema.get("value")
        if not isinstance(value, str):
            continue
        for match in PLACEHOLDER_RE.finditer(value):
            key, placeholder_type = match.group(1), match.group(2)
            if placeholder_type not in _ALL_TYPES:
                continue
            existing = found.get(key)
            if existing is not None and existing != placeholder_type:
                raise ValueError(
                    f"Conflicting placeholder types for '{key}' in scope '{scope_label}': "
                    f"'{existing}' vs '{placeholder_type}'"
                )
            found[key] = placeholder_type
    return found


def _sheet_name_placeholder(name: str) -> str | None:
    match = SHEET_NAME_PLACEHOLDER_RE.fullmatch(name)
    if not match:
        return None
    return match.group(1)


def _require_dict_payload(
    *, data: dict[str, Any], key: str, error_label: str
) -> dict[str, Any]:
    if key not in data:
        raise KeyError(
            f"Template requires {error_label} '{key}' but it was not provided in data"
        )
    payload = data[key]
    if not isinstance(payload, dict):
        raise TypeError(
            f"Data for {error_label} '{key}' must be an object/dict, got {type(payload).__name__}"
        )
    return payload


def _resolve_sheet_payloads(
    schema: WorkbookSchema, data: dict[str, Any]
) -> list[tuple[SheetSchema, str, dict[str, Any], dict[str, str]]]:
    if not isinstance(data, dict):
        raise TypeError(f"Data must be an object/dict, got {type(data).__name__}")

    resolved: list[tuple[SheetSchema, str, dict[str, Any], dict[str, str]]] = []
    output_names: set[str] = set()

    for sheet in schema["sheets"]:
        placeholders = _collect_sheet_placeholders(sheet)
        dynamic_key = _sheet_name_placeholder(sheet["name"])

        if dynamic_key is None:
            sheet_payload = _require_dict_payload(
                data=data, key=sheet["name"], error_label="sheet"
            )
            _validate_data(placeholders, sheet_payload, sheet["name"])
            if sheet["name"] in output_names:
                raise ValueError(
                    f"Duplicate output sheet name '{sheet['name']}' is not allowed"
                )
            output_names.add(sheet["name"])
            resolved.append((sheet, sheet["name"], sheet_payload, placeholders))
            continue

        group_payload = _require_dict_payload(
            data=data,
            key=dynamic_key,
            error_label="dynamic sheet group",
        )
        for output_name, sheet_payload in group_payload.items():
            if not isinstance(output_name, str):
                raise TypeError(
                    f"Dynamic sheet names under '{dynamic_key}' must be strings, "
                    f"got {type(output_name).__name__}"
                )
            if not isinstance(sheet_payload, dict):
                raise TypeError(
                    f"Data for dynamic sheet '{output_name}' under '{dynamic_key}' "
                    f"must be an object/dict, got {type(sheet_payload).__name__}"
                )
            _validate_data(placeholders, sheet_payload, output_name)
            if output_name in output_names:
                raise ValueError(
                    f"Duplicate output sheet name '{output_name}' is not allowed"
                )
            output_names.add(output_name)
            resolved.append((sheet, output_name, sheet_payload, placeholders))

    if not resolved:
        raise ValueError("No output sheets were resolved from template + input data")

    return resolved


def _substitute_scalars(value: str, data: dict[str, Any]) -> Any:
    """Replace scalar placeholders; whole-cell placeholders return raw values."""
    full_match = PLACEHOLDER_RE.fullmatch(value.strip())
    if full_match:
        key, placeholder_type = full_match.group(1), full_match.group(2)
        if placeholder_type in _SCALAR_TYPES and key in data:
            return data[key]

    def replacer(match: re.Match[str]) -> str:
        key, placeholder_type = match.group(1), match.group(2)
        if placeholder_type not in _SCALAR_TYPES or key not in data:
            return match.group(0)
        return str(data[key])

    return PLACEHOLDER_RE.sub(replacer, value)


def _infer_cell_type(value: Any) -> str:
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "string"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime.datetime, datetime.date)):
        return "date"
    return "string"


def _to_headers(df_or_headers: Any) -> list[str]:
    """Return header names from DataFrame/LazyFrame or list input."""
    module = getattr(type(df_or_headers), "__module__", "") or ""
    qualname = type(df_or_headers).__qualname__

    if "polars" in module:
        if qualname == "LazyFrame":
            df_or_headers = df_or_headers.collect()
        return [str(col) for col in df_or_headers.columns]

    if "pandas" in module and "DataFrame" in qualname:
        return [str(col) for col in df_or_headers.columns]

    if _is_parquet_source(df_or_headers):
        if df_or_headers.columns is not None:
            return [str(col) for col in df_or_headers.columns]
        import pyarrow.parquet as pq

        parquet_file = pq.ParquetFile(df_or_headers.path)
        return [str(col) for col in parquet_file.schema_arrow.names]

    if isinstance(df_or_headers, list):
        return [str(col) for col in df_or_headers]

    raise TypeError(
        f"Expected a polars/pandas DataFrame/LazyFrame, ParquetSource, or list of headers, got {type(df_or_headers).__name__}"
    )


def _parse_dims(dimensions: str) -> tuple[int, int, int, int]:
    """Parse 'A1:F9' to (min_col, min_row, max_col, max_row)."""
    parts = dimensions.split(":")
    start = parts[0]
    end = parts[1] if len(parts) > 1 else parts[0]
    start_col_letter, start_row = coordinate_from_string(start)
    end_col_letter, end_row = coordinate_from_string(end)
    return (
        column_index_from_string(start_col_letter),
        start_row,
        column_index_from_string(end_col_letter),
        end_row,
    )


# §4 Public Functions


def get_template_inputs(schema: WorkbookSchema) -> dict[str, Any]:
    """Return sheet-scoped required inputs including dynamic sheet groups."""
    found: dict[str, Any] = {}
    for sheet in schema["sheets"]:
        placeholders = _collect_sheet_placeholders(sheet)
        dynamic_key = _sheet_name_placeholder(sheet["name"])
        if dynamic_key is None:
            found[sheet["name"]] = placeholders
            continue

        existing = found.get(dynamic_key)
        if existing is None:
            found[dynamic_key] = {"*": placeholders}
            continue

        if not isinstance(existing, dict) or "*" not in existing:
            raise ValueError(
                f"Template contains incompatible scopes for key '{dynamic_key}' in sheet inputs"
            )
        found[dynamic_key] = {
            "*": _merge_placeholder_types(existing["*"], placeholders, dynamic_key)
        }

    return found

# §5 Entrypoints
