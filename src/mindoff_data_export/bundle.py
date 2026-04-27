from __future__ import annotations

import datetime
import json
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pyarrow.parquet as pq
from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
)

from .renderer import (
    PLACEHOLDER_RE,
    _DATAFRAME_TYPES,
    _SCALAR_TYPES,
    _infer_cell_type,
    _parse_dims,
    _resolve_sheet_payloads,
    _substitute_scalars,
    _to_headers,
    get_template_inputs,
)
from .schema import CellSchema, SheetSchema, WorkbookSchema

__all__ = [
    "ParquetSource",
    "ReportBundle",
    "compile_report_bundle",
    "load_report_bundle",
    "parquet_source",
]

# Â§1 Constants & Exceptions

BUNDLE_VERSION = "1.0"

# Â§2 Classes and Sub Classes


@dataclass(frozen=True)
class ParquetSource:
    """Disk-backed dataframe input for larger-than-RAM exports."""

    path: str
    columns: list[str] | None = None
    row_count: int | None = None


@dataclass(frozen=True)
class ReportBundle:
    """Canonical intermediate report artifact, backed by a directory."""

    manifest: dict[str, Any]
    report: dict[str, Any]
    path: str

    def write(self, bundle_path: str) -> None:
        _copy_bundle_dir(self, bundle_path)

    @classmethod
    def load(cls, bundle_path: str) -> "ReportBundle":
        return load_report_bundle(bundle_path)


# Â§3 Private Helper Functions


def _copy_bundle_dir(bundle: ReportBundle, bundle_path: str) -> None:
    source = Path(bundle.path).resolve()
    target = Path(bundle_path).resolve()
    if source == target:
        _write_bundle_metadata(bundle)
        return
    if target.exists() and not target.is_dir():
        raise ValueError(f"Report bundle path must be a directory: {bundle_path}")
    shutil.copytree(source, target, dirs_exist_ok=True)


def _write_bundle_metadata(bundle: ReportBundle) -> None:
    path = Path(bundle.path)
    path.mkdir(parents=True, exist_ok=True)
    (path / "manifest.json").write_text(
        json.dumps(bundle.manifest, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    (path / "report.json").write_text(
        json.dumps(bundle.report, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def _prepare_bundle_dir(bundle_path: str | None) -> Path:
    if bundle_path is None:
        path = Path(tempfile.mkdtemp(prefix="mindoff_report_bundle_"))
    else:
        path = Path(bundle_path)
    if path.exists() and not path.is_dir():
        raise ValueError(f"Report bundle path must be a directory: {path}")
    path.mkdir(parents=True, exist_ok=True)
    (path / "data").mkdir(exist_ok=True)
    return path


def _resolved_parquet_path(source: ParquetSource) -> Path:
    path = Path(source.path)
    if not path.exists():
        raise FileNotFoundError(f"Parquet source not found: {source.path}")
    if not path.is_file():
        raise ValueError(f"Parquet source must be a file: {source.path}")
    return path.resolve()


def _parquet_metadata(source: ParquetSource) -> tuple[list[str], int]:
    path = _resolved_parquet_path(source)
    parquet_file = pq.ParquetFile(path)
    available = [str(name) for name in parquet_file.schema_arrow.names]
    if source.columns is None:
        columns = available
    else:
        columns = [str(col) for col in source.columns]
        missing = [col for col in columns if col not in available]
        if missing:
            raise ValueError(
                f"Parquet source is missing requested columns: {', '.join(missing)}"
            )
    row_count = (
        source.row_count
        if source.row_count is not None
        else parquet_file.metadata.num_rows
    )
    if row_count < 0:
        raise ValueError(f"Parquet source row_count must be non-negative: {row_count}")
    return columns, int(row_count)


def _source_from_dataframe(value: Any) -> tuple[list[str], int]:
    module = getattr(type(value), "__module__", "") or ""
    qualname = type(value).__qualname__

    if "polars" in module:
        if qualname == "LazyFrame":
            value = value.collect()
        return [str(col) for col in value.columns], int(value.height)

    if "pandas" in module and "DataFrame" in qualname:
        return [str(col) for col in value.columns], int(len(value))

    raise TypeError(
        f"Expected a polars/pandas DataFrame/LazyFrame or ParquetSource, got {type(value).__name__}"
    )


def _write_source_file(
    *,
    value: Any,
    source_id: str,
    bundle_dir: Path,
    source_cache: dict[Path, dict[str, Any]],
) -> dict[str, Any]:
    if isinstance(value, ParquetSource):
        return _copy_parquet_source(
            value=value,
            source_id=source_id,
            bundle_dir=bundle_dir,
            source_cache=source_cache,
        )
    return _write_dataframe_source(value=value, source_id=source_id, bundle_dir=bundle_dir)


def _copy_parquet_source(
    *,
    value: ParquetSource,
    source_id: str,
    bundle_dir: Path,
    source_cache: dict[Path, dict[str, Any]],
) -> dict[str, Any]:
    input_path = _resolved_parquet_path(value)
    cached = source_cache.get(input_path)
    if cached is not None:
        return cached

    columns, row_count = _parquet_metadata(value)
    rel_path = f"data/{source_id}.parquet"
    output_path = bundle_dir / rel_path
    if input_path != output_path.resolve():
        with input_path.open("rb") as src, output_path.open("wb") as dst:
            shutil.copyfileobj(src, dst, length=1024 * 1024)

    record = {
        "id": source_id,
        "path": rel_path,
        "format": "parquet",
        "columns": columns,
        "rows": row_count,
        "file_backed": True,
    }
    source_cache[input_path] = record
    return record


def _write_dataframe_source(
    *, value: Any, source_id: str, bundle_dir: Path
) -> dict[str, Any]:
    columns, row_count = _source_from_dataframe(value)
    rel_path = f"data/{source_id}.parquet"
    output_path = bundle_dir / rel_path
    module = getattr(type(value), "__module__", "") or ""
    qualname = type(value).__qualname__

    if "polars" in module:
        if qualname == "LazyFrame":
            value = value.collect()
        value.write_parquet(output_path)
    elif "pandas" in module and "DataFrame" in qualname:
        value.to_parquet(output_path, index=False, engine="pyarrow")
    else:
        raise TypeError(
            f"Expected a polars/pandas DataFrame/LazyFrame, got {type(value).__name__}"
        )

    return {
        "id": source_id,
        "path": rel_path,
        "format": "parquet",
        "columns": columns,
        "rows": row_count,
        "file_backed": True,
    }


def _schema_value(value: Any) -> Any:
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return value


def _unique_id(used: set[str], raw: str) -> str:
    safe = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in raw)
    safe = safe.strip("._-") or "source"
    candidate = safe
    idx = 2
    while candidate in used:
        candidate = f"{safe}_{idx}"
        idx += 1
    used.add(candidate)
    return candidate


def _compile_sheet(
    *,
    sheet: SheetSchema,
    output_name: str,
    sheet_data: dict[str, Any],
    bundle_dir: Path,
    data_sources: list[dict[str, Any]],
    used_ids: set[str],
    source_cache: dict[Path, dict[str, Any]],
) -> dict[str, Any]:
    min_col, min_row, max_col, max_row = _parse_dims(sheet["dimensions"])
    static_cells: dict[str, CellSchema] = {}
    anchors: list[dict[str, Any]] = []

    for coord, cell in sheet["cells"].items():
        value = cell.get("value")
        if not isinstance(value, str):
            static_cells[coord] = cell
            continue

        full_match = PLACEHOLDER_RE.fullmatch(value.strip())
        if (
            full_match
            and full_match.group(2) in _DATAFRAME_TYPES
            and full_match.group(1) in sheet_data
        ):
            key = full_match.group(1)
            anchor_type = full_match.group(2)
            columns = _to_headers(sheet_data[key])
            anchor_id = _unique_id(used_ids, f"{output_name}__{key}__{anchor_type}")
            source_record: dict[str, Any] | None = None
            if anchor_type == "dataframe-content":
                source_record = _write_source_file(
                    value=sheet_data[key],
                    source_id=anchor_id,
                    bundle_dir=bundle_dir,
                    source_cache=source_cache,
                )
                if source_record not in data_sources:
                    data_sources.append(source_record)

            col_letter, row_idx = coordinate_from_string(coord)
            col_idx = column_index_from_string(col_letter)
            max_col = max(max_col, col_idx + max(len(columns) - 1, 0))
            anchors.append(
                {
                    "id": anchor_id,
                    "key": key,
                    "placeholder_type": anchor_type,
                    "coordinate": coord,
                    "start_row": row_idx,
                    "start_col": col_idx,
                    "columns": columns,
                    "source": source_record["path"] if source_record else None,
                    "source_format": source_record["format"] if source_record else None,
                    "cell": cell,
                }
            )
            continue

        new_value = _substitute_scalars(value, sheet_data)
        new_cell: dict[str, Any] = dict(cell)
        new_cell["value"] = _schema_value(new_value)
        if full_match and full_match.group(2) in _SCALAR_TYPES:
            new_cell["cell_type"] = _infer_cell_type(new_value)
        static_cells[coord] = new_cell  # type: ignore[assignment]

    dimensions = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    result = dict(sheet)
    result["name"] = output_name
    result["dimensions"] = dimensions
    result["cells"] = static_cells
    result["dataframe_anchors"] = anchors
    return result


# Â§4 Public Functions


def parquet_source(
    path: str, *, columns: list[str] | None = None, row_count: int | None = None
) -> ParquetSource:
    return ParquetSource(path=path, columns=columns, row_count=row_count)


def compile_report_bundle(
    template: WorkbookSchema,
    data: dict[str, Any],
    bundle_path: str | None = None,
) -> ReportBundle:
    """Validate *data* against *template* and produce a directory ReportBundle."""
    bundle_dir = _prepare_bundle_dir(bundle_path)
    data_sources: list[dict[str, Any]] = []
    used_ids: set[str] = set()
    source_cache: dict[Path, dict[str, Any]] = {}

    sheets = [
        _compile_sheet(
            sheet=sheet,
            output_name=output_name,
            sheet_data=sheet_data,
            bundle_dir=bundle_dir,
            data_sources=data_sources,
            used_ids=used_ids,
            source_cache=source_cache,
        )
        for sheet, output_name, sheet_data, _ in _resolve_sheet_payloads(template, data)
    ]

    report = {"version": BUNDLE_VERSION, "sheets": sheets, "assets": []}
    manifest = {
        "version": BUNDLE_VERSION,
        "bundle_format": "directory",
        "created_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "inputs": get_template_inputs(template),
        "sheets": [
            {
                "name": sheet["name"],
                "dataframe_anchors": [
                    anchor["id"] for anchor in sheet.get("dataframe_anchors", [])
                ],
            }
            for sheet in sheets
        ],
        "placeholder_resolution": {
            "scalars": "resolved in report.json",
            "dataframes": "stored as dataframe anchors pointing to data/*.parquet",
        },
        "dataframe_sources": data_sources,
        "assets": [],
        "output_capabilities": {"xlsx": True, "pdf": True, "image": False},
    }
    bundle = ReportBundle(manifest=manifest, report=report, path=str(bundle_dir))
    _write_bundle_metadata(bundle)
    return bundle


def load_report_bundle(bundle_path: str) -> ReportBundle:
    path = Path(bundle_path)
    if not path.exists():
        raise FileNotFoundError(f"Report bundle not found: {bundle_path}")
    if not path.is_dir():
        raise ValueError(f"Report bundle must be a directory: {bundle_path}")

    manifest_path = path / "manifest.json"
    report_path = path / "report.json"
    if not manifest_path.exists() or not report_path.exists():
        raise ValueError(
            "Malformed report bundle: expected manifest.json and report.json"
        )

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    report = json.loads(report_path.read_text(encoding="utf-8"))
    if manifest.get("version") != BUNDLE_VERSION:
        raise ValueError(
            f"Unsupported report bundle version: {manifest.get('version')!r}"
        )
    if manifest.get("bundle_format") != "directory":
        raise ValueError(
            f"Unsupported report bundle format: {manifest.get('bundle_format')!r}"
        )
    return ReportBundle(manifest=manifest, report=report, path=str(path))


# Â§5 Entrypoints
