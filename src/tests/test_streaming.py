from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

from mindoff_dataport import mo_dataport

# §1 Constants & Exceptions

polars = pytest.importorskip("polars", reason="polars not installed")

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _cell(coord: str, value):
    return {
        "coordinate": coord,
        "value": value,
        "cell_type": "string",
        "number_format": None,
        "font": {
            "name": "Calibri",
            "size": 11.0,
            "bold": False,
            "italic": False,
            "underline": None,
            "color": None,
        },
        "fill": {"bg_color": None},
        "alignment": {"horizontal": None, "vertical": None, "wrap_text": False},
        "borders": {
            "top": {"style": None, "color": None},
            "bottom": {"style": None, "color": None},
            "left": {"style": None, "color": None},
            "right": {"style": None, "color": None},
        },
        "merged": False,
        "merge_anchor": None,
    }


def _schema(
    cells: dict[str, dict], *, dims: str = "A1:C3", merges: list[str] | None = None
):
    return {
        "sheets": [
            {
                "name": "Sheet1",
                "dimensions": dims,
                "merged_regions": merges or [],
                "column_widths": {},
                "row_heights": {},
                "cells": cells,
            }
        ]
    }


def _export_streaming(schema, data, managed_tmp_dir: Path, **options):
    bundle = mo_dataport.compile(schema, data)
    out = managed_tmp_dir / "filled.xlsx"
    return mo_dataport.export(bundle, str(out), export_mode="streaming", **options)


# §4 Public Functions


def test_streaming_split_outputs_and_names(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1, 2, 3, 4, 5]})

    paths = _export_streaming(
        schema,
        {"Sheet1": {"rows": df}},
        managed_tmp_dir,
        max_rows_per_workbook=3,
    )

    assert isinstance(paths, list)
    assert len(paths) == 1
    zip_path = Path(paths[0])
    assert zip_path.name == "filled.zip"
    assert zip_path.exists()

    with ZipFile(zip_path) as zip_file:
        assert zip_file.namelist() == ["filled.part001.xlsx", "filled.part002.xlsx"]
        zip_file.extractall(path=managed_tmp_dir)

    wb1 = openpyxl.load_workbook(managed_tmp_dir / "filled.part001.xlsx", data_only=True)
    wb2 = openpyxl.load_workbook(managed_tmp_dir / "filled.part002.xlsx", data_only=True)
    assert [wb1["Sheet1"]["A1"].value, wb1["Sheet1"]["A2"].value, wb1["Sheet1"]["A3"].value] == [1, 2, 3]
    assert [wb2["Sheet1"]["A1"].value, wb2["Sheet1"]["A2"].value] == [4, 5]
    wb1.close()
    wb2.close()


def test_streaming_single_part_does_not_zip(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1, 2]})

    paths = _export_streaming(
        schema,
        {"Sheet1": {"rows": df}},
        managed_tmp_dir,
        max_rows_per_workbook=10,
    )

    assert len(paths) == 1
    assert Path(paths[0]).name == "filled.part001.xlsx"
    assert Path(paths[0]).exists()
    assert not (managed_tmp_dir / "filled.zip").exists()


def test_streaming_reads_lazyframe_source_in_batches(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    polars.DataFrame({"A": [1, 2, 3], "B": [4, 5, 6]}).write_parquet(source_path)
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")

    paths = _export_streaming(
        schema,
        {"Sheet1": {"rows": polars.scan_parquet(source_path).select(["B", "A"])}},
        managed_tmp_dir,
        streaming_chunk_rows=1,
        max_rows_per_workbook=10,
    )

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    assert [wb["Sheet1"]["A1"].value, wb["Sheet1"]["B1"].value] == [4, 1]
    assert [wb["Sheet1"]["A3"].value, wb["Sheet1"]["B3"].value] == [6, 3]
    wb.close()


def test_streaming_rejects_hug_mode(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})

    with pytest.raises(ValueError, match="does not support 'hug'"):
        _export_streaming(
            schema,
            {"Sheet1": {"rows": df}},
            managed_tmp_dir,
            column_width_mode="hug",
        )


def test_streaming_rejects_invalid_row_limits(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})

    with pytest.raises(ValueError, match="streaming_chunk_rows must be greater than 0"):
        _export_streaming(
            schema,
            {"Sheet1": {"rows": df}},
            managed_tmp_dir,
            streaming_chunk_rows=0,
        )
    with pytest.raises(ValueError, match="max_rows_per_workbook must be between"):
        _export_streaming(
            schema,
            {"Sheet1": {"rows": df}},
            managed_tmp_dir,
            max_rows_per_workbook=0,
        )


def test_streaming_rejects_unknown_export_mode(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="Unsupported export_mode"):
        mo_dataport.export(bundle, str(managed_tmp_dir / "out.xlsx"), export_mode="fast")


def test_streaming_rejects_multiple_content_anchors(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{first:dataframe-content}}"),
            "A3": _cell("A3", "{{second:dataframe-content}}"),
        },
        dims="A1:B3",
    )
    df = polars.DataFrame({"A": [1]})

    with pytest.raises(ValueError, match="one dataframe-content placeholder"):
        _export_streaming(
            schema,
            {"Sheet1": {"first": df, "second": df.clone()}},
            managed_tmp_dir,
        )


def test_streaming_rejects_missing_and_unsupported_source(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1]})}},
        bundle_path=str(managed_tmp_dir / "bundle"),
    )
    source = bundle.manifest["dataframe_sources"][0]
    out = managed_tmp_dir / "out.xlsx"

    source_path = Path(bundle.path) / source["path"]
    source_path.unlink()
    with pytest.raises(ValueError, match="missing dataframe source"):
        mo_dataport.export(bundle, str(out), export_mode="streaming")

    polars.DataFrame({"A": [1]}).write_parquet(source_path)
    source["format"] = "csv"
    with pytest.raises(ValueError, match="Unsupported dataframe source format"):
        mo_dataport.export(bundle, str(out), export_mode="streaming")


def test_streaming_writes_empty_dataframe_without_content_rows(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")

    paths = _export_streaming(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": []})}},
        managed_tmp_dir,
    )

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    assert wb["Sheet1"]["A1"].value is None
    wb.close()


def test_streaming_applies_fixed_and_even_dimensions(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:B2")
    sheet = schema["sheets"][0]
    sheet["column_widths"] = {"A": 33.0}
    sheet["row_heights"] = {"1": 24.0}
    paths = _export_streaming(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1], "B": [2]})}},
        managed_tmp_dir,
    )
    wb = openpyxl.load_workbook(paths[0])
    ws = wb["Sheet1"]
    assert ws.column_dimensions["A"].width == pytest.approx(33.0)
    assert ws.row_dimensions[1].height == pytest.approx(24.0)
    wb.close()

    even_schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:B2")
    even_schema["sheets"][0].update(
        {
            "column_width_mode": "even",
            "default_column_width": 21.0,
            "row_height_mode": "even",
            "default_row_height": 19.0,
        }
    )
    paths = _export_streaming(
        even_schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1], "B": [2]})}},
        managed_tmp_dir,
    )
    wb = openpyxl.load_workbook(paths[0])
    ws = wb["Sheet1"]
    assert ws.column_dimensions["A"].width == pytest.approx(21.0)
    assert ws.column_dimensions["B"].width == pytest.approx(21.0)
    assert ws.row_dimensions[1].height == pytest.approx(19.0)
    assert ws.row_dimensions[2].height == pytest.approx(19.0)
    wb.close()


def test_streaming_preserves_static_merged_regions(managed_tmp_dir: Path):
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["fill"] = {"bg_color": "FF00FF00"}
    shadow = _cell("B1", "Should not be written")
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema(
        {
            "A1": title,
            "B1": shadow,
            "A3": _cell("A3", "{{rows:dataframe-content}}"),
        },
        dims="A1:B3",
        merges=["A1:B1"],
    )
    df = polars.DataFrame({"A": [10], "B": [20]})

    paths = _export_streaming(schema, {"Sheet1": {"rows": df}}, managed_tmp_dir)

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert [str(region) for region in ws.merged_cells.ranges] == ["A1:B1"]
    assert ws["A1"].value == "Merged Title"
    assert ws["B1"].value is None
    assert ws["A1"].fill.fgColor.rgb == "FF00FF00"
    assert [ws["A3"].value, ws["B3"].value] == [10, 20]
    wb.close()


def test_streaming_rejects_merged_regions_that_intersect_content(managed_tmp_dir: Path):
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}")},
        dims="A1:A1",
        merges=["A1:B1"],
    )
    df = polars.DataFrame({"A": [1]})

    with pytest.raises(ValueError, match="intersect dataframe-content"):
        _export_streaming(schema, {"Sheet1": {"rows": df}}, managed_tmp_dir)


def test_streaming_anchor_style_is_cloned(managed_tmp_dir: Path):
    anchor = _cell("A1", "{{rows:dataframe-content}}")
    anchor["fill"] = {"bg_color": "FFFF0000"}
    anchor["number_format"] = "0.00"
    schema = _schema({"A1": anchor}, dims="A1:A1")
    df = polars.DataFrame({"A": [1.5]})

    paths = _export_streaming(schema, {"Sheet1": {"rows": df}}, managed_tmp_dir)
    wb = openpyxl.load_workbook(paths[0])
    cell = wb["Sheet1"]["A1"]
    assert cell.fill.fgColor.type == "rgb"
    assert cell.fill.fgColor.rgb == "FFFF0000"
    assert cell.number_format == "0.00"
    wb.close()


def test_streaming_expands_dynamic_sheet_names_in_order(managed_tmp_dir: Path):
    schema = {
        "sheets": [
            {
                "name": "{{sheet_1}}",
                "dimensions": "A1:A1",
                "merged_regions": [],
                "column_widths": {},
                "row_heights": {},
                "cells": {"A1": _cell("A1", "{{name:string}}")},
            }
        ]
    }

    paths = _export_streaming(
        schema,
        {
            "sheet_1": {
                "Sheet Name 1": {"name": "Alpha"},
                "Sheet Name 2": {"name": "Beta"},
            }
        },
        managed_tmp_dir,
        max_rows_per_workbook=10,
    )

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    assert wb.sheetnames == ["Sheet Name 1", "Sheet Name 2"]
    assert wb["Sheet Name 1"]["A1"].value == "Alpha"
    assert wb["Sheet Name 2"]["A1"].value == "Beta"
    wb.close()


def test_streaming_renders_repeat_records_in_one_sheet(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{line_items:dataframe}}"),
            "A4": _cell("A4", ""),
            "A5": _cell("A5", "{{reports:repeat-end}}"),
        },
        dims="A1:B5",
    )

    paths = _export_streaming(
        schema,
        {
            "Sheet1": {
                "reports": [
                    {
                        "customer_name": "Acme",
                        "line_items": polars.DataFrame({"sku": ["A"], "qty": [1]}),
                    },
                    {
                        "customer_name": "Globex",
                        "line_items": polars.DataFrame({"sku": ["G"], "qty": [2]}),
                    },
                    {
                        "customer_name": "Initech",
                        "line_items": polars.DataFrame({"sku": ["I"], "qty": [3]}),
                    },
                ]
            }
        },
        managed_tmp_dir,
        streaming_chunk_rows=1,
        max_rows_per_workbook=20,
    )

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert [ws["A1"].value, ws["A4"].value, ws["A7"].value] == [
        "Acme",
        "Globex",
        "Initech",
    ]
    assert [ws["A2"].value, ws["B2"].value, ws["A3"].value, ws["B3"].value] == [
        "sku",
        "qty",
        "A",
        1,
    ]
    assert [ws["A8"].value, ws["B8"].value, ws["A9"].value, ws["B9"].value] == [
        "sku",
        "qty",
        "I",
        3,
    ]
    wb.close()


def test_streaming_splits_repeat_records_across_parts(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{name:string}}"),
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:A3",
    )

    paths = _export_streaming(
        schema,
        {"Sheet1": {"reports": [{"name": "A"}, {"name": "B"}, {"name": "C"}]}},
        managed_tmp_dir,
        max_rows_per_workbook=2,
    )

    with ZipFile(paths[0]) as zip_file:
        assert zip_file.namelist() == ["filled.part001.xlsx", "filled.part002.xlsx"]
        zip_file.extractall(path=managed_tmp_dir)

    wb1 = openpyxl.load_workbook(managed_tmp_dir / "filled.part001.xlsx", data_only=True)
    wb2 = openpyxl.load_workbook(managed_tmp_dir / "filled.part002.xlsx", data_only=True)
    assert [wb1["Sheet1"]["A1"].value, wb1["Sheet1"]["A2"].value] == ["A", "B"]
    assert wb2["Sheet1"]["A1"].value == "C"
    wb1.close()
    wb2.close()


def test_pdf_renders_repeat_records(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{line_items:dataframe}}"),
            "A4": _cell("A4", "{{reports:repeat-end}}"),
        },
        dims="A1:B4",
    )
    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "reports": [
                    {
                        "customer_name": "Acme",
                        "line_items": polars.DataFrame({"sku": ["A"], "qty": [1]}),
                    },
                    {
                        "customer_name": "Globex",
                        "line_items": polars.DataFrame({"sku": ["G"], "qty": [2]}),
                    },
                ]
            }
        },
    )
    out = managed_tmp_dir / "repeat.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", streaming_chunk_rows=1)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_streaming_supports_repeat_merged_fixed_rows(managed_tmp_dir: Path):
    title = _cell("A2", "{{name:string}}")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    title["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thick", "color": "FF000000"},
        "left": {"style": "medium", "color": "FF000000"},
        "right": {"style": "dashed", "color": "FF000000"},
    }
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": title,
            "B2": shadow,
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:B3",
        merges=["A2:B2"],
    )

    paths = _export_streaming(
        schema,
        {"Sheet1": {"reports": [{"name": "Acme"}, {"name": "Globex"}]}},
        managed_tmp_dir,
    )

    wb = openpyxl.load_workbook(paths[0])
    ws = wb["Sheet1"]
    assert sorted(str(region) for region in ws.merged_cells.ranges) == [
        "A1:B1",
        "A2:B2",
    ]
    assert ws["A1"].value == "Acme"
    assert ws["A2"].value == "Globex"
    assert ws["A1"].border.top.style == "thin"
    assert ws["B1"].border.right.style == "dashed"
    wb.close()


def test_pdf_supports_repeat_merged_fixed_rows(managed_tmp_dir: Path):
    title = _cell("A2", "{{name:string}}")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": title,
            "B2": shadow,
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:B3",
        merges=["A2:B2"],
    )
    bundle = mo_dataport.compile(schema, {"Sheet1": {"reports": [{"name": "Acme"}]}})
    out = managed_tmp_dir / "merged-repeat.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", streaming_chunk_rows=2)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_supports_static_merge_before_repeat_section(managed_tmp_dir: Path):
    title = _cell("A1", "Repeat Sample")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["fill"] = {"bg_color": "FF003366"}
    shadow = _cell("B1", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema(
        {
            "A1": title,
            "B1": shadow,
            "A2": _cell("A2", "{{reports:repeat-start}}"),
            "A3": _cell("A3", "{{name:string}}"),
            "A4": _cell("A4", "{{reports:repeat-end}}"),
        },
        dims="A1:B4",
        merges=["A1:B1"],
    )
    bundle = mo_dataport.compile(schema, {"Sheet1": {"reports": [{"name": "Acme"}]}})
    out = managed_tmp_dir / "static-merged-title.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", streaming_chunk_rows=2)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_streaming_renders_sibling_repeat_sections(managed_tmp_dir: Path):
    first = _cell("A3", "{{customer_name:string}}")
    first["merged"] = True
    first["merge_anchor"] = "A3"
    first_shadow = _cell("B3", None)
    first_shadow["merged"] = True
    first_shadow["merge_anchor"] = "A3"
    second = _cell("A8", "{{region:string}}")
    second["merged"] = True
    second["merge_anchor"] = "A8"
    second_shadow = _cell("B8", None)
    second_shadow["merged"] = True
    second_shadow["merge_anchor"] = "A8"
    schema = _schema(
        {
            "A1": _cell("A1", "Title"),
            "A2": _cell("A2", "{{customers:repeat-start}}"),
            "A3": first,
            "B3": first_shadow,
            "A4": _cell("A4", "{{customers:repeat-end}}"),
            "A5": _cell("A5", "Between"),
            "A7": _cell("A7", "{{summaries:repeat-start}}"),
            "A8": second,
            "B8": second_shadow,
            "A9": _cell("A9", "{{totals:dataframe-content}}"),
            "A10": _cell("A10", "{{summaries:repeat-end}}"),
            "A11": _cell("A11", "Footer"),
        },
        dims="A1:B11",
        merges=["A3:B3", "A8:B8"],
    )

    paths = _export_streaming(
        schema,
        {
            "Sheet1": {
                "customers": [{"customer_name": "Acme"}, {"customer_name": "Globex"}],
                "summaries": [
                    {
                        "region": "North",
                        "totals": polars.DataFrame({"metric": ["Sales"], "value": [10]}),
                    },
                    {
                        "region": "South",
                        "totals": polars.DataFrame({"metric": ["Sales"], "value": [20]}),
                    },
                ],
            }
        },
        managed_tmp_dir,
        streaming_chunk_rows=1,
        max_rows_per_workbook=20,
    )

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert [
        ws["A1"].value,
        ws["A2"].value,
        ws["A3"].value,
        ws["A4"].value,
        ws["A5"].value,
        ws["A6"].value,
        ws["A7"].value,
        ws["A8"].value,
        ws["A9"].value,
        ws["A10"].value,
    ] == [
        "Title",
        "Acme",
        "Globex",
        "Between",
        None,
        "North",
        "Sales",
        "South",
        "Sales",
        "Footer",
    ]
    assert [ws["B7"].value, ws["B9"].value] == [10, 20]
    assert sorted(str(region) for region in ws.merged_cells.ranges) == [
        "A2:B2",
        "A3:B3",
        "A6:B6",
        "A8:B8",
    ]
    wb.close()


def test_pdf_renders_sibling_repeat_sections(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "Title"),
            "A2": _cell("A2", "{{customers:repeat-start}}"),
            "A3": _cell("A3", "{{customer_name:string}}"),
            "A4": _cell("A4", "{{customers:repeat-end}}"),
            "A5": _cell("A5", "Between"),
            "A6": _cell("A6", "{{summaries:repeat-start}}"),
            "A7": _cell("A7", "{{region:string}}"),
            "A8": _cell("A8", "{{summaries:repeat-end}}"),
        },
        dims="A1:A8",
    )
    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "customers": [{"customer_name": "Acme"}],
                "summaries": [{"region": "North"}],
            }
        },
    )
    out = managed_tmp_dir / "sibling-repeat.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", streaming_chunk_rows=2)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


# §5 Entrypoints
