import json

import openpyxl
import pytest
from openpyxl.styles import Border, Color, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils.cell import get_column_letter

from mindoff_dataport import extract_template
from mindoff_dataport.style_conversion import (
    argb_to_color,
    border_side_to_dict,
    dict_to_border_side,
    extract_theme_colors,
    normalize_color,
    resolve_theme_color,
)

# §1. Constants & Exceptions

# §2. Classes and Sub Classes

# §3. Private Helper Functions

_CUSTOM_THEME = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Custom">
  <a:themeElements>
    <a:clrScheme name="Custom">
      <a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
      <a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>
      <a:lt2><a:srgbClr val="EEECE1"/></a:lt2>
      <a:dk2><a:srgbClr val="1F497D"/></a:dk2>
      <a:accent1><a:srgbClr val="112233"/></a:accent1>
      <a:accent2><a:srgbClr val="445566"/></a:accent2>
      <a:accent3><a:srgbClr val="778899"/></a:accent3>
      <a:accent4><a:srgbClr val="AABBCC"/></a:accent4>
      <a:accent5><a:srgbClr val="DDEEFF"/></a:accent5>
      <a:accent6><a:srgbClr val="123456"/></a:accent6>
      <a:hlink><a:srgbClr val="0000FF"/></a:hlink>
      <a:folHlink><a:srgbClr val="800080"/></a:folHlink>
    </a:clrScheme>
  </a:themeElements>
</a:theme>"""

# §4. Public Functions


def test_extract_returns_workbook_schema(workbook_schema):
    assert "sheets" in workbook_schema
    assert len(workbook_schema["sheets"]) >= 1
    assert "theme_colors" in workbook_schema
    assert len(workbook_schema["theme_colors"]) >= 10


def test_sheet_has_required_keys(workbook_schema):
    sheet = workbook_schema["sheets"][0]
    for key in (
        "name",
        "dimensions",
        "merged_regions",
        "column_widths",
        "row_heights",
        "cells",
    ):
        assert key in sheet, f"Missing key: {key}"


def test_merged_regions_captured(workbook_schema):
    sheet = workbook_schema["sheets"][0]
    assert len(sheet["merged_regions"]) >= 2
    assert "A1:C2" in sheet["merged_regions"]
    assert "D1:E3" in sheet["merged_regions"]


def test_merge_anchor_cell_tagged(workbook_schema):
    cells = workbook_schema["sheets"][0]["cells"]
    assert cells["A1"]["merged"] is True
    assert cells["A1"]["merge_anchor"] == "A1"


def test_non_anchor_merged_cell_tagged(workbook_schema):
    cells = workbook_schema["sheets"][0]["cells"]
    assert cells["B1"]["merged"] is True
    assert cells["B1"]["merge_anchor"] == "A1"
    assert cells["B1"]["value"] is None


def test_non_merged_cell_not_tagged(workbook_schema):
    cells = workbook_schema["sheets"][0]["cells"]
    assert cells["A3"]["merged"] is False
    assert cells["A3"]["merge_anchor"] is None


def test_plain_string_cell(workbook_schema):
    cell = workbook_schema["sheets"][0]["cells"]["A3"]
    assert cell["value"] == "Label"
    assert cell["cell_type"] == "string"


def test_number_cell_with_format(workbook_schema):
    cell = workbook_schema["sheets"][0]["cells"]["B3"]
    assert cell["cell_type"] == "number"
    assert cell["number_format"] == "0.00"


def test_formula_cell(workbook_schema):
    cell = workbook_schema["sheets"][0]["cells"]["C3"]
    assert cell["cell_type"] == "formula"
    assert str(cell["value"]).startswith("=")


def test_date_cell(workbook_schema):
    cell = workbook_schema["sheets"][0]["cells"]["A4"]
    assert cell["cell_type"] == "date"
    assert "2024" in str(cell["value"])


def test_font_bold_italic(workbook_schema):
    font = workbook_schema["sheets"][0]["cells"]["B4"]["font"]
    assert font["bold"] is True
    assert font["italic"] is True
    assert font["color"] is not None


def test_font_strike_captured(workbook_schema):
    font = workbook_schema["sheets"][0]["cells"]["A6"]["font"]
    assert font["strike"] is True


def test_font_vert_align_superscript(workbook_schema):
    font = workbook_schema["sheets"][0]["cells"]["B6"]["font"]
    assert font["vert_align"] == "superscript"


def test_font_vert_align_subscript(workbook_schema):
    font = workbook_schema["sheets"][0]["cells"]["C6"]["font"]
    assert font["vert_align"] == "subscript"


def test_fill_solid_fg_color_captured(workbook_schema):
    fill = workbook_schema["sheets"][0]["cells"]["C4"]["fill"]
    assert fill["pattern_type"] == "solid"
    assert fill["fg_color"] is not None
    assert len(fill["fg_color"]) == 8


def test_fill_pattern_type_captured(workbook_schema):
    fill = workbook_schema["sheets"][0]["cells"]["A9"]["fill"]
    assert fill["pattern_type"] == "gray125"
    assert fill["fg_color"] is not None
    assert fill["bg_color"] is not None


def test_extract_theme_colors_from_custom_workbook(managed_tmp_dir):
    workbook = openpyxl.Workbook()
    workbook.loaded_theme = _CUSTOM_THEME
    sheet = workbook.active
    sheet["A1"].value = "Theme"
    sheet["A1"].fill = PatternFill(
        patternType="solid",
        fgColor=Color(theme=4, tint=0.5),
    )
    path = managed_tmp_dir / "theme.xlsx"
    workbook.save(path)

    schema = extract_template(str(path))

    assert schema["theme_colors"][4] == "FF112233"
    fill = schema["sheets"][0]["cells"]["A1"]["fill"]
    assert fill["fg_color"] == "theme:4:0.5"


def test_resolve_theme_color_uses_palette_and_tint():
    palette = extract_theme_colors(_CUSTOM_THEME)

    assert resolve_theme_color("theme:4:0.0", palette) == "FF112233"
    assert resolve_theme_color("theme:4:0.5", palette) == "FF889099"
    assert resolve_theme_color("theme:4:-0.5", palette) == "FF081119"
    assert resolve_theme_color("FFABCDEF", palette) == "FFABCDEF"
    assert resolve_theme_color("theme:not-a-number:0", palette) is None


def test_resolve_theme_color_rejects_invalid_shapes():
    assert resolve_theme_color("theme:2", DEFAULT_THEME := extract_theme_colors(None)) is None
    assert resolve_theme_color("theme:-1:0.0", DEFAULT_THEME) is None
    assert resolve_theme_color("theme:100:0.0", DEFAULT_THEME) is None
    assert resolve_theme_color("theme:2:0.0", ["BAD"]) is None


def test_extract_theme_colors_falls_back_for_invalid_xml_and_missing_scheme():
    assert extract_theme_colors("<broken") == extract_theme_colors(None)
    assert (
        extract_theme_colors(
            b'<?xml version="1.0"?><a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"></a:theme>'
        )
        == extract_theme_colors(None)
    )


def test_style_color_and_border_conversion_defensive_paths():
    assert normalize_color(Color(type="indexed", indexed=4)) == "FF0000FF"
    assert normalize_color(Color(type="indexed", indexed=999)) is None
    assert normalize_color(Color(type="auto")) is None
    assert normalize_color(Color(rgb="00000000")) is None
    assert normalize_color(Color(rgb="112233")) == "00112233"

    themed = argb_to_color("theme:4:0.25")
    assert themed is not None
    assert themed.type == "theme"
    assert themed.theme == 4

    plain = argb_to_color("FF112233")
    assert plain is not None
    assert plain.rgb == "FF112233"

    side = dict_to_border_side({"style": "thin", "color": "theme:1:0.0"})
    converted = border_side_to_dict(side)
    assert converted["style"] == "thin"
    assert converted["color"] == "theme:1:0.0"


def test_extract_theme_colors_falls_back_when_scheme_entries_are_missing():
    xml = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Broken">
  <a:themeElements>
    <a:clrScheme name="Broken">
      <a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
    </a:clrScheme>
  </a:themeElements>
</a:theme>"""
    assert extract_theme_colors(xml) == extract_theme_colors(None)


def test_borders_captured(workbook_schema):
    borders = workbook_schema["sheets"][0]["cells"]["A5"]["borders"]
    assert borders["top"]["style"] == "thick"
    assert borders["left"]["style"] == "medium"
    assert borders["right"]["style"] == "dashed"


def test_border_diagonal_captured(workbook_schema):
    borders = workbook_schema["sheets"][0]["cells"]["B8"]["borders"]
    assert borders["diagonal"]["style"] == "thin"
    assert borders["diagonal_up"] is True
    assert borders["diagonal_down"] is True


def test_border_start_end_captured(workbook_schema):
    borders = workbook_schema["sheets"][0]["cells"]["C8"]["borders"]
    assert borders["start"]["style"] == "medium"
    assert borders["end"]["style"] == "dashed"


def test_merged_region_border_is_captured_from_edges(managed_tmp_dir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:B2")
    sheet["A1"].value = "Merged"
    sheet["A1"].border = Border(top=Side(style="thin"), left=Side(style="medium"))
    sheet["B2"].border = Border(bottom=Side(style="thick"), right=Side(style="dashed"))
    path = managed_tmp_dir / "merged-border.xlsx"
    workbook.save(path)

    schema = extract_template(str(path))

    borders = schema["sheets"][0]["cells"]["A1"]["borders"]
    assert borders["top"]["style"] == "thin"
    assert borders["bottom"]["style"] == "thick"
    assert borders["left"]["style"] == "medium"
    assert borders["right"]["style"] == "dashed"


def test_alignment_captured(workbook_schema):
    alignment = workbook_schema["sheets"][0]["cells"]["B5"]["alignment"]
    assert alignment["wrap_text"] is True
    assert alignment["horizontal"] == "center"


def test_alignment_indent_captured(workbook_schema):
    alignment = workbook_schema["sheets"][0]["cells"]["A7"]["alignment"]
    assert alignment["indent"] == 2


def test_alignment_text_rotation_captured(workbook_schema):
    alignment = workbook_schema["sheets"][0]["cells"]["B7"]["alignment"]
    assert alignment["text_rotation"] == 45


def test_alignment_shrink_to_fit_captured(workbook_schema):
    alignment = workbook_schema["sheets"][0]["cells"]["C7"]["alignment"]
    assert alignment["shrink_to_fit"] is True


def test_alignment_reading_order_captured(workbook_schema):
    alignment = workbook_schema["sheets"][0]["cells"]["A8"]["alignment"]
    assert alignment["reading_order"] == 2


def test_column_widths_captured(workbook_schema):
    widths = workbook_schema["sheets"][0]["column_widths"]
    assert "A" in widths
    assert abs(widths["A"] - 20.0) < 0.5


def test_column_dimension_ranges_are_expanded(managed_tmp_dir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Range width"
    sheet.column_dimensions["C"].width = 31.25
    sheet.column_dimensions["C"].min = 3
    sheet.column_dimensions["C"].max = 12
    path = managed_tmp_dir / "range-width.xlsx"
    workbook.save(path)

    schema = extract_template(str(path))

    widths = schema["sheets"][0]["column_widths"]
    for col_idx in range(3, 13):
        assert widths[get_column_letter(col_idx)] == pytest.approx(31.25)


def test_sheet_gridline_visibility_captured(managed_tmp_dir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "No gridlines"
    path = managed_tmp_dir / "gridlines.xlsx"
    workbook.save(path)

    schema = extract_template(str(path))

    assert schema["sheets"][0]["show_gridlines"] is False


def test_row_heights_captured(workbook_schema):
    heights = workbook_schema["sheets"][0]["row_heights"]
    assert "1" in heights
    assert abs(heights["1"] - 30.0) < 0.5


def test_manual_row_and_column_page_breaks_are_extracted(managed_tmp_dir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Breaks"
    sheet.row_breaks.append(Break(id=2))
    sheet.col_breaks.append(Break(id=3))
    path = managed_tmp_dir / "page-breaks.xlsx"
    workbook.save(path)

    schema = extract_template(str(path))

    assert schema["sheets"][0]["row_page_breaks"] == [2]
    assert schema["sheets"][0]["column_page_breaks"] == [3]


def test_manual_page_breaks_are_normalized(managed_tmp_dir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Breaks"
    sheet.row_breaks.append(Break(id=5))
    sheet.row_breaks.append(Break(id=2))
    sheet.row_breaks.append(Break(id=5))
    sheet.row_breaks.append(Break(id=0))
    sheet.col_breaks.append(Break(id=4))
    sheet.col_breaks.append(Break(id=2))
    sheet.col_breaks.append(Break(id=4))
    path = managed_tmp_dir / "normalized-breaks.xlsx"
    workbook.save(path)

    schema = extract_template(str(path))

    assert schema["sheets"][0]["row_page_breaks"] == [2, 5]
    assert schema["sheets"][0]["column_page_breaks"] == [2, 4]


def test_schema_is_json_serializable(workbook_schema):
    json.dumps(workbook_schema)


# §5. Entrypoints
