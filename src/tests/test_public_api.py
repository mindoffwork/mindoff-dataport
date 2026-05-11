import mindoff_dataport

import openpyxl
import pytest

from mindoff_dataport import (
    compile_report_bundle,
    export_report_bundle,
    extract_template,
    get_template_inputs,
    mo_dataport,
    repeat_records,
)

# §1. Constants & Exceptions

# §2. Classes and Sub Classes

# §3. Private Helper Functions


def _cell(coord: str, value: str) -> dict:
    return {
        "coordinate": coord,
        "value": value,
        "cell_type": "string",
        "number_format": None,
        "font": {
            "name": "Calibri",
            "size": 11.0,
            "bold": False,
            "italic": False,
            "underline": None,
            "color": None,
        },
        "fill": {"bg_color": None},
        "alignment": {"horizontal": None, "vertical": None, "wrap_text": False},
        "borders": {
            "top": {"style": None, "color": None},
            "bottom": {"style": None, "color": None},
            "left": {"style": None, "color": None},
            "right": {"style": None, "color": None},
        },
        "merged": False,
        "merge_anchor": None,
    }


def _schema() -> dict:
    return {
        "sheets": [
            {
                "name": "Sheet1",
                "dimensions": "A1:A1",
                "merged_regions": [],
                "column_widths": {},
                "row_heights": {},
                "cells": {"A1": _cell("A1", "{{name:string}}")},
            }
        ]
    }


# §4. Public Functions


def test_mo_dataport_namespace_exposes_bundle_first_aliases():
    assert mo_dataport.extract is extract_template
    assert mo_dataport.inputs is get_template_inputs
    assert mo_dataport.compile is compile_report_bundle
    assert mo_dataport.export is export_report_bundle
    assert mo_dataport.repeat_records is repeat_records
    assert not hasattr(mindoff_dataport, "parquet_source")
    assert not hasattr(mo_dataport, "parquet_source")
    assert not hasattr(mo_dataport, "build")
    assert not hasattr(mo_dataport, "alter_schema")
    assert not hasattr(mindoff_dataport, "mode")


def test_mo_dataport_bundle_aliases_work(managed_tmp_dir):
    schema = _schema()

    inputs = mo_dataport.inputs(schema)
    assert inputs == {"Sheet1": {"name": "string"}}

    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})
    assert bundle.report["sheets"][0]["cells"]["A1"]["value"] == "Alice"

    output = managed_tmp_dir / "out.xlsx"
    mo_dataport.export(bundle, str(output))
    assert output.exists()


def test_export_rejects_streaming_engine_for_pdf(managed_tmp_dir):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="PDF export does not support streaming_engine"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "out.pdf"),
            format="pdf",
            streaming_engine="openpyxl",
        )


def test_export_rejects_xlsxwriter_for_fidelity(managed_tmp_dir):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="Fidelity XLSX export does not support"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "out.xlsx"),
            export_mode="fidelity",
            streaming_engine="xlsxwriter",
        )


def test_fidelity_defaults_to_openpyxl_engine(managed_tmp_dir):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})
    out = managed_tmp_dir / "fidelity-default.xlsx"

    mo_dataport.export(bundle, str(out), export_mode="fidelity")

    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb["Sheet1"]["A1"].value == "Alice"
    wb.close()


def test_fidelity_accepts_explicit_openpyxl_engine(managed_tmp_dir):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})
    out = managed_tmp_dir / "fidelity-openpyxl.xlsx"

    mo_dataport.export(
        bundle,
        str(out),
        export_mode="fidelity",
        streaming_engine="openpyxl",
    )

    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb["Sheet1"]["A1"].value == "Alice"
    wb.close()


@pytest.mark.parametrize("engine", [None, "openpyxl", "xlsxwriter"])
def test_streaming_allows_openpyxl_and_xlsxwriter_engines(managed_tmp_dir, engine):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})
    out = managed_tmp_dir / f"streaming-{engine or 'default'}.xlsx"
    options = {"export_mode": "streaming", "streaming_chunk_rows": 1}
    if engine is not None:
        options["streaming_engine"] = engine

    paths = mo_dataport.export(bundle, str(out), **options)

    assert isinstance(paths, list) and paths
    wb = openpyxl.load_workbook(paths[0], data_only=True)
    assert wb["Sheet1"]["A1"].value == "Alice"
    wb.close()


def test_pdf_export_uses_default_renderer_path(managed_tmp_dir):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})
    out = managed_tmp_dir / "out.pdf"

    mo_dataport.export(bundle, str(out), format="pdf")

    assert out.exists()
    assert out.stat().st_size > 0


def test_streaming_rejects_unknown_streaming_engine(managed_tmp_dir):
    bundle = mo_dataport.compile(_schema(), {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="Unsupported streaming_engine"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "out.xlsx"),
            export_mode="streaming",
            streaming_engine="fast",
        )


# §5. Entrypoints
