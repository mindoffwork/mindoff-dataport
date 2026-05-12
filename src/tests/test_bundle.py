from __future__ import annotations

import datetime
import json
from pathlib import Path

import openpyxl
import pytest
import reportlab
from reportlab.platypus import PageBreak

from mindoff_dataport import mo_dataport
from mindoff_dataport.bundle import load_report_bundle
from mindoff_dataport.pdf_renderer import (
    _FontResolver,
    _LazyFlowables,
    _apply_tint,
    _chunked_row_flowables,
    _chunked_row_tables,
    _column_widths,
    _data_source_map,
    _dataframe_pdf_rows,
    _fast_grid_plan,
    _hex_color,
    _paragraph,
    _repeat_record_rows,
    _row_chunk_table,
    _sheet_flowables,
    _table_style,
)
from mindoff_dataport.xlsx_renderer import _expanded_cells
from mindoff_dataport.xlsx_renderer import (
    _apply_merged_region_borders,
    _apply_streaming_dimensions,
    _delete_bundle_dir,
    _edge_border_schema_bounds,
    _merged_region_edge_schemas,
    _resolve_color_field,
    _resolved_borders,
    _resolved_fill,
    _resolved_font,
    _xlsxwriter_border_style,
    _xlsxwriter_color,
    _xlsxwriter_format_props,
    _xlsxwriter_logical_border_sides,
    _xlsxwriter_pattern,
    _xlsxwriter_vertical_alignment,
)

# §1. Constants & Exceptions

polars = pytest.importorskip("polars", reason="polars not installed")

# §2. Classes and Sub Classes

# §3. Private Helper Functions


def _cell(coord: str, value):
    return {
        "coordinate": coord,
        "value": value,
        "cell_type": "string",
        "number_format": None,
        "font": {
            "name": "Calibri",
            "size": 11.0,
            "bold": False,
            "italic": False,
            "underline": None,
            "strike": False,
            "vert_align": None,
            "color": None,
        },
        "fill": {"pattern_type": None, "fg_color": None, "bg_color": None},
        "alignment": {
            "horizontal": None,
            "vertical": None,
            "wrap_text": False,
            "indent": None,
            "relative_indent": None,
            "text_rotation": None,
            "shrink_to_fit": False,
            "reading_order": None,
        },
        "borders": {
            "top": {"style": None, "color": None},
            "bottom": {"style": None, "color": None},
            "left": {"style": None, "color": None},
            "right": {"style": None, "color": None},
            "start": {"style": None, "color": None},
            "end": {"style": None, "color": None},
            "horizontal": {"style": None, "color": None},
            "vertical": {"style": None, "color": None},
            "diagonal": {"style": None, "color": None},
            "diagonal_up": False,
            "diagonal_down": False,
            "outline": True,
        },
        "merged": False,
        "merge_anchor": None,
    }


def _schema(cells: dict[str, dict], *, dims: str = "A1:C3") -> dict:
    return {
        "sheets": [
            {
                "name": "Sheet1",
                "dimensions": dims,
                "merged_regions": [],
                "column_widths": {},
                "row_heights": {},
                "cells": cells,
            }
        ]
    }


def _schema_with_options(cells: dict[str, dict], *, dims: str = "A1:C3", **options):
    schema = _schema(cells, dims=dims)
    schema["sheets"][0].update(options)
    return schema


def _has_line_command(commands: list[tuple], command: str, width: float) -> bool:
    return any(
        item[:4] == (command, (0, 0), (1, 1), width)
        for item in commands
    )


def _vera_font_path() -> str:
    return str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")


# §4. Public Functions


def test_compile_creates_valid_report_bundle_directory(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{name:string}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:B2",
    )
    df = polars.DataFrame({"A": [1, 2], "B": [3, 4]})
    bundle_path = managed_tmp_dir / "report_bundle"

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"name": "Alice", "rows": df}},
        bundle_path=str(bundle_path),
    )

    assert bundle_path.exists()
    assert bundle_path.is_dir()
    assert (bundle_path / "manifest.json").exists()
    assert (bundle_path / "report.json").exists()
    assert bundle.report["sheets"][0]["cells"]["A1"]["value"] == "Alice"
    assert "A2" not in bundle.report["sheets"][0]["cells"]
    source_path = bundle.report["sheets"][0]["dataframe_anchors"][0]["source"]
    assert source_path.startswith("data/")
    assert (bundle_path / source_path).exists()

    manifest = json.loads((bundle_path / "manifest.json").read_text(encoding="utf-8"))
    report = json.loads((bundle_path / "report.json").read_text(encoding="utf-8"))

    assert manifest["bundle_format"] == "directory"
    assert manifest["output_capabilities"] == {
        "xlsx": True,
        "pdf": True,
        "image": False,
    }
    assert report["sheets"][0]["dataframe_anchors"][0]["columns"] == ["A", "B"]


def test_compile_creates_compact_repeat_section_bundle(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{line_items:dataframe}}"),
            "A4": _cell("A4", "{{reports:repeat-end}}"),
        },
        dims="A1:B4",
    )
    rows = polars.DataFrame({"sku": ["A", "B"], "qty": [1, 2]})

    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "reports": [
                    {"customer_name": "Acme", "line_items": rows},
                    {"customer_name": "Globex", "line_items": rows},
                ]
            }
        },
        bundle_path=str(managed_tmp_dir / "bundle"),
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["cells"] == {}
    assert "repeat_sections" in sheet
    section = sheet["repeat_sections"][0]
    assert section["key"] == "reports"
    assert len(section["records"]) == 2
    assert section["cell_templates"]
    assert section["records"][0]["cells"][0]["value"] == "Acme"
    assert "cell" not in section["records"][0]["cells"][0]
    assert len(bundle.manifest["dataframe_sources"]) == 1


def test_compile_source_backed_repeat_keeps_records_out_of_report(
    managed_tmp_dir: Path,
):
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{line_items:dataframe-content}}"),
            "A4": _cell("A4", "{{reports:repeat-end}}"),
        },
        dims="A1:B4",
    )
    employees = polars.DataFrame({"customer_name": [f"Name {idx}" for idx in range(100)]})
    rows = polars.DataFrame({"sku": ["A"], "qty": [1]})

    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "reports": mo_dataport.repeat_records(
                    employees.lazy(),
                    constants={"line_items": rows},
                )
            }
        },
        bundle_path=str(managed_tmp_dir / "bundle"),
    )

    section = bundle.report["sheets"][0]["repeat_sections"][0]
    assert section["record_count"] == 100
    assert section["record_source"].startswith("data/")
    assert "records" not in section
    assert len(section["record_bindings"]) == 1
    assert section["record_bindings"][0]["scalar_keys"] == ["customer_name"]
    assert len(section["dataframe_anchors"]) == 1
    assert "Name 99" not in json.dumps(bundle.report)


def test_compile_rejects_repeat_payload_that_is_not_list():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:A3",
    )

    with pytest.raises(TypeError, match="must be a list"):
        mo_dataport.compile(schema, {"Sheet1": {"reports": {"customer_name": "Acme"}}})


def test_compile_rejects_repeat_item_that_is_not_dict():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:A3",
    )

    with pytest.raises(TypeError, match="item 0.*must be an object/dict"):
        mo_dataport.compile(schema, {"Sheet1": {"reports": ["Acme"]}})


def test_compile_rejects_repeat_item_missing_required_value():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:A3",
    )

    with pytest.raises(KeyError, match="Sheet 'Sheet1.reports\\[0\\]' requires"):
        mo_dataport.compile(schema, {"Sheet1": {"reports": [{}]}})


def test_compile_preserves_static_content_after_repeat_end():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{customer_name:string}}"),
            "A3": _cell("A3", "{{reports:repeat-end}}"),
            "A4": _cell("A4", "Footer"),
        },
        dims="A1:A4",
    )

    bundle = mo_dataport.compile(schema, {"Sheet1": {"reports": [{"customer_name": "Acme"}]}})

    assert bundle.report["sheets"][0]["cells"]["A4"]["value"] == "Footer"


def test_compile_creates_ordered_sibling_repeat_sections():
    schema = _schema(
        {
            "A1": _cell("A1", "Title"),
            "A2": _cell("A2", "{{customers:repeat-start}}"),
            "A3": _cell("A3", "{{customer_name:string}}"),
            "A4": _cell("A4", "{{customers:repeat-end}}"),
            "A5": _cell("A5", "Between"),
            "A6": _cell("A6", "{{summaries:repeat-start}}"),
            "A7": _cell("A7", "{{region:string}}"),
            "A8": _cell("A8", "{{totals:dataframe-content}}"),
            "A9": _cell("A9", "{{summaries:repeat-end}}"),
            "A10": _cell("A10", "Footer"),
        },
        dims="A1:B10",
    )

    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "customers": [{"customer_name": "Acme"}],
                "summaries": [
                    {
                        "region": "North",
                        "totals": polars.DataFrame({"metric": ["Sales"], "value": [10]}),
                    }
                ],
            }
        },
    )

    sheet = bundle.report["sheets"][0]
    assert [section["key"] for section in sheet["repeat_sections"]] == [
        "customers",
        "summaries",
    ]
    assert sheet["cells"]["A1"]["value"] == "Title"
    assert sheet["cells"]["A5"]["value"] == "Between"
    assert sheet["cells"]["A10"]["value"] == "Footer"


def test_compile_stores_repeat_merges_as_relative_metadata():
    title = _cell("A2", "{{customer_name:string}}")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": title,
            "B2": shadow,
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:B3",
    )
    schema["sheets"][0]["merged_regions"] = ["A2:B2"]

    bundle = mo_dataport.compile(schema, {"Sheet1": {"reports": [{"customer_name": "Acme"}]}})

    section = bundle.report["sheets"][0]["repeat_sections"][0]
    assert bundle.report["sheets"][0]["merged_regions"] == []
    assert section["merged_regions"] == [
        {"min_row_offset": 0, "max_row_offset": 0, "min_col": 1, "max_col": 2}
    ]


def test_compile_rejects_repeat_merge_over_dataframe_content():
    anchor = _cell("A2", "{{line_items:dataframe-content}}")
    anchor["merged"] = True
    anchor["merge_anchor"] = "A2"
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": anchor,
            "B2": shadow,
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:B3",
    )
    schema["sheets"][0]["merged_regions"] = ["A2:B2"]

    with pytest.raises(ValueError, match="dataframe-content rows"):
        mo_dataport.compile(
            schema,
            {
                "Sheet1": {
                    "reports": [
                        {"line_items": polars.DataFrame({"A": [1], "B": [2]})}
                    ]
                }
            },
        )


def test_compile_without_bundle_path_creates_temp_directory():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})

    bundle = mo_dataport.compile(schema, {"Sheet1": {"rows": df}})

    path = Path(bundle.path)
    assert path.exists()
    assert path.is_dir()
    assert (path / "manifest.json").exists()
    paths = [source["path"] for source in bundle.manifest["dataframe_sources"]]
    assert paths == ["data/Sheet1__rows__dataframe-content.parquet"]
    assert (path / paths[0]).exists()


def test_compile_accepts_lazyframe_without_collecting_rows(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    polars.DataFrame({"A": [1, 2], "B": [3, 4]}).write_parquet(source_path)
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "bundle"

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.scan_parquet(source_path).select(["B", "A"])}},
        bundle_path=str(bundle_path),
    )

    source = bundle.manifest["dataframe_sources"][0]
    assert source["columns"] == ["B", "A"]
    assert source["rows"] == 2
    assert source["format"] == "parquet"
    assert (bundle_path / source["path"]).exists()


def test_report_bundle_write_copies_directory(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})
    target = managed_tmp_dir / "copied_bundle"

    bundle.write(str(target))
    copied = load_report_bundle(str(target))

    assert copied.report["sheets"][0]["cells"]["A1"]["value"] == "Alice"
    assert copied.manifest["bundle_format"] == "directory"


def test_compile_rejects_bundle_path_that_is_file(managed_tmp_dir: Path):
    target = managed_tmp_dir / "bundle-file"
    target.write_text("not a directory", encoding="utf-8")
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")

    with pytest.raises(ValueError, match="must be a directory"):
        mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}}, bundle_path=str(target))


def test_report_bundle_write_rejects_file_target(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})
    target = managed_tmp_dir / "bundle-file"
    target.write_text("not a directory", encoding="utf-8")

    with pytest.raises(ValueError, match="must be a directory"):
        bundle.write(str(target))


def test_load_rejects_missing_file_path_and_unsupported_metadata(managed_tmp_dir: Path):
    missing = managed_tmp_dir / "missing"
    file_path = managed_tmp_dir / "bundle-file"
    file_path.write_text("not a directory", encoding="utf-8")

    with pytest.raises(FileNotFoundError, match="Report bundle not found"):
        load_report_bundle(str(missing))
    with pytest.raises(ValueError, match="must be a directory"):
        load_report_bundle(str(file_path))

    bad_version = managed_tmp_dir / "bad-version"
    bad_version.mkdir()
    (bad_version / "manifest.json").write_text(
        json.dumps({"version": "0.0", "bundle_format": "directory"}),
        encoding="utf-8",
    )
    (bad_version / "report.json").write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported report bundle version"):
        load_report_bundle(str(bad_version))

    bad_format = managed_tmp_dir / "bad-format"
    bad_format.mkdir()
    (bad_format / "manifest.json").write_text(
        json.dumps({"version": "1.0", "bundle_format": "zip"}),
        encoding="utf-8",
    )
    (bad_format / "report.json").write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported report bundle format"):
        load_report_bundle(str(bad_format))


def test_compile_rejects_non_polars_dataframe_source():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")

    with pytest.raises(TypeError, match="expected a polars DataFrame or LazyFrame"):
        mo_dataport.compile(schema, {"Sheet1": {"rows": [{"A": 1}]}})


def test_compile_sanitizes_and_deduplicates_dataframe_source_ids(managed_tmp_dir: Path):
    schema = {
        "sheets": [
            {
                "name": "{{reports}}",
                "dimensions": "A1:A1",
                "merged_regions": [],
                "column_widths": {},
                "row_heights": {},
                "cells": {"A1": _cell("A1", "{{rows:dataframe-content}}")},
            }
        ]
    }
    df = polars.DataFrame({"A": [1]})

    bundle = mo_dataport.compile(
        schema,
        {"reports": {"North/West": {"rows": df}, "North?West": {"rows": df.clone()}}},
        bundle_path=str(managed_tmp_dir / "bundle"),
    )

    paths = [source["path"] for source in bundle.manifest["dataframe_sources"]]
    assert paths == [
        "data/North_West__rows__dataframe-content.parquet",
        "data/North_West__rows__dataframe-content_2.parquet",
    ]


def test_compile_serializes_date_scalar_in_report():
    schema = _schema({"A1": _cell("A1", "{{when:date}}")}, dims="A1:A1")

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"when": datetime.datetime(2024, 1, 2, 3, 4, 5)}},
    )

    assert bundle.report["sheets"][0]["cells"]["A1"]["value"] == "2024-01-02T03:04:05"


def test_xlsx_export_uses_bundle_data_without_expanding_report(managed_tmp_dir: Path):
    schema = _schema(
        {
            "A1": _cell("A1", "{{name:string}}"),
            "A2": _cell("A2", "{{headers:dataframe-header}}"),
            "A3": _cell("A3", "{{rows:dataframe-content}}"),
        },
        dims="A1:B3",
    )
    df = polars.DataFrame({"A": [1, 2], "B": [3, 4]})
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"name": "Alice", "headers": df, "rows": df}},
    )

    assert "A3" not in bundle.report["sheets"][0]["cells"]
    out = managed_tmp_dir / "out.xlsx"
    paths = mo_dataport.export(bundle, str(out), export_mode="streaming")

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert ws["A1"].value == "Alice"
    assert [ws["A2"].value, ws["B2"].value] == ["A", "B"]
    assert [ws["A3"].value, ws["B3"].value, ws["A4"].value, ws["B4"].value] == [
        1,
        3,
        2,
        4,
    ]
    wb.close()


def test_dataframe_placeholder_writes_headers_and_content(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe}}")}, dims="A1:B1")
    df = polars.DataFrame({"A": [1, 2], "B": [3, 4]}).lazy()
    bundle = mo_dataport.compile(schema, {"Sheet1": {"rows": df}})

    anchors = bundle.report["sheets"][0]["dataframe_anchors"]
    assert [anchor["placeholder_type"] for anchor in anchors] == [
        "dataframe-header",
        "dataframe-content",
    ]
    assert anchors[0]["start_row"] == 1
    assert anchors[1]["start_row"] == 2

    out = managed_tmp_dir / "combined.xlsx"
    paths = mo_dataport.export(bundle, str(out), export_mode="streaming")

    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert [ws["A1"].value, ws["B1"].value] == ["A", "B"]
    assert [ws["A2"].value, ws["B2"].value, ws["A3"].value, ws["B3"].value] == [
        1,
        3,
        2,
        4,
    ]
    wb.close()


def test_dataframe_header_only_does_not_write_source_file(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-header}}")}, dims="A1:B1")
    df = polars.DataFrame({"A": [1], "B": [2]}).lazy()

    bundle = mo_dataport.compile(schema, {"Sheet1": {"rows": df}})

    assert bundle.manifest["dataframe_sources"] == []
    anchors = bundle.report["sheets"][0]["dataframe_anchors"]
    assert len(anchors) == 1
    assert anchors[0]["placeholder_type"] == "dataframe-header"
    assert anchors[0]["source"] is None


def test_compile_stores_dataframe_column_layouts_and_expands_dimensions():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A"], "Amount": [10]})}},
        dataframe_options={
            "Sheet1": {
                "rows": {
                    "columns": {
                        "Name": {"occupation": 2, "alignment": "left"},
                        "Amount": {"occupation": 3, "alignment": "right"},
                        "Ignored": {"occupation": 9, "alignment": "center"},
                    }
                }
            }
        },
    )

    sheet = bundle.report["sheets"][0]
    anchor = sheet["dataframe_anchors"][0]
    assert sheet["dimensions"] == "A1:E1"
    assert anchor["column_layouts"] == [
        {
            "name": "Name",
            "start_col_offset": 0,
            "occupation": 2,
            "alignment": "left",
        },
        {
            "name": "Amount",
            "start_col_offset": 2,
            "occupation": 3,
            "alignment": "right",
        },
    ]


def test_compile_scopes_dataframe_column_layouts_by_placeholder_key():
    schema = _schema(
        {
            "A1": _cell("A1", "{{headers:dataframe-header}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:A2",
    )
    df = polars.DataFrame({"Amount": [10]})

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"headers": df, "rows": df}},
        dataframe_options={
            "Sheet1": {
                "headers": {"columns": {"Amount": {"occupation": 1, "alignment": "center"}}},
                "rows": {"columns": {"Amount": {"occupation": 3, "alignment": "right"}}},
            }
        },
    )

    header, content = bundle.report["sheets"][0]["dataframe_anchors"]
    assert header["column_layouts"][0]["occupation"] == 1
    assert header["column_layouts"][0]["alignment"] == "center"
    assert content["column_layouts"][0]["occupation"] == 3
    assert content["column_layouts"][0]["alignment"] == "right"


def test_compile_resolves_dynamic_row_page_breaks_after_dataframe_content():
    schema = _schema(
        {
            "A1": _cell("A1", "{{rows:dataframe-content}}"),
            "A2": _cell("A2", "Footer"),
        },
        dims="A1:A2",
    )
    schema["sheets"][0]["row_page_breaks"] = [1]

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2, 3]})}},
        dataframe_shift="vertical",
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["row_page_breaks"] == [1]
    assert sheet["resolved_row_page_breaks"] == [3]


def test_compile_resolves_dynamic_column_page_breaks_after_dataframe_expansion():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:C1")
    schema["sheets"][0]["column_page_breaks"] = [1]

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1], "B": [2]})}},
        dataframe_shift="horizontal",
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["column_page_breaks"] == [1]
    assert sheet["resolved_column_page_breaks"] == [2]


def test_compile_resolves_repeat_row_page_breaks_against_rendered_records():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
            "A4": _cell("A4", "After"),
            "A5": _cell("A5", "{{reports:repeat-end}}"),
        },
        dims="A1:A5",
    )
    schema["sheets"][0]["row_page_breaks"] = [2]

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"reports": [{"rows": polars.DataFrame({"A": [1, 2, 3]})}]}},
        dataframe_shift="vertical",
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["resolved_row_page_breaks"] == [3]


def test_compile_shifts_right_side_merge_away_from_dataframe_occupied_range(
    managed_tmp_dir: Path,
):
    title = _cell("B1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "B1"
    shadow = _cell("C1", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "B1"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "B1": title, "C1": shadow},
        dims="A1:C1",
    )
    schema["sheets"][0]["merged_regions"] = ["B1:C1"]

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A"]})}},
        dataframe_options={
            "Sheet1": {"rows": {"columns": {"Name": {"occupation": 2}}}}
        },
        dataframe_shift="horizontal",
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["merged_regions"] == ["C1:D1"]
    assert sheet["cells"]["C1"]["value"] == "Merged Title"

    out = managed_tmp_dir / "shifted-right.xlsx"
    paths = mo_dataport.export(bundle, str(out), export_mode="streaming")
    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert sorted(str(region) for region in ws.merged_cells.ranges) == [
        "A1:B1",
        "C1:D1",
    ]
    assert ws["A1"].value == "A"
    assert ws["C1"].value == "Merged Title"
    wb.close()


def test_compile_shifts_bottom_merge_below_dataframe_rows(managed_tmp_dir: Path):
    title = _cell("A2", "Totals")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "A2": title, "B2": shadow},
        dims="A1:B2",
    )
    schema["sheets"][0]["merged_regions"] = ["A2:B2"]

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A", "B"]})}},
        dataframe_shift="vertical",
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["merged_regions"] == ["A3:B3"]
    assert sheet["cells"]["A3"]["value"] == "Totals"

    out = managed_tmp_dir / "shifted-bottom.xlsx"
    paths = mo_dataport.export(bundle, str(out), export_mode="streaming")
    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert sorted(str(region) for region in ws.merged_cells.ranges) == ["A3:B3"]
    assert [ws["A1"].value, ws["A2"].value, ws["A3"].value] == ["A", "B", "Totals"]
    wb.close()


def test_compile_shifts_diagonal_merge_right_and_down(managed_tmp_dir: Path):
    title = _cell("B2", "Notes")
    title["merged"] = True
    title["merge_anchor"] = "B2"
    shadow = _cell("C2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "B2"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "B2": title, "C2": shadow},
        dims="A1:C2",
    )
    schema["sheets"][0]["merged_regions"] = ["B2:C2"]

    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A", "B"]})}},
        dataframe_options={
            "Sheet1": {"rows": {"columns": {"Name": {"occupation": 2}}}}
        },
    )

    sheet = bundle.report["sheets"][0]
    assert sheet["merged_regions"] == ["C3:D3"]
    assert sheet["cells"]["C3"]["value"] == "Notes"

    out = managed_tmp_dir / "shifted-diagonal.xlsx"
    paths = mo_dataport.export(bundle, str(out), export_mode="streaming")
    wb = openpyxl.load_workbook(paths[0], data_only=True)
    ws = wb["Sheet1"]
    assert sorted(str(region) for region in ws.merged_cells.ranges) == [
        "A1:B1",
        "A2:B2",
        "C3:D3",
    ]
    assert [ws["A1"].value, ws["A2"].value, ws["C3"].value] == ["A", "B", "Notes"]
    wb.close()


def test_compile_shifts_repeat_record_cells_around_dataframe_output():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "Customer: {{customer_name:string}}"),
            "A3": _cell("A3", "{{region:string}}"),
            "A4": _cell("A4", "{{line_items:dataframe-header}}"),
            "C4": _cell("C4", "{{note:string}}"),
            "A5": _cell("A5", "{{line_items:dataframe-content}}"),
            "A6": _cell("A6", "{{footer:string}}"),
            "A7": _cell("A7", "{{reports:repeat-end}}"),
        },
        dims="A1:C7",
    )

    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "reports": [
                    {
                        "customer_name": "Acme",
                        "region": "North",
                        "line_items": polars.DataFrame({"Item": ["A", "B"]}),
                        "note": "Side 1",
                        "footer": "Below 1",
                    },
                    {
                        "customer_name": "Globex",
                        "region": "South",
                        "line_items": polars.DataFrame({"Item": ["C"]}),
                        "note": "Side 2",
                        "footer": "Below 2",
                    },
                ]
            }
        },
        dataframe_options={
            "Sheet1": {"line_items": {"columns": {"Item": {"occupation": 2}}}}
        },
    )

    section = bundle.report["sheets"][0]["repeat_sections"][0]
    first = section["records"][0]
    second = section["records"][1]

    assert first["block_height"] == 6
    assert second["block_height"] == 5
    assert {(item["value"], item["row_offset"], item["start_col"]) for item in first["cells"]} == {
        ("Customer: Acme", 0, 1),
        ("North", 1, 1),
        ("Side 1", 2, 4),
        ("Below 1", 5, 1),
    }
    assert {(item["value"], item["row_offset"], item["start_col"]) for item in second["cells"]} == {
        ("Customer: Globex", 0, 1),
        ("South", 1, 1),
        ("Side 2", 2, 4),
        ("Below 2", 4, 1),
    }
    assert [anchor["start_row_offset"] for anchor in first["dataframe_anchors"]] == [2, 3]
    assert [anchor["start_col"] for anchor in first["dataframe_anchors"]] == [1, 1]


def test_pdf_dataframe_rows_use_shifted_template_merges():
    title = _cell("A2", "Totals")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "A2": title, "B2": shadow},
        dims="A1:B2",
    )
    schema["sheets"][0]["merged_regions"] = ["A2:B2"]
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A", "B"]})}},
    )

    rows = list(
        _dataframe_pdf_rows(
            bundle,
            _data_source_map(bundle),
            bundle.report["sheets"][0],
            batch_size=1,
        )
    )

    assert rows[2]["cells"][1]["value"] == "Totals"
    assert rows[2]["merges"] == [
        {"min_row_offset": 0, "max_row_offset": 0, "min_col": 1, "max_col": 2}
    ]


def test_pdf_chunks_keep_shifted_vertical_merges_together():
    title = _cell("A2", "Totals")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    shadow_cells = {}
    for coord in ["B2", "A3", "B3"]:
        shadow = _cell(coord, None)
        shadow["merged"] = True
        shadow["merge_anchor"] = "A2"
        shadow_cells[coord] = shadow
    schema = _schema(
        {
            "A1": _cell("A1", "{{rows:dataframe-content}}"),
            "A2": title,
            **shadow_cells,
        },
        dims="A1:B3",
    )
    schema["sheets"][0]["merged_regions"] = ["A2:B3"]
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A", "B"]})}},
    )
    sheet = bundle.report["sheets"][0]
    rows = _dataframe_pdf_rows(
        bundle,
        _data_source_map(bundle),
        sheet,
        batch_size=1,
    )

    tables = list(
        _chunked_row_tables(
            sheet,
            rows,
            1,
            2,
            500,
            _FontResolver(),
            streaming_chunk_rows=1,
        )
    )

    assert any(("SPAN", (0, 0), (1, 1)) in table._spanCmds for table in tables)


def test_compile_dataframe_shift_horizontal_only_rejects_vertical_collision():
    title = _cell("A2", "Totals")
    title["merged"] = True
    title["merge_anchor"] = "A2"
    shadow = _cell("B2", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A2"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "A2": title, "B2": shadow},
        dims="A1:B2",
    )
    schema["sheets"][0]["merged_regions"] = ["A2:B2"]

    with pytest.raises(ValueError, match="must not overlap dataframe output ranges"):
        mo_dataport.compile(
            schema,
            {"Sheet1": {"rows": polars.DataFrame({"Name": ["A", "B"]})}},
            dataframe_shift="horizontal",
        )


def test_compile_dataframe_shift_vertical_only_rejects_horizontal_collision():
    title = _cell("B1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "B1"
    shadow = _cell("C1", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "B1"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "B1": title, "C1": shadow},
        dims="A1:C1",
    )
    schema["sheets"][0]["merged_regions"] = ["B1:C1"]

    with pytest.raises(ValueError, match="must not overlap dataframe output ranges"):
        mo_dataport.compile(
            schema,
            {"Sheet1": {"rows": polars.DataFrame({"Name": ["A"]})}},
            dataframe_options={
                "Sheet1": {"rows": {"columns": {"Name": {"occupation": 2}}}}
            },
            dataframe_shift="vertical",
        )


def test_compile_dataframe_shift_none_uses_strict_overlap_validation():
    title = _cell("B1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "B1"
    shadow = _cell("C1", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "B1"
    schema = _schema(
        {"A1": _cell("A1", "{{rows:dataframe-content}}"), "B1": title, "C1": shadow},
        dims="A1:C1",
    )
    schema["sheets"][0]["merged_regions"] = ["B1:C1"]

    with pytest.raises(ValueError, match="must not overlap dataframe output ranges"):
        mo_dataport.compile(
            schema,
            {"Sheet1": {"rows": polars.DataFrame({"Name": ["A"]})}},
            dataframe_options={
                "Sheet1": {"rows": {"columns": {"Name": {"occupation": 2}}}}
            },
            dataframe_shift="none",
        )


def test_compile_rejects_invalid_dataframe_shift_mode():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")

    with pytest.raises(ValueError, match="dataframe_shift"):
        mo_dataport.compile(
            schema,
            {"Sheet1": {"rows": polars.DataFrame({"Name": ["A"]})}},
            dataframe_shift="diagonal",
        )


def test_compile_rejects_template_merge_covering_dataframe_anchor():
    anchor = _cell("A1", "{{rows:dataframe-content}}")
    anchor["merged"] = True
    anchor["merge_anchor"] = "A1"
    shadow = _cell("B1", None)
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema({"A1": anchor, "B1": shadow}, dims="A1:B1")
    schema["sheets"][0]["merged_regions"] = ["A1:B1"]

    with pytest.raises(ValueError, match="must not overlap dataframe output ranges"):
        mo_dataport.compile(schema, {"Sheet1": {"rows": polars.DataFrame({"Name": ["A"]})}})


def test_compile_rejects_invalid_dataframe_column_layout_options():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"Name": ["A"]})

    with pytest.raises(ValueError, match="positive integer"):
        mo_dataport.compile(
            schema,
            {"Sheet1": {"rows": df}},
            dataframe_options={
                "Sheet1": {"rows": {"columns": {"Name": {"occupation": 0}}}}
            },
        )
    with pytest.raises(ValueError, match="alignment"):
        mo_dataport.compile(
            schema,
            {"Sheet1": {"rows": df}},
            dataframe_options={
                "Sheet1": {
                    "rows": {"columns": {"Name": {"alignment": "justify"}}}
                }
            },
        )


def test_old_dataframe_headers_spelling_is_not_a_placeholder():
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-headers}}")}, dims="A1:A1")

    assert mo_dataport.inputs(schema) == {"Sheet1": {}}


def test_compile_deduplicates_same_lazyframe_object(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    polars.DataFrame({"A": [1], "B": [2]}).write_parquet(source_path)
    rows = polars.scan_parquet(source_path)
    schema = _schema(
        {
            "A1": _cell("A1", "{{headers:dataframe-header}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:B2",
    )

    bundle = mo_dataport.compile(schema, {"Sheet1": {"headers": rows, "rows": rows}})

    assert len(bundle.manifest["dataframe_sources"]) == 1
    content_anchor = [
        anchor
        for anchor in bundle.report["sheets"][0]["dataframe_anchors"]
        if anchor["placeholder_type"] == "dataframe-content"
    ][0]
    assert content_anchor["source"] == bundle.manifest["dataframe_sources"][0]["path"]


def test_compile_keeps_distinct_lazyframe_objects_separate(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    polars.DataFrame({"A": [1], "B": [2]}).write_parquet(source_path)
    schema = _schema(
        {
            "A1": _cell("A1", "{{first:dataframe-content}}"),
            "A3": _cell("A3", "{{second:dataframe-content}}"),
        },
        dims="A1:B3",
    )

    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "first": polars.scan_parquet(source_path),
                "second": polars.scan_parquet(source_path),
            }
        },
    )

    assert len(bundle.manifest["dataframe_sources"]) == 2


def test_export_accepts_bundle_path(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "report_bundle"
    mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}}, bundle_path=str(bundle_path))
    out = managed_tmp_dir / "out.xlsx"

    mo_dataport.export(str(bundle_path), str(out))

    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb["Sheet1"]["A1"].value == "Alice"
    wb.close()


def test_xlsx_export_preserves_gridline_visibility(managed_tmp_dir: Path):
    schema = _schema_with_options(
        {"A1": _cell("A1", "No gridlines")},
        dims="A1:A1",
        show_gridlines=False,
    )
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "no-gridlines.xlsx"

    mo_dataport.export(bundle, str(out))

    wb = openpyxl.load_workbook(out)
    assert wb["Sheet1"].sheet_view.showGridLines is False
    wb.close()


def test_xlsx_export_preserves_merged_region_border(managed_tmp_dir: Path):
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thick", "color": "FF000000"},
        "left": {"style": "medium", "color": "FF000000"},
        "right": {"style": "dashed", "color": "FF000000"},
    }
    shadow = _cell("B2", "")
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema({"A1": title, "B2": shadow}, dims="A1:B2")
    schema["sheets"][0]["merged_regions"] = ["A1:B2"]
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "merged-border.xlsx"

    mo_dataport.export(bundle, str(out))

    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    assert ws["A1"].border.top.style == "thin"
    assert ws["A2"].border.bottom.style == "thick"
    assert ws["A1"].border.left.style == "medium"
    assert ws["B1"].border.right.style == "dashed"
    wb.close()


def test_export_pdf_creates_nonempty_pdf(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})
    out = managed_tmp_dir / "out.pdf"

    mo_dataport.export(bundle, str(out), format="pdf")

    assert out.exists()
    assert out.stat().st_size > 0
    assert out.read_bytes().startswith(b"%PDF")


def test_export_image_is_reserved(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})

    with pytest.raises(NotImplementedError, match="image"):
        mo_dataport.export(bundle, str(managed_tmp_dir / "out.png"), format="image")


def test_pdf_export_renders_parquet_backed_dataframe(managed_tmp_dir: Path):
    source_path = managed_tmp_dir / "source.parquet"
    polars.DataFrame({"A": [1, 2], "B": [3, 4]}).write_parquet(source_path)
    schema = _schema(
        {
            "A1": _cell("A1", "{{headers:dataframe-header}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:B2",
    )
    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "headers": polars.scan_parquet(source_path).select(["B", "A"]),
                "rows": polars.scan_parquet(source_path).select(["B", "A"]),
            }
        },
    )
    out = managed_tmp_dir / "table.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", streaming_chunk_rows=1)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_dataframe_rows_emit_occupation_spans_and_alignment():
    schema = _schema(
        {
            "A1": _cell("A1", "{{headers:dataframe-header}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
        },
        dims="A1:A2",
    )
    df = polars.DataFrame(
        {"Employee Name": ["Alice"], "Department": ["Finance"], "Amount": [12]}
    )
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"headers": df, "rows": df}},
        dataframe_options={
            "Sheet1": {
                "headers": {
                    "columns": {
                        "Employee Name": {"occupation": 2, "alignment": "center"},
                        "Amount": {"occupation": 2, "alignment": "center"},
                    }
                },
                "rows": {
                    "columns": {
                        "Employee Name": {"occupation": 2, "alignment": "left"},
                        "Amount": {"occupation": 2, "alignment": "right"},
                    }
                },
            }
        },
    )
    sheet = bundle.report["sheets"][0]

    rows = list(
        _dataframe_pdf_rows(
            bundle,
            _data_source_map(bundle),
            sheet,
            batch_size=1,
        )
    )

    assert rows[0]["cells"][1]["value"] == "Employee Name"
    assert rows[0]["cells"][4]["value"] == "Amount"
    assert rows[1]["cells"][1]["value"] == "Alice"
    assert rows[1]["cells"][4]["value"] == 12
    assert rows[1]["cells"][1]["alignment"]["horizontal"] == "left"
    assert rows[1]["cells"][4]["alignment"]["horizontal"] == "right"
    assert rows[0]["merges"] == [
        {"min_row_offset": 0, "max_row_offset": 0, "min_col": 1, "max_col": 2},
        {"min_row_offset": 0, "max_row_offset": 0, "min_col": 4, "max_col": 5},
    ]
    table = _row_chunk_table(sheet, rows, 1, 5, 500, _FontResolver())
    assert ("SPAN", (0, 0), (1, 0)) in table._spanCmds
    assert ("SPAN", (3, 1), (4, 1)) in table._spanCmds


def test_pdf_table_uses_shared_expanded_xlsx_render_plan_for_dataframe_styles():
    header = _cell("A1", "{{headers:dataframe-header}}")
    header["fill"] = {"pattern_type": "solid", "fg_color": "FF1F4E79", "bg_color": None}
    header["font"] = dict(header["font"])
    header["font"]["color"] = "FFFFFFFF"
    header["alignment"] = dict(header["alignment"])
    header["alignment"]["vertical"] = "center"
    content = _cell("A2", "{{rows:dataframe-content}}")
    content["fill"] = {"pattern_type": "solid", "fg_color": "FFD9EAF7", "bg_color": None}
    content["font"] = dict(content["font"])
    content["font"]["color"] = "FF203040"
    content["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thin", "color": "FF000000"},
        "left": {"style": "thin", "color": "FF000000"},
        "right": {"style": "thin", "color": "FF000000"},
    }
    schema = _schema({"A1": header, "A2": content}, dims="A1:A2")
    df = polars.DataFrame({"Amount": [12]})
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"headers": df, "rows": df}},
        dataframe_options={
            "Sheet1": {
                "headers": {"columns": {"Amount": {"occupation": 2, "alignment": "center"}}},
                "rows": {"columns": {"Amount": {"occupation": 2, "alignment": "right"}}},
            }
        },
    )

    sheet, cells = _expanded_cells(
        bundle, bundle.report["sheets"][0], streaming_chunk_rows=1
    )
    commands = _table_style(sheet, cells, 1, 1, 2, 2, _FontResolver()).getCommands()

    assert sheet["merged_regions"] == ["A1:B1", "A2:B2"]
    assert cells[(1, 1)]["value"] == "Amount"
    assert cells[(1, 1)]["alignment"]["horizontal"] == "center"
    assert cells[(2, 1)]["value"] == 12
    assert cells[(2, 1)]["alignment"]["horizontal"] == "right"
    assert ("SPAN", (0, 0), (1, 0)) in commands
    assert ("SPAN", (0, 1), (1, 1)) in commands
    assert any(item[0] == "BACKGROUND" and item[1] == (0, 1) and item[2] == (1, 1) for item in commands)
    assert any(item[0] == "TEXTCOLOR" and item[1] == (0, 1) and item[2] == (1, 1) for item in commands)
    assert ("ALIGN", (0, 1), (1, 1), "RIGHT") in commands
    assert ("LEFTPADDING", (0, 1), (1, 1), 0) in commands


def test_pdf_column_widths_follow_xlsx_widths_with_deterministic_scaling():
    sheet = _schema(
        {"A1": _cell("A1", "A"), "B1": _cell("B1", "B")},
        dims="A1:B1",
    )["sheets"][0]
    sheet["column_widths"] = {"A": 20.0, "B": 10.0}

    assert _column_widths(sheet, {}, 1, 1, 2, 1, 1_000) == [140.0, 70.0]
    assert _column_widths(sheet, {}, 1, 1, 2, 1, 105) == [70.0, 35.0]


def test_pdf_fast_grid_plan_prunes_shifted_invisible_dataframe_columns():
    title = _cell("A1", "{{report_title:string}}")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    generated = _cell("A2", "Generated: {{generated_on:date}}")
    generated["merged"] = True
    generated["merge_anchor"] = "A2"
    rows = _cell("D2", "Rows: {{row_count:number}}")
    rows["merged"] = True
    rows["merge_anchor"] = "D2"
    header = _cell("A4", "{{bench_data:dataframe-header}}")
    content = _cell("A5", "{{bench_data:dataframe-content}}")
    schema = _schema(
        {
            "A1": title,
            "B1": _cell("B1", None),
            "C1": _cell("C1", None),
            "D1": _cell("D1", None),
            "E1": _cell("E1", None),
            "A2": generated,
            "B2": _cell("B2", None),
            "C2": _cell("C2", None),
            "D2": rows,
            "E2": _cell("E2", None),
            "A3": _cell("A3", None),
            "B3": _cell("B3", None),
            "C3": _cell("C3", None),
            "D3": _cell("D3", None),
            "E3": _cell("E3", None),
            "A4": header,
            "B4": _cell("B4", None),
            "C4": _cell("C4", None),
            "D4": _cell("D4", None),
            "E4": _cell("E4", None),
            "A5": content,
            "B5": _cell("B5", None),
            "C5": _cell("C5", None),
            "D5": _cell("D5", None),
            "E5": _cell("E5", None),
        },
        dims="A1:E5",
    )
    schema["sheets"][0]["merged_regions"] = ["A1:E1", "A2:C2", "D2:E2"]
    schema["sheets"][0]["column_widths"] = {
        "A": 12.0,
        "B": 28.0,
        "C": 18.0,
        "D": 14.0,
        "E": 16.0,
    }
    schema["sheets"][0]["row_heights"] = {
        "1": 36.0,
        "2": 22.0,
        "3": 8.0,
        "4": 22.0,
        "5": 18.0,
    }
    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "report_title": "Benchmark Report",
                "generated_on": datetime.date(2024, 1, 1),
                "row_count": 2,
                "bench_data": polars.DataFrame(
                    {
                        "id": [1, 2],
                        "name": ["A", "B"],
                        "category": ["C", "D"],
                        "value": [1.0, 2.0],
                        "date": ["2024-01-01", "2024-01-02"],
                    }
                ),
            }
        },
    )

    assert bundle.report["sheets"][0]["dimensions"].startswith("A1:I")
    plan = _fast_grid_plan(
        bundle,
        bundle.report["sheets"][0],
        column_width_mode=None,
        row_height_mode=None,
        default_column_width=None,
        default_row_height=None,
        available_width=500,
        repeat_dataframe_headers=False,
    )

    assert plan is not None
    assert (plan["min_col"], plan["max_col"]) == (1, 5)


def test_pdf_fast_grid_plan_falls_back_for_wrapped_text():
    cell = _cell("A1", "line one\nline two")
    cell["alignment"] = dict(cell["alignment"])
    cell["alignment"]["wrap_text"] = True
    bundle = mo_dataport.compile(_schema({"A1": cell}, dims="A1:A1"), {"Sheet1": {}})

    plan = _fast_grid_plan(
        bundle,
        bundle.report["sheets"][0],
        column_width_mode=None,
        row_height_mode=None,
        default_column_width=None,
        default_row_height=None,
        available_width=500,
        repeat_dataframe_headers=False,
    )

    assert plan is None


def test_pdf_fast_grid_plan_keeps_styled_blank_columns():
    styled_blank = _cell("C1", None)
    styled_blank["fill"] = {
        "pattern_type": "solid",
        "fg_color": "FFFF0000",
        "bg_color": None,
    }
    bundle = mo_dataport.compile(
        _schema({"A1": _cell("A1", "Visible"), "C1": styled_blank}, dims="A1:C1"),
        {"Sheet1": {}},
    )

    plan = _fast_grid_plan(
        bundle,
        bundle.report["sheets"][0],
        column_width_mode=None,
        row_height_mode=None,
        default_column_width=None,
        default_row_height=None,
        available_width=500,
        repeat_dataframe_headers=False,
    )

    assert plan is not None
    assert (plan["min_col"], plan["max_col"]) == (1, 3)


def test_pdf_export_streams_dataframe_occupation_chunks(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"Name": ["A", "B", "C"]})}},
        dataframe_options={
            "Sheet1": {"rows": {"columns": {"Name": {"occupation": 2}}}}
        },
    )
    out = managed_tmp_dir / "occupied.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", streaming_chunk_rows=1)

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_export_rejects_hug_sizing_for_dataframe_content(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1]})}},
    )

    with pytest.raises(ValueError, match="hug.*dataframe-content"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "hug.pdf"),
            format="pdf",
            column_width_mode="hug",
        )


def test_pdf_export_allows_row_hug_for_dataframe_content(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2]})}},
    )
    out = str(managed_tmp_dir / "row_hug.pdf")
    mo_dataport.export(bundle, out, format="pdf", row_height_mode="hug")
    assert Path(out).exists()


def test_pdf_repeat_dataframe_occupation_spans():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{name:string}}"),
            "A3": _cell("A3", "{{line_items:dataframe}}"),
            "A4": _cell("A4", ""),
            "A5": _cell("A5", "{{reports:repeat-end}}"),
        },
        dims="A1:A5",
    )
    bundle = mo_dataport.compile(
        schema,
        {
            "Sheet1": {
                "reports": [
                    {
                        "name": "Acme",
                        "line_items": polars.DataFrame({"sku": ["A"], "qty": [1]}),
                    }
                ]
            }
        },
        dataframe_options={
            "Sheet1": {
                "line_items": {
                    "columns": {
                        "sku": {"occupation": 2},
                        "qty": {"occupation": 2},
                    }
                }
            }
        },
    )
    sheet = bundle.report["sheets"][0]
    section = sheet["repeat_sections"][0]
    record = section["records"][0]

    rows = list(
        _repeat_record_rows(
            bundle,
            _data_source_map(bundle),
            record,
            block_height=section["block_height"],
            merges=[],
            cell_templates=section["cell_templates"],
            batch_size=1,
        )
    )

    table = _row_chunk_table(sheet, rows, 1, 4, 500, _FontResolver())
    assert ("SPAN", (0, 1), (1, 1)) in table._spanCmds
    assert ("SPAN", (2, 2), (3, 2)) in table._spanCmds


def test_fidelity_inherits_anchor_row_height_for_dataframe_rows(managed_tmp_dir: Path):
    schema = _schema({"A2": _cell("A2", "{{rows:dataframe-content}}")}, dims="A1:B2")
    schema["sheets"][0]["row_heights"] = {"2": 24.0}
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2], "B": [3, 4]})}},
    )
    out = managed_tmp_dir / "out.xlsx"

    mo_dataport.export(bundle, str(out))

    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    assert ws.row_dimensions[2].height == pytest.approx(24.0)
    assert ws.row_dimensions[3].height == pytest.approx(24.0)
    wb.close()


def test_fidelity_export_writes_resolved_page_breaks(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "Title")}, dims="A1:B2")
    schema["sheets"][0]["row_page_breaks"] = [1]
    schema["sheets"][0]["column_page_breaks"] = [1]
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "page-breaks.xlsx"

    mo_dataport.export(bundle, str(out), export_mode="fidelity")

    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    assert [item.id for item in ws.row_breaks.brk] == [1]
    assert [item.id for item in ws.col_breaks.brk] == [1]
    wb.close()


def test_pdf_export_handles_merges_and_basic_styles(managed_tmp_dir: Path):
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["font"] = dict(title["font"])
    title["font"]["bold"] = True
    title["fill"] = {"pattern_type": "solid", "fg_color": "FF003366", "bg_color": None}
    title["alignment"] = dict(title["alignment"])
    title["alignment"]["horizontal"] = "center"
    shadow = _cell("B1", "")
    shadow["merged"] = True
    shadow["merge_anchor"] = "A1"
    schema = _schema({"A1": title, "B1": shadow}, dims="A1:B1")
    schema["sheets"][0]["merged_regions"] = ["A1:B1"]
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "styled.pdf"

    mo_dataport.export(bundle, str(out), format="pdf")

    assert out.exists()
    assert out.stat().st_size > 0


def test_pdf_sheet_flowables_insert_manual_page_break_for_static_rows():
    schema = _schema(
        {
            "A1": _cell("A1", "Top"),
            "A2": _cell("A2", "Bottom"),
        },
        dims="A1:A2",
    )
    schema["sheets"][0]["row_page_breaks"] = [1]
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=10,
            font_resolver=_FontResolver(),
        )
    )

    assert any(isinstance(item, PageBreak) for item in flowables)


def test_pdf_sheet_flowables_insert_manual_page_break_for_streamed_dataframe_rows():
    schema = _schema(
        {
            "A1": _cell("A1", "{{rows:dataframe-content}}"),
            "A2": _cell("A2", "After"),
        },
        dims="A1:A2",
    )
    schema["sheets"][0]["row_page_breaks"] = [1]
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2, 3]})}},
        dataframe_shift="vertical",
    )

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=1,
            font_resolver=_FontResolver(),
        )
    )

    assert any(isinstance(item, PageBreak) for item in flowables)


def test_pdf_repeat_dataframe_headers_default_keeps_repeat_rows_disabled():
    schema = _schema(
        {
            "A1": _cell("A1", "{{rows:dataframe}}"),
        },
        dims="A1:A1",
    )
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2, 3]})}},
    )

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=1,
            font_resolver=_FontResolver(),
        )
    )
    tables = [item for item in flowables if hasattr(item, "repeatRows")]
    assert len(tables) >= 2
    assert all(table.repeatRows == 0 for table in tables)


def test_pdf_repeat_dataframe_headers_does_not_repeat_for_plain_chunk_splits():
    schema = _schema(
        {
            "A1": _cell("A1", "{{rows:dataframe}}"),
        },
        dims="A1:A1",
    )
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2, 3]})}},
    )

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=1,
            font_resolver=_FontResolver(),
            repeat_dataframe_headers=True,
        )
    )
    tables = [item for item in flowables if hasattr(item, "repeatRows")]
    assert len(tables) == 1
    assert tables[0].repeatRows == 1


def test_pdf_repeat_dataframe_headers_with_manual_break_keeps_pagebreak_and_header_repeat():
    sheet = _schema({"A1": _cell("A1", "x")}, dims="A1:A1")["sheets"][0]
    header_row = {
        "cells": {1: _cell("A1", "Col")},
        "merges": [],
        "is_dataframe_header_row": True,
    }
    rows = iter(
        [
            {
                "cells": {1: _cell("A1", "Col")},
                "merges": [],
                "row_idx": 1,
                "dataframe_header_sources": ["key:rows"],
                "is_dataframe_header_row": True,
            },
            {
                "cells": {1: _cell("A1", 1)},
                "merges": [],
                "row_idx": 2,
                "dataframe_content_sources": ["key:rows"],
                "dataframe_content_start_sources": ["key:rows"],
                "dataframe_header_rows_by_source": {"key:rows": header_row},
            },
            {
                "cells": {1: _cell("A1", 2)},
                "merges": [],
                "row_idx": 3,
                "dataframe_content_sources": ["key:rows"],
                "dataframe_content_start_sources": [],
                "dataframe_header_rows_by_source": {"key:rows": header_row},
            },
        ]
    )

    flowables = list(
        _chunked_row_flowables(
            sheet,
            rows,
            1,
            1,
            500,
            _FontResolver(),
            streaming_chunk_rows=10,
            page_breaks={2},
            repeat_dataframe_headers=True,
        )
    )
    assert any(isinstance(item, PageBreak) for item in flowables)
    tables = [item for item in flowables if hasattr(item, "repeatRows")]
    assert len(tables) >= 2
    assert tables[1].repeatRows == 1


def test_pdf_repeat_dataframe_headers_does_not_fake_header_without_header_anchor():
    schema = _schema(
        {
            "A1": _cell("A1", "{{rows:dataframe-content}}"),
        },
        dims="A1:A1",
    )
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": polars.DataFrame({"A": [1, 2, 3]})}},
    )

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=1,
            font_resolver=_FontResolver(),
            repeat_dataframe_headers=True,
        )
    )
    tables = [item for item in flowables if hasattr(item, "repeatRows")]
    assert len(tables) == 1
    assert tables[0].repeatRows == 0


def test_pdf_repeat_dataframe_headers_repeat_section_does_not_repeat_for_plain_chunk_splits():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{rows:dataframe}}"),
            "A3": _cell("A3", "{{reports:repeat-end}}"),
        },
        dims="A1:A3",
    )
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"reports": [{"rows": polars.DataFrame({"A": [1, 2, 3]})}]}},
    )

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=1,
            font_resolver=_FontResolver(),
            repeat_dataframe_headers=True,
        )
    )
    tables = [item for item in flowables if hasattr(item, "repeatRows")]
    assert len(tables) == 1
    assert tables[0].repeatRows == 1


def test_pdf_sheet_flowables_insert_manual_page_break_for_repeat_rows():
    schema = _schema(
        {
            "A1": _cell("A1", "{{reports:repeat-start}}"),
            "A2": _cell("A2", "{{rows:dataframe-content}}"),
            "A4": _cell("A4", "After"),
            "A5": _cell("A5", "{{reports:repeat-end}}"),
        },
        dims="A1:A5",
    )
    schema["sheets"][0]["row_page_breaks"] = [2]
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"reports": [{"rows": polars.DataFrame({"A": [1, 2, 3]})}]}},
        dataframe_shift="vertical",
    )

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=1,
            font_resolver=_FontResolver(),
        )
    )

    assert any(isinstance(item, PageBreak) for item in flowables)


def test_pdf_column_page_breaks_do_not_change_flow():
    schema = _schema({"A1": _cell("A1", "Only")}, dims="A1:A1")
    schema["sheets"][0]["column_page_breaks"] = [1]
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})

    flowables = list(
        _sheet_flowables(
            bundle,
            bundle.report["sheets"][0],
            available_width=500,
            column_width_mode=None,
            row_height_mode=None,
            default_column_width=None,
            default_row_height=None,
            streaming_chunk_rows=10,
            font_resolver=_FontResolver(),
        )
    )

    assert not any(isinstance(item, PageBreak) for item in flowables)


def test_pdf_table_style_applies_merged_region_border_to_span():
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thick", "color": "FF000000"},
        "left": {"style": "medium", "color": "FF000000"},
        "right": {"style": "dashed", "color": "FF000000"},
    }
    sheet = _schema({"A1": title}, dims="A1:B2")["sheets"][0]
    sheet["merged_regions"] = ["A1:B2"]

    commands = _table_style(
        sheet, {(1, 1): title}, 1, 1, 2, 2, _FontResolver()
    ).getCommands()

    assert _has_line_command(commands, "LINEABOVE", 0.5)
    assert _has_line_command(commands, "LINEBELOW", 1.5)
    assert _has_line_command(commands, "LINEBEFORE", 1.0)
    assert _has_line_command(commands, "LINEAFTER", 0.75)


def test_pdf_table_style_applies_merged_region_fill_to_span():
    title = _cell("A1", "Merged Title")
    title["merged"] = True
    title["merge_anchor"] = "A1"
    title["fill"] = {"pattern_type": "solid", "fg_color": "FF003366", "bg_color": None}
    sheet = _schema({"A1": title}, dims="A1:D1")["sheets"][0]
    sheet["merged_regions"] = ["A1:D1"]

    commands = _table_style(
        sheet, {(1, 1): title}, 1, 1, 4, 1, _FontResolver()
    ).getCommands()

    assert any(item[0] == "SPAN" and item[1] == (0, 0) and item[2] == (3, 0) for item in commands)
    assert any(item[0] == "BACKGROUND" and item[1] == (0, 0) and item[2] == (3, 0) for item in commands)


def test_pdf_table_style_does_not_draw_default_grid_for_empty_cells():
    cell = _cell("A1", "Visible")
    sheet = _schema({"A1": cell}, dims="A1:B2")["sheets"][0]

    commands = _table_style(
        sheet, {(1, 1): cell}, 1, 1, 2, 2, _FontResolver()
    ).getCommands()

    assert not any(item[0] == "GRID" for item in commands)
    assert not any(
        item[0] in {"LINEABOVE", "LINEBELOW", "LINEBEFORE", "LINEAFTER"}
        for item in commands
    )


def test_pdf_table_style_uses_registered_custom_font():
    cell = _cell("A1", "Custom Font")
    cell["font"] = dict(cell["font"])
    cell["font"]["name"] = "Vera"
    sheet = _schema({"A1": cell}, dims="A1:A1")["sheets"][0]
    resolver = _FontResolver({"Vera": _vera_font_path()})

    commands = _table_style(sheet, {(1, 1): cell}, 1, 1, 1, 1, resolver).getCommands()

    font_commands = [item for item in commands if item[0] == "FONTNAME"]
    assert font_commands
    assert str(font_commands[0][3]).startswith("Mindoff-Vera-regular-")


def test_pdf_lazy_flowables_accepts_split_insertion():
    flowables = _LazyFlowables(["first", "second"])

    flowables[0:0] = ["split"]
    flowables.insert(0, "retry")

    assert [flowables[0], flowables[1], flowables[2], flowables[3]] == [
        "retry",
        "split",
        "first",
        "second",
    ]


def test_pdf_export_accepts_custom_fonts(managed_tmp_dir: Path):
    cell = _cell("A1", "Custom Font")
    cell["font"] = dict(cell["font"])
    cell["font"]["name"] = "Vera"
    schema = _schema({"A1": cell}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "custom-font.pdf"

    mo_dataport.export(
        bundle,
        str(out),
        format="pdf",
        fonts={"Vera": {"regular": _vera_font_path()}},
    )

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_hex_color_resolves_theme_zero_to_white():
    color = _hex_color("theme:0:0.0")
    assert color is not None
    assert round(color.red * 255) == 255
    assert round(color.green * 255) == 255
    assert round(color.blue * 255) == 255


def test_pdf_hex_color_resolves_theme_one_to_black():
    color = _hex_color("theme:1:0.0")
    assert color is not None
    assert round(color.red * 255) == 0
    assert round(color.green * 255) == 0
    assert round(color.blue * 255) == 0


def test_pdf_apply_tint_lightens_color():
    r, g, b = _apply_tint((0, 0, 255), 0.5)
    assert r == 127
    assert g == 127
    assert b == 255


def test_pdf_apply_tint_darkens_color():
    r, g, b = _apply_tint((100, 100, 100), -0.5)
    assert r == 50
    assert g == 50
    assert b == 50


def test_pdf_paragraph_respects_wrap_text_false():
    cell = _cell("A1", "line one\nline two")
    result = _paragraph(cell, _FontResolver())
    assert "<br/>" not in result.text


def test_pdf_paragraph_respects_wrap_text_true():
    cell = _cell("A1", "line one\nline two")
    cell["alignment"] = dict(cell["alignment"])
    cell["alignment"]["wrap_text"] = True
    result = _paragraph(cell, _FontResolver())
    assert "<br/>" in result.text


def test_pdf_paragraph_applies_inline_markup_for_font_variants():
    cell = _cell("A1", "note")
    cell["font"] = dict(cell["font"])
    cell["font"]["underline"] = "single"
    cell["font"]["strike"] = True
    cell["font"]["vert_align"] = "superscript"
    result = _paragraph(cell, _FontResolver())
    assert "<u>" in result.text
    assert "<strike>" in result.text
    assert "<super>" in result.text


def test_pdf_table_style_rtl_start_end_borders_map_to_edges():
    cell = _cell("A1", "RTL")
    cell["alignment"] = dict(cell["alignment"])
    cell["alignment"]["reading_order"] = 2
    cell["borders"] = dict(cell["borders"])
    cell["borders"]["start"] = {"style": "medium", "color": "FF000000"}
    cell["borders"]["end"] = {"style": "thin", "color": "FF000000"}
    sheet = _schema({"A1": cell}, dims="A1:A1")["sheets"][0]
    commands = _table_style(sheet, {(1, 1): cell}, 1, 1, 1, 1, _FontResolver()).getCommands()
    assert any(item[0] == "LINEAFTER" for item in commands)
    assert any(item[0] == "LINEBEFORE" for item in commands)


def test_pdf_table_style_pattern_fill_prefers_bg_then_fg_fallback():
    cell = _cell("A1", "Pattern")
    cell["fill"] = {"pattern_type": "gray125", "fg_color": "FFFF0000", "bg_color": "FFFFFFFF"}
    sheet = _schema({"A1": cell}, dims="A1:A1")["sheets"][0]
    commands = _table_style(sheet, {(1, 1): cell}, 1, 1, 1, 1, _FontResolver()).getCommands()
    assert any(item[0] == "BACKGROUND" for item in commands)


def test_pdf_table_style_preserves_supported_cell_styling():
    cell = _cell("A1", "Styled")
    cell["fill"] = {"pattern_type": "solid", "fg_color": "FF123456", "bg_color": None}
    cell["alignment"] = dict(cell["alignment"])
    cell["alignment"]["horizontal"] = "right"
    cell["alignment"]["vertical"] = "center"
    cell["borders"] = dict(cell["borders"])
    cell["borders"]["bottom"] = {"style": "thin", "color": "FF654321"}
    sheet = _schema({"A1": cell}, dims="A1:A1")["sheets"][0]

    commands = _table_style(sheet, {(1, 1): cell}, 1, 1, 1, 1, _FontResolver()).getCommands()

    assert ("ALIGN", (0, 0), (0, 0), "RIGHT") in commands
    assert ("VALIGN", (0, 0), (0, 0), "MIDDLE") in commands
    assert any(item[0] == "BACKGROUND" for item in commands)
    assert any(item[0] == "LINEBELOW" for item in commands)


def test_xlsxwriter_helper_mappings_cover_known_and_unknown_values():
    assert _xlsxwriter_pattern("solid") == 1
    assert _xlsxwriter_pattern("gray125") == 17
    assert _xlsxwriter_pattern("unknown") is None

    assert _xlsxwriter_border_style("thin") == 1
    assert _xlsxwriter_border_style("double") == 6
    assert _xlsxwriter_border_style("unknown") is None

    assert _xlsxwriter_logical_border_sides({"reading_order": 2}) == {"start": "right", "end": "left"}
    assert _xlsxwriter_logical_border_sides({"reading_order": 1}) == {"start": "left", "end": "right"}
    assert _xlsxwriter_vertical_alignment("center") == "vcenter"
    assert _xlsxwriter_vertical_alignment("justify") == "vjustify"

    assert _xlsxwriter_color("FF112233") == "#112233"
    assert _xlsxwriter_color("theme:0:0.0") == "#FFFFFF"
    assert _xlsxwriter_color("theme:bad:0.0") is None


def test_xlsxwriter_format_props_covers_pattern_and_border_variants():
    cell = _cell("A1", 1)
    cell["font"] = dict(cell["font"])
    cell["font"]["color"] = "theme:0:0.0"
    cell["font"]["vert_align"] = "subscript"
    cell["fill"] = {
        "pattern_type": "gray125",
        "fg_color": "FF112233",
        "bg_color": "FF445566",
    }
    cell["alignment"] = dict(cell["alignment"])
    cell["alignment"]["reading_order"] = 2
    cell["alignment"]["vertical"] = "center"
    cell["borders"] = dict(cell["borders"])
    cell["borders"]["start"] = {"style": "thin", "color": "FF010203"}
    cell["borders"]["end"] = {"style": "medium", "color": "FF040506"}
    cell["borders"]["diagonal"] = {"style": "dashed", "color": "FF112233"}
    cell["borders"]["diagonal_up"] = True
    cell["borders"]["diagonal_down"] = True
    cell["number_format"] = "0.00"

    props = _xlsxwriter_format_props(cell, ["FFFFFFFF"] * 12)

    assert props["font_script"] == 2
    assert props["pattern"] == 17
    assert props["fg_color"] == "#112233"
    assert props["bg_color"] == "#445566"
    assert props["right"] == 1
    assert props["left"] == 2
    assert props["diag_type"] == 3
    assert props["num_format"] == "0.00"
    assert props["valign"] == "vcenter"


def test_xlsxwriter_format_props_drops_unresolvable_fill_colors():
    cell = _cell("A1", "x")
    cell["fill"] = {
        "pattern_type": "solid",
        "fg_color": "theme:99:0.0",
        "bg_color": "theme:99:0.0",
    }

    props = _xlsxwriter_format_props(cell, ["FFFFFFFF"])

    assert "pattern" not in props
    assert "fg_color" not in props
    assert "bg_color" not in props


def test_xlsx_resolve_color_helpers_cover_theme_and_passthrough():
    theme = ["FF010203"] * 12
    assert _resolve_color_field("theme:0:0.0", theme) == "FF010203"
    assert _resolve_color_field("FF112233", theme) == "FF112233"
    assert _resolve_color_field("theme:0:0.0", None) == "theme:0:0.0"

    fill = {"pattern_type": "solid", "fg_color": "theme:0:0.0", "bg_color": None}
    font = {"name": "Calibri", "size": 11.0, "color": "theme:0:0.0"}
    borders = {
        "left": {"style": "thin", "color": "theme:0:0.0"},
        "right": {"style": None, "color": None},
    }

    assert _resolved_fill(fill, theme)["fg_color"] == "FF010203"
    assert _resolved_font(font, theme)["color"] == "FF010203"
    assert _resolved_borders(borders, theme)["left"]["color"] == "FF010203"


def test_xlsx_merged_edge_schema_bounds_and_region_synthesis():
    borders = _cell("A1", "x")["borders"]
    borders["top"] = {"style": "thin", "color": "FF000000"}
    borders["left"] = {"style": "medium", "color": "FF000000"}
    borders["diagonal_up"] = True

    edge = _edge_border_schema_bounds(
        borders,
        min_row=1,
        max_row=2,
        min_col=1,
        max_col=2,
        row=1,
        col=1,
    )
    assert edge is not None
    assert edge["top"]["style"] == "thin"
    assert edge["left"]["style"] == "medium"
    assert edge["diagonal_up"] is True

    empty_borders = _cell("A1", "x")["borders"]
    assert (
        _edge_border_schema_bounds(
            empty_borders,
            min_row=1,
            max_row=1,
            min_col=1,
            max_col=1,
            row=1,
            col=1,
        )
        is None
    )

    anchor = _cell("A1", "Merged")
    anchor["borders"] = borders
    sheet = _schema({"A1": anchor}, dims="A1:B2")["sheets"][0]
    sheet["merged_regions"] = ["A1:B2", "C1:D2"]
    edges = _merged_region_edge_schemas(sheet)
    assert (1, 1) in edges
    assert edges[(1, 1)]["value"] == "Merged"


def test_xlsx_apply_merged_borders_uses_cells_override_and_skips_missing_anchor():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:B1")

    anchor = _cell("A1", "Title")
    anchor["borders"]["top"] = {"style": "thin", "color": "FF000000"}
    sheet = _schema({}, dims="A1:B1")["sheets"][0]
    sheet["merged_regions"] = ["A1:B1", "C1:D1"]
    _apply_merged_region_borders(ws, sheet, cells={(1, 1): anchor})
    assert ws["A1"].border.top.style == "thin"


def test_xlsx_apply_streaming_dimensions_handles_fixed_and_even_modes():
    wb = openpyxl.Workbook()
    ws = wb.active

    sheet_fixed = _schema({}, dims="A1:B2")["sheets"][0]
    sheet_fixed["column_width_mode"] = "fixed"
    sheet_fixed["row_height_mode"] = "fixed"
    sheet_fixed["column_widths"] = {"A": 20.0, "B": None}
    sheet_fixed["row_heights"] = {"1": 30.0, "2": None}
    _apply_streaming_dimensions(ws, sheet_fixed)
    assert ws.column_dimensions["A"].width == 20.0
    assert ws.row_dimensions[1].height == 30.0

    sheet_even = _schema({}, dims="A1:C3")["sheets"][0]
    sheet_even["column_width_mode"] = "even"
    sheet_even["row_height_mode"] = "even"
    sheet_even["default_column_width"] = 11.0
    sheet_even["default_row_height"] = 13.0
    _apply_streaming_dimensions(ws, sheet_even)
    assert ws.column_dimensions["C"].width == 11.0
    assert ws.row_dimensions[3].height == 13.0


def test_xlsx_delete_bundle_dir_safety_guards(managed_tmp_dir: Path):
    bundle_dir = managed_tmp_dir / "bundle-dir"
    bundle = mo_dataport.compile(
        _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1"),
        {"Sheet1": {"name": "Alice"}},
        bundle_path=str(bundle_dir),
    )
    assert bundle_dir.exists()
    _delete_bundle_dir(bundle)
    assert not bundle_dir.exists()

    file_path = managed_tmp_dir / "not-dir"
    file_path.write_text("x", encoding="utf-8")
    file_bundle = mo_dataport.compile(
        _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1"),
        {"Sheet1": {"name": "Bob"}},
    )
    object.__setattr__(file_bundle, "path", str(file_path))
    _delete_bundle_dir(file_bundle)

    malformed = managed_tmp_dir / "malformed"
    malformed.mkdir()
    malformed_bundle = mo_dataport.compile(
        _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1"),
        {"Sheet1": {"name": "Carol"}},
    )
    object.__setattr__(malformed_bundle, "path", str(malformed))
    with pytest.raises(ValueError, match="Refusing to delete malformed report bundle directory"):
        _delete_bundle_dir(malformed_bundle)


def test_pdf_export_accepts_bundle_path(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "report_bundle"
    mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}}, bundle_path=str(bundle_path))
    out = managed_tmp_dir / "out.pdf"

    mo_dataport.export(str(bundle_path), str(out), format="pdf")

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_export_auto_delete_removes_bundle_after_success(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "bundle"
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"name": "Alice"}},
        bundle_path=str(bundle_path),
    )
    out = managed_tmp_dir / "out.pdf"

    mo_dataport.export(bundle, str(out), format="pdf", auto_delete_bundle=True)

    assert out.exists()
    assert not bundle_path.exists()


def test_pdf_export_rejects_streaming_export_mode(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="PDF export does not support export_mode"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "out.pdf"),
            format="pdf",
            export_mode="streaming",
        )


def test_pdf_export_rejects_invalid_page_options(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {"name": "Alice"}})

    with pytest.raises(ValueError, match="Unsupported PDF page_size"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "bad-page.pdf"),
            format="pdf",
            page_size="tabloid",
        )
    with pytest.raises(ValueError, match="Unsupported PDF orientation"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "bad-orientation.pdf"),
            format="pdf",
            orientation="sideways",
        )
    with pytest.raises(ValueError, match="margin must be non-negative"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "bad-margin.pdf"),
            format="pdf",
            margin=-1,
        )


def test_pdf_export_supports_landscape_orientation(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "Landscape")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})
    out = managed_tmp_dir / "landscape.pdf"

    mo_dataport.export(
        bundle,
        str(out),
        format="pdf",
        page_size="A4",
        orientation="landscape",
        margin=12,
    )

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_export_rejects_invalid_font_configs(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "Custom Font")}, dims="A1:A1")
    bundle = mo_dataport.compile(schema, {"Sheet1": {}})

    with pytest.raises(TypeError, match="PDF font config must be a path or dict"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "bad-font-type.pdf"),
            format="pdf",
            fonts={"Broken": 42},
        )
    with pytest.raises(ValueError, match="file does not exist"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "missing-font.pdf"),
            format="pdf",
            fonts={"Missing": str(managed_tmp_dir / "missing.ttf")},
        )
    with pytest.raises(ValueError, match="requires a regular font file"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "missing-regular.pdf"),
            format="pdf",
            fonts={"Broken": {"bold": _vera_font_path()}},
        )


def test_load_rejects_malformed_bundle(managed_tmp_dir: Path):
    bad_path = managed_tmp_dir / "bad"
    bad_path.mkdir()
    (bad_path / "manifest.json").write_text("{}", encoding="utf-8")

    with pytest.raises(ValueError, match="Malformed report bundle"):
        load_report_bundle(str(bad_path))


def test_export_auto_delete_removes_bundle_after_success(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{name:string}}")}, dims="A1:A1")
    bundle_path = managed_tmp_dir / "bundle"
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"name": "Alice"}},
        bundle_path=str(bundle_path),
    )
    out = managed_tmp_dir / "out.xlsx"

    mo_dataport.export(bundle, str(out), auto_delete_bundle=True)

    assert out.exists()
    assert not bundle_path.exists()


def test_failed_export_preserves_bundle_directory(managed_tmp_dir: Path):
    schema = _schema({"A1": _cell("A1", "{{rows:dataframe-content}}")}, dims="A1:A1")
    df = polars.DataFrame({"A": [1]})
    bundle_path = managed_tmp_dir / "bundle"
    bundle = mo_dataport.compile(
        schema,
        {"Sheet1": {"rows": df}},
        bundle_path=str(bundle_path),
    )

    with pytest.raises(ValueError, match="Unsupported export_mode"):
        mo_dataport.export(
            bundle,
            str(managed_tmp_dir / "out.xlsx"),
            export_mode="fast",
            auto_delete_bundle=True,
        )

    assert bundle_path.exists()


# §5. Entrypoints
