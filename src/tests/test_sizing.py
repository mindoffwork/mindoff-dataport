"""Tests for column_width_mode and row_height_mode."""

import datetime

import openpyxl
import pytest

from mindoff_dataport import mode

# §1 Constants & Exceptions

# §2 Classes and Sub Classes

# §3 Private Helper Functions


def _minimal_cell(coord, value="Hello"):
    return {
        "coordinate": coord,
        "value": value,
        "cell_type": "string",
        "number_format": None,
        "font": {
            "name": "Calibri",
            "size": 11.0,
            "bold": False,
            "italic": False,
            "underline": None,
            "color": None,
        },
        "fill": {"bg_color": None},
        "alignment": {"horizontal": None, "vertical": None, "wrap_text": False},
        "borders": {
            "top": {"style": None, "color": None},
            "bottom": {"style": None, "color": None},
            "left": {"style": None, "color": None},
            "right": {"style": None, "color": None},
        },
        "merged": False,
        "merge_anchor": None,
    }


def _minimal_sheet(cells_dict, **extra):
    return {
        "name": "Sheet1",
        "dimensions": "A1:C3",
        "merged_regions": [],
        "column_widths": {},
        "row_heights": {},
        "cells": cells_dict,
        **extra,
    }


def _build_and_reload(schema, managed_tmp_dir):
    out = str(managed_tmp_dir / "out.xlsx")
    bundle = mode.compile(schema, {"Sheet1": {}})
    mode.export(bundle, out)
    return openpyxl.load_workbook(out)


# §4 Public Functions


def test_even_columns_applies_uniform_width(managed_tmp_dir):
    cells = {c: _minimal_cell(c) for c in ["A1", "B1", "C1"]}
    schema = {
        "sheets": [
            _minimal_sheet(cells, column_width_mode="even", default_column_width=25.0)
        ]
    }
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    for col in ["A", "B", "C"]:
        assert ws.column_dimensions[col].width == pytest.approx(25.0)
    wb.close()


def test_even_columns_default_width_when_not_specified(managed_tmp_dir):
    cells = {"A1": _minimal_cell("A1")}
    schema = {"sheets": [_minimal_sheet(cells, column_width_mode="even")]}
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    assert ws.column_dimensions["A"].width == pytest.approx(15.0)
    wb.close()


def test_hug_columns_wider_for_longer_content(managed_tmp_dir):
    cells = {
        "A1": _minimal_cell("A1", value="Hi"),
        "B1": _minimal_cell("B1", value="A very long header label here"),
    }
    schema = {"sheets": [_minimal_sheet(cells, column_width_mode="hug")]}
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width
    wb.close()


def test_hug_columns_bold_factor(managed_tmp_dir):
    plain = _minimal_cell("A1", value="SameText")
    bold = _minimal_cell("B1", value="SameText")
    bold["font"] = dict(bold["font"])
    bold["font"]["bold"] = True
    cells = {"A1": plain, "B1": bold}
    schema = {"sheets": [_minimal_sheet(cells, column_width_mode="hug")]}
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width
    wb.close()


def test_even_rows_applies_uniform_height(managed_tmp_dir):
    cells = {f"A{r}": _minimal_cell(f"A{r}") for r in range(1, 4)}
    schema = {
        "sheets": [
            _minimal_sheet(cells, row_height_mode="even", default_row_height=30.0)
        ]
    }
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    for row in range(1, 4):
        assert ws.row_dimensions[row].height == pytest.approx(30.0)
    wb.close()


def test_hug_rows_larger_for_bigger_font(managed_tmp_dir):
    small = _minimal_cell("A1", value="text")
    large = _minimal_cell("A2", value="text")
    large["font"] = dict(large["font"])
    large["font"]["size"] = 24.0
    cells = {"A1": small, "A2": large}
    schema = {"sheets": [_minimal_sheet(cells, row_height_mode="hug")]}
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    assert ws.row_dimensions[2].height > ws.row_dimensions[1].height
    wb.close()


def test_fixed_mode_uses_explicit_widths(managed_tmp_dir):
    cells = {"A1": _minimal_cell("A1")}
    schema = {
        "sheets": [
            {
                "name": "Sheet1",
                "dimensions": "A1:A1",
                "merged_regions": [],
                "cells": cells,
                "column_widths": {"A": 42.0},
                "row_heights": {},
            }
        ]
    }
    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active
    assert ws.column_dimensions["A"].width == pytest.approx(42.0)
    wb.close()


def test_xlsx_export_preserves_cell_styles_and_date_values(managed_tmp_dir):
    styled = _minimal_cell("A1", "Styled")
    styled["font"] = {
        "name": "Calibri",
        "size": 14.0,
        "bold": True,
        "italic": True,
        "underline": "single",
        "color": "FFFF0000",
    }
    styled["fill"] = {"bg_color": "FF00FF00"}
    styled["alignment"] = {
        "horizontal": "center",
        "vertical": "bottom",
        "wrap_text": True,
    }
    styled["borders"] = {
        "top": {"style": "thin", "color": "FF000000"},
        "bottom": {"style": "thick", "color": "FF000000"},
        "left": {"style": "medium", "color": "FF000000"},
        "right": {"style": "dashed", "color": "FF000000"},
    }
    date_cell = _minimal_cell("A2", "2024-01-02T03:04:05")
    date_cell["cell_type"] = "date"
    date_cell["number_format"] = "yyyy-mm-dd hh:mm"
    schema = {"sheets": [_minimal_sheet({"A1": styled, "A2": date_cell})]}

    wb = _build_and_reload(schema, managed_tmp_dir)
    ws = wb.active

    assert ws["A1"].font.bold is True
    assert ws["A1"].font.italic is True
    assert ws["A1"].font.underline == "single"
    assert ws["A1"].font.color.rgb == "FFFF0000"
    assert ws["A1"].fill.fgColor.rgb == "FF00FF00"
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A1"].alignment.vertical == "bottom"
    assert ws["A1"].alignment.wrap_text is True
    assert ws["A1"].border.top.style == "thin"
    assert ws["A1"].border.bottom.style == "thick"
    assert ws["A1"].border.left.style == "medium"
    assert ws["A1"].border.right.style == "dashed"
    assert ws["A2"].value == datetime.datetime(2024, 1, 2, 3, 4, 5)
    assert ws["A2"].number_format == "yyyy-mm-dd hh:mm"
    wb.close()


# §5 Entrypoints
